"""Approved administrative presentation and complete old/new scope bindings.

SPT form sheets, original captures and completion decisions are owned elsewhere.
"""
from collections import Counter
from copy import copy, deepcopy
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.chart import Reference
from openpyxl.utils import get_column_letter

RESOURCE = Path(__file__).resolve().parent / 'resources' / 'resume_format.xlsx'
LABELS = ('NPWP', 'Tahun', 'Wajib Pajak', 'Dasar Tahun', 'Status',
          'Total Pilihan SPT', 'SPT disedot status COMPLETE', 'Waktu mulai')
LEGACY_LABELS = ('NPWP', 'Tahun', 'Wajib Pajak', 'Dasar Tahun', 'Status',
                 'SPT dalam lingkup', 'SPT berhasil ditangkap', 'Waktu mulai')
INDEX_HEADERS = ('No.', 'Jenis Pajak', 'Jenis SPT', 'Masa Pajak', 'Model / Pembetulan',
                 'Status Daftar', 'Bukti Penerimaan', 'Tanggal Submit',
                 'Hasil Capture', 'Bagian', 'URL Sumber')
AUDIT_HEADERS = ('Status / Objek', 'Total sumber', 'Baris dibaca', 'Dipilih', 'Hasil', 'Keterangan')


def summary_scope(workbook):
    """Read one complete known layout; mixed headers/coordinates never qualify."""
    if not {'Ringkasan', 'Daftar SPT', 'Audit'}.issubset(workbook.sheetnames):
        raise ValueError('Workbook scope/NPWP has no supported administrative layout')
    summary, index, audit = (workbook[name] for name in ('Ringkasan', 'Daftar SPT', 'Audit'))
    candidates = [start for start, labels in ((2, LABELS), (4, LEGACY_LABELS))
                  if tuple(summary.cell(start + i, 1).value for i in range(8)) == labels]
    if len(candidates) != 1:
        raise ValueError('Workbook scope/NPWP has incomplete or mixed layout labels')
    start = candidates[0]
    audit_row = 3 if start == 2 else 4
    if (tuple(index.cell(4, c).value for c in range(1, 12)) != INDEX_HEADERS
            or tuple(audit.cell(audit_row, c).value for c in range(1, 7)) != AUDIT_HEADERS):
        raise ValueError('Workbook administrative headers cannot prove scope/NPWP')
    values = [summary.cell(start + i, 2).value for i in range(8)]
    npwp, year, _, basis, status, count, captured, _ = values
    if (not isinstance(npwp, str) or not re.fullmatch(r'[0-9]{16}', npwp)
            or summary.cell(start, 2).data_type != 's'):
        raise ValueError('Workbook NPWP is not literal sixteen-digit text')
    if type(year) is not int or not 1900 <= year <= 2100:
        raise ValueError('Workbook tax year is invalid')
    if basis not in ('Masa/Tahun Pajak', 'Tanggal Submit SPT'):
        raise ValueError('Workbook year basis is invalid')
    if (status not in ('COMPLETE', 'PARTIAL') or type(count) is not int or count < 0
            or type(captured) is not int or not 0 <= captured <= count
            or index.max_row != 4 + count):
        raise ValueError('Workbook completion/count scope is invalid')
    return {'npwp': npwp, 'year': year,
            'year_basis': 'period' if basis == 'Masa/Tahun Pajak' else 'submission',
            'complete': status == 'COMPLETE', 'record_count': count,
            'captured_count': captured, 'layout': 'operator-v1' if start == 2 else 'legacy'}


def _style(target, reference):
    # Style IDs are workbook-local. Copy the style components, never their IDs.
    for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
        setattr(target, attribute, copy(getattr(reference, attribute)))
    target.number_format = reference.number_format


def _literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'


def _reset_sheet(sheet, reference):
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))
    sheet._cells.clear()
    sheet._charts.clear()
    sheet.tables.clear()
    sheet.row_dimensions.clear()
    sheet.column_dimensions.clear()
    for key, dimension in reference.column_dimensions.items():
        sheet.column_dimensions[key] = copy(dimension)
    for key, dimension in reference.row_dimensions.items():
        sheet.row_dimensions[key] = copy(dimension)
    for attribute in ('sheet_format', 'sheet_properties', 'views', 'page_setup',
                      'page_margins', 'print_options'):
        setattr(sheet, attribute, deepcopy(getattr(reference, attribute)))
    sheet.freeze_panes = reference.freeze_panes
    for merged in reference.merged_cells.ranges:
        sheet.merge_cells(str(merged))


def apply_resume_layout(workbook, run):
    """Remap only administrative cells, retaining their literal values and links."""
    reference = load_workbook(RESOURCE)
    try:
        # Keep the already-bound record/PDF links, not stale reference links.
        data = {}
        for name in ('Daftar SPT', 'Audit'):
            data[name] = [(cell.row, cell.column, cell.value, cell.data_type, copy(cell.hyperlink))
                          for row in workbook[name] for cell in row
                          if cell.value is not None or cell.hyperlink is not None]
        for name in ('Ringkasan', 'Daftar SPT', 'Audit'):
            _reset_sheet(workbook[name], reference[name])
        summary = workbook['Ringkasan']
        source = reference['Ringkasan']
        records = run['records']
        values = [run['npwp'], run['year'], run.get('taxpayer_name', ''),
                  'Masa/Tahun Pajak' if run['year_basis'] == 'period' else 'Tanggal Submit SPT',
                  'COMPLETE' if run['complete'] else 'PARTIAL', len(records),
                  sum(bool(record.get('capture_complete')) for record in records),
                  run.get('started_at', '')]
        for row, (label, value) in enumerate(zip(LABELS, values), 2):
            for col, item in ((1, label), (2, value)):
                cell = summary.cell(row, col)
                _style(cell, source.cell(row, col))
                _literal(cell, item)
        for coordinate in ('A1', 'A13', 'B13'):
            cell = summary[coordinate]
            _style(cell, source[coordinate])
            _literal(cell, source[coordinate].value)
        # The single reference scope slot replaces the older multirow banner.
        if run.get('scope_description'):
            _style(summary['A11'], source['A18'])
            _literal(summary['A11'], 'Ruang lingkup')
            _style(summary['B11'], source['B4'])
            _literal(summary['B11'], run['scope_description'])
        counts = Counter(record['meta'].get('Jenis Pajak', 'Tidak diketahui') for record in records)
        for row, (tax, count) in enumerate(sorted(counts.items()), 14):
            for col, value in ((1, tax), (2, count)):
                cell = summary.cell(row, col)
                _style(cell, source.cell(14, col))
                _literal(cell, value)
        mapping_row = 14 + len(counts)
        reports = run.get('template_mapping', [])
        if reports:
            _style(summary.cell(mapping_row, 1), source['A18'])
            _style(summary.cell(mapping_row, 2), source['B18'])
            _literal(summary.cell(mapping_row, 1), source['A18'].value)
            _literal(summary.cell(mapping_row, 2),
                     'COMPLETE' if all(report['complete'] for report in reports) else 'PARTIAL')
        if counts:
            chart = deepcopy(source._charts[0])
            chart.series = []
            chart.add_data(Reference(summary, min_col=2, min_row=13, max_row=13 + len(counts)),
                           titles_from_data=True)
            chart.set_categories(Reference(summary, min_col=1, min_row=14, max_row=13 + len(counts)))
            summary.add_chart(chart, 'D13')
        index = workbook['Daftar SPT']
        source = reference['Daftar SPT']
        last = 4 + len(records)
        for row, col, value, kind, link in data['Daftar SPT']:
            if row > last:
                continue
            cell = index.cell(row, col)
            style_row = row if row <= 4 else 8 if row == last else 5
            _style(cell, source.cell(style_row, col))
            _literal(cell, value)
            cell.data_type = kind
            cell.hyperlink = link
        if records:
            table = deepcopy(next(iter(source.tables.values())))
            table.ref = f'A4:K{last}'
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref
            index.add_table(table)
        audit = workbook['Audit']
        source = reference['Audit']
        header_rows = {3}
        if reports:
            header_rows.add(4 + len(run.get('audit', [])) + len(run.get('errors', [])))
        if any(record.get('attachments') for record in records):
            header_rows.add(4 + len(run.get('audit', [])) + len(run.get('errors', []))
                            + (1 + len(reports) if reports else 0))
        for row, col, value, kind, link in data['Audit']:
            if row in (2, 3):
                continue
            row = row - 1 if row >= 4 else row
            cell = audit.cell(row, col)
            _style(cell, source.cell(3 if row in header_rows else 1 if row == 1 else 4, col))
            _literal(cell, value)
            cell.data_type = kind
            cell.hyperlink = link
        index.print_title_rows = '1:4'
        audit.print_title_rows = '1:3'
        for name, width in (('Ringkasan', 8), ('Daftar SPT', 11), ('Audit', 6)):
            sheet = workbook[name]
            sheet.print_area = f'A1:{get_column_letter(width)}{sheet.max_row}'
    finally:
        reference.close()
