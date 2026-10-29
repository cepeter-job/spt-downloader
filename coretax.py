"""Orkestrasi CLI Coretax dan publikasi Excel yang memeriksa kembali hasilnya.

Login CLI diselesaikan pengguna di browser. Validasi data bersama dipakai oleh
CLI dan worker desktop; batas keamanannya dijelaskan di docs/arsitektur.md.
"""
from spt_desktop import diagnostics
import argparse
from contextlib import contextmanager, ExitStack
from datetime import datetime
import json
import os
import shutil
import tempfile
from pathlib import Path
import sys
import time
from urllib.parse import urlparse

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from coretax_exporter import (validate_scope, validate_export_run, export_split_workbooks,
                             workbook_groups, detail_sheet_title, record_sheet_titles,
                             workbook_proof, read_workbook_proof, workbook_sheet_titles,
                             export_workbook)
from coretax_browser import SEARCH_URL, collect_inventory, collect_details, _atomic_checkpoint
from coretax_pdf import attachment_declarations, pdf_custody, copy_held_pdfs
from coretax_json_capture import (JSON_OUTPUT_LAYOUT, json_output_names, json_custody,
                                  copy_held_json, verify_existing_json)
from coretax_resume_layout import summary_scope
from coretax_mpn import mpn_filename, mpn_custody, require_mpn
from coretax_mpn_browser import collect_mpn
from coretax_mpn_export import (read_mpn_proof, verify_mpn_workbook,
                                read_resume_mpn, resume_mpn_declaration)

BASE = Path(__file__).resolve().parent


def verify_workbook(path, run, *, include_details=True, include_administration=True):
    """Cocokkan scope, status, jumlah baris, dan lembar dengan data run."""
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if not include_administration:
            proof = read_workbook_proof(wb)
            if proof is None or proof != workbook_proof(run):
                raise ValueError('Workbook scope/PDF proof differs from selected records')
            return set(wb.sheetnames)
        scope = summary_scope(wb)
        checks = (
            (scope['npwp'] == run['npwp'], 'NPWP mismatch in output'),
            (scope['year'] == run['year'], 'Tax year mismatch in output'),
            (scope['year_basis'] == run['year_basis'], 'Year basis mismatch in output'),
            (scope['complete'] == run['complete'], 'Incorrect completion status'),
            (scope['record_count'] == len(run['records']), 'Excel scope count mismatch'),
            (wb['Daftar SPT'].max_row == 4 + len(run['records']), 'Excel index count mismatch'),
            (len(wb.sheetnames) == (3 + sum(len(record_sheet_titles(i, r)) for i, r in enumerate(run['records'], 1)) if include_details else 3),
             'Incorrect worksheet count'),
        )
        for valid, message in checks:
            if not valid:
                raise ValueError(message)
        return set(wb.sheetnames)
    finally:
        wb.close()


def output_path(folder, name):
    """Tolak alias Windows dan komponen reparse sebelum membaca atau mengganti hasil."""
    if (not isinstance(name, str) or not name or '\\' in name or ':' in name
            or any(part in ('', '.', '..') or part.endswith((' ', '.')) for part in name.split('/'))
            or Path(name).is_absolute()):
        raise ValueError('Invalid workbook filename')
    root = Path(folder).absolute()
    path = root / name
    for component in [root, *[root.joinpath(*name.split('/')[:i]) for i in range(1, len(name.split('/'))+1)]]:
        if component.exists() or component.is_symlink():
            if component.is_symlink() or getattr(component.lstat(), 'st_file_attributes', 0) & 1024:
                raise ValueError('Invalid source workbook path: symbolic link or reparse point')
    if not path.resolve().is_relative_to(root.resolve()):
        raise ValueError('Invalid source workbook path')
    return path


def workbook_files(folder, exclude=None):
    folder = Path(folder).absolute()
    # Normalisasi hanya pembanding; path yang dibaca tetap mempertahankan symlink/"..".
    exclude = Path(os.path.abspath(exclude)) if exclude is not None else None
    found = set()
    stack = [folder]
    while stack:
        current = stack.pop()
        for entry in os.scandir(current):
            path = Path(entry.path)
            if exclude is not None and Path(os.path.abspath(path)) == exclude:
                continue
            if entry.is_symlink() or getattr(entry.stat(follow_symlinks=False), 'st_file_attributes', 0) & 1024:
                raise ValueError('Invalid source workbook path: reparse point')
            if entry.is_dir(follow_symlinks=False):
                stack.append(path)
            elif path.suffix.casefold() == '.xlsx':
                found.add(path.relative_to(folder).as_posix())
    return found


def _legacy_resume_proof(workbook, run):
    """Read the old administrative layout without inventing form or PDF metadata."""
    if not {'Ringkasan', 'Daftar SPT', 'Audit'}.issubset(workbook.sheetnames):
        raise ValueError('Retained workbook has no supported scope proof')
    scope = summary_scope(workbook)
    if any(scope[key] != run[key] for key in ('npwp', 'year', 'year_basis')):
        raise ValueError('Retained workbook scope/NPWP cannot authorize global index')
    rows = list(workbook['Daftar SPT'].iter_rows(min_row=5, max_col=11))
    forms = [s for s in workbook.sheetnames if s not in {'Ringkasan', 'Daftar SPT', 'Audit'}]
    locations = {"#'" + s.replace("'", "''") + "'!A1": index for index, s in enumerate(forms)}
    starts = []
    for row in rows:
        link = row[0].hyperlink
        if not link or link.target not in locations or row[8].value not in ('OK', 'PARTIAL'):
            raise ValueError('Retained legacy first-form link or capture status cannot be verified')
        starts.append(locations[link.target])
    if (not starts or starts[0] != 0 or starts != sorted(set(starts))
            or scope['record_count'] != len(rows)):
        raise ValueError('Retained legacy index count/order cannot be verified')
    records = []
    for row, first, end in zip(rows, starts, starts[1:] + [len(forms)]):
        records.append({'meta': {'Jenis Pajak': row[1].value, 'Masa Pajak': row[3].value,
                                 'Model SPT': row[4].value},
                        'capture_complete': row[8].value == 'OK',
                        'source_status': row[5].value or '',
                        'sections': [{'name': name} for name in forms[first:end]]})
    return {'npwp': run['npwp'], 'year': run['year'], 'year_basis': run['year_basis'],
            'records': records, 'complete': scope['complete']}


def folder_resume_run(run, folder):
    """Extend the index from verified retained form proofs, never inferred records."""
    selected = set(workbook_groups(run)) | {'resume.xlsx'}
    if mpn_filename(run) is not None:
        selected.add(mpn_filename(run))
    retained = sorted(workbook_files(folder) - selected)
    if not retained:
        return run
    combined = {**run, 'records': list(run['records']), 'complete': run['complete'],
                'scope_description': 'Indeks seluruh workbook terverifikasi dalam folder; mencakup hasil pengambilan sebelumnya.'}
    for name in retained:
        with output_path(folder, name).open('rb') as stream:
            wb = load_workbook(stream, data_only=False)
            try:
                mpn_proof = read_mpn_proof(wb)
                if mpn_proof is not None:
                    if (mpn_proof['scope'] != [run['npwp'], run['year'], run['year_basis']]
                            or name != f"MPN_{run['year']}.xlsx" or 'retained_mpn' in combined):
                        raise ValueError('Retained MPN scope/filename cannot authorize global index')
                    combined['retained_mpn'] = {'filename': name, 'status': mpn_proof['status'],
                                                'scope': mpn_proof['scope']}
                    combined['complete'] = combined['complete'] and mpn_proof['status'] == 'COMPLETE'
                    continue
                proof = read_workbook_proof(wb)
                if proof is None:
                    proof = _legacy_resume_proof(wb, run)
                if any(proof[k] != run[k] for k in ('npwp', 'year', 'year_basis')):
                    raise ValueError('Retained workbook scope cannot be verified for global index')
                if set(workbook_groups(proof)) != {name}:
                    raise ValueError('Retained workbook filename does not match its scope')
                combined['records'].extend(proof['records'])
                combined['complete'] = combined['complete'] and proof['complete']
            finally:
                wb.close()
    validate_export_run(combined)
    return combined


def write_folder_resume(run, existing, destination):
    combined = folder_resume_run(run, existing)
    targets = {}
    ordinals = {id(record): index for index, record in enumerate(combined['records'], 1)}
    selected_files = set(workbook_groups(run))
    for filename, records in workbook_groups(combined).items():
        selected = filename in selected_files
        title_plan = workbook_sheet_titles(records, presentation=selected)
        if not selected:
            # Rekaman lama mendeskripsikan sheet fisik yang sudah ada; aturan tampilan
            # baru tidak boleh mengubah target tautan pada workbook yang dipertahankan.
            retained = load_workbook(output_path(existing, filename), read_only=True)
            try:
                actual = set(retained.sheetnames)
                if read_workbook_proof(retained) is None:
                    actual -= {'Ringkasan', 'Daftar SPT', 'Audit'}
            finally:
                retained.close()
            if {title for titles in title_plan for title in titles} != actual:
                title_plan = workbook_sheet_titles(records, legacy=True, presentation=False)
            if {title for titles in title_plan for title in titles} != actual:
                raise ValueError('Retained workbook names differ from its physical proof')
        for record, titles in zip(records, title_plan):
            targets[ordinals[id(record)]] = (filename, titles[0])
    export_workbook(combined, destination, include_details=False, detail_targets=targets)
    verify_workbook(destination, combined, include_details=False)


def _verify_workbook_outputs(run, folder, *, allow_extra=False):
    """Verifikasi set Excel dan tautan pembetulan, termasuk saat Python memakai -O.

    Pemeriksaan ini memakai exception eksplisit agar optimasi interpreter tidak
    menghilangkan syarat publikasi hasil CLI maupun desktop.
    """
    folder = Path(folder)
    validate_export_run(run)
    grouped = workbook_groups(run)
    selected_files = set(grouped)
    expected = {'resume.xlsx', *grouped}
    mpn_name = mpn_filename(run)
    if mpn_name is not None:
        expected.add(mpn_name)
    declared = run.get('output_files', list(expected))
    if (not isinstance(declared, list) or
            any(not isinstance(n, str) or not n.endswith('.xlsx') for n in declared)):
        raise ValueError('Invalid workbook filename declaration')
    if len(declared) != len(set(declared)) or set(declared) != expected:
        raise ValueError('Workbook filename/count mismatch')
    for name in declared:
        output_path(folder, name)
    actual = workbook_files(folder)
    if (not allow_extra and actual != expected) or not expected.issubset(actual):
        raise ValueError('Workbook file set differs from the selected scope')
    if any((folder / name).is_symlink() or not (folder / name).is_file() for name in expected):
        raise ValueError('Invalid source workbook path')
    index_run = folder_resume_run(run, folder) if allow_extra else run
    verify_workbook(folder / 'resume.xlsx', index_run, include_details=False)
    sheet_names = {
        filename: verify_workbook(folder / filename, {**run, 'records': records}, include_administration=False)
        for filename, records in grouped.items()
    }
    if mpn_name is not None:
        verify_mpn_workbook(run, output_path(folder, mpn_name))
    grouped = workbook_groups(index_run)
    for filename in grouped.keys() - sheet_names.keys():
        wb = load_workbook(output_path(folder, filename), read_only=True)
        try:
            sheet_names[filename] = set(wb.sheetnames)
            if read_workbook_proof(wb) is None:
                sheet_names[filename] -= {'Ringkasan', 'Daftar SPT', 'Audit'}
        finally:
            wb.close()
    first_titles = {}
    for filename, records in grouped.items():
        selected = filename in selected_files
        titles = workbook_sheet_titles(records, presentation=selected)
        if {title for group in titles for title in group} != sheet_names[filename]:
            titles = workbook_sheet_titles(records, legacy=True, presentation=selected)
        if {title for group in titles for title in group} != sheet_names[filename]:
            raise ValueError('Workbook names do not match a supported physical form layout')
        first_titles.update({id(record): group[0] for record, group in zip(records, titles)})
    # Mode read-only openpyxl tidak menyediakan hyperlink; baca indeks secara normal.
    resume = load_workbook(folder / 'resume.xlsx')
    try:
        if read_resume_mpn(resume) != resume_mpn_declaration(index_run):
            raise ValueError('MPN resume declaration differs from its physical workbook')
        for ordinal, record in enumerate(index_run['records'], 1):
            link = resume['Daftar SPT'].cell(ordinal + 4, 1).hyperlink
            expected_file = next(name for name, records in grouped.items() if any(r is record for r in records))
            if not link or link.target != expected_file or not isinstance(link.location, str):
                raise ValueError('Missing or incorrect type-workbook link')
            linked_sheet = link.location.rsplit('!', 1)[0][1:-1].replace("''", "'")
            expected_sheet = first_titles[id(record)]
            expected_location = "'" + expected_sheet.replace("'", "''") + "'!A1"
            if linked_sheet not in sheet_names[expected_file] or link.location != expected_location:
                raise ValueError('Resume links to an incorrect detail sheet or amendment')
    finally:
        resume.close()


@diagnostics.traced('verify')
def verify_outputs(run, folder, *, allow_extra=False):
    """Verify declared held PDF bytes as well as the existing workbook scope."""
    with pdf_custody(run, folder), json_custody(run, folder):
        _verify_workbook_outputs(run, folder, allow_extra=allow_extra)


def _check_output_scope(run, folder, staging=None):
    """Pastikan checkpoint dan header Excel lama membuktikan scope yang sama."""
    verify_existing_json(run, folder)
    scope_file = folder / 'run.json'
    if scope_file.is_symlink():
        raise FileExistsError('Output scope cannot be a symbolic link')
    previous = None
    if scope_file.exists():
        previous = json.loads(scope_file.read_text(encoding='utf-8'))
        scope_keys = ('npwp', 'year', 'year_basis')
        if (not isinstance(previous, dict) or
                any(previous.get(key) != run.get(key) for key in scope_keys)):
            raise FileExistsError('Output scope differs from existing run.json; use a new output folder')
    expected = {'resume.xlsx', *workbook_groups(run)}
    if mpn_filename(run) is not None:
        expected.add(mpn_filename(run))
    folded = {name.casefold(): name for name in expected}
    for relative in workbook_files(folder, exclude=staging):
        path = output_path(folder, relative)
        if previous is None:
            raise FileExistsError('Existing workbook has no verified output scope; use a new output folder')
        if relative.casefold() in folded and relative != folded[relative.casefold()]:
            raise FileExistsError('Existing workbook filename collision; use a new output folder')
        if relative not in expected:
            raise FileExistsError('Existing workbooks differ from selected output scope; use a new output folder')
        if path.is_symlink() or not path.is_file():
            raise FileExistsError('Invalid existing workbook path; use a new output folder')
        # run.json juga menjadi checkpoint dan bisa lebih baru daripada set Excel.
        # Header Excel lama harus membuktikan scope secara independen.
        try:
            workbook = load_workbook(path, read_only=True)
            try:
                mpn_proof = read_mpn_proof(workbook)
                proof = None if mpn_proof is not None else read_workbook_proof(workbook)
                if mpn_proof is not None:
                    matches = (relative == f"MPN_{run['year']}.xlsx"
                               and mpn_proof['scope'] == [run['npwp'], run['year'], run['year_basis']])
                elif proof is not None:
                    matches = (all(proof[key] == run[key] for key in ('npwp', 'year', 'year_basis'))
                               and set(workbook_groups(proof)) == {relative})
                else:
                    scope = summary_scope(workbook)
                    matches = all(scope[key] == run[key] for key in ('npwp', 'year', 'year_basis'))
            finally:
                workbook.close()
        except Exception as error:
            raise FileExistsError('Existing workbook scope cannot be verified; use a new output folder') from error
        if not matches:
            raise FileExistsError('Existing workbook scope differs; use a new output folder')


def _publish_output(staging, folder, names):
    """Pulihkan penggantian terdahulu bila publikasi berkas berikutnya gagal.

    Backup membantu saat exception atau berkas dikunci Excel. Ini bukan jaminan
    transaksi atomik ketika proses dimatikan atau listrik padam.
    """
    backup = staging / 'backup'
    backup.mkdir()
    for name in names:
        existing = output_path(folder, name)
        if existing.exists():
            (backup / name).parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(existing, backup / name)
    attempted = []
    try:
        for name in names:
            target = output_path(folder, name)
            target.parent.mkdir(parents=True, exist_ok=True)
            attempted.append(name)
            (staging / name).replace(target)
    except BaseException as error:
        diagnostics.record('publish.rollback', phase='rollback', error=error, level='WARNING')
        try:
            for name in reversed(attempted):
                if (staging / name).exists():
                    continue  # Rename gagal sebelum memakai sumber dari staging.
                saved = backup / name
                if saved.exists():
                    saved.replace(folder / name)
                else:
                    (folder / name).unlink(missing_ok=True)
        except BaseException as recovery_error:
            # Kunci eksternal atau interupsi juga bisa menghalangi rollback.
            raise RuntimeError(f'Output rollback blocked; recovery files retained at {staging}') from recovery_error
        raise


@diagnostics.traced('export')
def save_output(run, folder, *, json_source=None):
    """Validasi, tulis ke staging, baca ulang, lalu publikasikan seluruh set."""
    folder = Path(folder)
    validate_export_run(run)
    folder.mkdir(parents=True, exist_ok=True)
    _check_output_scope(run, folder)
    staging = Path(tempfile.mkdtemp(prefix='.spt-export-', dir=folder))
    retain_recovery = False
    try:
        current = {**run, 'output_attachments': attachment_declarations(run),
                   'json_output_layout': JSON_OUTPUT_LAYOUT}
        current['output_json'] = json_output_names(current)
        evidence_source = Path(json_source) if json_source is not None else folder
        with mpn_custody(current, evidence_source), pdf_custody(current, folder) as (attachments, streams):
            # Unduhan sudah menempatkan PDF sumber privat. Tahan stream terverifikasi
            # selama publikasi workbook/checkpoint; ekspor ulang Excel tidak mengganti PDF.
            copy_held_pdfs(attachments, streams, staging)
            with json_custody(run, evidence_source,
                    published=json_source is not None and 'output_json' in run,
                    target_layout=JSON_OUTPUT_LAYOUT) as (json_items, json_streams):
                copy_held_json(json_items, json_streams, staging)
            current['template_mapping'] = []
            files = export_split_workbooks(current, staging, mapping_reports=current['template_mapping'])
            saved_run = {**current, 'output_files': list(files)}
            verify_outputs(saved_run, staging)
            (staging / 'run.json').write_text(json.dumps(saved_run, ensure_ascii=False, indent=2), encoding='utf-8')
            _check_output_scope(run, folder, staging)
            # Simpan staging/backup sejak publikasi dimulai agar interupsi atau
            # rollback yang terhalang tetap menyisakan bahan pemulihan.
            retain_recovery = True
            _publish_output(staging, folder, [*current['output_json'], *files, 'run.json'])
            retain_recovery = False
            run['output_files'] = list(files)
            run['output_attachments'] = current['output_attachments']
            run['output_json'] = current['output_json']
            run['json_output_layout'] = current['json_output_layout']
            run['template_mapping'] = current.get('template_mapping', [])
    finally:
        if not retain_recovery:
            shutil.rmtree(staging)
    return folder / 'resume.xlsx'


def validate_debug_endpoint(value):
    """Terima hanya origin loopback eksplisit tanpa kredensial atau rute URL."""
    message = '--reuse-browser must be a localhost HTTP debugging origin without credentials, path or query'
    if not isinstance(value, str) or any(ord(c) <= 32 or ord(c) == 127 for c in value):
        raise ValueError(message)
    try:
        endpoint = urlparse(value)
        port = endpoint.port
    except ValueError as error:
        raise ValueError(message) from error
    if (endpoint.scheme not in ('http', 'https') or
            endpoint.hostname not in ('127.0.0.1', 'localhost', '::1') or
            endpoint.username is not None or endpoint.password is not None or
            endpoint.path not in ('', '/') or endpoint.params or endpoint.query or endpoint.fragment or
            endpoint.netloc.endswith(':') or port == 0):
        raise ValueError(message)
    return value


def authenticate(page, noninteractive):
    page.goto(SEARCH_URL, wait_until='domcontentloaded', timeout=60000)
    if page.locator('#ObjectPermitNumber').is_visible():
        return
    page.bring_to_front()
    if noninteractive:
        # Tunggu render kontrol pencarian SPA; halaman login tidak pernah dilewati.
        page.locator('#ObjectPermitNumber').wait_for(state='visible', timeout=20000)
        return
    print('\nSign in directly in the browser window. Complete password, CAPTCHA and MFA there.')
    print('Do not enter any password or verification code into this script.')
    input('After sign-in, press Enter here to continue: ')
    page.goto(SEARCH_URL, wait_until='domcontentloaded', timeout=60000)
    page.locator('#ObjectPermitNumber').wait_for(state='visible', timeout=45000)


class _DetailSource(argparse.Action):
    def __call__(self, parser, namespace, values, option_string=None):
        setattr(namespace, self.dest, values)
        namespace.detail_source_explicit = True


def parser():
    cli = argparse.ArgumentParser(description='Read-only authorized Coretax SPT export by exact NPWP and selected tax year.')
    cli.add_argument('--npwp', help='Exactly 16 digits; no normalization or separators')
    cli.add_argument('--year', help='Four-digit tax year')
    cli.add_argument('--year-basis', choices=['period', 'submission'], default='period', help='Default: tax period/year, not submission date')
    cli.add_argument('--output', type=Path, help='Output folder; default is a timestamped local folder')
    cli.add_argument('--resume', type=Path, help='Resume a previous output folder from verified checkpoints')
    cli.add_argument('--from-json', type=Path, help='Rebuild Excel offline from an existing run.json')
    cli.add_argument('--browser', choices=['msedge', 'chrome'], default='msedge')
    cli.add_argument('--profile', type=Path, default=BASE / 'browser-profile', help='Dedicated persistent profile, never your personal browser profile')
    cli.add_argument('--reuse-browser', nargs='?', const='http://127.0.0.1:9235', help='Connect to an already opened local automation browser')
    cli.add_argument('--non-interactive', action='store_true', help='Requires NPWP/year and an authenticated session; never prompts for credentials')
    cli.add_argument('--inventory-only', action='store_true', help='Capture SPT list only; output is explicitly PARTIAL for full-detail scope')
    cli.add_argument('--detail-limit', type=int, help='Test-only detail limit; output stays PARTIAL if any return remains unexported')
    cli.set_defaults(detail_source_explicit=False)
    cli.add_argument('--detail-source', choices=['dom', 'json-first'], default='dom', action=_DetailSource,
                     help='DOM by default; opt-in JSON hybrid with whole-record DOM fallback')
    cli.add_argument('--max-pages', type=int, default=1000, help='Safety cap per list/table; reaching it is reported as PARTIAL')
    return cli


@contextmanager
def _cli_cleanup():
    """Close every owned resource without replacing an active error or Stop."""
    stack = ExitStack()
    failed = interrupted = False

    def register(callback, *args):
        def close():
            nonlocal failed, interrupted
            try:
                callback(*args)
            except KeyboardInterrupt as error:
                interrupted = True
                diagnostics.record('operation.cancelled', error=error, phase='cli', status='CANCELLED')
            except Exception as error:
                # Each external close is containment only; it cannot authorize
                # capture/publication or prevent the remaining owned cleanup.
                failed = True
                diagnostics.record('operation.failed', error=error, level='WARNING', phase='cli', status='ERROR')
        stack.callback(close)

    try:
        yield register
    finally:
        active_failure = sys.exception() is not None
        stack.close()
        if interrupted:
            raise KeyboardInterrupt('CLI interrupted during owned cleanup')
        if failed and not active_failure:
            raise RuntimeError('CLI owned resource cleanup failed')


def _close_cli_page(page):
    if not page.is_closed():
        page.close()


def _failed_cli_run(run, folder, *, interrupted=False):
    """Preserve a validated PARTIAL checkpoint; never retry workbook export."""
    if run is not None:
        try:
            run['complete'] = False
            run.setdefault('errors', []).append({'object': 'Run', 'phase': 'details',
                'message': ('Interrupted. Resume this output folder to retry incomplete returns.'
                            if interrupted else 'Capture/export failed. Inspect diagnostics and original checkpoints.')})
            validate_export_run(run)
            _check_output_scope(run, folder)
            _atomic_checkpoint(folder / 'run.json', run)
        except KeyboardInterrupt as error:
            interrupted = True
            diagnostics.record('operation.cancelled', error=error, phase='cli', status='CANCELLED')
        except Exception as error:
            # Best effort only. Invalid scope or failed storage preserves the
            # last durable bytes and the original terminal result.
            diagnostics.record('operation.failed', error=error, level='WARNING', phase='cli', status='PARTIAL')
            print('Checkpoint update unavailable; existing checkpoint files retained.', file=sys.stderr)
    if interrupted:
        print('Interrupted; existing checkpoints retained.', file=sys.stderr)
        return 130
    print('BLOCKED: capture/export failed; inspect diagnostics and original checkpoints.', file=sys.stderr)
    return 1


def _main(argv=None):
    cli = parser()
    args = cli.parse_args(argv)
    if args.max_pages < 1 or args.detail_limit is not None and args.detail_limit < 1:
        cli.error('Page and detail limits must be positive')
    diagnostics.record('operation.started', operation='offline' if args.from_json else 'resume' if args.resume else 'export')
    if args.from_json:
        run = json.loads(args.from_json.read_text(encoding='utf-8'))
        validate_scope(run['npwp'], run['year'])
        folder = args.output or args.from_json.parent
        path = save_output(run, folder, json_source=args.from_json.parent)
        print(f'Resume rebuilt and verified: {path.resolve()}')
        print('SPT files: ' + ', '.join(name for name in run['output_files'] if name != 'resume.xlsx'))
        return 0 if run.get('complete') else 2
    if args.resume:
        from coretax_resume import load_resume_inventory
        folder = args.resume.absolute()
        try:
            run, args.detail_source = load_resume_inventory(folder,
                detail_source=args.detail_source if args.detail_source_explicit else None)
        except (OSError, ValueError) as error:
            cli.error(str(error))
        if args.npwp is not None and args.npwp != run['npwp'] or args.year is not None and str(args.year) != str(run['year']):
            cli.error('Resume scope must match the checkpoint exactly')
        npwp, year = validate_scope(run['npwp'], run['year'])
    else:
        if args.non_interactive and (args.npwp is None or args.year is None):
            cli.error('--non-interactive requires --npwp and --year, or --resume')
        npwp = args.npwp if args.npwp is not None else input('NPWP (16 digits, no separators): ')
        entered_year = args.year if args.year is not None else input('Tax year (e.g. 2025): ')
        npwp, year = validate_scope(npwp, entered_year)
        stamp = datetime.now().strftime('%Y%m%d-%H%M%S-%f')
        folder = (args.output or BASE / 'output' / f'{npwp}_{year}_{stamp}').resolve()
        run = None
    print(f'Scope: NPWP {npwp} | year {year} | basis {run["year_basis"] if run else args.year_basis}')
    print('READ-ONLY: no submission, cancellation, PDF-generation or email actions.')
    if args.reuse_browser:
        try:
            validate_debug_endpoint(args.reuse_browser)
        except ValueError as error:
            cli.error(str(error))
    start = time.monotonic()
    owned_context = None
    work_page = None
    output_path = None
    try:
        with _cli_cleanup() as register:
            manager = sync_playwright()
            # Initialization may acquire the transport before raising or Stop.
            # The contained callback also tolerates an uninitialized manager.
            register(manager.__exit__, None, None, None)
            playwright = manager.__enter__()
            if args.reuse_browser:
                browser = playwright.chromium.connect_over_cdp(args.reuse_browser, timeout=20000)
                context = browser.contexts[0]
            else:
                owned_context = playwright.chromium.launch_persistent_context(str(args.profile.resolve()), channel=args.browser,
                    headless=False, no_viewport=True, service_workers='block',
                    locale='id-ID', timezone_id='Asia/Jakarta',
                    args=['--no-first-run', '--no-default-browser-check'])
                register(owned_context.close)
                owned_context._spt_financial_route_policy = 'service_workers_blocked'
                context = owned_context
            work_page = context.new_page()
            register(_close_cli_page, work_page)
            work_page.set_viewport_size({'width': 1600, 'height': 1000})
            authenticate(work_page, args.non_interactive)
            if run is None:
                run = collect_inventory(work_page, npwp, year, args.year_basis, folder, max_pages=args.max_pages)
            run['started_at'] = run.get('started_at') or datetime.now().astimezone().isoformat()
            run['complete'] = False
            if args.inventory_only:
                run['errors'].append({'object': 'Scope', 'message': 'Inventory-only run. Return-detail content not captured.'})
            else:
                require_mpn(run, args.detail_source)
                def progress(current_run, ordinal):
                    if ordinal % 5 == 0:
                        current_run['complete'] = False
                        save_output(current_run, folder)
                collect_details(work_page, run, folder, max_pages=args.max_pages, detail_limit=args.detail_limit,
                                on_progress=progress, detail_source=args.detail_source)
                collect_mpn(work_page, run, folder, max_pages=args.max_pages, detail_source=args.detail_source)
            run['finished_at'] = datetime.now().astimezone().isoformat()
            run['duration_seconds'] = round(time.monotonic() - start, 2)
            output_path = save_output(run, folder)
    except KeyboardInterrupt as error:
        diagnostics.record('operation.cancelled', error=error, status='CANCELLED')
        return _failed_cli_run(run, folder, interrupted=True)
    except Exception as error:
        diagnostics.record('operation.failed', error=error, level='ERROR', status='ERROR')
        return _failed_cli_run(run, folder)
    count = sum(r.get('capture_complete', False) for r in run['records'])
    print(f"{'COMPLETE' if run['complete'] else 'PARTIAL'}: {count}/{len(run['records'])} return details verified")
    print(f'Resume Excel: {output_path.resolve()}')
    print('SPT files: ' + ', '.join(name for name in run['output_files'] if name != 'resume.xlsx'))
    print(f'Checkpoints: {folder}')
    print(f"Duration: {run['duration_seconds']} seconds")
    return 0 if run['complete'] else 2


def main(argv=None):
    with diagnostics.session(diagnostics.default_root(), 'cli'):
        try:
            code = _main(argv)
        except BaseException as error:
            diagnostics.record('operation.cancelled' if isinstance(error, KeyboardInterrupt) else 'operation.failed',
                               error=error, level='ERROR')
            raise
        diagnostics.record('operation.finished', exit_code=code,
                           status={0:'COMPLETE', 2:'PARTIAL', 130:'CANCELLED'}.get(code, 'ERROR'))
        return code


if __name__ == '__main__':
    raise SystemExit(main())
