"""MPN is an independent workbook, never an SPT row or public raw archive."""
from copy import deepcopy
import importlib
import sys

from openpyxl import load_workbook
import pytest

from test_mpn import scope_run, payment_source


def mpn_run(*, amount='-1250.50', complete=True):
    from coretax_mpn import require_mpn, project_payment
    run = scope_run()
    value = require_mpn(run, 'dom')
    header, detail = payment_source(amount=amount)
    value.update(status='COMPLETE' if complete else 'PARTIAL', inventory_complete=True,
                 details_complete=complete, headers=1, details=1,
                 rows=[project_payment(header, detail, year=2025)])
    run['complete'] = complete
    return run


def exporter():
    return importlib.import_module('coretax_mpn_export')


def test_mpn_round_trip_keeps_negative_amount_identifiers_and_literal_reference(tmp_path):
    run = mpn_run()
    path = tmp_path / 'MPN_2025.xlsx'
    exporter().export_mpn_workbook(run, path)
    exporter().verify_mpn_workbook(run, path)
    wb = load_workbook(path)
    try:
        assert wb.sheetnames == ['MPN', 'Audit']
        cells = [cell for row in wb['MPN'] for cell in row]
        assert any(cell.value == -1250.5 and cell.data_type == 'n' for cell in cells)
        assert any(cell.value == '=SYNTHETIC-REFERENCE' and cell.data_type == 's' for cell in cells)
        assert any(cell.value == '0000000000000007' and cell.data_type == 's' for cell in cells)
        proof = exporter().read_mpn_proof(wb)
        assert proof['rows'] == run['mpn']['rows']
        assert proof['scope'] == [run['npwp'], 2025, 'period']
    finally:
        wb.close()


@pytest.mark.parametrize('amount', ['-1234567890123456', '-1e-309'])
def test_unsafe_excel_numbers_stay_original_text_with_presentation_audit(tmp_path, amount):
    run = mpn_run(amount=amount)
    path = tmp_path / 'MPN_2025.xlsx'
    exporter().export_mpn_workbook(run, path)
    wb = load_workbook(path)
    try:
        assert any(cell.value == amount and cell.data_type == 's' for row in wb['MPN'] for cell in row)
        assert any(cell.value == 'LITERAL / PRECISION' for row in wb['Audit'] for cell in row)
        assert exporter().read_mpn_proof(wb)['rows'][0]['amount'] == amount
    finally:
        wb.close()


def test_physical_mpn_cells_must_match_proof_before_publication(tmp_path):
    run = mpn_run()
    path = tmp_path / 'MPN_2025.xlsx'
    exporter().export_mpn_workbook(run, path)
    wb = load_workbook(path)
    for row in wb['MPN']:
        for cell in row:
            if cell.value == -1250.5:
                cell.value = 1250.5
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match='MPN'):
        exporter().verify_mpn_workbook(run, path)


def test_full_empty_spt_job_publishes_mpn_and_one_relative_resume_link(tmp_path):
    import coretax
    from spt_desktop import output
    run = mpn_run()
    private = tmp_path / 'private'
    coretax.save_output(run, private)
    assert set(run['output_files']) == {'resume.xlsx', 'MPN_2025.xlsx'}
    final = output.publish_workbooks(run, private, tmp_path / 'Downloads', 'synthetic-mpn')
    coretax.verify_outputs(run, final)
    assert not list(final.rglob('*.json'))
    wb = load_workbook(final / 'resume.xlsx')
    try:
        assert wb['Ringkasan']['B7'].value == 0
        assert wb['Daftar SPT'].max_row == 4
        links = [cell.hyperlink for row in wb['Ringkasan'] for cell in row if cell.hyperlink]
        assert len(links) == 1 and links[0].target == 'MPN_2025.xlsx'
        assert links[0].location == "'MPN'!A1"
        assert exporter().read_resume_mpn(wb)['filename'] == 'MPN_2025.xlsx'
        rows = [row for row in wb['Audit'] if row[0].value == 'MPN' and row[5].hyperlink]
        assert len(rows) == 1
        assert [cell.value for cell in rows[0]][:5] == ['MPN', 1, 1, 1, 'COMPLETE']
        assert rows[0][5].hyperlink.target == 'MPN_2025.xlsx'
    finally:
        wb.close()


def test_mpn_partial_preserves_verified_spt_and_cannot_promote_complete(tmp_path):
    from support_runs import fixture_run
    import coretax
    run = fixture_run()
    run['mpn'] = mpn_run(complete=False)['mpn']
    private = tmp_path / 'private'
    coretax.save_output(run, private)
    coretax.verify_outputs(run, private)
    assert (private / 'MPN_2025.xlsx').is_file()
    assert len(run['output_files']) == 7  # Five literal period/model files, resume, MPN.
    wb = load_workbook(private / 'resume.xlsx')
    try:
        assert wb['Ringkasan']['B6'].value == 'PARTIAL'
        assert wb['Ringkasan']['B7'].value == 5
        assert exporter().read_resume_mpn(wb)['status'] == 'PARTIAL'
    finally:
        wb.close()


def test_legacy_refresh_preserves_mpn_bytes_and_its_resume_reference(tmp_path):
    import coretax
    from spt_desktop import output
    run = mpn_run()
    private = tmp_path / 'first'
    coretax.save_output(run, private)
    downloads = tmp_path / 'Downloads'
    final = output.publish_workbooks(run, private, downloads, 'synthetic-first')
    before = (final / 'MPN_2025.xlsx').read_bytes()
    legacy = scope_run()
    legacy['complete'] = False
    coretax.save_output(legacy, tmp_path / 'second')
    output.publish_workbooks(legacy, tmp_path / 'second', downloads, 'synthetic-second', overwrite=True)
    assert (final / 'MPN_2025.xlsx').read_bytes() == before
    wb = load_workbook(final / 'resume.xlsx')
    try:
        assert exporter().read_resume_mpn(wb)['filename'] == 'MPN_2025.xlsx'
    finally:
        wb.close()
    coretax.verify_outputs(legacy, final, allow_extra=True)


def test_changed_private_mpn_evidence_blocks_publication_before_downloads(monkeypatch, tmp_path):
    import coretax
    from coretax_mpn import MpnSafetyError
    from spt_desktop import output
    from test_mpn_capture_protocol import prepare
    browser, _ = prepare(monkeypatch)
    private = tmp_path / 'private'
    private.mkdir()
    run = scope_run()
    browser.collect_mpn(object(), run, private)
    coretax.save_output(run, private)
    path = private / run['mpn']['evidence'][0]['path']
    path.write_bytes(path.read_bytes() + b' ')
    downloads = tmp_path / 'Downloads'
    with pytest.raises(MpnSafetyError):
        output.publish_workbooks(run, private, downloads, 'synthetic-changed')
    assert not downloads.exists()


def mpn_replacement(tmp_path):
    import coretax
    from support_runs import fixture_run
    from spt_desktop import output
    run = fixture_run()
    run['mpn'] = mpn_run()['mpn']
    downloads = tmp_path / 'Downloads'
    coretax.save_output(run, tmp_path / 'old-source')
    final = output.publish_workbooks(run, tmp_path / 'old-source', downloads, 'old-mpn')
    selected = deepcopy(run)
    selected['records'] = selected['records'][2:3]
    selected['mpn'] = mpn_run(amount='-2500.75')['mpn']
    source = tmp_path / 'new-source'
    coretax.save_output(selected, source)
    return selected, source, downloads, final


def test_mpn_replacement_requires_consent_and_preserves_unselected_spt(tmp_path):
    from spt_desktop import output
    from test_desktop_overwrite_safety import workbook_bytes
    run, source, downloads, final = mpn_replacement(tmp_path)
    before = workbook_bytes(final)
    snapshot = output.existing_outputs(run, downloads)
    assert set(snapshot['files']) == {'resume.xlsx', 'MPN_2025.xlsx',
                                     'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx'}
    with pytest.raises(FileExistsError, match='consent|overwrite'):
        output.publish_workbooks(run, source, downloads, 'denied-mpn')
    assert workbook_bytes(final) == before
    output.publish_workbooks(run, source, downloads, 'accepted-mpn', overwrite=True, expected_outputs=snapshot)
    assert (final / 'MPN_2025.xlsx').read_bytes() != before['MPN_2025.xlsx']
    for name in set(before) - set(snapshot['files']):
        assert (final / name).read_bytes() == before[name]


def test_mpn_consent_snapshot_refuses_changed_workbook_bytes(tmp_path):
    from spt_desktop import output
    from test_desktop_overwrite_safety import workbook_bytes
    run, source, downloads, final = mpn_replacement(tmp_path)
    snapshot = output.existing_outputs(run, downloads)
    path = final / 'MPN_2025.xlsx'
    path.write_bytes(path.read_bytes() + b'SYNTHETIC CONCURRENT CHANGE')
    before = workbook_bytes(final)
    with pytest.raises(ValueError, match='changed|consent|snapshot'):
        output.publish_workbooks(run, source, downloads, 'changed-mpn', overwrite=True, expected_outputs=snapshot)
    assert workbook_bytes(final) == before


def test_failure_after_installing_mpn_rolls_back_all_workbook_bytes(tmp_path, monkeypatch):
    from pathlib import Path
    from spt_desktop import output
    from test_desktop_overwrite_safety import workbook_bytes
    run, source, downloads, final = mpn_replacement(tmp_path)
    before = workbook_bytes(final)
    real_move = output._move_no_replace
    failed = []
    def move(source_path, destination):
        result = real_move(source_path, destination)
        if Path(destination) == final / 'MPN_2025.xlsx' and not failed:
            failed.append(True)
            raise OSError('Synthetic failure after MPN installation')
        return result
    monkeypatch.setattr(output, '_move_no_replace', move)
    with pytest.raises(OSError, match='Synthetic failure'):
        output.publish_workbooks(run, source, downloads, 'rollback-mpn', overwrite=True)
    assert failed and workbook_bytes(final) == before


@pytest.mark.parametrize('boundary', ['displaced', 'installed', 'index'])
@pytest.mark.skipif(sys.platform != 'win32', reason='Process-death recovery requires native Windows custody')
def test_process_death_recovers_mpn_spt_and_resume_as_one_publication(tmp_path, boundary):
    import subprocess
    import sys
    from test_publication_crash_recovery import CRASH_CHILD
    from test_desktop_overwrite_safety import workbook_bytes
    from spt_desktop.publication_recovery import recover_publication
    _, source, downloads, final = mpn_replacement(tmp_path)
    before = workbook_bytes(final)
    child_code = CRASH_CHILD.replace(
        "selected = final / 'SPT Masa 21' / 'SPT_21_Januari 2025_Normal.xlsx'",
        "selected = final / 'MPN_2025.xlsx'")
    child = subprocess.run([sys.executable, '-c', child_code, str(source), str(downloads),
                            str(final), boundary], capture_output=True, text=True, timeout=90)
    assert child.returncode == 73, child.stderr
    assert list(source.glob('publication-recovery-*.json'))
    assert recover_publication(source, downloads) == final
    assert workbook_bytes(final) == before
    assert recover_publication(source, downloads) == final
    assert not list(downloads.rglob('*.json'))
