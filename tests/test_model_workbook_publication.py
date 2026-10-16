"""Synthetic model-separated publication and old-layout migration contracts."""
from copy import deepcopy
import json

import pytest
from openpyxl import load_workbook

import coretax
import coretax_exporter as exporter
from spt_desktop import output
from support_financial_publication import PDF, CHANGED_PDF, declare_pdf
from support_runs import model_run


from support_financial_publication import (
    legacy_export, physical_bytes
)

def resume_targets(folder):
    workbook = load_workbook(folder / 'resume.xlsx')
    try:
        return [workbook['Daftar SPT'].cell(row, 1).hyperlink.target
                for row in range(5, workbook['Daftar SPT'].max_row + 1)]
    finally:
        workbook.close()


@pytest.mark.parametrize('selected_index', [0, 1, 2])
def test_replacement_changes_only_selected_model_workbook_and_bound_pdf(tmp_path, selected_index):
    source = tmp_path / 'original-private-job'
    old_run = model_run()
    for record in old_run['records']:
        declare_pdf(record, source)
    coretax.save_output(old_run, source)
    downloads = tmp_path / 'Downloads'
    final = output.publish_workbooks(old_run, source, downloads, 'synthetic-old')
    targets = resume_targets(final)
    assert len(set(targets)) == 3
    before = physical_bytes(final)
    selected = deepcopy(old_run)
    selected['records'] = [selected['records'][selected_index]]
    selected.pop('output_files', None)
    selected.pop('output_attachments', None)
    replacement = tmp_path / 'replacement-private-job'
    declare_pdf(selected['records'][0], replacement, CHANGED_PDF)
    coretax.save_output(selected, replacement)
    snapshot = output.existing_outputs(selected, downloads)
    pdf = selected['records'][0]['attachments'][0]['relative_path']
    assert set(snapshot['selected_files']) == {'resume.xlsx', targets[selected_index], pdf}
    assert set(snapshot['files']) == {'resume.xlsx', targets[selected_index], pdf}
    with pytest.raises(FileExistsError, match='overwrite|consent'):
        output.publish_workbooks(selected, replacement, downloads, 'synthetic-no',
                                 overwrite=False, expected_outputs=snapshot)
    assert physical_bytes(final) == before
    result = output.publish_workbooks(selected, replacement, downloads, 'synthetic-yes',
                                      overwrite=True, expected_outputs=snapshot)
    assert result == final and (final / pdf).read_bytes() == CHANGED_PDF
    for index, record in enumerate(old_run['records']):
        if index == selected_index:
            continue
        relative = record['attachments'][0]['relative_path']
        assert (final / targets[index]).read_bytes() == before[targets[index]]
        assert (final / relative).read_bytes() == before[relative] == PDF
    # The global index includes both refreshed and retained model workbooks.
    assert resume_targets(final)[0] == targets[selected_index]
    assert set(resume_targets(final)) == set(targets)
    assert len(resume_targets(final)) == len(targets)
    coretax.verify_outputs(selected, final, allow_extra=True)


@pytest.mark.parametrize('field,value', [('Model SPT', 'Pembetulan 99'), ('Masa Pajak', 'Februari 2025')])
def test_existing_filename_cannot_authorize_changed_literal_model_or_period(tmp_path, field, value):
    run = model_run(models=('Normal',))
    source = tmp_path / 'original-private-job'
    coretax.save_output(run, source)
    downloads = tmp_path / 'Downloads'
    final = output.publish_workbooks(run, source, downloads, 'synthetic-scope')
    path = final / resume_targets(final)[0]
    workbook = load_workbook(path)
    try:
        proof = json.loads(workbook.properties.identifier)
        proof['records'][0]['meta'][field] = value
        workbook.properties.identifier = json.dumps(proof)
        workbook.save(path)
    finally:
        workbook.close()
    before = physical_bytes(final)
    with pytest.raises(ValueError, match='scope|model|period|filename'):
        output.existing_outputs(run, downloads)
    assert physical_bytes(final) == before


@pytest.mark.parametrize('format_name', ['forms', 'legacy'])
def test_old_combined_outputs_are_preserved_and_new_layout_uses_a_new_folder(tmp_path, format_name):
    downloads = tmp_path / 'Downloads'
    run, old = legacy_export(downloads, format_name=format_name)
    before = physical_bytes(old)
    source = tmp_path / 'new-private-job'
    for record in run['records']:
        declare_pdf(record, source, CHANGED_PDF)
    coretax.save_output(run, source)
    assert output.existing_outputs(run, downloads) is None
    final = output.publish_workbooks(run, source, downloads, 'synthetic-new')
    assert final != old and final.parent == old.parent
    assert physical_bytes(old) == before
    assert len(set(resume_targets(final))) == 3
    for record in run['records']:
        relative = record['attachments'][0]['relative_path']
        assert (final / relative).read_bytes() == CHANGED_PDF
        assert (old / relative).read_bytes() == PDF
    coretax.verify_outputs(run, final)


@pytest.mark.parametrize('format_name', ['forms', 'legacy'])
def test_old_layout_migration_does_not_ignore_wrong_taxpayer_scope(tmp_path, format_name):
    downloads = tmp_path / 'Downloads'
    run, old = legacy_export(downloads, format_name=format_name)
    path = old / 'resume.xlsx'
    workbook = load_workbook(path)
    try:
        workbook['Ringkasan']['B4'] = '0000000000000002'
        workbook.save(path)
    finally:
        workbook.close()
    before = physical_bytes(old)
    with pytest.raises(ValueError, match='scope|NPWP'):
        output.existing_outputs(run, downloads)
    assert physical_bytes(old) == before


def test_corporate_pdf_proof_follows_each_model_file_not_a_fixed_badan_filename(tmp_path):
    run = model_run()
    source = tmp_path / 'private-job'
    for index, record in enumerate(run['records']):
        declare_pdf(record, source, PDF if index == 0 else CHANGED_PDF)
    coretax.save_output(run, source)
    targets = resume_targets(source)
    assert len(set(targets)) == 3
    for record, filename in zip(run['records'], targets):
        workbook = load_workbook(source / filename)
        try:
            proof = exporter.read_workbook_proof(workbook)
            assert [item['record_id'] for item in proof['records']] == [record['record_id']]
            assert proof['records'][0]['attachments'] == record['attachments']
            relative = record['attachments'][0]['relative_path']
            assert (source / relative).is_file()
        finally:
            workbook.close()
    downloads = tmp_path / 'Downloads'
    final = output.publish_workbooks(run, source, downloads, 'synthetic-pdfs')
    snapshot = output.existing_outputs(run, downloads)
    assert snapshot['target'] == str(final)
    assert set(snapshot['selected_files']) == {'resume.xlsx', *targets,
        *(record['attachments'][0]['relative_path'] for record in run['records'])}
    coretax.verify_outputs(run, final)


def test_selected_model_cannot_replace_different_source_text_with_same_sanitized_filename(tmp_path):
    run = model_run(models=('Pembetulan/1',))
    source = tmp_path / 'old-private-job'
    coretax.save_output(run, source)
    downloads = tmp_path / 'Downloads'
    final = output.publish_workbooks(run, source, downloads, 'synthetic-collision')
    before = physical_bytes(final)
    selected = model_run(models=('Pembetulan:1',))
    with pytest.raises(ValueError, match='collision|model|scope'):
        output.existing_outputs(selected, downloads)
    assert physical_bytes(final) == before
