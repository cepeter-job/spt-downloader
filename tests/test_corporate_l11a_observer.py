"""Owned View, passive L11-A response custody and saved exports; runners only."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

import coretax
import coretax_browser as driver
from coretax_corporate import observe_forms
from coretax_json import JsonSafetyError, adapt_corporate_detail
from coretax_json_capture import DetailJsonCapture, checkpoint_bodies
from l11a_observer_fixture import Fixture, ASSET
from test_corporate_l11a_native import DETAIL, GRIDS, OWNERS, GROUPS
from test_json_hybrid_native import browser
import support_json_corpus as corpus


def open_selected(fixture, capture):
    page, popup = driver.open_return_readonly(fixture.page, fixture.page.locator('tbody tr'),
        corpus.TAXPAYER_ID, corpus.RECORD_ID, json_capture=capture)
    assert page is fixture.page and popup is False
    page.wait_for_function('window.primaryDone===true')
    return page


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
@pytest.mark.parametrize('transport', ['singletons', 'pair-reordered'])
def test_owned_l11a_view_keeps_original_pages_catalogs_checkpoint_and_compact_saved_form(
        browser, tmp_path, monkeypatch, model, transport):
    fixture = Fixture(browser, monkeypatch, model=model, transport=transport)
    try:
        with DetailJsonCapture(fixture.context, fixture.page, fixture.record, corpus.SCOPE, tmp_path) as capture:
            page = open_selected(fixture, capture)
            forms = observe_forms(page, fixture.record, corpus.SCOPE, capture, approved_digest=fixture.pins[ASSET])
            assert forms is not None, 'Native L11-A owned observer is not implemented'
            assert [form['name'] for form in forms] == ['Induk', 'L11-A']
            proof = forms[1]['projection']
            assert proof['kind'] == 'corporate-l11a-v1' and proof['builds'] == fixture.pins
            assert [t['component'] for t in proof['tables']] == OWNERS
            assert [t['source_structure'] for t in proof['tables']] == [
                {'form': 'L11-A', 'component': owner, 'groups': groups} for owner, groups in zip(OWNERS, GROUPS)]
            manifest = capture.drain()
            assert manifest['complete'] and manifest['component_lifetime']['same_document'] is True
            bodies = list(capture.iter_bodies())
            assert len(bodies) == 6 + len(fixture.reference_raw)
            assert next(p['body'] for p in bodies if p['context']['url'] == corpus.PRIMARY_URL) == fixture.raw
            for source in bodies:
                url = source['context']['url']
                if url in GRIDS.values():
                    assert source['body'] == fixture.grid_raw[(url, source['context']['request_json']['First'])]
                elif url in fixture.reference_raw: assert source['body'] == fixture.reference_raw[url]
            before = deepcopy((bodies, forms, fixture.record))
            result = adapt_corporate_detail(iter(bodies), fixture.record, corpus.SCOPE, forms)
            assert result['eligible'] and result['coverage']['complete'], result['reason']
            assert (bodies, forms, fixture.record) == before
            fixture.record['detail_url'] = DETAIL
            replay = adapt_corporate_detail(checkpoint_bodies(tmp_path, manifest, fixture.record, corpus.SCOPE),
                                           fixture.record, corpus.SCOPE, forms)
            assert replay['sections'] == result['sections'] and replay['provenance'] == result['provenance']
            assert page.evaluate('window.businessReads') == page.evaluate('window.forbiddenClicks') == 0
        section = next(s for s in result['sections'] if s['name'] == 'L11-A')
        assert [t['body_row_count'] for t in section['tables']] == [12, 2, 2, 2, 2]
        assert len(section['fields']) == 12 and all(t['footer'] == [] for t in section['tables'])
        fixture.record.update(capture_complete=True, sections=result['sections'])
        run = {'npwp': corpus.NPWP, 'year': 2025, 'year_basis': 'period', 'inventory_complete': True,
               'complete': True, 'errors': [], 'records': [fixture.record]}
        coretax.save_output(run, tmp_path / 'export')
        coretax.verify_outputs(run, tmp_path / 'export')
        name, = [n for n in run['output_files'] if n != 'resume.xlsx']
        assert model in name
        book = load_workbook(tmp_path / 'export' / name)
        try:
            sheet = book['L11-A']
            values = [cell.value for row in sheet for cell in row]
            assert values.count('0007') == 14 and values.count('0005') == 2
            assert values.count('0025') == values.count('SYNTHETIC PERSON') == 2
            assert sheet.max_row < 60 and 'JUMLAH' not in values
            assert 'A1:D1' in {str(area) for area in sheet.merged_cells.ranges}
        finally: book.close()
        mapping = next(r for r in run['template_mapping'] if r['form'] == 'L11-A')
        assert not mapping['complete'] and mapping['unmapped_fields'] == 12
        assert mapping['unmapped_rows'] == mapping['unmapped_tables'] == 0
        assert fixture.unexpected == []
        for index, url in enumerate(GRIDS.values()): assert fixture.requests.count((url, 'POST')) == (2 if index == 0 else 1)
        for url in fixture.reference_raw: assert fixture.requests.count((url, 'GET')) == 1
        for url in fixture.assets: assert fixture.requests.count((url, 'GET')) == 2
    finally: fixture.close()


@pytest.mark.parametrize('defect', ['external-context', 'timezone', 'locale', 'vendor-pin',
    'owner', 'header', 'footer', 'wrong-group', 'missing-outer-group', 'misplaced-grid',
    'wrong-ui-owner', 'table-outside-p-table', 'duplicate-owner', 'hidden-owner',
    'missing-grid', 'failed-page', 'missing-reference', 'changed-header-after-page', 'changed-group-after-page',
    'changed-offline-header-after-page', 'changed-offline-group-after-page',
    'corporate-fetch-only', 'corporate-preload-only', 'main-fetch-only', 'main-preload-only',
    'vendor-fetch-only', 'vendor-preload-only'])
def test_unproved_owned_l11a_provider_owner_structure_or_dependency_cannot_grant_coverage(
        browser, tmp_path, monkeypatch, defect):
    fixture = Fixture(browser, monkeypatch, defect=defect)
    try:
        with DetailJsonCapture(fixture.context, fixture.page, fixture.record, corpus.SCOPE, tmp_path) as capture:
            page = open_selected(fixture, capture)
            forms = observe_forms(page, fixture.record, corpus.SCOPE, capture, approved_digest=fixture.pins[ASSET])
            capture.drain()
            if defect.startswith('changed-'):
                assert page.evaluate('window.l11aNextDone') is True
                assert fixture.requests.count((list(GRIDS.values())[0], 'POST')) == 2
                selector = OWNERS[3] if defect.startswith('changed-offline-') else OWNERS[0]
                if 'header' in defect:
                    assert 'SYNTHETIC CHANGED' in page.locator(selector + ' thead th').first.inner_text()
                else:
                    assert 'SYNTHETIC CHANGED' in page.locator(selector).evaluate('(e)=>e.closest("p-accordiontab").querySelector(".p-accordion-header-text").textContent')
            if forms is not None:
                result = adapt_corporate_detail(capture.iter_bodies(), fixture.record, corpus.SCOPE, forms)
                assert not result['eligible'] and result['sections'] == []
            assert page.evaluate('window.businessReads') == page.evaluate('window.forbiddenClicks') == 0
            assert fixture.unexpected == []
    finally: fixture.close()


def test_owned_l11a_changed_page_count_is_fatal_after_observation_even_if_binding_refused(browser, tmp_path, monkeypatch):
    fixture = Fixture(browser, monkeypatch, defect='changed-count')
    try:
        with DetailJsonCapture(fixture.context, fixture.page, fixture.record, corpus.SCOPE, tmp_path) as capture:
            page = open_selected(fixture, capture)
            page.locator('a[role=tab]').filter(has_text='L11-A').click()
            page.wait_for_function('window.l11aReplies===4', timeout=10000)
            forms = observe_forms(page, fixture.record, corpus.SCOPE, capture, approved_digest=fixture.pins[ASSET])
            if not page.evaluate('window.l11aNextDone'):
                # An unsupported observer may return immediately after visiting
                # Induk. Reopen the inspected owner before issuing its Next.
                page.locator('a[role=tab]').filter(has_text='L11-A').click()
                page.locator(OWNERS[0] + ' .p-paginator-next').click()
            page.wait_for_function('window.l11aNextDone===true', timeout=10000)
            capture.drain()
            bodies = list(capture.iter_bodies())
            grid = list(GRIDS.values())[0]
            assert len([p for p in bodies if p['context']['url'] == grid]) == 2
            with pytest.raises(JsonSafetyError, match='snapshot|count'):
                adapt_corporate_detail(iter(bodies), fixture.record, corpus.SCOPE, forms or [])
    finally: fixture.close()


@pytest.mark.parametrize('phase', ['entry', 'after-next'])
def test_owned_l11a_stop_propagates_without_publication_and_releases_document(browser, tmp_path, monkeypatch, phase):
    fixture = Fixture(browser, monkeypatch)
    stopped = []
    try:
        with DetailJsonCapture(fixture.context, fixture.page, fixture.record, corpus.SCOPE, tmp_path) as capture:
            page = open_selected(fixture, capture)
            def stop():
                if phase == 'entry' or fixture.requests.count((list(GRIDS.values())[0], 'POST')) >= 2:
                    stopped.append(True)
                    raise driver.CaptureCancelled('SYNTHETIC L11A STOP')
            with pytest.raises(driver.CaptureCancelled):
                observe_forms(page, fixture.record, corpus.SCOPE, capture,
                              approved_digest=fixture.pins[ASSET], check_cancel=stop)
            assert stopped == [True] and not list(tmp_path.rglob('manifest.json'))
            if phase == 'after-next': assert fixture.requests.count((list(GRIDS.values())[0], 'POST')) == 2
            assert page.evaluate('window.businessReads') == page.evaluate('window.forbiddenClicks') == 0
            assert fixture.unexpected == []
        assert capture._component_document is None
    finally: fixture.close()


def test_owned_l11a_business_read_traps_cover_offline_active_and_recap_containers(browser, tmp_path, monkeypatch):
    fixture = Fixture(browser, monkeypatch)
    try:
        with DetailJsonCapture(fixture.context, fixture.page, fixture.record, corpus.SCOPE, tmp_path) as capture:
            page = open_selected(fixture, capture)
            page.evaluate('show(1)')
            page.wait_for_function('window.l11aReplies===4')
            for selector in (OWNERS[0] + ' table', OWNERS[3] + ' table', '#panel-1'):
                for key in ('innerText', 'textContent'):
                    assert page.locator(selector).evaluate('(e,key)=>{try{void e[key];return false;}catch(_){return true;}}', key)
            assert page.evaluate('window.businessReads') == 6
            assert page.locator(OWNERS[0] + ' th').first.inner_text() == 'No.'
    finally: fixture.close()
