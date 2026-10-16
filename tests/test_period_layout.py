"""Kontrak tata letak keluaran dengan data sintetis, bukan catatan wajib pajak."""
from copy import deepcopy
import pytest
from openpyxl import load_workbook
import coretax
import coretax_exporter as exporter
from support_runs import fixture_run


def periodic_run():
    run = fixture_run()
    amendment = deepcopy(run['records'][2])
    amendment.update(key='21-january-amendment', record_id='21-january-amendment')
    amendment['meta']['Model SPT'] = 'Pembetulan ke-1'
    run['records'].append(amendment)
    ppn = deepcopy(run['records'][2])
    ppn.update(key='ppn-january', record_id='ppn-january')
    ppn['meta']['Jenis Pajak'] = 'PPN'
    run['records'].append(ppn)
    return run


def test_each_period_and_amendment_has_its_own_workbook(tmp_path):
    files = exporter.export_split_workbooks(periodic_run(), tmp_path)
    assert set(files) == {'resume.xlsx',
        'SPT_Badan_Januari - Desember 2025_Normal.xlsx',
        'SPT_Badan_Januari - Desember 2025_Amendment 001.xlsx',
        'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx',
        'SPT Masa 21/SPT_21_Januari 2025_Pembetulan ke-1.xlsx',
        'SPT Masa 21/SPT_21_Februari 2025_Normal.xlsx',
        'SPT Masa Unifikasi/SPT_Unifikasi_Januari 2025_Normal.xlsx',
        'SPT Masa PPN/SPT_PPN_Januari 2025_Normal.xlsx'}
    for model in ('Normal', 'Pembetulan ke-1'):
        workbook = load_workbook(files[f'SPT Masa 21/SPT_21_Januari 2025_{model}.xlsx'])
        try:
            assert workbook.sheetnames == ['Induk']
            proof = exporter.read_workbook_proof(workbook)
            assert [record['meta']['Model SPT'] for record in proof['records']] == [model]
            assert all(sheet['A1'].value == 'Induk SPT PPh Pasal 21/26' for sheet in workbook)
            assert not any(cell.hyperlink and 'resume.xlsx' in (cell.hyperlink.target or '')
                           for sheet in workbook for row in sheet for cell in row)
        finally:
            workbook.close()


def test_corporate_models_have_separate_files_with_all_observed_forms(tmp_path):
    run = fixture_run()
    run['records'] = run['records'][:2]
    for record in run['records']:
        record['sections'].append({'name': 'L1-D', 'fields': [{'label':'Amount','value':'1.234,50'}], 'tables':[], 'text':''})
    files = exporter.export_split_workbooks(run, tmp_path)
    assert len(files) == 3
    for record in run['records']:
        model = record['meta']['Model SPT']
        workbook = load_workbook(files[f'SPT_Badan_Januari - Desember 2025_{model}.xlsx'])
        try:
            assert workbook.sheetnames == ['Induk', 'L1-D']
            assert len({name.casefold() for name in workbook.sheetnames}) == 2
            proof = exporter.read_workbook_proof(workbook)
            assert [item['meta']['Model SPT'] for item in proof['records']] == [model]
            for sheet in workbook:
                assert not any(merged.min_row >= 6 for merged in sheet.merged_cells.ranges)
        finally:
            workbook.close()


def test_nested_corporate_sections_keep_full_source_identity(tmp_path):
    run = fixture_run()
    run['records'] = run['records'][:2]
    expected = {}
    for record in run['records']:
        record['sections'] = []
        model = record['meta']['Model SPT']
        expected[model] = []
        for name in ['Induk / Rincian', 'L1-D / Rincian']:
            marker = f"SYNTHETIC {record['record_id']} / {name}"
            record['sections'].append({'name': name, 'fields': [{'label': 'Marker', 'value': marker}],
                                       'tables': [], 'text': ''})
            expected[model].append((name, marker))
    files = exporter.export_split_workbooks(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    for model, entries in expected.items():
        workbook = load_workbook(files[f'SPT_Badan_Januari - Desember 2025_{model}.xlsx'])
        try:
            assert len(workbook.worksheets) == len(entries)
            all_markers = {marker for items in expected.values() for _, marker in items}
            for sheet, (name, marker) in zip(workbook, entries):
                values = [cell.value for row in sheet for cell in row]
                assert name in values
                assert [value for value in values if value in all_markers] == [marker]
        finally:
            workbook.close()


def test_fields_and_tables_are_filterable_with_numeric_and_raw_values(tmp_path):
    run = fixture_run()
    run['records'] = run['records'][:1]
    run['records'][0]['sections'][0].update(fields=[{'label':'Pendapatan','field':'Revenue','value':'1.234,50'},
        {'label':'NPWP','field':'NPWP','value':'0000000000000001'}], tables=[{'name':'Details','headers':['Kode','Nilai'],
        'rows':[['001','1.000,00'],['001','1.000,00']], 'complete':True}])
    exporter.export_split_workbooks(run,tmp_path)
    wb = load_workbook(tmp_path/'SPT_Badan_Januari - Desember 2025_Normal.xlsx')
    try:
        ws=wb.worksheets[0]
        assert ws.tables and ws.freeze_panes
        cells=[c for row in ws for c in row]
        assert sum(c.value=='001' for c in cells)==2
        assert any(c.value=='1.234,50' and c.data_type=='s' for c in cells)
        assert any(c.value==1234.5 and c.data_type=='n' for c in cells)
        assert all(c.data_type=='s' for c in cells if c.value=='0000000000000001')
    finally:
        wb.close()


def test_scope_verification_supports_nested_workbooks(tmp_path):
    run=periodic_run()
    coretax.save_output(run,tmp_path)
    coretax.verify_outputs(run,tmp_path)
    assert len(list(tmp_path.rglob('*.xlsx')))==8


def test_period_path_collision_rejected_before_writing(tmp_path):
    run=fixture_run()
    run['records']=run['records'][2:4]
    run['records'][0]['meta']['Masa Pajak']='Januari/2025'
    run['records'][1]['meta']['Masa Pajak']='Januari:2025'
    with pytest.raises(ValueError,match='collision'):
        exporter.export_split_workbooks(run,tmp_path)
    assert not list(tmp_path.rglob('*.xlsx'))


def test_long_form_captions_keep_unique_short_names_and_source_text(tmp_path):
    run=fixture_run()
    run['records']=run['records'][:2]
    names=['A'*100+' first','A'*100+' second','A'*100+' / same','A'*100+' : same']
    for record in run['records']:
        record['sections']=[{'name':name,'fields':[{'label':'Marker','value':name}],
            'tables':[],'text':''} for name in names]
    files=exporter.export_split_workbooks(run,tmp_path)
    assert len(files)==3
    for model in ('Normal', 'Amendment 001'):
        workbook=load_workbook(files[f'SPT_Badan_Januari - Desember 2025_{model}.xlsx'])
        try:
            titles=workbook.sheetnames
            assert len(titles)==4 and len({title.casefold() for title in titles})==4
            assert all(len(title)<=31 for title in titles)
            for name in names:
                assert any(cell.value==name for sheet in workbook for row in sheet for cell in row)
        finally:
            workbook.close()


def test_broken_symlink_component_is_rejected_before_resolve(tmp_path,monkeypatch):
    from pathlib import Path
    missing=tmp_path/'missing-link'
    original=Path.is_symlink
    monkeypatch.setattr(Path,'is_symlink',lambda p:p==missing or original(p))
    with pytest.raises(ValueError,match='symbolic|reparse'):
        coretax.output_path(tmp_path,'missing-link/result.xlsx')


@pytest.mark.parametrize('name',['.hidden.xlsx','resume.pdf','resume'])
def test_unexpected_output_declaration_never_authorizes_a_generic_path(tmp_path,name):
    run=fixture_run()
    run['output_files']=['resume.xlsx',name]
    with pytest.raises(ValueError,match='filename|count|scope'):
        coretax.verify_outputs(run,tmp_path)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize(('value', 'expected'), [
    ('123.456.789.012.345,67', None),
    ('999.999.999.999.999,99', None),
    ('10.000.000.000.000.001', None),
    ('123.456.789.012,34', 123456789012.34),
    ('1.234,50', 1234.5),
    ('-1.234,50', -1234.5),
    ('1.000.000.000.000.000,00', 1e15),
])
def test_optional_numeric_view_does_not_lose_decimal_precision(value, expected):
    assert exporter.numeric_analysis(value, 'Nilai') == expected


def test_derived_table_headers_are_unique_without_changing_source_captions(tmp_path):
    run=fixture_run()
    run['records']=run['records'][:1]
    headers=['Nilai','Nilai','Nilai (2)','Numerik 1 (Nilai)']
    run['records'][0]['sections'][0]['tables']=[{'name':'SYNTHETIC CAPTIONS','headers':headers,
        'rows':[['1.000,00','2.000,00','3.000,00','4.000,00']], 'complete':True}]
    files=exporter.export_split_workbooks(run,tmp_path)
    wb=load_workbook(files['SPT_Badan_Januari - Desember 2025_Normal.xlsx'])
    try:
        ws=wb.worksheets[0]
        for table in ws.tables.values():
            labels=[col.name.casefold() for col in table.tableColumns]
            assert len(labels)==len(set(labels))
        assert all(any(c.value==name for row in ws for c in row) for name in headers)
    finally:
        wb.close()
