"""Explicit presentation omissions retain custody and reject lookalikes."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from support_masa_binding import render_pph21_binding as render_template_form
from support_templates import cells_with, section
from support_pph21_dom import appendix_table
from test_pph21_template import REFERENCE
from support_pph21_numeric import money

IDENTITY = 'A. IDENTITAS PEMOTONG'
DECLARATION = 'D. PERNYATAAN DAN TANDA TANGAN'


def omitted_fields(form):
    if form == 'Induk':
        items = [('Tin', '1. NPWP', IDENTITY, '0000000000000001'),
                 ('Name', '2. Nama', IDENTITY, 'SYNTHETIC TAXPAYER'),
                 ('Address', '3. Alamat', IDENTITY, 'SYNTHETIC ADDRESS'),
                 ('PhoneNumber', '4. Nomor Telepon', IDENTITY, '000000'),
                 ('MyCheckBox', 'DITANDATANGANI OLEH / Wajib Pajak / Kuasa / NAMA / KODE QR', DECLARATION, False),
                 ('DeclarationBy', 'Wajib Pajak', DECLARATION, 'TAXPAYER'),
                 ('DeclaratorName', 'NAMA', DECLARATION, 'SYNTHETIC SIGNER')]
    else:
        groups = {'L-IA': ['LIST-IA'], 'L-IB': ['BPA1', 'BPA2'],
                  'L-II': ['BPA1', 'BPA2'], 'L-III': ['BP21', 'BP26']}[form]
        items = [('Tin', 'Nomor Identitas WP', 'HEADER', '0000000000000001'),
                 ('TaxPeriod', 'MASA PAJAK (MM-YYYY)', 'HEADER', '01-2025')]
        items.extend(('', 'Unnamed field', group, '100') for group in groups)
    return [dict(field=name, label=label, group=group, value=value) for name, label, group, value in items]


@pytest.mark.parametrize('form', ['Induk', 'L-IA', 'L-IB', 'L-II', 'L-III'])
def test_known_out_of_layout_controls_are_omitted_but_source_is_unchanged(form):
    fields = omitted_fields(form)
    sections = [section(form, fields=fields)]
    original = deepcopy(sections)
    book = load_workbook(REFERENCE)
    try:
        report = render_template_form(book[form], form, sections)
        assert report is not None and report['complete'] is True
        assert sections == original
        assert not cells_with(book, 'Data sumber tambahan / belum dipetakan')
        assert not cells_with(book, 'SYNTHETIC TAXPAYER')
    finally:
        book.close()


@pytest.mark.parametrize('field', [
    {'field': 'Name', 'label': '2. Nama', 'group': 'UNKNOWN', 'value': 'SYNTHETIC'},
    {'field': 'Name', 'label': 'AMOUNT', 'group': IDENTITY, 'value': '123,00'},
    {'field': '', 'label': 'Unnamed field', 'group': 'BPA2', 'value': '999,25'},
    {'field': 'Unknown', 'label': 'Nomor Identitas WP', 'group': 'HEADER', 'value': 'SYNTHETIC'},
])
def test_omission_exceptions_do_not_cover_unrecognized_semantics(field):
    form = 'Induk' if field['field'] == 'Name' else 'L-IB'
    book = load_workbook(REFERENCE)
    try:
        report = render_template_form(book[form], form, [section(form, fields=[field])])
        assert report is not None and report['complete'] is False
        assert report['unmapped_fields'] == 1
        assert not cells_with(book, 'Data sumber tambahan / belum dipetakan')
    finally:
        book.close()


def test_bp21_subtotals_are_not_misbound_to_bp26_slots():
    first, _ = appendix_table('L-III', 'BP21', 1)
    footer = [[], ['', 'JUMLAH PENGHASILAN BRUTO UNTUK PASAL 21 DAN PAJAK PENGHASILAN PASAL 21 YANG DITANGGUNG OLEH PEMERINTAH', '991,00', '17,00', ''],
              ['', 'JUMLAH UNTUK PASAL 21 DAN PENGHASILAN BRUTO PASAL 21 SERTA PAJAK PENGHASILAN YANG DIPOTONG', '887,00', '-23,00', '']]
    first['footer'] = footer
    first['footer_spans'] = [[], [4, 4, 1, 1, 4], [4, 4, 1, 1, 4]]
    first['rows'].extend(deepcopy(footer))
    second, totals = appendix_table('L-III', 'BP26', 0, with_totals=True)
    sections = [section('L-III', tables=[first, second])]
    original = deepcopy(sections)
    book = load_workbook(REFERENCE)
    try:
        report = render_template_form(book['L-III'], 'L-III', sections)
        assert report is not None and report['complete'] is True
        assert sections == original
        for row, amounts in totals.items():
            assert tuple(book['L-III'].cell(row, col).value for col in (7, 8)) == tuple(map(money, amounts))
        assert not cells_with(book, '991,00') and not cells_with(book, '887,00')
    finally:
        book.close()
