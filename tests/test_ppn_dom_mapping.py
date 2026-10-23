"""PPN-only DOM release regressions; values are synthetic."""

from pathlib import Path


import pytest


from openpyxl import load_workbook


import coretax_exporter as exporter


from support_templates import run_for, section, table, assert_format


RESOURCES = Path(__file__).resolve().parents[1] / 'resources'


def exported(tmp_path, tax, sections):
    run = run_for(sections)
    run['records'][0]['meta']['Jenis Pajak'] = tax
    from support_masa_binding import reference_geometry
    with reference_geometry():
        files = exporter.export_split_workbooks(run, tmp_path)
    return load_workbook(tmp_path / next(f for f in files if f != 'resume.xlsx'))


@pytest.mark.parametrize('tax,observed,reference,resource', [
    *[('PPN', v[0]+'-'+v[1], v, 'ppn_format.xlsx') for v in ['A1','A2','B1','B2','B3']],
])
def test_observed_alias_selects_template(tmp_path, tax, observed, reference, resource):
    wb = exported(tmp_path, tax, [section(observed)])
    ref = load_workbook(RESOURCES / resource)
    try:
        assert wb.sheetnames == [observed]
        assert wb[observed]['A1'].value == ref[reference]['A1'].value
        assert_format(wb[observed]['A1'], ref[reference]['A1'])
        assert set(map(str, wb[observed].merged_cells.ranges)) == set(map(str, ref[reference].merged_cells.ranges))
    finally:
        wb.close(); ref.close()


PPN_LABEL = 'PPN kurang atau (lebih) bayar (III.A - III.B - III.C - III.D)'


def test_ppn_scalar_field_maps_to_amount_slot(tmp_path):
    wb = exported(tmp_path, 'PPN', [section('INDUK', fields=[{'label':PPN_LABEL,'field':'','value':'-17,00'}])])
    try:
        assert wb['INDUK']['E45'].value == -17
        assert wb['INDUK']['E45'].data_type == 'n'
    finally:
        wb.close()


def test_ppn_duplicate_scalar_not_arbitrarily_selected(tmp_path):
    fields = [{'label':PPN_LABEL,'field':'','value':v} for v in ['17,00','19,00']]
    wb = exported(tmp_path, 'PPN', [section('Induk', fields=fields)])
    try:
        assert wb['Induk']['E45'].value is None
        assert not ({c.value for cells in wb['Induk'] for c in cells} & {'17,00','19,00'})
    finally:
        wb.close()


def test_ppn_repeated_label_uses_group_column_not_order(tmp_path):
    label = 'Jumlah PPN dan PPnBM yang dipungut'
    fields = [{'label':label,'field':'','value':'19,00','group':'VIII. PEMUNGUTAN PPN ATAU PPN DAN PPNBM OLEH PIHAK LAIN','column_label':'PPN (Rupiah)'},
              {'label':label,'field':'','value':'17,00','group':'VII. PEMUNGUTAN PPN ATAU PPN DAN PPNBM OLEH PEMUNGUT PPN','column_label':'PPN (Rupiah)'}]
    wb = exported(tmp_path, 'PPN', [section('Induk', fields=fields)])
    try:
        assert wb['Induk']['G69'].value == 17
        assert wb['Induk']['G75'].value == 19
        assert wb['Induk']['H69'].value is None
    finally:
        wb.close()


def test_ppn_appendix_dom_headers_and_extra_source_column(tmp_path):
    ref = load_workbook(RESOURCES/'ppn_format.xlsx')
    headers = [ref['A2'].cell(4,c).value for c in range(1,11)]
    ref.close()
    headers[4] += ' (dd-MM-yyyy)'
    headers.append('Masa Pajak')
    values = ['1','SYNTHETIC BUYER','0000000000000001','SYN-INVOICE','01-01-2025','17','19','23','0','SYN-REPLACED-INVOICE','EXTRA SYNTHETIC PROOF']
    wb = exported(tmp_path, 'PPN', [section('A-2',tables=[table('Tabel 1',[values],headers=headers)])])
    try:
        assert wb['A-2']['B5'].value == 'SYNTHETIC BUYER'
        assert wb['A-2']['E5'].value == '01-01-2025'
        assert wb['A-2']['H5'].value == 23
        assert not any(c.value == 'EXTRA SYNTHETIC PROOF' for cells in wb['A-2'] for c in cells)
    finally:
        wb.close()


def test_resume_reports_unmapped_template_evidence(tmp_path):
    run = run_for([section('Induk', fields=[{'label':'UNSUPPORTED SYNTHETIC AMOUNT','field':'','value':'17'}])])
    run['records'][0]['meta']['Jenis Pajak'] = 'PPN'
    exporter.export_split_workbooks(run,tmp_path)
    wb = load_workbook(tmp_path/'resume.xlsx')
    try:
        assert wb['Ringkasan']['B23'].value == 'PARTIAL'
        assert any(c.value == 'Pemetaan template' for cells in wb['Audit'] for c in cells)
        hints = [c.value for cells in wb['Audit'] for c in cells
                 if isinstance(c.value, str) and 'jumlah field/tabel/baris belum dipetakan' in c.value]
        assert hints and all('checkpoint' in value and 'di bawah form' not in value for value in hints)
    finally:
        wb.close()
