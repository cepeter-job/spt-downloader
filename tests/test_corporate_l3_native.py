"""Pinned L3 original-packet/footer/catalog integration contracts; runners only."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

import coretax
import coretax_browser as driver
from coretax_json import JsonSafetyError, adapt_corporate_detail
import support_json_corpus as corpus
from test_corporate_l7_native import bindings as old_bindings, selected
from test_corporate_l11c_native import L11C_PINS
from test_corporate_l1d_native import DETAIL, packet as fresh_packet
from test_corporate_l3_projection import ROOT, HEADERS, SPANS, LABELS, inputs


GRID = corpus.ORIGIN + '/returnsheet/api/loadcit/l3-table-b-grid'
FOOTER = GRID + '/footer'
REFERENCES = {kind: corpus.ORIGIN + '/referencedata/api/currentreferencedata/id-ID/' + kind
              for kind in LABELS}
OWNERS = ['rsh-citr-l3-grid1', 'rshshr-citr-l3-grid2']
GROUPS = ['A. PENGHASILAN DARI LUAR NEGERI', 'B. PPh YANG DIPOTONG/DIPUNGUT PIHAK LAIN']


def request(first):
    return {'ReturnSheetRecordId': corpus.RECORD_ID, 'Language': 'id-ID', 'First': first,
            'Rows': 10, 'SortField': 'WithholdingSlipsDate', 'SortOrder': 1,
            'Filters': [], 'LanguageId': 'id-ID'}


def envelope(model='Normal', form=None):
    root = corpus.current_form()
    root['MainFormData']['IsIncomeTaxOverseas'] = 1
    root[ROOT] = inputs()[0] if form is None else form
    return {'IsSuccessful': True, 'Payload': {'TIN': corpus.NPWP, 'PeriodYear': 2025,
        'IsDraft': True, 'ReturnSheetStatusCode': 'SUBMITTED', 'IsIncomeTaxOverseas': 1,
        'ReturnSheetModel': 'NORMAL' if model == 'Normal' else 'AMENDMENT',
        'PreviousFormData': None, 'LastSubmittedFormData': None,
        'CurrentFormData': corpus.wire(root)}}


def catalog(kind):
    return [{'ReferenceDataType': kind, 'language': 'id-ID', 'Details': [
        {'Code': code, 'CodeDescription': text, 'ValidFrom': None, 'ValidTo': None}
        for code, text in LABELS[kind].items()]}]


def packet(value, sequence, endpoint, body, *, method='POST', status=200):
    item = corpus.packet(value, sequence, endpoint=endpoint, request=body, status=status)
    item['context'].update(frame_url=DETAIL, method=method,
                          component_lifetime=deepcopy(fresh_packet()['context']['component_lifetime']))
    return item


def packets(model='Normal', *, form=None):
    _, pages, footer = inputs()
    values = [fresh_packet(envelope(model, form)),
        packet({'IsSuccessful': True, 'Payload': footer}, 2, FOOTER,
               {'ReturnSheetRecordId': corpus.RECORD_ID})]
    for index, page in enumerate(pages):
        values.append(packet({'IsSuccessful': True, 'Payload': {
            'TotalRecords': page['total'], 'Data': page['data']}}, index + 3, GRID, request(page['offset'])))
    values.extend(packet({'IsSuccessful': True, 'Payload': catalog(kind)}, index + 5,
                         url, None, method='GET') for index, (kind, url) in enumerate(REFERENCES.items()))
    return values


def bindings():
    return [old_bindings()[0], {'name': 'L3', 'paths': [[ROOT]], 'grids': [],
        'responses': [{'name': 'L3/domestic', 'endpoint': GRID, 'requests': [request(0), request(10)]}],
        'dependencies': [{'endpoint': FOOTER, 'request_json': {'ReturnSheetRecordId': corpus.RECORD_ID}}],
        'projection': {'kind': 'corporate-l3-v1', 'builds': dict(L11C_PINS), 'route': DETAIL,
            'locale': 'id-ID', 'timezone': 'Asia/Jakarta',
            'references': [{'endpoint': url, 'method': 'GET', 'sequences': [index + 5]}
                           for index, url in enumerate(REFERENCES.values())],
            'tables': [{'component': owner, 'headers': list(HEADERS[index]),
                'footer_spans': deepcopy(SPANS[index]),
                'source_structure': {'form': 'L3', 'component': owner, 'groups': [GROUPS[index]]}}
                for index, owner in enumerate(OWNERS)]}}]


@pytest.mark.parametrize('kind', list(REFERENCES))
def test_l3_each_catalog_uses_sorted_consumer_label_and_original_selected_path(kind):
    raw, observed = packets(), bindings()
    index = list(REFERENCES).index(kind)
    value = catalog(kind)
    first = deepcopy(value[0]['Details'][0])
    first['CodeDescription'] = 'SYNTHETIC ZULU'
    second = {**first, 'CodeDescription': 'SYNTHETIC ALPHA'}
    value[0]['Details'] = [first, second]
    sequence = index + 5
    raw[index + 4] = packet({'IsSuccessful': True, 'Payload': value}, sequence,
                            REFERENCES[kind], None, method='GET')
    before = deepcopy(raw)
    result = adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, observed)
    assert result['eligible'], result['reason']
    tables = next(s for s in result['sections'] if s['name'] == 'L3')['tables']
    destinations = {'COUNTRY_CODE': (0, 2), 'CURRENCY': (0, 7),
                    'FOREIGN_INCOME_TYPE_CIT': (0, 4), 'TAX_TYPE_BY_OTHER_PARTIES': (1, 3)}
    table, column = destinations[kind]
    assert tables[table]['rows'][0][column] == 'SYNTHETIC ALPHA'
    destination = {'table': 'L3/' + ('overseas' if table == 0 else 'domestic'), 'row': 0, 'column': column}
    evidence = result['provenance']['projections'][0]['values']
    assert any(e['response_sequence'] == sequence and e['container'] == ['Payload']
               and e['path'] == [0, 'Details', 1, 'CodeDescription']
               and e['value'] == 'SYNTHETIC ALPHA' and e['destination'] == destination for e in evidence)
    assert raw == before


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
@pytest.mark.parametrize('order', ['source', 'response-completion'])
def test_l3_original_packets_reach_final_saved_form_with_authoritative_footer_and_provenance(tmp_path, model, order):
    raw, observed, record = packets(model), bindings(), selected(model)
    if order == 'response-completion':
        raw = [raw[0], raw[3], *raw[4:], raw[2], raw[1]]
    before = deepcopy((raw, observed, record))
    result = adapt_corporate_detail(iter(raw), record, corpus.SCOPE, observed)
    assert result['eligible'] and result['coverage']['complete'], result['reason']
    assert (raw, observed, record) == before
    section, = [s for s in result['sections'] if s['name'] == 'L3']
    assert [t['body_row_count'] for t in section['tables']] == [2, 12]
    assert section['tables'][1]['footer'] == [['JUMLAH', '50', '40', ''],
        ['KREDIT PAJAK LUAR NEGERI', '', '1', ''], ['JUMLAH KREDIT PAJAK', '', '41', '']]
    assert [r[0] for r in section['tables'][1]['rows'][:12]] == [str(i) for i in range(1, 11)] + ['1', '2']
    assert section['tables'][1]['rows'][0][2] == '0000000000000002'
    assert {f['field']: f['value'] for f in section['fields']} == {
        'OverseasRefundPreviousYear': '1', 'OverseasTotalCalculatedCurrentYear': '888'}
    evidence, = [p['values'] for p in result['provenance']['projections'] if p['form'] == 'L3']
    assert any(e['response_sequence'] == 2 and e['container'] == ['Payload']
               and e['path'] == ['IncomeTax'] and e['type'] == 'number' and e['value'] == '40'
               for e in evidence)
    assert any(e['response_sequence'] == 4 and e['container'] == ['Payload']
               and e['path'] == ['Data', 0, 'WithholdingSlipsNumber'] and e['value'] == '000001'
               for e in evidence)
    for sequence, kind in enumerate(REFERENCES, 5):
        assert any(e['response_sequence'] == sequence and e['path'] == [0, 'Details', 0, 'CodeDescription']
                   and e['value'] in LABELS[kind].values() for e in evidence)
    record.update(capture_complete=True, sections=result['sections'])
    run = {'npwp': corpus.NPWP, 'year': 2025, 'year_basis': 'period', 'inventory_complete': True,
           'complete': True, 'errors': [], 'records': [record]}
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    name, = [n for n in run['output_files'] if n != 'resume.xlsx']
    assert model in name
    book = load_workbook(tmp_path / name)
    try:
        sheet = book['L3']
        totals = [row[5] for row in sheet if row[0].value in (
            'JUMLAH', 'KREDIT PAJAK LUAR NEGERI', 'JUMLAH KREDIT PAJAK')]
        assert [cell.value for cell in totals[-3:]] == [40, 1, 41]
        assert all('Rp' in cell.number_format and '0.00' not in cell.number_format for cell in totals[-3:])
        assert sheet.max_row < 40 and sheet['D3'].value == '02-12-2025'
        assert not any(cell.value == 'Data sumber tambahan / belum dipetakan' for row in sheet for cell in row)
    finally:
        book.close()
    mapping, = [r for r in run['template_mapping'] if r['form'] == 'L3']
    assert mapping['unmapped_fields'] == 2 and not mapping['complete']
    assert mapping['unmapped_tables'] == mapping['unmapped_rows'] == 0
    induk, = [r for r in run['template_mapping'] if r['form'] == 'Induk']
    assert not induk['complete']  # L3 cannot certify the other form's mapping.


@pytest.mark.parametrize('defect', ['missing-lifetime', 'pin', 'route', 'locale', 'timezone', 'owner',
    'group', 'header', 'span', 'missing-table', 'induk-after', 'missing-previous', 'model',
    'missing-page', 'missing-footer', 'failed-footer', 'null-footer', 'footer-shape',
    'page-result-only', 'page-index-shape', 'missing-reference', 'unknown-code',
    'reference-sequence', 'language', 'language-id', 'sort', 'filters', 'extra-footer-request',
    'outer-flag-missing', 'current-flag-missing', 'flags-disabled', 'flags-conflict', 'boolean-flag',
    'wrong-catalog-type', 'extra-catalog-type', 'cross-endpoint-sequence',
    'sort-order-bool', 'sort-order-string', 'missing-request-key', 'extra-request-key'])
def test_unproved_l3_packet_or_owner_dependency_keeps_whole_record_fallback(defect):
    raw, observed = packets(), bindings()
    proof = observed[1]['projection']
    if defect == 'missing-lifetime':
        for item in raw: del item['context']['component_lifetime']
    elif defect == 'pin': proof['builds'][next(iter(proof['builds']))] = '0' * 64
    elif defect == 'route': proof['route'] = DETAIL.replace('?view=true', '?history=true')
    elif defect in ('locale', 'timezone'): proof[defect] = 'SYNTHETIC UNPROVED'
    elif defect == 'owner': proof['tables'][0]['component'] = 'synthetic-foreign-owner'
    elif defect == 'group': proof['tables'][1]['source_structure']['groups'] = [GROUPS[0]]
    elif defect == 'header': proof['tables'][0]['headers'][0] = 'SYNTHETIC UNPROVED'
    elif defect == 'span': proof['tables'][1]['footer_spans'] = [[8]] * 3
    elif defect == 'missing-table': proof['tables'].pop()
    elif defect == 'induk-after': observed.reverse()
    elif defect in ('missing-previous', 'model'):
        value = envelope()
        if defect == 'missing-previous': del value['Payload']['PreviousFormData']
        else: value['Payload']['ReturnSheetModel'] = 'AMENDMENT'
        raw[0] = fresh_packet(value)
    elif defect == 'missing-page': raw.pop(3)
    elif defect == 'missing-footer': raw.pop(1)
    elif defect in ('failed-footer', 'null-footer', 'footer-shape'):
        payload = None if defect == 'null-footer' else {**inputs()[2], **({'Unproved': 0} if defect == 'footer-shape' else {})}
        raw[1] = packet({'IsSuccessful': defect != 'failed-footer', 'Payload': payload}, 2, FOOTER,
                        {'ReturnSheetRecordId': corpus.RECORD_ID}, status=500 if defect == 'failed-footer' else 200)
    elif defect in ('page-result-only', 'page-index-shape'):
        page = inputs()[1][0]
        payload = {'TotalRecords': 12, 'Data': page['data']}
        if defect == 'page-result-only': payload = {'TotalRecords': 12, 'Result': {'Data': page['data']}}
        else: payload['Index'] = None  # Initial supported consumer schema stays exact.
        raw[2] = packet({'IsSuccessful': True, 'Payload': payload}, 3, GRID, request(0))
    elif defect == 'missing-reference':
        raw.pop()
        proof['references'].pop()
    elif defect == 'unknown-code':
        form = inputs()[0]
        form['IncomeOverseas'][0]['CountryCode'] = 'SYNTHETIC UNBOUND'
        raw[0] = fresh_packet(envelope(form=form))
    elif defect == 'reference-sequence': proof['references'][0]['sequences'] = [999]
    elif defect in ('language', 'language-id', 'sort', 'filters'):
        key = {'language': 'Language', 'language-id': 'LanguageId', 'sort': 'SortField', 'filters': 'Filters'}[defect]
        value = [{'PropertyName': 'Name', 'Value': 'SYNTHETIC'}] if defect == 'filters' else 'SYNTHETIC UNPROVED'
        for item in raw:
            if item['context']['url'] == GRID: item['context']['request_json'][key] = deepcopy(value)
        for req in observed[1]['responses'][0]['requests']: req[key] = deepcopy(value)
    elif defect == 'extra-footer-request':
        raw[1]['context']['request_json']['Unproved'] = False
        observed[1]['dependencies'][0]['request_json']['Unproved'] = False
    elif defect in ('outer-flag-missing', 'current-flag-missing', 'flags-disabled', 'flags-conflict', 'boolean-flag'):
        value = envelope()
        if defect == 'outer-flag-missing': del value['Payload']['IsIncomeTaxOverseas']
        else:
            root = corpus.current_form()
            root[ROOT] = inputs()[0]
            if defect != 'current-flag-missing':
                root['MainFormData']['IsIncomeTaxOverseas'] = 0 if defect != 'boolean-flag' else True
            if defect == 'flags-disabled': value['Payload']['IsIncomeTaxOverseas'] = 0
            if defect == 'boolean-flag': value['Payload']['IsIncomeTaxOverseas'] = True
            value['Payload']['CurrentFormData'] = corpus.wire(root)
        raw[0] = fresh_packet(value)
    elif defect in ('wrong-catalog-type', 'extra-catalog-type'):
        kind, url = next(iter(REFERENCES.items()))
        value = catalog(kind)
        if defect == 'wrong-catalog-type': value[0]['ReferenceDataType'] = 'CURRENCY'
        else: value.append(catalog('CURRENCY')[0])
        raw[4] = packet({'IsSuccessful': True, 'Payload': value}, 5, url, None, method='GET')
    elif defect == 'cross-endpoint-sequence':
        proof['references'][0]['sequences'], proof['references'][1]['sequences'] = [6], [5]
    elif defect in ('sort-order-bool', 'sort-order-string', 'missing-request-key', 'extra-request-key'):
        requests = [item['context']['request_json'] for item in raw if item['context']['url'] == GRID]
        requests += observed[1]['responses'][0]['requests']
        for req in requests:
            if defect == 'missing-request-key': del req['Rows']
            elif defect == 'extra-request-key': req['SYNTHETIC_UNPROVED'] = False
            else: req['SortOrder'] = True if defect == 'sort-order-bool' else 'asc'
    before = deepcopy((raw, observed))
    result = adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, observed)
    assert not result['eligible'] and result['sections'] == []
    assert (raw, observed) == before


@pytest.mark.parametrize('bound', [False, True])
@pytest.mark.parametrize('defect', ['changed-footer', 'foreign-footer-id',
    'foreign-grid-id', 'foreign-frame', 'authorization', 'privacy', 'clock', 'lifetime'])
def test_l3_known_source_contradictions_remain_fatal_before_any_semantic_fallback(bound, defect):
    raw = packets()
    if defect == 'changed-footer':
        raw.append(packet({'IsSuccessful': True, 'Payload': {**inputs()[2], 'IncomeTax': 41}},
                          9, FOOTER, {'ReturnSheetRecordId': corpus.RECORD_ID}))
    elif defect in ('foreign-footer-id', 'foreign-grid-id'):
        raw[1 if defect == 'foreign-footer-id' else 2]['context']['request_json']['ReturnSheetRecordId'] = corpus.FOREIGN_ID
    elif defect == 'foreign-frame': raw[1]['context']['frame_url'] = DETAIL.replace(corpus.RECORD_ID, corpus.FOREIGN_ID)
    elif defect == 'authorization': raw[1]['context']['status'] = 403
    elif defect == 'privacy':
        raw[1] = packet({'IsSuccessful': True, 'Payload': {'AccessToken': 'SYNTHETIC PRIVATE BOUNDARY'}},
                        2, FOOTER, {'ReturnSheetRecordId': corpus.RECORD_ID})
    elif defect == 'clock':
        raw[1]['context']['clock_interval'] = {'attempt': corpus.ATTEMPT, 'start_ms': 0,
                                             'end_ms': 1, 'timezone': 'Asia/Jakarta'}
    elif defect == 'lifetime': raw[1]['context']['component_lifetime']['same_document'] = False
    category = ('snapshot' if defect.startswith('changed-') else
        'contradicts selected return' if defect in ('foreign-footer-id', 'foreign-grid-id') else
        'frame' if defect == 'foreign-frame' else 'authorization' if defect == 'authorization' else
        'Authentication material|privacy|forbidden' if defect == 'privacy' else
        'clock' if defect == 'clock' else 'lifetime')
    with pytest.raises(JsonSafetyError, match=category):
        adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, bindings() if bound else [])


@pytest.mark.parametrize('bound', [False, True])
@pytest.mark.parametrize('kind', list(REFERENCES))
@pytest.mark.parametrize('defect', ['snapshot', 'locale'])
def test_every_l3_catalog_preserves_fatal_snapshot_and_locale_checks_before_fallback(bound, kind, defect):
    raw = packets()
    index = list(REFERENCES).index(kind)
    value = catalog(kind)
    if defect == 'snapshot': value[0]['Details'][0]['CodeDescription'] = 'SYNTHETIC CHANGED'
    else: value[0]['language'] = 'en-US'
    item = packet({'IsSuccessful': True, 'Payload': value}, 9 if defect == 'snapshot' else index + 5,
                  REFERENCES[kind], None, method='GET')
    if defect == 'snapshot': raw.append(item)
    else: raw[index + 4] = item
    with pytest.raises(JsonSafetyError, match=defect):
        adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, bindings() if bound else [])


@pytest.mark.parametrize('endpoint', [GRID, FOOTER])
def test_known_l3_request_identity_is_fatal_before_observer_or_build_eligibility(endpoint):
    observations = [{'url': endpoint, 'method': 'POST',
                     'request_json': {'ReturnSheetRecordId': corpus.FOREIGN_ID}}]
    with pytest.raises(JsonSafetyError, match='contradicts selected return'):
        driver._check_corporate_request_identities(observations, selected(), corpus.SCOPE)
