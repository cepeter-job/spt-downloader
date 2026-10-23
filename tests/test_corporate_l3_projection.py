"""Pinned L3 pure source/auxiliary-footer contracts; runners only."""
from copy import deepcopy
import json

import pytest
from openpyxl import load_workbook

import coretax
import coretax_corporate_projection as projection
from coretax_json import _UnsupportedMetadata, _corporate_json
from support_templates import section, source_run


HEADERS = [['NO.', 'PEMOTONG PAJAK / NAMA', 'PEMOTONG PAJAK / NEGARA',
    'TANGGAL TRANSAKSI/PEMBAYARAN PPh', 'JENIS PENGHASILAN', 'PENGHASILAN NETO (RUPIAH)',
    'PPh YANG DIBAYAR/DIPOTONG/TERUTANG DI LUAR NEGERI / NILAI (Rp)',
    'PPh YANG DIBAYAR/DIPOTONG/TERUTANG DI LUAR NEGERI / MATA UANG',
    'PPh YANG DIBAYAR/DIPOTONG/TERUTANG DI LUAR NEGERI / NILAI DALAM MATA UANG ASING',
    'JUMLAH KREDIT PAJAK YANG DAPAT DIPERHITUNGKAN (Rp)'],
    ['NO.', 'NAMA PEMOTONG/PEMUNGUT', 'NPWPW PEMOTONG/PEMUNGUT', 'Jenis Pajak',
     'DASAR PENGENAAN PAJAK (Rupiah)', 'PPH YANG DIPOTONG/DIPUNGUT (Rupiah)',
     'BUKTI POTONG/SSP/SSPCP - NOMOR', 'BUKTI POTONG/SSP/SSPCP - TANGGAL']]
SPANS = [[[1, 4, 1, 1, 2, 1]], [[4, 1, 1, 2]] * 3]
LABELS = {'COUNTRY_CODE': {'ZZ': 'SYNTHETIC COUNTRY'},
          'CURRENCY': {'XXX': 'SYNTHETIC CURRENCY'},
          'FOREIGN_INCOME_TYPE_CIT': {'SYNTHETIC-INCOME': 'SYNTHETIC INCOME'},
          'TAX_TYPE_BY_OTHER_PARTIES': {'SYNTHETIC-TAX': 'SYNTHETIC TAX'}}
ROOT = 'ListIncomeTaxWithheldByOtherParties'


def inputs():
    overseas = {'Name': 'SYNTHETIC OVERSEAS', 'CountryCode': 'ZZ',
        'DateOfTransaction': '2025-12-01T23:30:00.000Z', 'IncomeCode': 'SYNTHETIC-INCOME',
        'NetIncome': 3, 'PayableRupiah': 2, 'PayableCurrency': 'XXX',
        'AmountForeignCurrency': 12, 'CreditCalculated': 1}
    form = {'IncomeOverseas': [deepcopy(overseas), deepcopy(overseas)],
        'OverseasTotalNetIncome': 6, 'OverseasTotalTaxPayable': 4, 'OverseasTotalTaxCredit': 2,
        'OverseasRefundPreviousYear': 1, 'OverseasTotalCalculatedCurrentYear': 888,
        # The captured successful footer legitimately overwrites these controls.
        'OtherPartiesTotalTaxBase': 999, 'OtherPartiesIncomeTaxWithheld': 999,
        'OtherPartiesIncomeTaxWithheldUsd': 999, 'OtherPartiesOverseasIncomeTax': 777,
        'OtherPartiesTotalCredit': 999, 'OtherPartiesTotalCreditUsd': 999}
    row = {'RecordId': 'SYNTHETIC-ROW', 'Name': 'SYNTHETIC DOMESTIC',
        'TIN': '0000000000000002', 'TaxType': 'SYNTHETIC-TAX', 'TaxBase': 5, 'IncomeTax': 3,
        'WithholdingSlipsNumber': '000001', 'WithholdingSlipsDate': '2025-12-01'}
    pages = [{'offset': 0, 'rows': 10, 'total': 12, 'data': [deepcopy(row) for _ in range(10)]},
             {'offset': 10, 'rows': 10, 'total': 12, 'data': [deepcopy(row) for _ in range(2)]}]
    return form, pages, {'TaxBase': 50, 'IncomeTax': 40, 'IncomeTaxUsd': 0}


def project(form, pages, footer, *, footer_wire=None, **overrides):
    function = getattr(projection, 'project_l3', None)
    assert callable(function), 'L3 pure source projector is not implemented'
    options = {'headers': deepcopy(HEADERS), 'footer_spans': deepcopy(SPANS),
               'tax_type': 'ICT_RCIT', 'locale': 'id-ID', 'timezone_name': 'Asia/Jakarta',
               'labels': deepcopy(LABELS)}
    options.update(overrides)
    captured = _corporate_json(json.dumps(form))
    observed_pages = [{**p, 'data': _corporate_json(json.dumps(p['data']))} for p in pages]
    captured_footer = _corporate_json(json.dumps(footer) if footer_wire is None else footer_wire)
    before = deepcopy((captured, observed_pages, captured_footer, options))
    try:
        return function(captured, pages=observed_pages, footer_payload=captured_footer, **options)
    finally:
        assert (captured, observed_pages, captured_footer, options) == before


def test_l3_uses_actual_current_offline_and_paged_fields_and_authoritative_footer():
    form, pages, footer = inputs()
    before = deepcopy((form, pages, footer))
    result = project(form, pages, footer)
    assert (form, pages, footer) == before
    a, b = result['tables']
    assert [a['body_row_count'], b['body_row_count']] == [2, 12]
    assert a['rows'][0] == ['1', 'SYNTHETIC OVERSEAS', 'SYNTHETIC COUNTRY', '02-12-2025',
        'SYNTHETIC INCOME', '3', '2', 'SYNTHETIC CURRENCY', '12,00', '1']
    assert a['rows'][1] == ['2', *a['rows'][0][1:]]
    assert a['footer'] == [['', 'JUMLAH', '6', '4', '', '2']]
    assert a['footer_spans'] == SPANS[0]
    assert b['rows'][0] == ['1', 'SYNTHETIC DOMESTIC', '0000000000000002',
        'SYNTHETIC TAX', '5', '3', '000001', '01-12-2025']
    assert [r[0] for r in b['rows'][:12]] == [str(i) for i in range(1, 11)] + ['1', '2']
    assert b['footer'] == [['JUMLAH', '50', '40', ''],
        ['KREDIT PAJAK LUAR NEGERI', '', '1', ''], ['JUMLAH KREDIT PAJAK', '', '41', '']]
    assert b['footer_spans'] == SPANS[1]
    assert all(t['complete'] and t['source_total'] == t['body_row_count'] for t in (a, b))
    assert {(f['field'], f['label'], f['value']) for f in result['fields']} == {
        ('OverseasRefundPreviousYear',
         'PENGEMBALIAN/PENGURANGAN PAJAK PENGHASILAN LUAR NEGERI (PASAL 24) YANG TELAH DIKREDITKAN UNTUK TAHUN SEBELUMNYA', '1'),
        ('OverseasTotalCalculatedCurrentYear',
         'JUMLAH PAJAK PENGHASILAN YANG DIBAYAR DI LUAR NEGERI YANG DAPAT DIPERHITUNGKAN DALAM TAHUN BERJALAN', '888')}
    evidence = result['provenance']
    assert any(e['source'] == 'current' and e['path'] == [ROOT, 'IncomeOverseas', 1, 'Name']
               and e['value'] == 'SYNTHETIC OVERSEAS' for e in evidence)
    assert any(e['source'] == 'grid' and e['page_index'] == 1
               and e['path'] == ['Data', 0, 'WithholdingSlipsNumber'] and e['value'] == '000001'
               for e in evidence)
    assert any(e['source'] == 'footer' and e['path'] == ['IncomeTax']
               and e['type'] == 'number' and e['value'] == '40'
               for e in evidence)
    # Final credits derive from the model's actual operands, not saved888/777/999.
    assert any(e['source'] == 'current' and e['path'] == [ROOT, 'OverseasRefundPreviousYear']
               and e['type'] == 'number' and e['value'] == '1' for e in evidence)
    for field in ('OverseasTotalCalculatedCurrentYear', 'OtherPartiesOverseasIncomeTax',
                  'OtherPartiesTotalTaxBase', 'OtherPartiesIncomeTaxWithheld',
                  'OtherPartiesIncomeTaxWithheldUsd', 'OtherPartiesTotalCredit', 'OtherPartiesTotalCreditUsd'):
        assert any(e['source'] == 'current' and e['path'] == [ROOT, field]
                   and e['type'] == 'number' and e['value'] == str(form[field]) for e in evidence)


@pytest.mark.parametrize('lexeme', ['40.0', '4e1'])
def test_l3_preserves_exact_footer_numeric_lexemes_while_rendering_the_same_value(lexeme):
    form, pages, footer = inputs()
    result = project(form, pages, footer,
        footer_wire='{"TaxBase":50,"IncomeTax":' + lexeme + ',"IncomeTaxUsd":0}')
    assert result['tables'][1]['footer'][0] == ['JUMLAH', '50', '40', '']
    assert any(e['source'] == 'footer' and e['path'] == ['IncomeTax']
               and e['type'] == 'number' and e['value'] == lexeme for e in result['provenance'])


@pytest.mark.parametrize('case', ['a-empty', 'b-null', 'b-empty'])
def test_l3_keeps_distinct_offline_and_active_grid_empty_date_consumers(case):
    form, pages, footer = inputs()
    if case == 'a-empty':
        form['IncomeOverseas'][0]['DateOfTransaction'] = ''
        assert project(form, pages, footer)['tables'][0]['rows'][0][3] == ''
    elif case == 'b-null':
        pages[0]['data'][0]['WithholdingSlipsDate'] = None
        assert project(form, pages, footer)['tables'][1]['rows'][0][7] == ''
    else:
        pages[0]['data'][0]['WithholdingSlipsDate'] = ''
        with pytest.raises(_UnsupportedMetadata):
            project(form, pages, footer)


@pytest.mark.parametrize('empty', [None, []])
def test_l3_explicit_empty_bodies_need_complete_zero_page_and_reported_totals(empty):
    form, _, _ = inputs()
    form.update(IncomeOverseas=empty, OverseasTotalNetIncome=0, OverseasTotalTaxPayable=0,
                OverseasTotalTaxCredit=0, OverseasRefundPreviousYear=None)
    result = project(form, [{'offset': 0, 'rows': 10, 'total': 0, 'data': []}],
                     {'TaxBase': 0, 'IncomeTax': 0, 'IncomeTaxUsd': 0})
    assert [t['body_row_count'] for t in result['tables']] == [0, 0]
    assert result['tables'][0]['footer'] == [['', 'JUMLAH', '0', '0', '', '0']]
    assert result['tables'][1]['footer'][-1] == ['JUMLAH KREDIT PAJAK', '', '0', '']


@pytest.mark.parametrize('defect', ['tax-type', 'locale', 'timezone', 'header', 'span',
    'missing-root-field', 'missing-overseas', 'string-money', 'bool-money', 'fraction-money',
    'unsafe-money', 'unsafe-intermediate', 'reported-mismatch', 'unsafe-refund',
    'missing-footer', 'footer-string', 'footer-bool', 'unsafe-credit', 'missing-page',
    'gap', 'changed-total', 'unsafe-offset', 'missing-actual-income-tax', 'singular-slip',
    'missing-label', 'numeric-code', 'invalid-date', 'year0001'])
def test_unproved_l3_source_arithmetic_pages_or_metadata_remains_unsupported(defect):
    form, pages, footer = inputs()
    overrides = {}
    if defect in ('tax-type', 'locale'):
        overrides[defect.replace('-', '_')] = 'SYNTHETIC-UNPROVED'
    elif defect == 'timezone':
        overrides['timezone_name'] = 'UTC'
    elif defect == 'header':
        overrides['headers'] = deepcopy(HEADERS)
        overrides['headers'][1][5] = 'SYNTHETIC UNUSED ALIAS'
    elif defect == 'span':
        overrides['footer_spans'] = [[[4, 1, 1, 2, 1]], SPANS[1]]
    elif defect == 'missing-root-field':
        del form['OverseasRefundPreviousYear']
    elif defect == 'missing-overseas':
        del form['IncomeOverseas']
    elif defect in ('string-money', 'bool-money', 'fraction-money', 'unsafe-money'):
        form['IncomeOverseas'][0]['NetIncome'] = {
            'string-money': '3', 'bool-money': True, 'fraction-money': 0.1,
            'unsafe-money': 1000000000000000}[defect]
    elif defect == 'unsafe-intermediate':
        for value in (999999999999999, 1, -1):
            row = deepcopy(form['IncomeOverseas'][0])
            row['NetIncome'] = value
            if value == 999999999999999:
                form['IncomeOverseas'] = []
            form['IncomeOverseas'].append(row)
        form.update(OverseasTotalNetIncome=999999999999999, OverseasTotalTaxPayable=6,
                    OverseasTotalTaxCredit=3)
    elif defect == 'reported-mismatch':
        form['OverseasTotalNetIncome'] = 7
    elif defect == 'unsafe-refund':
        form['OverseasRefundPreviousYear'] = -999999999999999
    elif defect == 'missing-footer':
        footer = None
    elif defect in ('footer-string', 'footer-bool'):
        footer['IncomeTax'] = '40' if defect == 'footer-string' else True
    elif defect == 'unsafe-credit':
        footer['IncomeTax'] = 999999999999999
    elif defect == 'missing-page':
        pages.pop()
    elif defect == 'gap':
        pages[1]['offset'] = 11
    elif defect == 'changed-total':
        pages[1]['total'] = 13
    elif defect == 'unsafe-offset':
        pages[0]['offset'] = False
    elif defect == 'missing-actual-income-tax':
        row = pages[0]['data'][0]
        row['IncomeTaxWithheld'] = row.pop('IncomeTax')
    elif defect == 'singular-slip':
        row = pages[0]['data'][0]
        row['WithholdingSlipNumber'] = row.pop('WithholdingSlipsNumber')
    elif defect == 'missing-label':
        overrides['labels'] = deepcopy(LABELS)
        overrides['labels']['CURRENCY'] = {}
    elif defect == 'numeric-code':
        form['IncomeOverseas'][0]['CountryCode'] = 0
    elif defect == 'invalid-date':
        form['IncomeOverseas'][0]['DateOfTransaction'] = '2025-02-30'
    elif defect == 'year0001':
        form['IncomeOverseas'][0]['DateOfTransaction'] = '0001-01-01'
    before = deepcopy((form, pages, footer, overrides))
    with pytest.raises(_UnsupportedMetadata):
        project(form, pages, footer, **overrides)
    assert (form, pages, footer, overrides) == before


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
def test_l3_final_saved_export_preserves_all_footers_and_truthful_unresolved_fields(tmp_path, model):
    result = project(*inputs())
    run = source_run('PPh Badan', [section('L3', fields=result['fields'], tables=result['tables'])])
    run['records'][0]['meta']['Model SPT'] = model
    original_fields = deepcopy(result['fields'])
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    assert run['records'][0]['sections'][0]['fields'] == original_fields
    name, = [n for n in run['output_files'] if n != 'resume.xlsx']
    assert model in name
    book = load_workbook(tmp_path / name)
    try:
        sheet = book['L3']
        assert sheet['D3'].value == '02-12-2025'
        assert sheet['I3'].value == '12,00'  # Number2 display remains literal source rendering.
        captions = [(cell.row, cell.value) for row in sheet for cell in row if cell.column == 1
                    and cell.value in ('JUMLAH', 'KREDIT PAJAK LUAR NEGERI', 'JUMLAH KREDIT PAJAK')]
        assert [label for _, label in captions] == ['JUMLAH', 'JUMLAH',
            'KREDIT PAJAK LUAR NEGERI', 'JUMLAH KREDIT PAJAK']
        btotal, overseas, credit = [row for row, _ in captions[1:]]
        assert sheet.cell(btotal, 5).value == 50 and sheet.cell(btotal, 6).value == 40
        assert sheet.cell(overseas, 5).value is None and sheet.cell(overseas, 6).value == 1
        assert sheet.cell(credit, 6).value == 41
        assert sheet.cell(btotal - 1, 3).value == '0000000000000002'
        assert sheet.cell(btotal - 1, 7).value == '000001'
        assert 'Rp' in sheet.cell(btotal, 6).number_format
        assert '0.00' not in sheet.cell(btotal, 6).number_format
        assert not any(cell.value == 'Data sumber tambahan / belum dipetakan' for row in sheet for cell in row)
    finally:
        book.close()
    report, = [r for r in run['template_mapping'] if r['form'] == 'L3']
    assert report['unmapped_fields'] == 2 and not report['complete']
    assert report['unmapped_rows'] == report['unmapped_tables'] == 0
