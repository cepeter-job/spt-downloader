"""Synthetic user-approved rate column and clean Induk regressions."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

import coretax_exporter as exporter
from support_templates import section, source_run


HEADERS = [
    'NO.', 'NIK/NPWP', 'NAMA', 'NOMOR BUKTI POTONG', 'TANGGAL BUKTI POTONG',
    'JENIS PAJAK', 'KODE OBJEK PAJAK', 'OBJEK PAJAK', 'DASAR PENGENAAN PAJAK (Rp)',
    'TINGKAT (%)', 'PAJAK PENGHASILAN', 'FASILITAS PERPAJAKAN',
    'UANG PERSEDIAAN / PEMBAYARAN LANGSUNG (UNTUK WP INSTANSI PEMERINTAH DENGAN DANA APBN)',
    'NITKU / NOMOR IDENTITAS SUBUNIT ORGANISASI', 'STATUS', 'KAP-KJS',
]


def export(tmp_path, sections):
    run = source_run('PPh Unifikasi', sections)
    original = deepcopy(run)
    reports = []
    paths = exporter.export_split_workbooks(run, tmp_path, mapping_reports=reports)
    assert run == original
    path, = [path for path in paths if path != 'resume.xlsx']
    return load_workbook(tmp_path / path), reports


@pytest.mark.parametrize('count', [0, 26, 41])
def test_rate_column_precedes_tax_without_overwriting_money(tmp_path, count):
    body = [
        [str(i + 1), '0000000000000002', 'SYNTHETIC NAME', f'SYNTHETIC-{i}',
         '02-01-2025', 'PPh Pasal 23', '23-100-01', 'SYNTHETIC OBJECT',
         '1.234,50', '2,00', '24,69', '', 'SYNTHETIC PAYMENT METHOD',
         '0000000000000002000001', 'Normal', '411124-100']
        for i in range(count)
    ]
    table = {'name': 'Tabel 1', 'group': 'TABEL I. BPPU', 'headers': HEADERS,
             'rows': body, 'body_row_count': count, 'source_total': count,
             'complete': True, 'footer': [], 'footer_spans': []}
    book, reports = export(tmp_path, [section('DAFTAR-I', tables=[table])])
    try:
        sheet = book['DAFTAR-I']
        assert sheet['J4'].value == 'TINGKAT (%)'
        assert sheet['K4'].value == 'PAJAK PENGHASILAN (Rp)'
        assert 'A1:P1' in {str(merge) for merge in sheet.merged_cells.ranges}
        for row in range(5, count + 5):
            assert sheet.cell(row, 9).value == 1234.5
            assert sheet.cell(row, 10).value == '2,00'
            assert sheet.cell(row, 10).data_type == 's'
            assert 'Rp' not in sheet.cell(row, 10).number_format
            assert sheet.cell(row, 11).value == 24.69
            assert sheet.cell(row, 11).data_type == 'n'
            assert sheet.cell(row, 14).value == '0000000000000002000001'
        # Supplemental payment-method evidence remains; rate is now on the form.
        assert sum(c.value == '2,00' for cells in sheet for c in cells) == count
        assert reports and reports[0]['complete']
    finally:
        book.close()


def test_known_induk_extra_identity_block_is_not_duplicated_below_form(tmp_path):
    fields = [
        {'field': 'TaxIdentificationNumber', 'label': 'NPWP',
         'group': 'A. IDENTITAS PEMOTONG', 'value': ''},
        {'field': 'FullName', 'label': 'Nama Lengkap',
         'group': 'A. IDENTITAS PEMOTONG', 'value': ''},
        {'field': 'Address', 'label': 'Alamat',
         'group': 'A. IDENTITAS PEMOTONG', 'value': ''},
        {'field': 'PhoneNumber', 'label': 'Nomor Telepon',
         'group': 'A. IDENTITAS PEMOTONG', 'value': ''},
        {'field': 'StatementTIN', 'label': 'NPWP',
         'group': 'C. PERNYATAAN DAN TANDA TANGAN', 'value': ''},
        {'field': 'StatementFullName', 'label': 'Nama Tanda Tangan',
         'group': 'C. PERNYATAAN DAN TANDA TANGAN', 'value': ''},
    ]
    book, reports = export(tmp_path, [section('INDUK', fields=fields)])
    try:
        assert not any(c.value is not None for cells in book['INDUK'].iter_rows(min_row=27) for c in cells)
        # Display omission must not erase unresolved mapping diagnostics.
        assert reports and reports[0]['unmapped_fields'] == len(fields)
        assert reports[0]['complete'] is False
    finally:
        book.close()


def test_unknown_induk_field_remains_unresolved_visible_evidence(tmp_path):
    field = {'field': 'UnexpectedAmount', 'label': 'NPWP',
             'group': 'A. IDENTITAS PEMOTONG', 'value': 'SYNTHETIC-UNRESOLVED'}
    book, reports = export(tmp_path, [section('INDUK', fields=[field])])
    try:
        assert reports and reports[0]['complete'] is False
        assert any(c.value == 'SYNTHETIC-UNRESOLVED' for row in book['INDUK'] for c in row)
    finally:
        book.close()
