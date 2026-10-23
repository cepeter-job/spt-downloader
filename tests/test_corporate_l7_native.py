"""Synthetic packet/native/workbook L7 contracts; execution belongs to CI."""
from copy import deepcopy
import hashlib
import importlib
import importlib.util
import json

import pytest

import coretax
import coretax_browser as driver
import coretax_exporter as exporter
from coretax_json import JsonSafetyError, adapt_corporate_detail
import support_json_corpus as corpus
from test_corporate_l7_projection import ROOT, HEADERS, SPANS, raw_source
from test_json_hybrid_native import ASSET, BUILD, NativeFixture, browser, PDF_URL


MAIN = corpus.ORIGIN + '/returnsheets/id-ID/main.cccfdb820c20596c.js'
MAIN_BUILD = b'/* SYNTHETIC MAIN BUILD -- no portal code or taxpayer data */'
PINS = {ASSET: 'fcd6d4a9244d6562a294675e802732562a9510621fbb9abbe651fa4de143048d',
        MAIN: '696ac66dc2897c937c7df7b466382dffdf25651837a20e0a4d4235787f3c9afd'}
FRAME = (corpus.ORIGIN + '/returnsheets/id-ID/corporate-income-tax-return/' +
         '/'.join((corpus.TAXPAYER_ID, corpus.RECORD_ID, 'ICT_RCIT', corpus.AGGREGATE_ID)))
REQUEST = {'RecordId': corpus.RECORD_ID, 'OnManualClick': False}


def selected(model='Normal'):
    record = corpus.selected_record()
    record['meta'].update({'Masa Pajak': 'Januari - Desember 2025', 'Model SPT': model})
    record['source_fields'].update(ReturnSheetStatusCode='SUBMITTED',
                                  ReturnSheetModel='NORMAL' if model == 'Normal' else 'AMENDMENT')
    return record


def envelope(model='Normal'):
    form = corpus.current_form()
    form['MainFormData']['NetIncomeBeforeTax'] = corpus.Number('1000')
    form[ROOT] = raw_source()
    return {'IsSuccessful': True, 'ErrorCode': 0, 'Payload': {
        'TIN': corpus.NPWP, 'PeriodYear': 2025, 'IsDraft': True,
        'ReturnSheetStatusCode': 'SUBMITTED',
        'ReturnSheetModel': 'NORMAL' if model == 'Normal' else 'AMENDMENT',
        'PreviousFormData': None, 'LastSubmittedFormData': None,
        'CurrentFormData': corpus.wire(form)}}


def marker():
    return {'kind': 'corporate-l7-v1', 'builds': dict(PINS), 'route': FRAME, 'locale': 'id-ID',
            'tables': [{'component': 'rsh-citr-l7-grid1', 'headers': list(HEADERS),
                        'footer_spans': deepcopy(SPANS)}]}


def bindings():
    return [{'name': 'Induk', 'paths': [['MainFormData']],
             'primary': {'endpoint': corpus.PRIMARY_URL, 'request_json': dict(REQUEST)}, 'grids': []},
            {'name': 'L7', 'paths': [[ROOT]], 'grids': [], 'projection': marker()}]


def packet(value=None):
    result = corpus.packet(envelope() if value is None else value, request=dict(REQUEST))
    result['context']['frame_url'] = FRAME
    return result


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
def test_l7_packet_semantics_preserve_literal_model_source_and_saved_approved_layout(tmp_path, model):
    record, raw, observed = selected(model), packet(envelope(model)), bindings()
    before = deepcopy((record, raw, observed))
    result = adapt_corporate_detail(iter([raw]), record, corpus.SCOPE, observed)
    assert result['eligible'] and result['coverage']['complete'], result['reason']
    section, = [section for section in result['sections'] if section['name'] == 'L7']
    table, = section['tables']
    assert table['headers'] == HEADERS
    assert table['rows'][0][:3] == ['1', '02023', '12.345']
    assert table['rows'][1][:3] == ['2', '2025', '1.000']
    assert table['rows'][2:] == table['footer']
    assert result['provenance']['draft_flags'] == {'outer': True, 'inner': False}
    assert result['provenance']['projections'][0]['form'] == 'L7'
    assert (record, raw, observed) == before
    record.update(capture_complete=True, sections=[section])
    run = {'npwp': corpus.NPWP, 'year': 2025, 'year_basis': 'period',
           'inventory_complete': True, 'complete': True, 'errors': [], 'records': [record]}
    assert_saved_l7(run, tmp_path, model)


def assert_saved_l7(run, folder, model):
    from openpyxl import load_workbook
    before, reports = deepcopy(run), []
    paths = exporter.export_split_workbooks(run, folder, mapping_reports=reports)
    assert run == before
    coretax.verify_outputs(run, folder)
    name, = [name for name in paths if name != 'resume.xlsx']
    assert model in name
    book = load_workbook(folder / name)
    try:
        sheet = book['L7']
        assert sheet['B3'].value == '02023' and sheet['B4'].value == '2025'
        assert sheet['C3'].value == 12345 and sheet['C4'].value == 1000
        assert sheet['D13'].value == 200 and sheet['I13'].value == 1200
        assert sheet['H2'].value == HEADERS[7] and sheet.max_row == 13
        assert 'A13:C13' in {str(area) for area in sheet.merged_cells.ranges}
        assert all('Rp' in sheet[cell].number_format and '*' in sheet[cell].number_format
                   for cell in ('C3', 'I4', 'D13', 'I13'))
    finally:
        book.close()
    report, = [report for report in reports if report['form'] == 'L7']
    assert report['complete'] and report['unmapped_rows'] == report['unmapped_tables'] == 0


@pytest.mark.parametrize('defect', ['pin', 'header', 'span', 'extra-table', 'unknown-marker',
    'route-history', 'locale', 'status-missing', 'draft-status', 'current-null', 'bad-previous',
    'prior-only', 'total-mismatch', 'unknown-root'])
def test_l7_unreviewed_lifecycle_or_structure_refuses_entire_semantic_record(defect):
    value, observed = envelope(), bindings()
    proof = observed[1]['projection']
    if defect == 'pin':
        proof['builds'][MAIN] = '0' * 64
    elif defect == 'header':
        proof['tables'][0]['headers'][1] = 'SYNTHETIC UNREVIEWED COLUMN'
    elif defect == 'span':
        proof['tables'][0]['footer_spans'] = [[9]]
    elif defect == 'extra-table':
        proof['tables'].append(deepcopy(proof['tables'][0]))
    elif defect == 'unknown-marker':
        proof['kind'] = 'SYNTHETIC UNKNOWN'
    elif defect == 'route-history':
        proof['route'] = FRAME.replace('corporate-income-tax-return/', 'corporate-income-tax-return-history/')
    elif defect == 'locale':
        proof['locale'] = 'en-US'
    elif defect == 'status-missing':
        del value['Payload']['ReturnSheetStatusCode']
    elif defect == 'draft-status':
        value['Payload']['ReturnSheetStatusCode'] = 'DRAFT'
    elif defect == 'bad-previous':
        value['Payload']['PreviousFormData'] = 'SYNTHETIC INVALID JSON'
    else:
        form = corpus.current_form()
        form['MainFormData']['NetIncomeBeforeTax'] = corpus.Number('1000')
        form[ROOT] = raw_source()
        if defect == 'current-null':
            form[ROOT] = None
        elif defect == 'prior-only':
            value['Payload']['PreviousFormData'] = corpus.wire({ROOT: raw_source()})
            del form[ROOT]
        elif defect == 'total-mismatch':
            form[ROOT]['TotalFiscalLossCompensationYearMinus4'] = corpus.Number('999')
        elif defect == 'unknown-root':
            form[ROOT]['SYNTHETIC_EXTRA'] = '0003'
        value['Payload']['CurrentFormData'] = corpus.wire(form)
    # DRAFT is deliberately unbound in selected metadata here: unsupported,
    # rather than a known selected-vs-response identity contradiction.
    record = selected()
    record['source_fields'].pop('ReturnSheetStatusCode')
    result = adapt_corporate_detail([packet(value)], record, corpus.SCOPE, observed)
    assert result['eligible'] is False and result['sections'] == []


@pytest.mark.parametrize('defect', ['route-tax', 'route-aggregate', 'selected-status'])
@pytest.mark.parametrize('observed', [[], bindings()], ids=['before-fallback', 'bound'])
def test_known_corporate_route_and_status_contradictions_are_fatal_before_fallback(defect, observed):
    raw = packet()
    if defect == 'route-tax':
        raw['context']['frame_url'] = FRAME.replace('/ICT_RCIT/', '/ICT_ROGCIT/')
    elif defect == 'route-aggregate':
        raw['context']['frame_url'] = FRAME.rsplit('/', 1)[0] + '/' + corpus.FOREIGN_ID
    else:
        value = envelope()
        value['Payload']['ReturnSheetStatusCode'] = 'DRAFT'
        raw = packet(value)
    with pytest.raises(JsonSafetyError):
        adapt_corporate_detail([raw], selected(), corpus.SCOPE, deepcopy(observed))


def l7_html():
    return ('<rshshr-citr-l7-form><rsh-citr-l7-grid1><p-table><table><thead>'
        '<tr><th rowspan="3">NO.</th><th colspan="2">LABA (RUGI) NETTO FISKAL</th>'
        '<th colspan="6">KOMPENSASI KERUGIAN FISKAL</th></tr>'
        '<tr><th rowspan="2">TAHUN/BAGIAN TAHUN PAJAK</th><th rowspan="2">NILAI (Rp)</th>'
        '<th>Y-4</th><th>Y-3</th><th>Y-2</th><th>Y-1</th><th>2025</th><th>Y+1</th></tr>'
        '<tr><th>NILAI (Rp)</th><th>NILAI (Rp)</th><th>NILAI (Rp)</th><th>NILAI (Rp)</th>'
        '<th>TAHUN PAJAK INI - NILAI (Rp)</th><th>TAHUN PAJAK BERJALAN - NILAI (Rp)</th></tr>'
        '</thead><tbody><tr>' + '<td>999999</td>' * 9 + '</tr></tbody>'
        '<tfoot><tr><td colspan="3">JUMLAH</td>' + '<td>999999</td>' * 6 +
        '</tr></tfoot></table></p-table></rsh-citr-l7-grid1></rshshr-citr-l7-form>')


class L7Fixture(NativeFixture):
    def __init__(self, browser, tmp_path, monkeypatch, *, model='Normal', popup=True, trap=True,
                 changed_main=False, changed_header=False):
        monkeypatch.setattr(corpus, 'FRAME_URL', FRAME)
        self.trap, self.changed_main, self.changed_header = trap, changed_main, changed_header
        super().__init__(browser, tmp_path, monkeypatch, popup=popup,
                         extra_form=('L7', 'rshshr-citr-l7-form'))
        record = selected(model)
        self.record['meta'] = record['meta']
        self.record['source_fields'] = record['source_fields']
        self.raw = corpus.wire(envelope(model)).encode()
        if importlib.util.find_spec('coretax_corporate_projection') is not None:
            projection = importlib.import_module('coretax_corporate_projection')
            monkeypatch.setattr(projection, 'BUILD_DIGESTS', {
                ASSET: hashlib.sha256(BUILD).hexdigest(), MAIN: hashlib.sha256(MAIN_BUILD).hexdigest()})

    def route(self, route):
        if route.request.url == MAIN:
            self.requests.append({'url': MAIN, 'method': route.request.method, 'body': None})
            route.fulfill(content_type='application/javascript',
                          body=MAIN_BUILD + (b'/* CHANGED */' if self.changed_main else b''))
            return
        super().route(route)

    def detail_html(self):
        source = super().detail_html()
        source = source.replace('<rshshr-citr-l7-form><label>UNSUPPORTED FIXTURE FORM</label>'
                                '</rshshr-citr-l7-form>', l7_html())
        source = source.replace('p.querySelectorAll("p-table")',
                                'p.querySelectorAll("p-table[data-fixture-config]")')
        source = source.replace('</head>', '<script src="' + MAIN + '"></script></head>')
        if self.changed_header:
            source = source.replace('TAHUN/BAGIAN TAHUN PAJAK', 'SYNTHETIC CHANGED HEADER')
        if self.trap:
            source = source.replace('</body>', '''<script>
                window.businessReads=0;
                for(const [proto,key] of [[HTMLInputElement.prototype,'value'],
                    [HTMLInputElement.prototype,'checked'],[HTMLElement.prototype,'innerText'],
                    [Node.prototype,'textContent']]){
                  const d=Object.getOwnPropertyDescriptor(proto,key);
                  Object.defineProperty(proto,key,{...d,get(){
                    if(this.matches?.('input,tbody,tbody *,tfoot td,tfoot td *')){
                      window.businessReads++;throw Error('DOM business value read');
                    }return d.get.call(this);
                  }});
                }
                </script></body>''')
        return source


@pytest.mark.parametrize('popup,model', [(True, 'Normal'), (False, 'Amendment 002')])
def test_native_l7_uses_no_business_dom_values_and_exports_source_bound_workbook(
        browser, tmp_path, monkeypatch, popup, model):
    fixture = L7Fixture(browser, tmp_path, monkeypatch, popup=popup, model=model)
    monkeypatch.setattr(driver, 'capture_detail', lambda *args, **kwargs:
                        pytest.fail('L7 native source unexpectedly entered whole DOM fallback'))
    try:
        result = fixture.collect()
        assert result['complete'] and result['errors'] == [], result['errors']
        record = result['records'][0]
        assert record['detail_source']['used'] == 'json'
        section, = [section for section in record['sections'] if section['name'] == 'L7']
        assert section['tables'][0]['headers'] == HEADERS
        assert record['attachments'][0]['status'] == 'COMPLETE'
        assert len([r for r in fixture.requests if r['url'] == PDF_URL]) == 1
        assert not any('/loadcit/' in r['url'] for r in fixture.requests)
        assert fixture.unexpected == []
        detail, = [page for page in fixture.context.pages if page.url == FRAME]
        assert detail.evaluate('window.businessReads') == 0
        assert_saved_l7(result, tmp_path, model)
    finally:
        fixture.close()


@pytest.mark.parametrize('defect', ['changed-main', 'changed-header'])
def test_changed_l7_source_or_header_uses_whole_dom(browser, tmp_path, monkeypatch, defect):
    fixture = L7Fixture(browser, tmp_path, monkeypatch, trap=False,
                        changed_main=defect == 'changed-main', changed_header=defect == 'changed-header')
    try:
        result = fixture.collect()
        assert result['complete'] and result['errors'] == [], result['errors']
        assert result['records'][0]['detail_source']['used'] == 'dom'
        assert fixture.unexpected == []
    finally:
        fixture.close()


def test_l7_native_fixture_is_independently_valid_before_native_support(browser, tmp_path, monkeypatch):
    fixture = L7Fixture(browser, tmp_path, monkeypatch, trap=False)
    try:
        fixture.page.goto(corpus.LIST_URL, wait_until='domcontentloaded')
        with fixture.context.expect_page() as opened:
            fixture.page.locator('#ViewButton').click()
        detail = opened.value
        detail.wait_for_url(FRAME)
        detail.wait_for_load_state('domcontentloaded')
        detail.get_by_role('tab', name='L7', exact=True).click()
        assert detail.locator('rsh-citr-l7-grid1 tbody td').first.inner_text() == '999999'
        assert detail.locator('rsh-citr-l7-grid1 tfoot td').first.get_attribute('colspan') == '3'
        detail.wait_for_function("() => performance.getEntriesByType('resource').some(e=>e.name.endsWith('/main.cccfdb820c20596c.js'))")
        assert fixture.unexpected == []
    finally:
        fixture.close()
