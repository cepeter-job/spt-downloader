"""Synthetic bounded stream use and unchanged physical proof refusals."""
from copy import deepcopy
from pathlib import PurePosixPath

from openpyxl import load_workbook
import pytest

from test_mpn_publication import exporter, mpn_run


def rows_run(count, amount='-1250.50'):
    run = mpn_run(amount=amount)
    source = run['mpn']['rows'][0]
    run['mpn'].update(headers=1 if count else 0, details=count, rows=[
        {**deepcopy(source), 'record_id': f'00000000-0000-4000-8000-{index+10:012d}'}
        for index in range(count)])
    return run


@pytest.mark.parametrize('count', [0, 24])
@pytest.mark.parametrize('amount', ['-1250.50', '-1234567890123456'])
def test_read_only_mpn_proof_uses_bounded_worksheet_streams(tmp_path, monkeypatch, count, amount):
    import zipfile
    run = rows_run(count, amount)
    path = tmp_path / 'MPN_2025.xlsx'
    exporter().export_mpn_workbook(run, path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    opened = []
    original = zipfile.ZipFile.open
    def counted(archive, name, *args, **kwargs):
        member = name.filename if isinstance(name, zipfile.ZipInfo) else name
        if PurePosixPath(member).parts[:2] == ('xl', 'worksheets'):
            opened.append(member)
        return original(archive, name, *args, **kwargs)
    monkeypatch.setattr(zipfile.ZipFile, 'open', counted)
    try:
        proof = exporter().read_mpn_proof(workbook)
        assert proof['rows'] == run['mpn']['rows']
        assert len(opened) <= 2
        assert set(opened) == {'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'}
    finally:
        workbook.close()


@pytest.mark.parametrize('changed', ['amount', 'literal_type', 'audit_status'])
def test_late_physical_changes_remain_fatal_under_sequential_verification(tmp_path, changed):
    run = rows_run(24)
    path = tmp_path / 'MPN_2025.xlsx'
    exporter().export_mpn_workbook(run, path)
    workbook = load_workbook(path)
    if changed == 'amount':
        workbook['MPN']['K28'] = -1
    elif changed == 'literal_type':
        # Keep the same apparent text while changing a literal into a formula.
        workbook['MPN']['L28'] = '=SYNTHETIC-REFERENCE'
    else:
        workbook['Audit']['B5'] = 'PARTIAL'
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match='MPN'):
        exporter().verify_mpn_workbook(run, path)
