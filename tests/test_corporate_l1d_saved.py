"""Independent final-export regression for source-proven 5322/group geometry."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

import coretax
from test_corporate_l1d_projection import ORDER, HEADERS, SPANS


def tables():
    body = []
    for code in ORDER:
        if not code.isdecimal():
            body.append(['', code, *([''] * 8)])
        else:
            values = ['6' if code == '4500' else '7' if code == '4700' else '0'] * 7
            if code == '5322':
                values[0] = '77'
            body.append([code, 'SYNTHETIC ' + code, *values[:6], '', values[6]])
    parts = [(body, ['4800', 'Laba (Rugi) Sebelum Pajak', *(['13'] * 6), '', '13'])]
    values = {'AC1101': '16', 'AC1522': '5', 'AC1700': '11', 'AC2102': '100',
              'AC3102': '11', 'AC2999': '100', 'AC3299': '11', 'AC3300': '111'}
    for descriptors in _BALANCE_ROWS:
        rows, footer = [], None
        for location, code, caption, control in descriptors:
            row = [code or '', caption, '' if control is None else values.get(control, '0')]
            if location == 'footer':
                footer = row
            else:
                rows.append(row)
        parts.append((rows, footer))
    return [{'name': f'L1-D/grid{i+1}', 'headers': deepcopy(HEADERS[i]),
        'rows': rows + [footer], 'footer': [footer], 'footer_spans': deepcopy(SPANS[i]),
        'body_row_count': len(rows), 'source_total': len(rows), 'complete': True,
        'source_structure': {'form': 'L1-D', 'component': f'rsh-citr-l1d-grid{i+1}',
            'groups': ['A. LAPORAN LABA RUGI' if i == 0 else 'B. LAPORAN POSISI KEUANGAN']}}
        for i, (rows, footer) in enumerate(parts)]


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
def test_full_l1d_final_save_moves_known_groups_for_5322_and_keeps_totals(tmp_path, model):
    sections = [{'name': 'L1-D', 'fields': [], 'tables': tables(), 'text': ''}]
    before = deepcopy(sections)
    record = {'record_id': 'SYNTHETIC-L1D', 'meta': {'Jenis Pajak': 'PPh Badan',
        'Masa Pajak': 'Januari - Desember 2025', 'Model SPT': model},
        'capture_complete': False, 'sections': sections}
    run = {'npwp': '0000000000000001', 'year': 2025, 'year_basis': 'period',
        'inventory_complete': True, 'complete': False, 'errors': [], 'records': [record]}
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    assert sections == before
    report, = [r for r in run['template_mapping'] if r['form'] == 'L1-D']
    assert report['complete'] and report['unmapped_rows'] == report['unmapped_tables'] == 0
    name, = [name for name in run['output_files'] if name != 'resume.xlsx']
    assert model in name
    workbook = load_workbook(tmp_path / name)
    try:
        sheet = workbook['L1-D']
        assert sheet.max_row == 98 and sheet['A1'].value == 'KODE AKUN'
        assert sheet['A20'].value == '5322' and sheet['C20'].value == 77
        assert sheet['A37'].value == '4800' and sheet['J37'].value == 13
        assert sheet['A98'].value == '3300' and sheet['C98'].value == 111
        merges = {str(region) for region in sheet.merged_cells.ranges}
        for row in (2, 8, 24, 30):
            assert f'A{row}:B{row}' in merges
        captions = ['Pendapatan:', 'Beban Usaha', 'Pendapatan Non Usaha', 'Beban Non Usaha']
        assert [sheet.cell(row, 1).value for row in (2, 8, 24, 30)] == captions
        depreciation, = [row for row in sheet.iter_rows() if row[1].value ==
                         'Dikurangi: Akumulasi penyusutan aset tetap dan inventaris']
        assert depreciation[2].value == -5
        for cell in ('C20', 'C37', 'J37', 'C98'):
            assert 'Rp' in sheet[cell].number_format and '0.00' not in sheet[cell].number_format
    finally:
        workbook.close()


# Static captions independently read from the pinned source, not generated output.
_BALANCE_ROWS = [[('body', None, 'Aset Lancar', None),
  ('body', '1101', 'Kas dan Setara Kas', 'AC1101'),
  ('body', '1122', 'Piutang Usaha - Pihak Ketiga', 'AC1122'),
  ('body', '1123', 'Piutang Usaha - Pihak yang Mempunyai Hubungan Istimewa', 'AC1123'),
  ('body', '1124', 'Piutang Lainnya - Pihak Ketiga', 'AC1124'),
  ('body', '1125', 'Piutang Lainnya - Pihak yang Mempunyai Hubungan Istimewa', 'AC1125'),
  ('body', '1131', 'Cadangan Kerugian Penurunan Nilai - Aset Lancar', 'AC1131'),
  ('body', '1181', 'Aset Kontrak', 'AC1181'),
  ('body', '1200', 'Investasi', 'AC1200'),
  ('body', '1401', 'Persediaan', 'AC1401'),
  ('body', '1421', 'Beban Dibayar di Muka', 'AC1421'),
  ('body', '1422', 'Uang Muka', 'AC1422'),
  ('body', '1423', 'Pajak Dibayar di Muka', 'AC1423'),
  ('body', '1499', 'Aset lancar lainnya', 'AC1499'),
  ('body', None, 'Aset Tidak Lancar', None),
  ('body', '1501', 'Piutang Jangka Panjang', 'AC1501'),
  ('body', '1521', 'Aset tetap dan inventaris', 'AC1521'),
  ('body', '1522', 'Dikurangi: Akumulasi penyusutan aset tetap dan inventaris', 'AC1522'),
  ('body',
   '1551',
   'Investasi pada Perusahaan Asosiasi, Ventura Bersama dan Anak Perusahaan',
   'AC1551'),
  ('body', '1599', 'Investasi Jangka Panjang Lainnya', 'AC1599'),
  ('body', '1600', 'Aset Tak Berwujud', 'AC1600'),
  ('body', '1601', 'Dikurangi : Akumulasi Amortisasi-Aset Tak Berwujud', 'AC1601'),
  ('body', '1611', 'Aktiva Pajak Tangguhan', 'AC1611'),
  ('body', '1651', 'Klaim atas pengembalian pajak', 'AC1651'),
  ('body', '1658', 'Cadangan Kerugian Penurunan Nilai - Aset Tidak Lancar', 'AC1658'),
  ('body', '1698', 'Aset Tidak Lancar Lainnya', 'AC1698'),
  ('footer', '1700', 'Jumlah Aset', 'AC1700')],
 [('body', None, 'Liabilitas Jangka Pendek', None),
  ('body', '2102', 'Utang Usaha - Pihak Ketiga', 'AC2102'),
  ('body', '2103', 'Utang Usaha - Pihak yang Mempunyai Hubungan Istimewa', 'AC2103'),
  ('body', '2111', 'Utang Bunga', 'AC2111'),
  ('body', '2186', 'Liabilitas Kontrak', 'AC2186'),
  ('body', '2187', 'Liabilitas Sewa Jangka Pendek', 'AC2187'),
  ('body', '2191', 'Utang Pajak', 'AC2191'),
  ('body', '2192', 'Utang Dividen', 'AC2192'),
  ('body', '2195', 'Beban yang Masih Harus Dibayar', 'AC2195'),
  ('body', '2201', 'Utang Bank Jangka Pendek', 'AC2201'),
  ('body', '2202', 'Utang Jangka Panjang yang Jatuh Tempo dalam Satu Tahun', 'AC2202'),
  ('body', '2203', 'Pendapatan Diterima di Muka', 'AC2203'),
  ('body', '2228', 'Liabilitas Jangka Pendek Lainnya', 'AC2228'),
  ('body', None, 'Liabilitas Jangka Panjang', None),
  ('body', '2301', 'Utang Bank Jangka Panjang', 'AC2301'),
  ('body', '2303', 'Utang Jangka Panjang-Pihak Ketiga', 'AC2303'),
  ('body', '2304', 'Utang Jangka Panjang - Pihak yang Mempunyai Hubungan Istimewa', 'AC2304'),
  ('body', '2312', 'Liabilitas Sewa Jangka Panjang', 'AC2312'),
  ('body', '2321', 'Liabilitas Pajak Tangguhan', 'AC2321'),
  ('body', '2322', 'Liabilitas Imbalan Kerja', 'AC2322'),
  ('body', '2998', 'Liabilitas Jangka Panjang Lainnya', 'AC2998'),
  ('body', '2999', 'Jumlah Liabilitas', 'AC2999'),
  ('body', None, 'Ekuitas', None),
  ('body', '3102', 'Modal Saham', 'AC3102'),
  ('body', '3120', 'Tambahan Modal Disetor', 'AC3120'),
  ('body', '3200', 'Laba Ditahan', 'AC3200'),
  ('body', '3297', 'Pendapatan Komprehensif Lainnya', 'AC3297'),
  ('body', '3298', 'Ekuitas Lainnya', 'AC3298'),
  ('body', '3299', 'Jumlah Ekuitas', 'AC3299'),
  ('footer', '3300', 'Jumlah Liabilitas dan Ekuitas', 'AC3300')]]
