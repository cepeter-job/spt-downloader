"""The operator layout is exercised using invented scope and zero source rows."""

from openpyxl import load_workbook
import pytest


def empty_run():
    return {'npwp': '0000000000000001', 'year': 2025, 'year_basis': 'period',
            'taxpayer_name': 'Synthetic template taxpayer', 'records': [],
            'audit': [], 'errors': [], 'complete': False, 'started_at': 'synthetic'}


@pytest.mark.parametrize('name', ['Ringkasan', 'Daftar SPT', 'Audit'])
def test_saved_resume_column_defaults_do_not_border_empty_space(tmp_path, name):
    from copy import copy
    from coretax_exporter import export_workbook
    from coretax_resume_layout import RESOURCE
    path = tmp_path / 'resume.xlsx'
    export_workbook(empty_run(), path, include_details=False)
    actual, reference = load_workbook(path), load_workbook(RESOURCE)
    try:
        for key, expected in reference[name].column_dimensions.items():
            column = actual[name].column_dimensions[key]
            assert column.width == expected.width
            assert (column.min, column.max) == (expected.min, expected.max)
            assert all(getattr(column.border, edge).style is None
                       for edge in ('left', 'right', 'top', 'bottom'))
            assert column.number_format == 'General'
            assert copy(column.font) == copy(expected.font)
    finally:
        actual.close()
        reference.close()


def test_saved_resume_row_defaults_preserve_the_reference_font_and_border(tmp_path):
    from copy import copy
    from coretax_exporter import export_workbook
    from coretax_resume_layout import RESOURCE
    path = tmp_path / 'resume.xlsx'
    export_workbook(empty_run(), path, include_details=False)
    actual, reference = load_workbook(path), load_workbook(RESOURCE)
    try:
        for row in (1, 3, 10):
            dimension = actual['Audit'].row_dimensions[row]
            expected = reference['Audit'].row_dimensions[row]
            assert copy(dimension.font) == copy(expected.font)
            assert all(getattr(dimension.border, edge).style is None
                       for edge in ('left', 'right', 'top', 'bottom'))
    finally:
        actual.close()
        reference.close()


@pytest.mark.parametrize('axis,key', [('column_dimensions', 'B'), ('row_dimensions', 7)])
def test_reset_sheet_transfers_dimension_styles_between_distinct_workbooks(tmp_path, axis, key):
    from copy import copy
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from coretax_resume_layout import _reset_sheet
    source, target = Workbook(), Workbook()
    reference, sheet = source.active, target.active
    dimension = getattr(reference, axis)[key]
    dimension.font = Font(name='Arial', bold=True, color='123456')
    dimension.fill = PatternFill('solid', fgColor='ABCDEF')
    dimension.border = Border(bottom=Side(style='double', color='654321'))
    dimension.alignment = Alignment(horizontal='right', wrap_text=True)
    dimension.protection = Protection(locked=False, hidden=True)
    dimension.number_format = '0.0000 "synthetic"'
    dimension.hidden, dimension.outlineLevel = True, 2
    if axis == 'column_dimensions':
        dimension.width, dimension.min, dimension.max = 23.5, 2, 4
    else:
        dimension.height = 31.5
    # Deliberately allocate incompatible destination IDs before transferring.
    sheet['A1'].font = Font(name='Courier New', italic=True)
    sheet['A1'].border = Border(top=Side(style='thick'))
    sheet['A1'].fill = PatternFill('solid', fgColor='FEDCBA')
    sheet['A1'].alignment = Alignment(horizontal='left')
    sheet['A1'].number_format = '0.0 "destination"'
    expected = {name: copy(getattr(dimension, name)) for name in
                ('font', 'fill', 'border', 'alignment', 'protection', 'number_format')}
    _reset_sheet(sheet, reference)
    path = tmp_path / 'dimensions.xlsx'
    target.save(path)
    restored = load_workbook(path)
    try:
        actual = getattr(restored.active, axis)[key]
        for name, value in expected.items():
            assert copy(getattr(actual, name)) == value
        assert actual.hidden and actual.outlineLevel == 2
        if axis == 'column_dimensions':
            assert (actual.width, actual.min, actual.max) == (23.5, 2, 4)
        else:
            assert actual.height == 31.5
        assert getattr(sheet, axis)[key].parent is sheet
        assert dimension.parent is reference
    finally:
        restored.close()
        source.close()
        target.close()


def test_resume_uses_the_operator_reference_coordinates_without_sample_rows(tmp_path):
    from coretax_exporter import export_workbook
    run = empty_run()
    path = tmp_path / 'resume.xlsx'
    export_workbook(run, path, include_details=False)
    wb = load_workbook(path)
    try:
        assert wb.sheetnames == ['Ringkasan', 'Daftar SPT', 'Audit']
        summary = wb['Ringkasan']
        assert summary['B2'].value == run['npwp']
        assert summary['B2'].data_type == 's'
        assert summary['B3'].value == run['year']
        assert summary['B4'].value == run['taxpayer_name']
        assert summary['B5'].value == 'Masa/Tahun Pajak'
        assert summary['B6'].value == 'PARTIAL'
        assert summary['B7'].value == summary['B8'].value == 0
        assert summary.column_dimensions['A'].width == 30
        assert summary.column_dimensions['B'].width == 58
        assert 'A1:H1' in summary.merged_cells
        assert summary.freeze_panes is None
        assert summary.print_area.endswith('$H$18')
        index = wb['Daftar SPT']
        assert index.max_row == 4
        assert index.freeze_panes is None
        assert index.print_area.endswith('$K$8')
        assert index['A5'].hyperlink is None
        assert wb['Audit']['A3'].value == 'Status / Objek'
        assert wb['Audit']['A1'].value == 'AUDIT | Kelengkapan'
        assert wb['Audit'].freeze_panes == 'A4'
    finally:
        wb.close()


def test_new_resume_scope_is_verified_independently_from_display_status(tmp_path):
    from coretax_exporter import export_workbook
    from coretax import verify_workbook
    run = empty_run()
    path = tmp_path / 'resume.xlsx'
    export_workbook(run, path, include_details=False)
    verify_workbook(path, run, include_details=False)
    wb = load_workbook(path)
    wb['Ringkasan']['B2'] = '0000000000000002'
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match='NPWP'):
        verify_workbook(path, run, include_details=False)


def test_resource_contains_only_static_cells_and_no_retained_operational_parts():
    from coretax_resume_layout import RESOURCE
    import zipfile
    import xml.etree.ElementTree as ET
    wb = load_workbook(RESOURCE)
    try:
        populated = {name: {cell.coordinate for row in wb[name] for cell in row
                            if cell.value is not None} for name in wb.sheetnames}
        assert populated['Ringkasan'] == {'A1', *{f'A{i}' for i in range(2, 10)},
                                         'A13', 'B13', 'A18'}
        assert populated['Daftar SPT'] == {'A1', 'A2', *{f'{c}4' for c in 'ABCDEFGHIJK'}}
        assert populated['Audit'] == {'A1', *{f'{c}{r}' for c in 'ABCDEF' for r in (3, 10)}}
        assert not any(cell.hyperlink for sheet in wb for row in sheet for cell in row)
    finally:
        wb.close()
    with zipfile.ZipFile(RESOURCE) as archive:
        assert not any('printerSettings' in name or 'sharedStrings' in name
                       or name.startswith('docProps/') for name in archive.namelist())
        for name in archive.namelist():
            root = ET.fromstring(archive.read(name))
            assert not any(item.get('TargetMode') == 'External' for item in root.iter())
            if name.startswith('xl/charts/'):
                assert not any(item.tag.rsplit('}', 1)[-1] in ('strCache', 'numCache')
                               for item in root.iter())


@pytest.mark.parametrize('count', [1, 4, 9])
def test_dynamic_rows_chart_and_table_follow_only_actual_records(tmp_path, count):
    from coretax_exporter import export_workbook
    run = empty_run()
    run['records'] = [{'meta': {'Jenis Pajak': f'Synthetic tax {i}', 'Masa Pajak': '2025',
                              'Model SPT': 'Normal'}, 'sections': [], 'capture_complete': False}
                      for i in range(count)]
    run['template_mapping'] = [{'file': 'synthetic.xlsx', 'form': 'Induk', 'complete': False}]
    path = tmp_path / 'resume.xlsx'
    export_workbook(run, path, include_details=False)
    wb = load_workbook(path)
    try:
        summary, index, audit = (wb[name] for name in ('Ringkasan', 'Daftar SPT', 'Audit'))
        assert summary['B7'].value == count
        assert [summary.cell(row, 2).value for row in range(14, 14 + count)] == [1] * count
        assert summary.cell(14 + count, 1).value == 'mapping template sedot terbaru'
        assert summary.cell(14 + count, 2).value == 'PARTIAL'
        chart = summary._charts[0]
        assert chart.anchor._from.col == 3 and chart.anchor._from.row == 12
        assert (chart.anchor.ext.cx, chart.anchor.ext.cy) == (6480000, 3240000)
        # ChartSpace.style is persisted but not restored by openpyxl's reader.
        import zipfile
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read('xl/charts/chart1.xml'))
            assert root.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}style').get('val') == '2'
        expected = "'Ringkasan'!$B$14" + (f':$B${13 + count}' if count > 1 else '')
        assert chart.series[0].val.numRef.f == expected
        assert chart.series[0].val.numRef.numCache is None
        assert index.max_row == count + 4
        assert index.tables['SPTIndex'].ref == f'A4:K{count + 4}'
        assert index['B5'].border.left.style == 'thin'
        from coretax_resume_layout import RESOURCE
        reference = load_workbook(RESOURCE)
        try:
            assert index.row_dimensions[count + 4].height == reference['Daftar SPT'].row_dimensions[5].height
            assert index.sheet_format.defaultRowHeight == reference['Daftar SPT'].sheet_format.defaultRowHeight
        finally:
            reference.close()
        assert audit['B4'].value == 'Field belum dipetakan'
        assert audit['B5'].value == 0
    finally:
        wb.close()


def test_hidden_administrative_sheet_cannot_prove_a_visible_export_scope(tmp_path):
    from coretax_exporter import export_workbook
    from coretax_resume_layout import summary_scope
    path = tmp_path / 'resume.xlsx'
    export_workbook(empty_run(), path, include_details=False)
    wb = load_workbook(path)
    try:
        wb['Ringkasan'].sheet_state = 'hidden'
        with pytest.raises(ValueError, match='scope|NPWP|visible'):
            summary_scope(wb)
    finally:
        wb.close()


def legacy_administration():
    from openpyxl import Workbook
    from coretax_resume_layout import LEGACY_LABELS, INDEX_HEADERS, AUDIT_HEADERS
    wb = Workbook()
    wb.active.title = 'Ringkasan'
    wb.create_sheet('Daftar SPT')
    wb.create_sheet('Audit')
    values = ['0000000000000001', 2025, 'Synthetic legacy taxpayer',
              'Masa/Tahun Pajak', 'PARTIAL', 0, 0, 'synthetic']
    for row, (label, value) in enumerate(zip(LEGACY_LABELS, values), 4):
        wb['Ringkasan'].cell(row, 1, label)
        wb['Ringkasan'].cell(row, 2, value)
    for col, value in enumerate(INDEX_HEADERS, 1):
        wb['Daftar SPT'].cell(4, col, value)
    for col, value in enumerate(AUDIT_HEADERS, 1):
        wb['Audit'].cell(4, col, value)
    return wb


def test_complete_legacy_labels_remain_readable_without_rewriting():
    from coretax_resume_layout import summary_scope
    wb = legacy_administration()
    try:
        scope = summary_scope(wb)
        assert scope['layout'] == 'legacy' and scope['npwp'] == '0000000000000001'
        assert scope['year'] == 2025 and scope['record_count'] == 0
        assert wb['Ringkasan']['A4'].value == 'NPWP'
        assert wb['Ringkasan']['B2'].value is None
        wb['Ringkasan']['A5'] = 'Status'
        with pytest.raises(ValueError, match='scope|NPWP'):
            summary_scope(wb)
    finally:
        wb.close()
