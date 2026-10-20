from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import coretax_exporter as app
from openpyxl import load_workbook
import pytest


def test_workbook_preserves_npwp_as_text(tmp_path):
    assert callable(getattr(app, 'export_workbook', None)), 'Workbook exporter is not implemented'
    run = {'npwp': '0000000000000001', 'year': 2025, 'year_basis': 'period',
           'taxpayer_name': 'Authorized taxpayer', 'records': [], 'audit': [], 'errors': [],
           'complete': False, 'started_at': '2026-10-09T10:00:00+07:00'}
    path = tmp_path / 'report.xlsx'
    app.export_workbook(run, path)
    wb = load_workbook(path)
    assert wb['Ringkasan']['B4'].value == run['npwp']
    assert wb['Ringkasan']['B4'].data_type == 's'
    assert wb['Ringkasan']['B5'].value == 2025
    assert wb['Ringkasan']['B8'].value == 'PARTIAL'


def test_detail_workbook_is_navigable_and_treats_source_formulas_as_text(tmp_path):
    run = {'npwp': '0000000000000001', 'year': 2025, 'year_basis': 'period',
           'taxpayer_name': 'Test fixture only', 'complete': True, 'started_at': 'fixture',
           'scope_description': 'SHORT TEST: exactly one SPT; maximum one page per detail table.',
           'audit': [{'status': 'Dilaporkan', 'total': 1, 'scanned': 1, 'selected': 1}], 'errors': [],
           'records': [{'key': 'fixture', 'meta': {'Jenis Pajak': 'PPh Badan', 'Masa Pajak': '2025', 'Model SPT': 'Normal'},
                        'source_status': 'Dilaporkan', 'detail_url': 'https://coretax.intranet.pajak.go.id/returnsheets/id-ID/test',
                        'capture_complete': True, 'sections': [{'name': 'Induk', 'text': 'Source fixture',
                           'fields': [{'label': 'Value', 'field': 'Value', 'value': '=1+1'}],
                           'tables': [{'name': 'Fixture', 'headers': ['Kode', 'Nilai'], 'rows': [['001', '1.000,00'], ['001', '1.000,00']], 'complete': True}]}]}]}
    path = tmp_path / 'detail.xlsx'
    app.export_workbook(run, path)
    wb = load_workbook(path)
    assert 'Daftar SPT' in wb.sheetnames, 'SPT index missing'
    assert 'Audit' in wb.sheetnames
    assert wb['Ringkasan']['B13'].value == run['scope_description']
    index = wb['Daftar SPT']
    assert index['A5'].hyperlink is not None, 'Detail navigation missing'
    assert index.freeze_panes == 'D5'
    detail = wb[wb.sheetnames[3]]
    source_cells = [c for row in detail for c in row if c.value == '=1+1']
    assert source_cells and all(c.data_type == 's' for c in source_cells)
    assert sum(c.value == '001' for row in detail for c in row) == 2, 'Duplicate source rows lost'
    assert detail.page_setup.orientation == 'landscape'
    assert detail.print_area


def test_legacy_tax_period_and_duplicate_status_exposures():
    assert not app.in_scope({'Masa Pajak': '03032016'}, 2025, 'period')
    assert app.in_scope({'Masa Pajak': '03032016'}, 2016, 'period')
    assert callable(getattr(app, 'deduplicate_inventory', None)), 'Cross-status deduplication missing'
    raw = [{'record_id': 'normal', 'key': 'submitted/normal', 'meta': {'Masa Pajak': '2025', 'Model SPT': 'Normal'}, 'source_status': 'Submitted', 'list_url': 'submitted', 'list_page': 1, 'row_position': 0},
           {'record_id': 'normal', 'key': 'archive/normal', 'meta': {'Masa Pajak': '2025', 'Model SPT': 'Normal'}, 'source_status': 'Archive', 'list_url': 'archive', 'list_page': 1, 'row_position': 0},
           {'record_id': 'amendment', 'key': 'submitted/amendment', 'meta': {'Masa Pajak': '2025', 'Model SPT': 'Pembetulan ke-1'}, 'source_status': 'Submitted', 'list_url': 'submitted', 'list_page': 1, 'row_position': 1}]
    unique = app.deduplicate_inventory(raw)
    assert len(unique) == 2, 'Normal return and amendment must remain separate'
    assert len(unique[0]['source_occurrences']) == 2


def test_scope_uses_tax_period_not_submission_date():
    assert callable(getattr(app, 'validate_scope', None)), 'Scope validation missing'
    assert app.validate_scope('0000000000000001', '2025') == ('0000000000000001', 2025)
    record = {'Masa Pajak': 'Desember 2025', 'Tanggal Submit SPT': '20-01-2026'}
    assert app.in_scope(record, 2025, 'period')
    assert not app.in_scope(record, 2025, 'submission')
    assert app.in_scope(record, 2026, 'submission')
    for npwp in ['0.000.000-000.0001', '000000000000001', ' 0000000000000001']:
        import pytest
        with pytest.raises(ValueError):
            app.validate_scope(npwp, '2025')
    assert app.in_scope({'Masa Pajak': '2024/2025'}, 2025, 'period')
    assert not app.in_scope({'Masa Pajak': 'Desember 2024'}, 2025, 'period')


def period_conflict_fixture():
    from datetime import datetime, timedelta, timezone
    from copy import deepcopy
    first = {'record_id': 'synthetic-same-id', 'key': 'submitted/synthetic',
             'source_status': 'SPT Dilaporkan', 'list_url': 'synthetic-submitted',
             'list_page': 1, 'row_position': 0,
             'meta': {'Masa Pajak': '01012024', 'Model SPT': 'Normal',
                      'Jenis Pajak': 'PPN', 'Nomor Bukti Penerimaan': '000001'}}
    second = deepcopy(first)
    second.update(key='archive/synthetic', source_status='Archieved Return Sheets',
                  list_url='synthetic-archive')
    second['meta']['Masa Pajak'] = 'Januari 2024'
    catalog = {'ReferenceDataType': 'PERIOD', 'language': 'id-ID',
               'Checksum': 'synthetic-period-v1',
               'ExpiryDate': (datetime.now(timezone.utc) + timedelta(days=1)).isoformat(),
               'Details': [{'Code': '01012024', 'CodeDescription': 'Januari 2024'}]}
    return [first, second], catalog


def source_proven_dedup(records, catalog):
    from inspect import signature
    assert 'period_catalog' in signature(app.deduplicate_inventory).parameters, 'Source-proven period equivalence is missing'
    return app.deduplicate_inventory(records, period_catalog=catalog, year=2024)


@pytest.mark.parametrize('reverse', [False, True])
def test_catalog_proven_period_alias_preserves_both_literal_occurrences(reverse):
    from copy import deepcopy
    records, catalog = period_conflict_fixture()
    if reverse:
        records.reverse()
    before = deepcopy(records)
    unique = source_proven_dedup(records, catalog)
    assert len(unique) == 1
    assert unique[0]['meta'] == before[0]['meta']
    assert unique[0]['list_url'] == before[0]['list_url']
    assert [o['meta'] for o in unique[0]['source_occurrences']] == [r['meta'] for r in before]
    proof = unique[0]['period_equivalence']
    assert proof['code'] == '01012024' and proof['label'] == 'Januari 2024'
    assert len(proof['catalog_sha256']) == len(proof['catalog_binding_sha256']) == 64
    assert records == before
    unique[0]['source_occurrences'][0]['meta']['Masa Pajak'] = 'mutated test copy'
    assert records == before


@pytest.mark.parametrize('problem', ['absent', 'expired', 'wrong_language', 'wrong_label',
                                    'duplicate_code', 'duplicate_label', 'ambiguous_code_label',
                                    'changed_model', 'changed_receipt', 'third_period'])
def test_period_alias_never_masks_unproved_or_other_metadata_conflicts(problem):
    from copy import deepcopy
    records, catalog = period_conflict_fixture()
    if problem == 'absent':
        catalog = None
    elif problem == 'expired':
        catalog['ExpiryDate'] = '2000-01-01T00:00:00Z'
    elif problem == 'wrong_language':
        catalog['language'] = 'en-US'
    elif problem == 'wrong_label':
        catalog['Details'][0]['CodeDescription'] = 'Februari 2024'
    elif problem == 'duplicate_code':
        catalog['Details'].append(dict(catalog['Details'][0]))
    elif problem == 'duplicate_label':
        catalog['Details'].append({'Code': 'OTHER', 'CodeDescription': 'Januari 2024'})
    elif problem == 'ambiguous_code_label':
        catalog['Details'].append({'Code': 'Januari 2024', 'CodeDescription': 'Other source label 2024'})
    elif problem == 'changed_model':
        records[1]['meta']['Model SPT'] = 'Pembetulan 1'
    elif problem == 'changed_receipt':
        records[1]['meta']['Nomor Bukti Penerimaan'] = '000002'
    else:
        records.append(deepcopy(records[1]))
        records[-1]['meta']['Masa Pajak'] = 'Februari 2024'
    with pytest.raises(ValueError, match='Conflicting source metadata'):
        source_proven_dedup(records, catalog)


def test_period_alias_does_not_merge_distinct_return_ids():
    records, catalog = period_conflict_fixture()
    records[1]['record_id'] = 'synthetic-distinct-amendment'
    unique = source_proven_dedup(records, catalog)
    assert len(unique) == 2
    assert all('period_equivalence' not in r for r in unique)
