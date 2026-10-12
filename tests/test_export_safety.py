"""Export safety regressions using synthetic data only."""
from copy import deepcopy
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import coretax_exporter as app


def fixture_run():
    return {
        'npwp': '0000000000000001', 'year': 2025, 'year_basis': 'period',
        'taxpayer_name': 'SYNTHETIC FIXTURE ONLY', 'complete': True,
        'audit': [], 'errors': [],
        'records': [{
            'record_id': 'synthetic-return', 'source_status': 'Fixture',
            'meta': {'Jenis Pajak': 'PPh Badan', 'Masa Pajak': '2025',
                     'Model SPT': 'Normal', 'Tanggal Submit SPT': '20-01-2026'},
            'capture_complete': True,
            'sections': [{'name': 'Induk', 'text': '', 'tables': [],
                          'fields': [{'label': 'Synthetic field', 'value': '001'}]}],
        }],
    }


@pytest.mark.parametrize('value', ['SYNTHETIC-' + 'x' * 32768,
                                  'SYNTHETIC\x0bCONTROL', 'SYNTHETIC\ufffeXML',
                                  9007199254740993, 1.2345678901234567,
                                  float('nan'), float('inf'),
                                  1234567890123456, 1.234567890123456,
                                  1e-320, 5e-324, -1e-320, -0.0, 10**24],
                         ids=['long-text', 'xml-control', 'xml-noncharacter',
                              'large-integer', 'precise-float', 'nan', 'infinity',
                              'excel-integer-precision', 'excel-float-precision',
                              'subnormal', 'smallest-subnormal', 'negative-subnormal',
                              'negative-zero', 'integer-binary-rounding'])
def test_unrepresentable_source_values_leave_existing_workbook_untouched(tmp_path, value):
    run = fixture_run()
    run['records'][0]['sections'][0]['fields'][0]['value'] = value
    path = tmp_path / 'existing.xlsx'
    original = b'SYNTHETIC EXISTING OUTPUT'
    path.write_bytes(original)

    with pytest.raises(ValueError, match='represent.*Excel|Excel.*represent'):
        app.export_workbook(run, path)

    assert path.read_bytes() == original


@pytest.mark.parametrize('value', [0, 0.0, 123456789012345, 1234567890123450, 10**20,
                                  1.23456789012345, 1.25, -2.5, 3e-308])
def test_supported_excel_numbers_round_trip_unchanged(tmp_path, value):
    run = fixture_run()
    run['records'][0]['sections'][0]['fields'][0]['value'] = value
    path = tmp_path / 'report.xlsx'
    app.export_workbook(run, path)

    wb = load_workbook(path)
    try:
        assert wb.worksheets[3]['F6'].value == value
        assert wb.worksheets[3]['F6'].data_type == 'n'
    finally:
        wb.close()


def test_split_preflight_checks_later_type_before_writing_any_workbook(tmp_path):
    run = fixture_run()
    other = deepcopy(run['records'][0])
    other['record_id'] = 'synthetic-second-type'
    other['meta']['Jenis Pajak'] = 'PPh Unifikasi'
    other['sections'][0]['fields'][0]['value'] = 'x' * 32768
    run['records'].append(other)

    with pytest.raises(ValueError, match='represent.*Excel|Excel.*represent'):
        app.export_split_workbooks(run, tmp_path)

    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize('label', ['model-status', 'spt-period'])
def test_split_preflight_checks_combined_source_labels_before_writing(tmp_path, label):
    run = fixture_run()
    other = deepcopy(run['records'][0])
    other['record_id'] = 'synthetic-second-type'
    other['meta']['Jenis Pajak'] = 'PPh Unifikasi'
    if label == 'model-status':
        other['meta']['Model SPT'] = 'M' * 20000
        other['source_status'] = 'S' * 20000
    else:
        other['meta']['Jenis Surat Pemberitahuan Pajak'] = 'T' * 20000
        other['meta']['Masa Pajak'] = ('P' * 20000) + ' 2025'
    run['records'].append(other)

    with pytest.raises(ValueError, match='represent.*Excel|Excel.*represent'):
        app.export_split_workbooks(run, tmp_path)

    assert list(tmp_path.iterdir()) == []


def test_split_preflight_checks_generated_scope_text_before_writing(tmp_path):
    run = fixture_run()
    run['scope_description'] = 'S' * 32600
    other = deepcopy(run['records'][0])
    other['record_id'] = 'synthetic-second-type'
    other['meta']['Jenis Pajak'] = 'T' * 200
    run['records'].append(other)

    with pytest.raises(ValueError, match='represent.*Excel|Excel.*represent'):
        app.export_split_workbooks(run, tmp_path)

    assert list(tmp_path.iterdir()) == []


def test_long_page_context_round_trips_without_truncation(tmp_path):
    run = fixture_run()
    source = ('SYNTHETIC context\n=1+1\n' * 5000) + 'FINAL SOURCE MARKER'
    run['records'][0]['sections'][0]['text'] = source
    path = tmp_path / 'report.xlsx'
    app.export_workbook(run, path)

    wb = load_workbook(path)
    try:
        detail = wb.worksheets[3]
        heading = next(cell.row for row in detail for cell in row
                       if cell.value == 'Teks halaman / konteks sumber')
        chunks = [detail.cell(row, 1).value for row in range(heading + 1, detail.max_row + 1)
                  if detail.cell(row, 1).value is not None]
        assert ''.join(chunks) == source
        assert all(detail.cell(heading + index + 1, 1).data_type == 's'
                   for index in range(len(chunks)))
    finally:
        wb.close()


@pytest.mark.parametrize('scope', [
    ['0000000000000002', 2025, 'period'],
    ['0000000000000001', 2024, 'period'],
    ['0000000000000001', 2025, 'submission'],
])
def test_mismatched_record_scope_cannot_be_exported(tmp_path, scope):
    run = fixture_run()
    run['records'][0]['scope'] = scope

    with pytest.raises(ValueError, match='scope'):
        app.export_split_workbooks(run, tmp_path)

    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize('basis,source', [
    ('period', {'Masa Pajak': '2024'}),
    ('submission', {'Tanggal Submit SPT': '20-01-2026'}),
])
def test_record_year_must_match_selected_basis(tmp_path, basis, source):
    run = fixture_run()
    run['year_basis'] = basis
    run['records'][0]['meta'].update(source)

    with pytest.raises(ValueError, match='year|scope'):
        app.export_split_workbooks(run, tmp_path)

    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize('field,value', [('npwp', '123'), ('npwp', 1),
                                        ('year', 2101), ('year_basis', 'unknown')])
def test_export_validates_top_level_scope(tmp_path, field, value):
    run = fixture_run()
    run[field] = value

    with pytest.raises(ValueError):
        app.export_split_workbooks(run, tmp_path)

    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize('contradiction', ['errors', 'capture', 'table', 'complete-string', 'capture-string'])
def test_inconsistent_completion_is_rejected_before_export(tmp_path, contradiction):
    run = fixture_run()
    if contradiction == 'errors':
        run['errors'] = [{'message': 'SYNTHETIC capture failure'}]
    elif contradiction == 'capture':
        run['records'][0]['capture_complete'] = False
    elif contradiction == 'table':
        run['records'][0]['sections'][0]['tables'] = [
            {'name': 'Synthetic partial table', 'headers': ['Value'],
             'rows': [['001']], 'complete': False}]
    elif contradiction == 'complete-string':
        run['complete'] = 'false'
    else:
        run['records'][0]['capture_complete'] = 'false'

    with pytest.raises(ValueError, match='complet|COMPLETE'):
        app.export_split_workbooks(run, tmp_path)

    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize('inventory_complete', [False, 'false', 0, 1, None])
def test_explicit_inventory_completion_must_support_complete_output(tmp_path, inventory_complete):
    run = fixture_run()
    run['inventory_complete'] = inventory_complete
    path = tmp_path / 'report.xlsx'
    original = b'SYNTHETIC EXISTING OUTPUT'
    path.write_bytes(original)

    with pytest.raises(ValueError, match='inventory|COMPLETE'):
        app.export_workbook(run, path)

    assert path.read_bytes() == original


def test_explicit_incomplete_inventory_can_be_exported_as_partial(tmp_path):
    run = fixture_run()
    run['inventory_complete'] = False
    run['complete'] = False
    path = tmp_path / 'report.xlsx'
    app.export_workbook(run, path)

    wb = load_workbook(path)
    try:
        assert wb['Ringkasan']['B8'].value == 'PARTIAL'
    finally:
        wb.close()


@pytest.mark.parametrize('complete', [True, False])
def test_duplicate_explicit_return_identifier_cannot_inflate_output_count(tmp_path, complete):
    run = fixture_run()
    run['complete'] = complete
    run['records'].append(deepcopy(run['records'][0]))

    with pytest.raises(ValueError, match='[Dd]uplicate.*identifier'):
        app.export_split_workbooks(run, tmp_path)

    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize('source_url', [
    'file:///synthetic-only/run.exe', r'\\synthetic.invalid\share\run.exe',
    'javascript:alert(1)', 'ms-excel:ofe|u|https://synthetic.invalid/fixture.xlsx',
])
def test_non_web_source_urls_are_literal_text_without_active_links(tmp_path, source_url):
    run = fixture_run()
    run['records'][0]['detail_url'] = source_url
    path = tmp_path / 'report.xlsx'
    app.export_workbook(run, path)

    wb = load_workbook(path)
    try:
        source_cells = [wb['Daftar SPT']['K5'], wb.worksheets[3]['A3']]
        assert all(cell.value == source_url for cell in source_cells)
        assert all(cell.hyperlink is None for cell in source_cells)
    finally:
        wb.close()


def test_web_source_link_and_literal_formula_text_are_preserved(tmp_path):
    run = fixture_run()
    source_url = 'https://coretax.intranet.pajak.go.id/returnsheets/synthetic-only'
    run['records'][0]['detail_url'] = source_url
    source_value = '=HYPERLINK("file:///synthetic-only", "001")'
    run['records'][0]['sections'][0]['fields'][0]['value'] = source_value
    original = deepcopy(run)
    path = tmp_path / 'report.xlsx'
    app.export_workbook(run, path)

    wb = load_workbook(path)
    try:
        assert wb['Daftar SPT']['K5'].hyperlink.target == source_url
        assert wb.worksheets[3]['A3'].hyperlink.target == source_url
        field = wb.worksheets[3]['F6']
        assert field.value == source_value
        assert field.data_type == 's'
        assert run == original
    finally:
        wb.close()


def test_duplicate_identifier_with_changed_submission_date_is_not_silently_merged():
    original = fixture_run()['records'][0]
    other = deepcopy(original)
    other['meta']['Tanggal Submit SPT'] = '21-01-2026'

    with pytest.raises(ValueError, match='Conflicting source metadata'):
        app.deduplicate_inventory([original, other])
