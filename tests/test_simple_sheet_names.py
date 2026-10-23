"""Synthetic SPT tab-name contracts; execution is GitHub-runner-only."""
from copy import deepcopy
import json
import re

import pytest
from openpyxl import load_workbook

import coretax
import coretax_exporter as exporter
from support_runs import fixture_run, unifikasi_induk_table


def form_run(names=('Induk', 'L-IA'), *, amendments=False):
    run = fixture_run()
    run['records'] = run['records'][:2 if amendments else 1]
    for index, record in enumerate(run['records']):
        record['meta']['Model SPT'] = 'Normal' if index == 0 and amendments else 'Pembetulan 2'
        record['sections'] = [
            {'name': name, 'fields': [{'label': 'Marker', 'value': f'SYNTHETIC {index} / {name}'}],
             'tables': [], 'text': ''}
            for name in names
        ]
        for section in record['sections']:
            if section['name'].casefold() == 'induk':
                section['fields'][0].update(label='2. Nama', field='Name', group='A. IDENTITAS WAJIB PAJAK')
    return run


def test_single_return_tabs_are_form_names_without_generated_prefixes(tmp_path):
    run = form_run()
    files = exporter.export_split_workbooks(run, tmp_path)
    workbook = load_workbook(files['SPT_Badan_Januari - Desember 2025_Pembetulan 2.xlsx'])
    try:
        assert workbook.sheetnames == ['Induk', 'L-IA']
        assert [sheet['A1'].value for sheet in workbook] == ['INDUK', 'L-IA']
        proof = exporter.read_workbook_proof(workbook)
        assert proof['records'][0]['meta']['Model SPT'] == 'Pembetulan 2'
        assert [section['name'] for section in proof['records'][0]['sections']] == ['Induk', 'L-IA']
    finally:
        workbook.close()


@pytest.mark.parametrize('tax,filename', [
    ('PPh Badan', 'SPT_Badan_Januari - Desember 2025_Normal.xlsx'),
    ('PPh Pasal 21/26', 'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx'),
    ('PPh Unifikasi', 'SPT Masa Unifikasi/SPT_Unifikasi_Januari 2025_Normal.xlsx'),
    ('PPN', 'SPT Masa PPN/SPT_PPN_Januari 2025_Normal.xlsx'),
])
def test_amendments_keep_all_forms_in_separate_file_namespaces(tmp_path, tax, filename):
    run = form_run(amendments=True)
    for record in run['records']:
        record['meta']['Jenis Pajak'] = tax
        if tax != 'PPh Badan':
            record['meta']['Masa Pajak'] = 'Januari 2025'
        if tax == 'PPh Pasal 21/26':
            record['sections'][0]['fields'][0].update(
                label='Status Penyampaian', field='SubmitInProgressStatus', group='HEADER')
            appendix = record['sections'][1]
            marker = appendix['fields'][0]['value']
            appendix.update(fields=[], tables=[{'name': 'Tabel 1', 'group': 'LIST-IA',
                                               'headers': ['NAMA'], 'rows': [[marker]], 'complete': True}])
    if tax == 'PPh Unifikasi':
        for index, record in enumerate(run['records']):
            record['sections'][0]['tables'] = [unifikasi_induk_table(100 + index)]
    files = exporter.export_split_workbooks(run, tmp_path)
    expected = [filename, filename.replace('_Normal.xlsx', '_Pembetulan 2.xlsx')]
    assert set(files) == {'resume.xlsx', *expected}
    captions = ['INDUK', 'L-IA']
    if tax == 'PPh Pasal 21/26':
        captions = ['Induk SPT PPh Pasal 21/26',
            ('DAFTAR PEMOTONGAN BULANAN PAJAK PENGHASILAN PASAL 21 BAGI PEGAWAI TETAP DAN '
            'PENSIUNAN YANG MENERIMA UANG TERKAIT PENSIUN SECARA BERKALA SERTA BAGI PEGAWAI '
            'NEGERI SIPIL, ANGGOTA TENTARA NASIONAL INDONESIA, ANGGOTA KEPOLISIAN REPUBLIK '
            'INDONESIA, PEJABAT NEGARA, DAN PENSIUNANNYA')]
    elif tax == 'PPN':
        captions = ['Induk SPT PPN', 'L-IA']
    elif tax == 'PPh Unifikasi':
        captions = ['Induk SPT PPh Unifikasi', 'L-IA']
    for index, (filename, record) in enumerate(zip(expected, run['records'])):
        workbook = load_workbook(files[filename])
        try:
            assert workbook.sheetnames == ['Induk', 'L-IA']
            assert [sheet['A1'].value for sheet in workbook] == captions
            assert all(not re.match(r'[0-9]{3}\.[0-9]{2} ', title) for title in workbook.sheetnames)
            markers = [cell.value for sheet in workbook for row in sheet for cell in row
                       if isinstance(cell.value, str) and cell.value.startswith('SYNTHETIC ')]
            marker_forms = ['L-IA'] if tax in ('PPN', 'PPh Unifikasi') else ['Induk', 'L-IA']
            assert markers == [f'SYNTHETIC {index} / {form}' for form in marker_forms]
            if tax == 'PPh Unifikasi':
                assert workbook['Induk']['C10'].value == 100 + index
                assert [cell.value for sheet in workbook for row in sheet for cell in row
                        if cell.value in (100, 101)] == [100 + index]
            proof = exporter.read_workbook_proof(workbook)
            assert [item['meta']['Model SPT'] for item in proof['records']] == [record['meta']['Model SPT']]
            assert [item['record_id'] for item in proof['records']] == [record['record_id']]
        finally:
            workbook.close()
    coretax.verify_outputs(run, tmp_path)


def test_case_insensitive_form_collisions_preserve_source_case_and_each_form(tmp_path):
    run = form_run(('Induk', 'induk', 'INDUK'))
    files = exporter.export_split_workbooks(run, tmp_path)
    workbook = load_workbook(files['SPT_Badan_Januari - Desember 2025_Pembetulan 2.xlsx'])
    try:
        assert workbook.sheetnames == ['Induk', 'induk (2)', 'INDUK (3)']
        assert [sheet['A1'].value for sheet in workbook] == ['INDUK'] * 3
        proof = exporter.read_workbook_proof(workbook)
        assert [section['name'] for section in proof['records'][0]['sections']] == ['Induk', 'induk', 'INDUK']
        assert [workbook[name]['B2'].value for name in workbook.sheetnames] == [
            f'SYNTHETIC 0 / {form}' for form in ('Induk', 'induk', 'INDUK')]
    finally:
        workbook.close()


def test_sanitized_collisions_do_not_merge_different_source_forms(tmp_path):
    run = form_run(('L/IA', 'L:IA', 'L-IA'))
    files = exporter.export_split_workbooks(run, tmp_path)
    workbook = load_workbook(files['SPT_Badan_Januari - Desember 2025_Pembetulan 2.xlsx'])
    try:
        assert workbook.sheetnames == ['L-IA', 'L-IA (2)', 'L-IA (3)']
        assert [sheet['A1'].value for sheet in workbook] == ['L/IA', 'L:IA', 'L-IA']
        exporter.read_workbook_proof(workbook)
    finally:
        workbook.close()


def test_truncation_reserves_suffix_space_and_preserves_full_source_captions(tmp_path):
    names = ('A' * 40 + ' first', 'A' * 40 + ' second', 'A' * 40 + ' third')
    run = form_run(names)
    files = exporter.export_split_workbooks(run, tmp_path)
    workbook = load_workbook(files['SPT_Badan_Januari - Desember 2025_Pembetulan 2.xlsx'])
    try:
        assert workbook.sheetnames == ['A' * 31, 'A' * 27 + ' (2)', 'A' * 27 + ' (3)']
        assert [sheet['A1'].value for sheet in workbook] == list(names)
        assert len({title.casefold() for title in workbook.sheetnames}) == 3
        exporter.read_workbook_proof(workbook)
    finally:
        workbook.close()


def test_literal_numbered_form_name_is_not_overwritten_by_an_earlier_suffix(tmp_path):
    run = form_run(('Induk',), amendments=True)
    third = deepcopy(run['records'][1])
    third.update(key='synthetic-third', record_id='synthetic-third')
    third['meta']['Model SPT'] = 'Pembetulan 3'
    third['sections'] = [{'name': 'Induk (2)', 'fields': [
        {'label': 'Marker', 'value': 'SYNTHETIC THIRD FORM'}], 'tables': [], 'text': ''}]
    run['records'].append(third)
    # Explicitly reconstruct the old combined writer to retain its collision contract.
    path = tmp_path / 'synthetic-legacy-combined.xlsx'
    exporter.export_workbook(run, path, include_administration=False)
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ['Induk', 'Induk (2)', 'Induk (2) (2)']
        assert [sheet['A1'].value for sheet in workbook] == ['INDUK', 'INDUK', 'Induk (2)']
        assert 'SYNTHETIC THIRD FORM' in [cell.value for row in workbook.worksheets[2] for cell in row]
        exporter.read_workbook_proof(workbook)
    finally:
        workbook.close()


def test_resume_links_resolve_to_the_right_return_not_only_an_existing_duplicate_form(tmp_path):
    run = form_run(amendments=True)
    coretax.save_output(run, tmp_path)
    resume = load_workbook(tmp_path / 'resume.xlsx')
    try:
        links = [resume['Daftar SPT'].cell(row, 1).hyperlink for row in (5, 6)]
        assert [link.target for link in links] == [
            'SPT_Badan_Januari - Desember 2025_Normal.xlsx',
            'SPT_Badan_Januari - Desember 2025_Pembetulan 2.xlsx']
        assert [link.location for link in links] == ["'Induk'!A1", "'Induk'!A1"]
        for link, marker in zip(links, ['SYNTHETIC 0 / Induk', 'SYNTHETIC 1 / Induk']):
            target = load_workbook(tmp_path / link.target)
            try:
                name = link.location.rsplit('!', 1)[0].strip("'")
                assert marker in [cell.value for row in target[name] for cell in row]
            finally:
                target.close()
    finally:
        resume.close()
    coretax.verify_outputs(run, tmp_path)


def test_each_period_workbook_starts_its_own_form_name_namespace(tmp_path):
    run = fixture_run()
    run['records'] = run['records'][2:4]
    files = exporter.export_split_workbooks(run, tmp_path)
    expected = ['SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx', 'SPT Masa 21/SPT_21_Februari 2025_Normal.xlsx']
    assert set(files) == {'resume.xlsx', *expected}
    for filename in expected:
        workbook = load_workbook(files[filename])
        try:
            assert workbook.sheetnames == ['Induk']
            exporter.read_workbook_proof(workbook)
        finally:
            workbook.close()
    coretax.verify_outputs(run, tmp_path)


def test_legacy_administrative_writer_reserves_its_index_and_audit_names(tmp_path):
    run = form_run(('Audit', 'Ringkasan', 'Daftar SPT'))
    path = tmp_path / 'synthetic-legacy.xlsx'
    exporter.export_workbook(run, path)
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ['Ringkasan', 'Daftar SPT', 'Audit',
                                       'Audit (2)', 'Ringkasan (2)', 'Daftar SPT (2)']
        assert [sheet['A5'].value for sheet in workbook.worksheets[3:]] == ['Audit', 'Ringkasan', 'Daftar SPT']
        assert workbook['Daftar SPT']['A5'].hyperlink.target == "#'Audit (2)'!A1"
    finally:
        workbook.close()


def test_pre_change_forms_only_v1_workbook_remains_verifiable(tmp_path):
    run = form_run(amendments=True)
    path = tmp_path / 'SPT_Badan.xlsx'
    exporter.export_workbook(run, path, include_administration=False)
    workbook = load_workbook(path)
    try:
        old_names = ['001.01 Normal Induk', '001.02 Normal L-IA',
                     '002.01 Pembetulan 2 Induk', '002.02 Pembetulan 2 L-IA']
        for sheet, name in zip(workbook.worksheets, old_names):
            sheet.title = name
        proof = json.loads(workbook.properties.identifier)
        proof['version'] = 1
        workbook.properties.identifier = json.dumps(proof, ensure_ascii=False, sort_keys=True)
        workbook.save(path)
    finally:
        workbook.close()
    workbook = load_workbook(path)
    try:
        proof = exporter.read_workbook_proof(workbook)
        assert workbook.sheetnames == old_names
        assert proof['npwp'] == '0000000000000001' and proof['year'] == 2025
        assert [record['meta']['Model SPT'] for record in proof['records']] == ['Normal', 'Pembetulan 2']
    finally:
        workbook.close()


@pytest.mark.parametrize('mutation', ['rename', 'missing', 'extra', 'hidden', 'reorder'])
def test_simple_tab_names_do_not_weaken_physical_form_proof(tmp_path, mutation):
    run = form_run()
    files = exporter.export_split_workbooks(run, tmp_path)
    path = files['SPT_Badan_Januari - Desember 2025_Pembetulan 2.xlsx']
    workbook = load_workbook(path)
    try:
        if mutation == 'rename':
            workbook.worksheets[0].title = 'SYNTHETIC UNDECLARED FORM'
        elif mutation == 'missing':
            workbook.remove(workbook.worksheets[0])
        elif mutation == 'extra':
            workbook.create_sheet('SYNTHETIC UNDECLARED FORM')
        elif mutation == 'hidden':
            workbook.worksheets[0].sheet_state = 'hidden'
        else:
            workbook.move_sheet(workbook.worksheets[1], offset=-1)
        workbook.save(path)
    finally:
        workbook.close()
    workbook = load_workbook(path)
    try:
        with pytest.raises(ValueError):
            exporter.read_workbook_proof(workbook)
    finally:
        workbook.close()
