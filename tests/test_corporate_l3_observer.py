"""Actual owned L3 SPA, passive custody and saved export contracts; runners only."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

import coretax
import coretax_browser as driver
from coretax_corporate import observe_forms
from coretax_json import JsonSafetyError, adapt_corporate_detail
from coretax_json_capture import DetailJsonCapture, checkpoint_bodies
from l3_observer_fixture import Fixture, ASSET
from test_corporate_l3_native import DETAIL, GRID, FOOTER, REFERENCES, OWNERS, GROUPS
from test_json_hybrid_native import browser
import support_json_corpus as corpus


def open_selected(fixture, capture):
    page, popup = driver.open_return_readonly(fixture.page, fixture.page.locator('tbody tr'),
        corpus.TAXPAYER_ID, corpus.RECORD_ID, json_capture=capture)
    assert page is fixture.page and popup is False
    page.wait_for_function('window.primaryDone===true')
    assert page.locator('rshshr-citr-main-form').is_visible()
    return page


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
@pytest.mark.parametrize('footer_order', ['early', 'late'])
def test_owned_l3_view_to_original_packets_checkpoint_and_compact_saved_form(
        browser, tmp_path, monkeypatch, model, footer_order):
    fixture = Fixture(browser, monkeypatch, model=model, footer_order=footer_order)
    try:
        with DetailJsonCapture(fixture.context, fixture.page, fixture.record, corpus.SCOPE, tmp_path) as capture:
            page = open_selected(fixture, capture)
            forms = observe_forms(page, fixture.record, corpus.SCOPE, capture, approved_digest=fixture.pins[ASSET])
            assert forms is not None, 'Native L3 owned observer is not implemented'
            assert [form['name'] for form in forms] == ['Induk', 'L3']
            proof = forms[1]['projection']
            assert proof['kind'] == 'corporate-l3-v1' and proof['builds'] == fixture.pins
            assert [t['component'] for t in proof['tables']] == OWNERS
            assert [t['source_structure'] for t in proof['tables']] == [
                {'form': 'L3', 'component': owner, 'groups': [GROUPS[index]]}
                for index, owner in enumerate(OWNERS)]
            manifest = capture.drain()
            assert manifest['complete'] and manifest['component_lifetime']['same_document'] is True
            bodies = list(capture.iter_bodies())
            assert len(bodies) == 8 and next(p['body'] for p in bodies if p['context']['url'] == corpus.PRIMARY_URL) == fixture.raw
            assert next(p['body'] for p in bodies if p['context']['url'] == FOOTER) == fixture.footer_raw
            for packet in bodies:
                url = packet['context']['url']
                if url == GRID:
                    assert packet['body'] == fixture.grid_raw[packet['context']['request_json']['First']]
                if url in REFERENCES.values(): assert packet['body'] == fixture.reference_raw[url]
            before = deepcopy((bodies, forms, fixture.record))
            result = adapt_corporate_detail(iter(bodies), fixture.record, corpus.SCOPE, forms)
            assert result['eligible'] and result['coverage']['complete'], result['reason']
            assert (bodies, forms, fixture.record) == before
            fixture.record['detail_url'] = DETAIL
            replay = adapt_corporate_detail(checkpoint_bodies(tmp_path, manifest, fixture.record, corpus.SCOPE),
                                           fixture.record, corpus.SCOPE, forms)
            assert replay['sections'] == result['sections'] and replay['provenance'] == result['provenance']
            assert page.evaluate('window.businessReads') == page.evaluate('window.forbiddenClicks') == 0
        section, = [s for s in result['sections'] if s['name'] == 'L3']
        assert [t['body_row_count'] for t in section['tables']] == [2, 12]
        assert section['tables'][1]['footer'][-1] == ['JUMLAH KREDIT PAJAK', '', '41', '']
        fixture.record.update(capture_complete=True, sections=result['sections'])
        run = {'npwp': corpus.NPWP, 'year': 2025, 'year_basis': 'period', 'inventory_complete': True,
               'complete': True, 'errors': [], 'records': [fixture.record]}
        coretax.save_output(run, tmp_path / 'export')
        coretax.verify_outputs(run, tmp_path / 'export')
        name, = [n for n in run['output_files'] if n != 'resume.xlsx']
        assert model in name
        book = load_workbook(tmp_path / 'export' / name)
        try:
            sheet = book['L3']
            assert sheet['D3'].value == '02-12-2025' and sheet.max_row < 40
            cells = [cell for row in sheet for cell in row]
            assert sum(cell.value == '000001' for cell in cells) == 12
            assert sum(cell.value == '0000000000000002' for cell in cells) == 12
            credit, = [sheet.cell(cell.row, 6) for cell in cells if cell.value == 'JUMLAH KREDIT PAJAK']
            assert credit.value == 41 and 'Rp' in credit.number_format and '0.00' not in credit.number_format
        finally:
            book.close()
        mapping, = [r for r in run['template_mapping'] if r['form'] == 'L3']
        assert not mapping['complete'] and mapping['unmapped_fields'] == 2
        assert mapping['unmapped_rows'] == mapping['unmapped_tables'] == 0
        assert fixture.unexpected == []
        assert fixture.requests.count((GRID, 'POST')) == 2 and fixture.requests.count((FOOTER, 'POST')) == 1
        first_grid, footer = fixture.requests.index((GRID, 'POST')), fixture.requests.index((FOOTER, 'POST'))
        assert (footer < first_grid) is (footer_order == 'early')
        for url in fixture.assets: assert fixture.requests.count((url, 'GET')) == 2
        for url in REFERENCES.values(): assert fixture.requests.count((url, 'GET')) == 1
    finally:
        fixture.close()


@pytest.mark.parametrize('defect', ['external-context', 'timezone', 'locale', 'vendor-pin', 'owner', 'header', 'span',
    'wrong-group', 'misplaced-grid', 'wrong-ui-owner', 'table-outside-p-table', 'duplicate-owner',
    'hidden-owner', 'missing-reference', 'missing-footer', 'failed-footer',
    'corporate-fetch-only', 'corporate-preload-only', 'main-fetch-only', 'main-preload-only',
    'vendor-fetch-only', 'vendor-preload-only'])
def test_unproved_owned_l3_structure_provider_or_dependency_cannot_grant_native_coverage(
        browser, tmp_path, monkeypatch, defect):
    fixture = Fixture(browser, monkeypatch, defect=defect)
    try:
        with DetailJsonCapture(fixture.context, fixture.page, fixture.record, corpus.SCOPE, tmp_path) as capture:
            page = open_selected(fixture, capture)
            forms = observe_forms(page, fixture.record, corpus.SCOPE, capture, approved_digest=fixture.pins[ASSET])
            capture.drain()
            if forms is not None:
                result = adapt_corporate_detail(capture.iter_bodies(), fixture.record, corpus.SCOPE, forms)
                assert not result['eligible'] and result['sections'] == []
            assert page.evaluate('window.businessReads') == page.evaluate('window.forbiddenClicks') == 0
            assert fixture.unexpected == []
    finally:
        fixture.close()


def test_late_conflicting_l3_footer_remains_fatal_after_owned_observation(browser, tmp_path, monkeypatch):
    fixture = Fixture(browser, monkeypatch, defect='conflicting-footer')
    try:
        with DetailJsonCapture(fixture.context, fixture.page, fixture.record, corpus.SCOPE, tmp_path) as capture:
            page = open_selected(fixture, capture)
            page.locator('a[role=tab]').filter(has_text='L3').click()
            forms = observe_forms(page, fixture.record, corpus.SCOPE, capture, approved_digest=fixture.pins[ASSET])
            page.wait_for_function('window.l3FooterReplies===2', timeout=10000)
            capture.drain()
            assert fixture.footer_count == 2
            bodies = list(capture.iter_bodies())
            footers = [p['body'] for p in bodies if p['context']['url'] == FOOTER]
            assert len(footers) == 2 and footers[0] == fixture.footer_raw and footers[1] != footers[0]
            with pytest.raises(JsonSafetyError, match='snapshot'):
                adapt_corporate_detail(iter(bodies), fixture.record, corpus.SCOPE, forms or [])
    finally:
        fixture.close()


def test_owned_l3_stop_propagates_and_releases_document_without_manifest(browser, tmp_path, monkeypatch):
    fixture = Fixture(browser, monkeypatch)
    try:
        with DetailJsonCapture(fixture.context, fixture.page, fixture.record, corpus.SCOPE, tmp_path) as capture:
            page = open_selected(fixture, capture)
            def stop(): raise driver.CaptureCancelled('SYNTHETIC STOP')
            with pytest.raises(driver.CaptureCancelled):
                observe_forms(page, fixture.record, corpus.SCOPE, capture,
                              approved_digest=fixture.pins[ASSET], check_cancel=stop)
            assert not list(tmp_path.rglob('manifest.json'))
        assert capture._component_document is None
    finally:
        fixture.close()


def test_l3_business_read_traps_cover_enclosing_panel_and_table(browser, tmp_path, monkeypatch):
    fixture = Fixture(browser, monkeypatch)
    try:
        with DetailJsonCapture(fixture.context, fixture.page, fixture.record, corpus.SCOPE, tmp_path) as capture:
            page = open_selected(fixture, capture)
            page.evaluate('show(1)')
            page.wait_for_function('window.l3Done===true')
            for selector in ('rsh-citr-l3-grid1 table', 'rshshr-citr-l3-grid2 tfoot', '#panel-1'):
                for key in ('innerText', 'textContent'):
                    assert page.locator(selector).evaluate('''(e,key)=>{
                        try{void e[key];return false;}catch(_){return true;}
                    }''', key)
            assert page.evaluate('window.businessReads') == 6
            assert page.locator('rsh-citr-l3-grid1 th').first.inner_text() == 'NO.'
    finally:
        fixture.close()
