"""Pinned current-only L11-A pure projection contracts; runners only."""
from copy import deepcopy
import json

import pytest
from openpyxl import load_workbook

import coretax
import coretax_corporate_projection as projection
from coretax_json import _UnsupportedMetadata, _corporate_json
from support_templates import section, source_run


ROOT = 'L11AFormData'
GROUP = ('IV.B RINCIAN PENGGANTIAN ATAU IMBALAN DALAM BENTUK NATURA DAN/ATAU KENIKMATAN '
         'YANG DIBERIKAN BERKENAAN DENGAN PELAKSANAAN PEKERJAAN DI DAERAH TERTENTU')
HEADERS = [
    ['No.', 'NOMOR IDENTITAS PENERIMA', 'NAMA PENERIMA', 'ALAMAT PENERIMA', 'TANGGAL',
     'JENIS BIAYA', 'NILAI (Rp.)', 'KETERANGAN', 'PPh YANG DIPOTONG/DIPUNGUT (Rp.)Rp.)',
     'NOMOR BUKTI POTONG'],
    ['TANGGAL PEMBERIAN ENTERTAINMENT', 'TEMPAT PEMBERIAN ENTERTAINMENT',
     'ALAMAT PEMBERIAN ENTERTAINMENT', 'JENIS PEMBERIAN ENTERTAINMENT',
     'NILAI PEMBERIAN ENTERTAINMENTRp.)', 'NAMA RELASI USAHA', 'JABATAN RELASI USAHA',
     'NAMA PERUSAHAAN RELASI USAHA', 'JENIS USAHA RELASI USAHA', 'DESKRIPSI'],
    ['No.', 'NOMOR IDENTITAS', 'NAMA DEBITUR', 'ALAMAT', 'PLAFON PIUTANG',
     'PIUTANG YANG NYATA-NYATA TIDAK DAPAT DITAGIH', 'METODE PEMBEBANAN',
     'JENIS DOKUMEN PEMBUKTIAN PEMENUHAN PERSYARATAN'],
    ['NO.', 'JENIS HARTA BERWUJUD', 'TAHUN PEROLEHAN', 'NILAI PEROLEHAN',
     'PENYUSUTAN / S.D. TAHUN LALU', 'PENYUSUTAN / PENYUSUTAN TAHUN INI',
     'PENYUSUTAN / S.D. TAHUN INI'],
    ['No.', 'NOMOR IDENTITAS', 'NAMA DEBITUR', 'ALAMAT',
     'NILAI KREDIT KURANG LANCAR AWAL TAHUN BUKU', 'NILAI KREDIT KURANG LANCAR AKHIR TAHUN BUKU',
     'JUMLAH BUNGA PADA TAHUN BUKU (AKRUAL)', 'KATEGORI']]
COSTS = ['ResidencyCost', 'HealthServiceCost', 'EducationCost', 'WorshipCost',
         'TransportationCost', 'SportsCost']
CAPTIONS = ['Alamat Lokasi', 'Nomor Surat Keputusan Penetapan Daerah Tertentu',
    'Tanggal Surat Keputusan Penetapan Daerah Tertentu',
    'Nomor Surat Keputusan Perpanjangan Penetapan Daerah Tertentu',
    'Tanggal Surat Keputusan Perpanjangan Penetapan Daerah Tertentu',
    'Biaya yang dikeluarkan untuk tempat tinggal, termasuk perumahan untuk pegawai dan keluarganya',
    'Biaya yang dikeluarkan untuk pelayanan kesehatan',
    'Biaya yang dikeluarkan untuk pendidikan bagi pegawai dan keluarganya',
    'Biaya yang dikeluarkan untuk peribadatan',
    'Biaya yang dikeluarkan untuk pengangkutan bagi pegawai dan keluarganya',
    'Biaya yang dikeluarkan untuk olahraga bagi pegawai dan keluarganya, tidak termasuk golf, '
    'power boating, pacuan kuda, dan terbang layang', 'Jumlah Biaya yang dikeluarkan']
RECAP = ['LocationAddress', 'NumberOfDecreeOfSpecificAreasDetermination',
    'DateOfDecreeOfSpecificAreasDetermination', 'NumberOfDecreeOfExtensificationOfSpecificAreasDetermination',
    'DateOfDecreeOfExtensificationOfSpecificAreasDetermination', *COSTS, 'TotalAmountofCosts']
TYPES = ['TYPE_OF_PROMOTION_COST', 'TANGIBLE_ASSET_INCLUDING_BUILDINGS',
         'NON_PERFORMING_LOAN_CATEGORY', 'DEDUCTION_METHOD', 'TYPE_OF_FULFILLMENT']
LABELS = {kind: {'SYN_CODE': 'SYNTHETIC LABEL ' + str(index)} for index, kind in enumerate(TYPES)}
ENDPOINTS = [f'l11a-table-{index}-grid' for index in (1, 2, 3, 5)]


def inputs():
    recap = dict(zip(RECAP, ['SYNTHETIC LOCATION', '0001', '2025-12-01T23:30:00.000Z',
        '0002', '2025-12-02', 1, 2, 3, 4, 5, 6, 21]))
    facility = {'TypeOfTangibleAssets': 'SYN_CODE', 'YearOfAcquisition': '0025',
        'AcquisitionValue': 1000, 'UpToPreviousYear': 2, 'CurrentYear': 3, 'UpToCurrentYear': 5}
    rows = [
        {'RecordId': 'SYNTHETIC-ROW', 'Index': '0007', 'RecipientIdentityNumber': '00000002',
         'RecipientName': 'SYNTHETIC RECIPIENT', 'RecipientAddress': 'SYNTHETIC ADDRESS',
         'RecipientDate': '2025-12-01T23:30:00.000Z', 'RecipientFormAndTypeOfPromotionCost': 'SYN_CODE',
         'RecipientAmountOfPromotionCost': 5, 'RecipientDescription': 'SYNTHETIC DESCRIPTION',
         'IncomeTax': 0, 'WithholdingSlipNumber': '00001'},
        {'RecordId': 'SYNTHETIC-ROW', 'EntertainmentDate': '2025-12-01',
         'EntertainmentPlace': 'SYNTHETIC PLACE', 'EntertainmentAddress': 'SYNTHETIC ADDRESS',
         'EntertainmentType': 'SYNTHETIC TYPE', 'EntertainmentAmount': 7,
         'PartnerName': 'SYNTHETIC PERSON', 'PartnerPosition': 'SYNTHETIC POSITION',
         'PartnerCompanyName': 'SYNTHETIC COMPANY', 'PartnerTypeOfBusiness': 'SYNTHETIC BUSINESS',
         'Description': 'SYNTHETIC DESCRIPTION'},
        {'RecordId': 'SYNTHETIC-ROW', 'Index': '0007', 'IdentityNumber': '00000002',
         'Name': 'SYNTHETIC DEBTOR', 'Address': 'SYNTHETIC ADDRESS', 'DebtCeiling': 13,
         'UncollectibleDebt': 11, 'DeductionMethod': 'SYN_CODE', 'TypeOfFulfillment': 'SYN_CODE'},
        {'RecordId': 'SYNTHETIC-ROW', 'Index': '0005', 'IdentityNumber': '00000002',
         'Name': 'SYNTHETIC DEBTOR', 'Address': 'SYNTHETIC ADDRESS', 'AmountBeginning': 13,
         'AmountEnd': 11, 'AmountInterestCurrentYear': 2, 'Category': 'SYN_CODE'}]
    pages = {endpoint: [{'offset': 0, 'rows': 10, 'total': 2, 'data': [deepcopy(row), deepcopy(row)]}]
             for endpoint, row in zip(ENDPOINTS, rows)}
    pages[ENDPOINTS[0]] = [
        {'offset': 0, 'rows': 10, 'total': 12, 'data': [deepcopy(rows[0]) for _ in range(10)]},
        {'offset': 10, 'rows': 10, 'total': 12, 'data': [deepcopy(rows[0]) for _ in range(2)]}]
    return {'Facilities': [deepcopy(facility), deepcopy(facility)], 'Recapitulation': recap}, pages


def project(form, pages, *, form_wire=None, **overrides):
    function = getattr(projection, 'project_l11a', None)
    assert callable(function), 'L11-A pure source projector is not implemented'
    options = {'previous': None, 'headers': deepcopy(HEADERS), 'footer_spans': [[], [], [], [], []],
        'tax_type': 'ICT_RCIT', 'locale': 'id-ID', 'timezone_name': 'Asia/Jakarta', 'labels': deepcopy(LABELS)}
    options.update(overrides)
    current = _corporate_json(json.dumps(form) if form_wire is None else form_wire)
    observed = {endpoint: [{**p, 'data': _corporate_json(json.dumps(p['data']))} for p in values]
                for endpoint, values in pages.items()}
    before = deepcopy((current, observed, options))
    try:
        return function(current, pages=observed, **options)
    finally:
        assert (current, observed, options) == before


def test_l11a_preserves_actual_five_tables_source_indices_and_all_unmapped_recap_values():
    form, pages = inputs()
    result = project(form, pages)
    tables = result['tables']
    assert [t['body_row_count'] for t in tables] == [12, 2, 2, 2, 2]
    assert all(t['source_total'] == t['body_row_count'] and t['complete'] for t in tables)
    assert all(t['footer'] == t['footer_spans'] == [] and len(t['rows']) == t['body_row_count'] for t in tables)
    assert tables[0]['rows'][0] == ['0007', '00000002', 'SYNTHETIC RECIPIENT', 'SYNTHETIC ADDRESS',
        '02-12-2025', 'SYNTHETIC LABEL 0', '5', 'SYNTHETIC DESCRIPTION', '0', '00001']
    assert tables[0]['rows'][0] == tables[0]['rows'][10]  # Index is source, not page ordinal.
    assert tables[1]['rows'][0][0] == '01-12-2025' and tables[1]['rows'][0][5:7] == [
        'SYNTHETIC PERSON', 'SYNTHETIC POSITION']
    assert tables[2]['rows'][0] == ['0007', '00000002', 'SYNTHETIC DEBTOR',
        'SYNTHETIC ADDRESS', '13', '11', 'SYNTHETIC LABEL 3', 'SYNTHETIC LABEL 4']
    assert tables[3]['rows'] == [['1', 'SYNTHETIC LABEL 1', '0025', '1.000', '2', '3', '5'],
                                ['2', 'SYNTHETIC LABEL 1', '0025', '1.000', '2', '3', '5']]
    assert tables[4]['rows'][0][-1] == 'SYNTHETIC LABEL 2'
    expected_values = ['SYNTHETIC LOCATION', '0001', '02-12-2025', '0002', '02-12-2025',
                       '1', '2', '3', '4', '5', '6', '21']
    assert [(f['group'], f['field'], f['label'], f['value']) for f in result['fields']] == [
        (GROUP, key, caption, value) for key, caption, value in zip(RECAP, CAPTIONS, expected_values)]
    evidence = result['provenance']
    assert any(e['source'] == 'grid' and e['endpoint'] == ENDPOINTS[0] and e['page_index'] == 1
               and e['path'] == ['Data', 0, 'Index'] and e['value'] == '0007' for e in evidence)
    assert any(e['source'] == 'current' and e['path'] == [ROOT, 'Facilities', 1, 'YearOfAcquisition']
               and e['value'] == '0025' for e in evidence)
    for key in RECAP:
        assert any(e['source'] == 'current' and e['path'] == [ROOT, 'Recapitulation', key]
                   and e['value'] == (str(form['Recapitulation'][key])) for e in evidence)


@pytest.mark.parametrize('lexeme', ['21.0', '2.1e1'])
def test_l11a_copies_reconciled_total_preserving_original_numeric_token(lexeme):
    form, pages = inputs()
    wire = json.dumps(form).replace('"TotalAmountofCosts": 21', '"TotalAmountofCosts": ' + lexeme)
    result = project(form, pages, form_wire=wire)
    assert result['fields'][-1]['value'] == '21'
    assert any(e['path'] == [ROOT, 'Recapitulation', 'TotalAmountofCosts']
               and e['type'] == 'number' and e['value'] == lexeme for e in result['provenance'])


@pytest.mark.parametrize('cost', COSTS)
@pytest.mark.parametrize('case', ['null', 'negative-clamp', 'unsafe'])
def test_l11a_every_cost_operand_follows_ordered_nullish_sum_and_final_clamp(cost, case):
    form, pages = inputs()
    recap = form['Recapitulation']
    recap.update({name: 0 for name in COSTS})
    recap[cost] = None if case == 'null' else -3 if case == 'negative-clamp' else 1000000000000000
    recap['TotalAmountofCosts'] = 0
    if case == 'unsafe':
        with pytest.raises(_UnsupportedMetadata):
            project(form, pages)
    else:
        result = project(form, pages)
        assert result['fields'][-1]['value'] == '0'
        assert result['fields'][RECAP.index(cost)]['value'] == ('' if case == 'null' else '-3')
        assert any(e['path'] == [ROOT, 'Recapitulation', cost]
                   and e['type'] == ('null' if case == 'null' else 'number')
                   and e['value'] == (None if case == 'null' else '-3') for e in result['provenance'])


def test_l11a_clamps_only_the_final_six_cost_sum():
    form, pages = inputs()
    form['Recapitulation'].update(dict(zip(COSTS, [-3, 5, 0, 0, 0, 0])))
    form['Recapitulation']['TotalAmountofCosts'] = 2
    result = project(form, pages)
    assert [result['fields'][RECAP.index(cost)]['value'] for cost in COSTS] == ['-3', '5', '0', '0', '0', '0']
    assert result['fields'][-1]['value'] == '2'


@pytest.mark.parametrize('empty', [None, []])
def test_l11a_explicit_empty_facilities_and_zero_pages_do_not_invent_footers(empty):
    form, pages = inputs()
    form['Facilities'] = empty
    pages = {key: [{'offset': 0, 'rows': 10, 'total': 0, 'data': []}] for key in pages}
    result = project(form, pages)
    assert all(t['body_row_count'] == 0 and t['rows'] == t['footer'] == [] for t in result['tables'])
    assert len(result['fields']) == 12


def test_l11a_active_null_money_and_date_remain_blank_and_numeric_partner_name_uses_number2():
    form, pages = inputs()
    pages[ENDPOINTS[0]][0]['data'][0].update(RecipientAmountOfPromotionCost=None, RecipientDate=None)
    pages[ENDPOINTS[1]][0]['data'][0]['PartnerName'] = 7
    result = project(form, pages)
    assert result['tables'][0]['rows'][0][4] == result['tables'][0]['rows'][0][6] == ''
    assert result['tables'][1]['rows'][0][5] == '7,00'


@pytest.mark.parametrize('defect', ['previous-empty-object', 'previous-facilities', 'previous-list',
    'tax-type', 'locale', 'timezone', 'header', 'footer', 'missing-root-field', 'extra-root-field',
    'missing-recap', 'extra-recap', 'reported-total', 'cost-string', 'cost-bool', 'cost-fraction',
    'unsafe-intermediate', 'missing-page', 'gap', 'changed-total', 'short-page', 'unsafe-offset',
    'missing-grid', 'extra-grid', 'missing-index', 'legacy-alias', 'missing-record-id',
    'unknown-row-field', 'money-string', 'money-bool', 'money-fraction', 'money-unsafe',
    'missing-label', 'code-bool', 'facility-key', 'facility-year-number', 'recap-date-null',
    'recap-date-empty', 'recap-date-invalid', 'body-date-empty', 'body-date-invalid', 'year0001'])
def test_unproved_l11a_predecessor_source_arithmetic_pages_or_metadata_remains_unsupported(defect):
    form, pages = inputs()
    recap, options = form['Recapitulation'], {}
    row = pages[ENDPOINTS[0]][0]['data'][0]
    if defect.startswith('previous-'):
        options['previous'] = {'previous-empty-object': {}, 'previous-facilities': {'Facilities': []},
                               'previous-list': []}[defect]
    elif defect in ('tax-type', 'locale', 'timezone'):
        options[{'tax-type': 'tax_type', 'locale': 'locale', 'timezone': 'timezone_name'}[defect]] = 'UNPROVED'
    elif defect == 'header':
        options['headers'] = deepcopy(HEADERS)
        options['headers'][0][4] = 'SYNTHETIC UNPROVED'
    elif defect == 'footer':
        options['footer_spans'] = [[[10]], [], [], [], []]
    elif defect == 'missing-root-field': del form['Facilities']
    elif defect == 'extra-root-field': form['Unproved'] = None
    elif defect == 'missing-recap': del recap['LocationAddress']
    elif defect == 'extra-recap': recap['Unproved'] = None
    elif defect == 'reported-total': recap['TotalAmountofCosts'] = 22
    elif defect in ('cost-string', 'cost-bool', 'cost-fraction'):
        recap['ResidencyCost'] = {'cost-string': '1', 'cost-bool': True, 'cost-fraction': 0.1}[defect]
    elif defect == 'unsafe-intermediate':
        recap.update(dict(zip(COSTS, [999999999999999, 1, -1, 0, 0, 0])))
        recap['TotalAmountofCosts'] = 999999999999999
    elif defect == 'missing-page': pages[ENDPOINTS[0]].pop()
    elif defect == 'gap': pages[ENDPOINTS[0]][1]['offset'] = 11
    elif defect == 'changed-total': pages[ENDPOINTS[0]][1]['total'] = 13
    elif defect == 'short-page': pages[ENDPOINTS[0]][0]['data'].pop()
    elif defect == 'unsafe-offset': pages[ENDPOINTS[0]][0]['offset'] = False
    elif defect == 'missing-grid': del pages[ENDPOINTS[1]]
    elif defect == 'extra-grid': pages['SYNTHETIC-UNPROVED'] = []
    elif defect == 'missing-index': del row['Index']
    elif defect == 'legacy-alias': row['Date'] = row.pop('RecipientDate')
    elif defect == 'missing-record-id': del row['RecordId']
    elif defect == 'unknown-row-field': row['Unproved'] = None
    elif defect in ('money-string', 'money-bool', 'money-fraction', 'money-unsafe'):
        row['RecipientAmountOfPromotionCost'] = {'money-string': '5', 'money-bool': True,
            'money-fraction': 0.1, 'money-unsafe': 1000000000000000}[defect]
    elif defect == 'missing-label':
        options['labels'] = deepcopy(LABELS)
        options['labels'][TYPES[0]] = {}
    elif defect == 'code-bool': row['RecipientFormAndTypeOfPromotionCost'] = True
    elif defect == 'facility-key': del form['Facilities'][0]['UpToPreviousYear']
    elif defect == 'facility-year-number': form['Facilities'][0]['YearOfAcquisition'] = 2025
    elif defect.startswith('recap-date-'):
        recap['DateOfDecreeOfSpecificAreasDetermination'] = {
            'recap-date-null': None, 'recap-date-empty': '', 'recap-date-invalid': '2025-02-30'}[defect]
    elif defect.startswith('body-date-'):
        row['RecipientDate'] = '' if defect == 'body-date-empty' else '2025-02-30'
    elif defect == 'year0001': row['RecipientDate'] = '0001-01-01'
    with pytest.raises(_UnsupportedMetadata):
        project(form, pages, **options)


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
def test_l11a_final_saved_compact_export_keeps_five_tables_and_truthful_recap_audit(tmp_path, model):
    result = project(*inputs())
    run = source_run('PPh Badan', [section('L11-A', fields=result['fields'], tables=result['tables'])])
    run['records'][0]['meta']['Model SPT'] = model
    before = deepcopy(result['fields'])
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    assert run['records'][0]['sections'][0]['fields'] == before
    name, = [n for n in run['output_files'] if n != 'resume.xlsx']
    assert model in name
    book = load_workbook(tmp_path / name)
    try:
        sheet = book['L11-A']
        values = [[cell.value for cell in row] for row in sheet]
        all_values = [value for row in values for value in row]
        assert all_values.count('0007') == 14 and all_values.count('0005') == 2
        assert all_values.count('00000002') == 16 and all_values.count('00001') == 12
        assert all_values.count('0025') == 2 and all_values.count('SYNTHETIC PERSON') == 2
        assert '02-12-2025' in all_values and '01-12-2025' in all_values
        assert 'JUMLAH' not in all_values and 'Data sumber tambahan / belum dipetakan' not in all_values
        assert 'A1:D1' in {str(area) for area in sheet.merged_cells.ranges}
        facilities = [row for row in sheet if len(row) >= 4 and row[1].value == 'SYNTHETIC LABEL 1']
        assert len(facilities) == 2 and all(row[2].value == '0025' and row[3].value == '1.000'
                                         and row[3].number_format == 'General' for row in facilities)
        entertainment = [row[4] for row in sheet if len(row) >= 6 and row[5].value == 'SYNTHETIC PERSON']
        assert len(entertainment) == 2 and all(cell.value == 7 and 'Rp' in cell.number_format
                                             and '0.00' not in cell.number_format for cell in entertainment)
        assert sheet.max_row < 60  # Final compaction remains active.
    finally:
        book.close()
    report, = [r for r in run['template_mapping'] if r['form'] == 'L11-A']
    assert not report['complete'] and report['unmapped_fields'] == 12
    assert report['unmapped_tables'] == report['unmapped_rows'] == 0
