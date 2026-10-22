"""Requested presentation deltas preserve data, totals and unresolved evidence."""
from copy import copy
from pathlib import Path

import pytest
from openpyxl import load_workbook

from support_templates import section, source_appendix, export_ppn, assert_format
from test_unifikasi_clean_percent import export, table


@pytest.mark.parametrize('first_count', [0, 4, 32, 43])
@pytest.mark.parametrize('second_count', [0, 7])
def test_daftar_i_keeps_one_borderless_separator_before_second_caption(tmp_path, first_count, second_count):
    book, reports = export(tmp_path, [section('DAFTAR-I', tables=[
        table(count=first_count), table('BPNR', count=second_count)])])
    try:
        sheet = book['DAFTAR-I']
        spacer = 9 + max(0, first_count - 4)
        assert sheet.cell(spacer + 1, 2).value == 'BPNR'
        assert sheet.cell(spacer + 2, 10).value == 'TINGKAT (%)'
        for cell in sheet[spacer]:
            assert cell.value is None
            assert all(getattr(cell.border, edge).style is None for edge in ('left', 'right', 'top', 'bottom'))
        for row in range(5, 5 + first_count):
            assert sheet.cell(row, 2).value == '0000000000000002'
            assert sheet.cell(row, 10).value == 0.025
            assert sheet.cell(row, 11).value == 30.86
        for row in range(spacer + 3, spacer + 3 + second_count):
            assert sheet.cell(row, 2).value == '0000000000000002'
            assert sheet.cell(row, 10).value == 0.025
            assert sheet.cell(row, 11).value == 30.86
        assert reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('count', [0, 1, 3, 4, 5, 8])
@pytest.mark.parametrize('source_spacer', [False, True])
def test_lampiran_i_reserves_one_empty_row_before_captured_total(tmp_path, count, source_spacer):
    reference = load_workbook(Path(__file__).resolve().parents[1] / 'resources/unifikasi_format.xlsx')
    try:
        sheet = reference['Lampiran I']
        headers = [cell.value for cell in sheet[4]]
        label = sheet['A10'].value
        body_reference = sheet['B5']
        total_reference = sheet['M10']
        heights = {row: sheet.row_dimensions[row].height for row in range(5, 10)}
    finally:
        reference.close()
    rows = [[str(index + 1)] + ['SYNTHETIC'] * 16 for index in range(count)]
    for row in rows:
        row[1], row[10], row[12] = '0000000000000002', '1.234,50', '24,69'
    footer = ([[]] if source_spacer else []) + [[label, '987,65', '']]
    source = {'name': 'Tabel 1', 'group': 'TABEL I. ATC', 'headers': headers,
              'rows': rows + footer, 'body_row_count': count, 'source_total': count,
              'complete': True, 'footer': footer,
              'footer_spans': ([[]] if source_spacer else []) + [[12, 1, 4]]}
    book, reports = export(tmp_path, [section('LAMPIRAN-I', tables=[source])])
    try:
        sheet = book['LAMPIRAN-I']
        total = 6 + count
        assert sheet.max_row == total
        assert sheet.cell(total, 1).value == label
        assert sheet.cell(total, 13).value == 987.65
        assert_format(sheet.cell(total, 13), total_reference)
        assert all(cell.value is None for cell in sheet[total - 1])
        for row in range(5, 5 + count):
            assert sheet.cell(row, 2).value == '0000000000000002'
            assert sheet.cell(row, 11).value == 1234.5
            assert sheet.cell(row, 13).value == 24.69
            assert_format(sheet.cell(row, 2), body_reference)
            assert sheet.row_dimensions[row].height == heights[min(row, 5) if row > 9 else row]
        assert reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('count', [0, 11, 64])
def test_ppn_a2_last_reserved_name_cell_is_centered_without_other_style_changes(tmp_path, count):
    source, _ = source_appendix('A2', 4, 5, 16, (6, 7, 8, 9), count)
    book, _, _ = export_ppn(tmp_path, [section('A-2', tables=[source])])
    try:
        sheet = book['A-2']
        assert sheet['B15'].alignment.horizontal == 'center'
        assert copy(sheet['B15'].alignment) == copy(sheet['B14'].alignment)
        assert copy(sheet['B15'].border) == copy(sheet['B14'].border)
        assert sheet['B15'].number_format == sheet['B14'].number_format
        if count >= 11:
            assert sheet['B15'].value == 'SYNTHETIC-10-2'
    finally:
        book.close()
