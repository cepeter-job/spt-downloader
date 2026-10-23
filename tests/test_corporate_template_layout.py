"""Corporate binding against reference geometry, before final row compaction.

These tests keep stable reference coordinates to isolate semantic binding,
growth and style contracts. test_corporate_compact_export exercises the actual
exporter without suppressing the final presentation phase.
"""
from copy import deepcopy
from pathlib import Path
import hashlib
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Border, Color, Side
import coretax
import coretax_exporter as exporter
from support_templates import assert_layout, section, source_run

RESOURCE = Path(__file__).resolve().parents[1] / 'resources/corporate_format.xlsx'
FORMS = ['L1-D', 'L2', 'L3', 'L4', 'L6', 'L7', 'L9', 'L10-A', 'L10-C',
         'L11-A', 'L11-B', 'L11-C']
APPROVED_ACCOUNTING_RANGES = {
    'L2': ['E13:E21', 'G13:G22', 'J13:J22', 'G3:G9', 'I3:I9'],
    'L3': ['E8:E4662', 'F8:F4664'],
    'L6': ['B2:B8'], 'L7': ['C1:I13'],
    'L9': ['D5:E187', 'H5:H188', 'D193:E193', 'G193:H194',
           'H226:H228', 'H259:H261'],
    'L10-A': ['G3:G33'], 'L11-A': ['E8:E292'],
    'L11-B': ['B2:B7', 'D12:P13', 'B18:N23'],
}


def export(tmp_path, sections, *, year=2025, capture_complete=True):
    run = source_run('PPh Badan', sections)
    run['year'] = year
    run['records'][0]['meta']['Masa Pajak'] = f'Januari - Desember {year}'
    run['records'][0]['capture_complete'] = capture_complete
    before = deepcopy(run)
    reports = []
    with patch('coretax_badan_template._compact_layout',
               side_effect=lambda sheet, form, deleted, der_row, averages: (der_row, averages)):
        paths = exporter.export_split_workbooks(run, tmp_path, mapping_reports=reports)
    assert run == before
    name, = [name for name in paths if name != 'resume.xlsx']
    coretax.verify_outputs(run, tmp_path)
    return load_workbook(tmp_path / name), reports


def test_binding_stage_preserves_reference_geometry_and_approved_styles(tmp_path):
    book, _ = export(tmp_path, [section(name) for name in FORMS])
    reference = load_workbook(RESOURCE)
    try:
        assert book.sheetnames == reference.sheetnames == FORMS
        # Explicitly approved number-format/border exceptions; all other style and
        # layout attributes still compare against the untouched reference.
        accounting = reference['L1-D']['C44'].number_format.replace('0.00', '0')
        for sheet in reference:
            for row in sheet:
                for cell in row:
                    if not isinstance(cell, MergedCell) and 'Rp' in cell.number_format and '*' in cell.number_format:
                        cell.number_format = cell.number_format.replace('0.00', '0')
        for form, ranges in APPROVED_ACCOUNTING_RANGES.items():
            for area in ranges:
                for row in reference[form][area]:
                    for cell in row:
                        if not isinstance(cell, MergedCell):
                            cell.number_format = accounting
        for area in ('H3:H9', 'F1:F21'):
            for row in reference['L2'][area]:
                for cell in row:
                    if not isinstance(cell, MergedCell):
                        cell.number_format = '0.00%'
        for address in ('F4', 'G4', 'J4'):
            reference['L3'][address].number_format = '_-"Rp"* #,##0_-;\\-"Rp"* #,##0_-;_-"Rp"* "-"??_-;_-@_-'
        for row in (226, 227, 228):
            reference['L9'].cell(row, 8).number_format = accounting.replace('0.00', '0')
        thin = Side(style='thin', color=Color(indexed=64))
        for row in reference['L6']['A2:B8']:
            for cell in row:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        der = reference['L11-B']
        der['A26'] = 'PERHITUNGAN DER'
        der['B26'] = 'Jumlah Saldo Rata-Rata Utang'
        der['B27'] = 'Jumlah Saldo Rata-Rata Modal'
        der['D26'].number_format = '0.00'
        for name in FORMS:
            assert_layout(book[name], reference[name])
            assert book[name].max_row == reference[name].max_row
            for cells in reference[name]:
                for cell in cells:
                    assert book[name][cell.coordinate].value == cell.value
    finally:
        book.close()
        reference.close()


def test_resource_has_accounting_blanks_without_taxpayer_values():
    assert hashlib.sha256(RESOURCE.read_bytes()).hexdigest() == (
        '47770d95a613ed496c8f121cbe39f41ea17947855e709ab0f931adbe9ca58ab0')
    reference = load_workbook(RESOURCE)
    try:
        assert reference.sheetnames == FORMS
        sheet = reference['L1-D']
        for row in range(44, 102):
            if row == 70:  # Empty separator between the two balance-sheet tables.
                assert sheet.cell(row, 3).value is None
                continue
            if row == 71:  # The repeated NILAI header is not an amount slot.
                assert sheet.cell(row, 3).value == 'NILAI'
                continue
            cell = sheet.cell(row, 3)
            assert cell.value is None
            assert 'Rp' in cell.number_format and '*' in cell.number_format
        for name in FORMS:
            assert not any(cell.data_type == 'f' or isinstance(cell.value, (int, float))
                           for cells in reference[name] for cell in cells)
    finally:
        reference.close()


@pytest.mark.parametrize('count', [1, 32])
def test_unique_table_binds_headers_and_grows_without_changing_identifier(tmp_path, count):
    reference = load_workbook(RESOURCE)
    try:
        headers = [cell.value for cell in reference['L10-A'][2]]
    finally:
        reference.close()
    rows = [['SYNTHETIC NAME', '0000000000000002', 'SYNTHETIC COUNTRY',
             'SYNTHETIC RELATION', 'SYNTHETIC ACTIVITY', 'SYNTHETIC TRANSACTION',
             '1.234,50', 'SYNTHETIC METHOD', 'SYNTHETIC REASON'] for _ in range(count)]
    table = {'name': 'Synthetic source table', 'headers': headers, 'rows': rows,
             'body_row_count': count, 'source_total': count, 'complete': True,
             'footer': [], 'footer_spans': []}
    book, reports = export(tmp_path, [section('L10-A', tables=[table])])
    try:
        sheet = book['L10-A']
        assert sheet['A2'].value == 'Nama'
        assert sheet.max_row == max(33, count + 2)
        for row in range(3, count + 3):
            assert sheet.cell(row, 2).value == '0000000000000002'
            assert sheet.cell(row, 7).value == 1234.5
            assert sheet.cell(row, 7).data_type == 'n'
        assert reports and reports[0]['complete']
    finally:
        book.close()


def test_ambiguous_balance_table_stays_unresolved_without_raw_display(tmp_path):
    table = {'name': 'Synthetic ambiguous table', 'headers': ['KODE AKUN', 'NAMA AKUN', 'NILAI'],
             'rows': [['0001', 'SYNTHETIC UNBOUND', 'Rp. | 1.234,50']],
             'body_row_count': 1, 'source_total': 1, 'complete': True,
             'footer': [], 'footer_spans': []}
    book, reports = export(tmp_path, [section('L1-D', tables=[table])])
    try:
        sheet = book['L1-D']
        assert sheet['C44'].value is None and sheet['C73'].value is None
        assert not any(cell.value == 'Rp. | 1.234,50' for cells in sheet for cell in cells)
        assert reports and reports[0]['complete'] is False
        assert reports[0]['unmapped_rows'] == 1
    finally:
        book.close()


def test_balance_roles_map_accounting_and_captured_total_without_recalculation(tmp_path):
    tables = []
    for component, marker, amount, total in (
            ('rsh-citr-l1d-grid3', 'SYNTHETIC LIABILITY', 'Rp. | 2.000,00', 'Rp. | 9.000,00'),
            ('rsh-citr-l1d-grid2', 'SYNTHETIC ASSET', 'Rp. | 1.234,50', 'Rp. | 8.000,00')):
        label = 'Jumlah Aset' if component.endswith('grid2') else 'Jumlah Liabilitas dan Ekuitas'
        footer = [['', label, total]]
        tables.append({'name': 'Same literal table name',
                       'headers': ['KODE AKUN', 'NAMA AKUN', 'NILAI'],
                       'rows': [['', marker, ''], ['0001', 'SYNTHETIC DETAIL', amount]] + footer,
                       'body_row_count': 2, 'source_total': 2, 'complete': True,
                       'footer': footer, 'footer_spans': [[1, 1, 1]],
                       'source_structure': {'form': 'L1-D', 'component': component,
                                            'groups': ['B. LAPORAN POSISI KEUANGAN']}})
    book, reports = export(tmp_path, [section('L1-D', tables=tables)])
    try:
        sheet = book['L1-D']
        assert sheet['A43'].value == 'SYNTHETIC ASSET'
        assert sheet['A72'].value == 'SYNTHETIC LIABILITY'
        assert sheet['A44'].value == sheet['A73'].value == '0001'
        for coordinate, value in [('C44', 1234.5), ('C73', 2000), ('C69', 8000), ('C101', 9000)]:
            assert sheet[coordinate].value == value
            assert sheet[coordinate].data_type == 'n'
            assert 'Rp' in sheet[coordinate].number_format and '*' in sheet[coordinate].number_format
        assert reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('form,slots', [
    ('L6', [(f'A{row}', f'B{row}') for row in range(2, 9)]),
    ('L11-B', [(f'A{row}', f'B{row}') for row in range(2, 8)]),
    ('L9', [(f'A{row}', f'H{row}') for row in (226, 227, 228, 259, 260, 261)]),
])
def test_unique_scalar_labels_fill_their_numeric_slots(tmp_path, form, slots):
    reference = load_workbook(RESOURCE)
    try:
        fields = [{'label': reference[form][label].value, 'field': 'synthetic_scalar',
                   'value': 'Rp. | 1.234,50'} for label, _ in slots]
    finally:
        reference.close()
    book, reports = export(tmp_path, [section(form, fields=fields)])
    try:
        for _, destination in slots:
            assert book[form][destination].value == 1234.5
            assert book[form][destination].data_type == 'n'
        assert reports[0]['complete'] is (form != 'L11-B')
        if form == 'L11-B':
            assert reports[0]['derived_reason'] == 'missing_average'
            assert reports[0]['unmapped_fields'] == reports[0]['unmapped_tables'] == 0
    finally:
        book.close()


def test_conflicting_scalar_duplicates_remain_unresolved_without_last_write_wins(tmp_path):
    label = '1. PENGHASILAN YANG MENJADI DASAR PENGHITUNGAN ANGSURAN'
    fields = [{'label': label, 'value': value} for value in ('1.234,50', '2.345,60')]
    book, reports = export(tmp_path, [section('L6', fields=fields)])
    try:
        assert book['L6']['B2'].value is None
        assert reports[0]['unmapped_fields'] == 2
        assert not reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('duplicate', [False, True])
def test_l9_reordered_roles_bind_each_group_and_reject_duplicate_snapshots(tmp_path, duplicate):
    specs = [
        ('grid1', 'HARTA BERWUJUD', group, header, total)
        for group, header, total in [('KELOMPOK 1', 4, 188), ('KELOMPOK 2', 192, 194),
                                     ('KELOMPOK 3', 198, 200), ('KELOMPOK 4', 204, 206),
                                     ('LAINNYA', 210, 212)]]
    specs += [('grid2', 'BANGUNAN', 'PERMANEN', 216, 218),
              ('grid2', 'BANGUNAN', 'TIDAK PERMANEN', 222, 224)]
    specs += [('grid3', 'HARTA TIDAK BERWUJUD', group, header, total)
              for group, header, total in [('KELOMPOK 1', 231, 233), ('KELOMPOK 2', 237, 239),
                                           ('KELOMPOK 3', 243, 245), ('KELOMPOK 4', 249, 251),
                                           ('LAINNYA', 255, 257)]]
    reference = load_workbook(RESOURCE)
    tables = []
    try:
        headers = [cell.value for cell in reference['L9'][4]]
    finally:
        reference.close()
    for index, (grid, outer, inner, _, _) in enumerate(specs):
        footer = [['TOTAL', '99,00', '']]
        tables.append({'name': 'Same literal table name', 'headers': headers,
                       'rows': [['0001', f'SYNTHETIC ASSET {index}', '01/2025',
                                 '1.234,50', '1.000,00', 'GL', 'GL', '34,50', '']] + footer,
                       'complete': True, 'body_row_count': 1, 'source_total': 1,
                       'footer': footer, 'footer_spans': [[7, 1, 1]],
                       'source_structure': {'form': 'L9', 'component': 'rshshr-citr-l9-form-' + grid,
                                            'groups': [outer, inner]}})
    if duplicate:
        tables.append(deepcopy(tables[0]))
    book, reports = export(tmp_path, [section('L9', tables=list(reversed(tables)))])
    try:
        for index, (_, _, _, header, total) in enumerate(specs):
            if duplicate and index == 0:
                assert book['L9'].cell(header + 1, 2).value is None
                continue
            assert book['L9'].cell(header + 1, 2).value == f'SYNTHETIC ASSET {index}'
            assert book['L9'].cell(header + 1, 4).value == 1234.5
            assert book['L9'].cell(total, 8).value == 99
        assert reports[0]['complete'] is (not duplicate)
        assert reports[0]['unmapped_tables'] == (2 if duplicate else 0)
    finally:
        book.close()


def test_l9_scalars_follow_labels_through_multiple_table_growths(tmp_path):
    specs = [('grid1', 'HARTA BERWUJUD', 'KELOMPOK 1', 184, 189),
             ('grid2', 'BANGUNAN', 'PERMANEN', 3, 221),
             ('grid2', 'BANGUNAN', 'TIDAK PERMANEN', 3, 229)]
    scalar_rows = (226, 227, 228, 259, 260, 261)
    reference = load_workbook(RESOURCE)
    try:
        headers = [cell.value for cell in reference['L9'][4]]
        fields = [{'label': reference['L9'].cell(row, 1).value,
                   'value': '1.234,50'} for row in scalar_rows]
    finally:
        reference.close()
    tables = []
    for grid, outer, inner, count, _ in specs:
        footer = [['TOTAL', '99,00', '']]
        rows = [['0001', f'SYNTHETIC ASSET {index}', '01/2025',
                 '1.234,50', '1.000,00', 'GL', 'GL', '34,50', '']
                for index in range(count)]
        tables.append({'name': 'Synthetic growth table', 'headers': headers,
                       'rows': rows + footer, 'body_row_count': count,
                       'source_total': count, 'complete': True,
                       'footer': footer, 'footer_spans': [[7, 1, 1]],
                       'source_structure': {'form': 'L9',
                                            'component': 'rshshr-citr-l9-form-' + grid,
                                            'groups': [outer, inner]}})
    book, reports = export(tmp_path, [section('L9', fields=fields, tables=list(reversed(tables)))])
    try:
        sheet = book['L9']
        for _, _, _, _, total_row in specs:
            assert sheet.cell(total_row, 8).value == 99
        # One added row in group 1 and two in each building block shift both
        # captions and values; fixed pre-growth coordinates are not the contract.
        for original, field in zip(scalar_rows, fields):
            shifted = original + 5
            assert sheet.cell(shifted, 1).value == field['label']
            assert sheet.cell(shifted, 8).value == 1234.5
            assert sheet.cell(shifted, 8).data_type == 'n'
            assert f'A{shifted}:G{shifted}' in sheet.merged_cells
        assert reports[0]['complete']
    finally:
        book.close()


def test_l1d_year_and_identifier_remain_literal_in_reference_identity_slots(tmp_path):
    fields = [{'label': 'Tahun Pajak', 'value': '2025'},
              {'label': 'Nomor Identitas WP', 'value': '0000000000000001'}]
    book, reports = export(tmp_path, [section('L1-D', fields=fields)])
    try:
        assert book['L1-D']['B2'].value == '2025'
        assert book['L1-D']['B3'].value == '0000000000000001'
        assert book['L1-D']['B3'].data_type == 's'
        assert reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('year,header_year', [(2024, 2024), (2025, 2025), (2025, 2024)])
def test_l7_uses_only_current_captured_year_header_bound_to_record_period(tmp_path, year, header_year):
    reference = load_workbook(RESOURCE)
    try:
        headers = [cell.value for cell in reference['L7'][2]]
    finally:
        reference.close()
    caption = f'KOMPENSASI KERUGIAN FISKAL / {header_year} / TAHUN PAJAK INI - NILAI (Rp)'
    headers[7] = caption
    table = {'name': 'Synthetic L7', 'headers': headers,
             'rows': [['0001', str(year - 1), '1.234,50', '0', '0', '0', '0', '5.000,00', '0']],
             'complete': True, 'body_row_count': 1, 'source_total': 1,
             'footer': [], 'footer_spans': []}
    book, reports = export(tmp_path, [section('L7', tables=[table])], year=year)
    try:
        sheet = book['L7']
        if year == header_year:
            assert sheet['H2'].value == caption
            assert sheet['A3'].value == '0001'
            assert sheet['B3'].value == str(year - 1)
            assert sheet['H3'].value == 5000
            assert reports[0]['complete']
        else:
            assert sheet['H2'].value is None and sheet['H3'].value is None
            assert reports[0]['unmapped_tables'] == 1
            assert not reports[0]['complete']
    finally:
        book.close()


def captured_table(form, header, rows, footer, spans, **extra):
    reference = load_workbook(RESOURCE)
    try:
        headers = [cell.value for cell in reference[form][header]]
        while headers and headers[-1] is None:
            headers.pop()
    finally:
        reference.close()
    return {'name': 'Synthetic captured table', 'headers': headers,
            'rows': rows + footer, 'body_row_count': len(rows),
            'source_total': len(rows), 'complete': True,
            'footer': footer, 'footer_spans': spans, **extra}


@pytest.mark.parametrize('grid,header,first,count,total,label', [
    (1, 5, 6, 34, 40, 'Laba (Rugi) Sebelum Pajak'),
    (2, 42, 43, 26, 69, 'Jumlah Aset'),
    (3, 71, 72, 29, 101, 'Jumlah Liabilitas dan Ekuitas'),
])
def test_l1d_full_blocks_preserve_footer_account_codes(tmp_path, grid, header, first, count, total, label):
    reference = load_workbook(RESOURCE)
    width = 10 if grid == 1 else 3
    try:
        rows = []
        for index in range(count):
            merged = f'A{first + index}:B{first + index}' in reference['L1-D'].merged_cells
            rows.append(['' if merged else '0001', f'SYNTHETIC ACCOUNT {index}']
                        + ['0'] * (width - 2))
        if grid == 1:
            for row in rows:
                row[8] = ''
    finally:
        reference.close()
    footer = [['0099', label] + ['1.234,50'] * (width - 2)]
    if grid == 1:
        footer[0][8] = ''
    table = captured_table('L1-D', header, rows, footer, [[1] * width],
                           source_structure={'form': 'L1-D',
                               'component': f'rsh-citr-l1d-grid{grid}',
                               'groups': ['A. LAPORAN LABA RUGI' if grid == 1 else
                                          'B. LAPORAN POSISI KEUANGAN']})
    book, reports = export(tmp_path, [section('L1-D', tables=[table])])
    try:
        sheet = book['L1-D']
        assert sheet.cell(first, 1).value == 'SYNTHETIC ACCOUNT 0'
        assert sheet.cell(total - 1, 2).value == f'SYNTHETIC ACCOUNT {count - 1}'
        assert sheet.cell(total, 1).value == '0099'
        assert sheet.cell(total, 1).data_type == 's'
        assert sheet.cell(total, 2).value == label
        assert sheet.cell(total, 3).value == 1234.5
        assert reports[0]['complete']
        assert sheet.max_row == 101
    finally:
        book.close()


@pytest.mark.parametrize('form,header,count,target,footer,spans,amounts', [
    ('L2', 2, 6, 9, [['', 'JUMLAH', '1.234,50', '100,00', '9.876,50']],
     [[1, 5, 1, 1, 1]], {'G9': 1234.5, 'H9': 1, 'I9': 9876.5}),
    ('L2', 12, 9, 22, [['', 'JUMLAH', '1.234,50', '', '9.876,50', '']],
     [[1, 5, 1, 2, 1, 2]], {'G22': 1234.5, 'J22': 9876.5}),
    ('L4', 8, 1, 10, [['', 'JUMLAH', '1.234,50']],
     [[1, 3, 1]], {'E10': 1234.5}),
])
def test_footer_label_after_blank_prefix_binds_merged_total(tmp_path, form, header, count, target, footer, spans, amounts):
    table = captured_table(form, header, [], footer, spans)
    rows = [[''] * len(table['headers']) for _ in range(count)]
    rows[0][0] = '0001'
    table.update(rows=rows + footer, body_row_count=count, source_total=count)
    book, reports = export(tmp_path, [section(form, tables=[table])])
    try:
        sheet = book[form]
        assert sheet.cell(header + 1, 1).value == '0001'
        assert sheet.cell(target, 1).value == 'JUMLAH'
        for cell, value in amounts.items():
            assert sheet[cell].value == value
        assert reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('count', [4654, 4656])
def test_l3_all_details_and_three_captured_totals_map_without_recalculation(tmp_path, count):
    rows = [[str(index), 'SYNTHETIC WITHHOLDER', '0000000000000002', 'SYNTHETIC TAX',
             '1.000,00', '20,00', f'000{index}', '01-01-2025'] for index in range(count)]
    footer = [['JUMLAH', '8.000,00', '900,00', ''],
              ['KREDIT PAJAK LUAR NEGERI', '', '100,00', ''],
              ['JUMLAH KREDIT PAJAK', '', '1.000,00', '']]
    table = captured_table('L3', 7, rows, footer, [[4, 1, 1, 2]] * 3)
    book, reports = export(tmp_path, [section('L3', tables=[table])])
    try:
        sheet = book['L3']
        last = 7 + count
        assert sheet['A8'].value == '0' and sheet.cell(last, 1).value == str(count - 1)
        assert sheet['C8'].value == '0000000000000002'
        assert sheet.cell(last, 7).value == f'000{count - 1}'
        assert sheet.cell(last + 1, 5).value == 8000
        assert sheet.cell(last + 1, 6).value == 900
        assert sheet.cell(last + 2, 1).value == 'KREDIT PAJAK LUAR NEGERI'
        assert sheet.cell(last + 2, 6).value == 100
        assert sheet.cell(last + 3, 6).value == 1000
        assert f'A{last + 3}:D{last + 3}' in sheet.merged_cells
        assert sheet.max_row == last + 3 and reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('form,last', [
    ('L2', 22), ('L3', 4664), ('L4', 10), ('L6', 8), ('L7', 13),
    ('L10-A', 33), ('L10-C', 3), ('L11-A', 305), ('L11-B', 32),
])
def test_requested_clean_forms_omit_raw_display_but_keep_truthful_audit(tmp_path, form, last):
    fields = [{'label': 'Synthetic unresolved field', 'value': '0009'}]
    table = {'name': 'Synthetic unknown table', 'headers': ['Unknown'],
             'rows': [['=SYNTHETIC()']], 'complete': True}
    book, reports = export(tmp_path, [section(form, fields=fields, tables=[table],
                                             text='Synthetic redundant page text')])
    try:
        sheet = book[form]
        assert sheet.max_row == last
        assert not any(cell.value == 'Data sumber tambahan / belum dipetakan'
                       for cells in sheet for cell in cells)
        assert reports[0]['unmapped_fields'] == 1
        assert reports[0]['unmapped_tables'] == 1 and reports[0]['unmapped_rows'] == 1
        assert reports[0]['complete'] is False
        resume = load_workbook(tmp_path / 'resume.xlsx')
        try:
            audit_rows = [row for row in resume['Audit'].iter_rows(values_only=True)
                          if row[0] == 'Pemetaan template']
            assert len(audit_rows) == 1
            assert audit_rows[0][1:5] == (1, 1, 1, 'PARTIAL')
            assert f' / {form}:' in audit_rows[0][5]
        finally:
            resume.close()
        if form in ('L6', 'L7'):
            cells = [sheet.cell(row, column) for row in (range(2, 9) if form == 'L6' else range(1, 14))
                     for column in ([2] if form == 'L6' else range(3, 10))]
            for cell in cells:
                # C13 belongs to the merged JUMLAH label and has no value slot.
                if cell.coordinate != 'C13':
                    assert 'Rp' in cell.number_format and '*' in cell.number_format
    finally:
        book.close()


@pytest.mark.parametrize('fault', ['unknown_label', 'conflicting_prefix', 'wrong_span', 'extra_footer'])
def test_unsupported_footer_stays_raw_without_partial_table_writes(tmp_path, fault):
    footer, spans = [['', 'JUMLAH', '1.234,50']], [[1, 3, 1]]
    if fault == 'unknown_label':
        footer[0][1] = 'SYNTHETIC UNKNOWN TOTAL'
    elif fault == 'conflicting_prefix':
        footer[0][0] = 'SYNTHETIC CONFLICT'
    elif fault == 'wrong_span':
        spans[0] = [1, 2, 1]
    else:
        footer *= 2
        spans *= 2
    table = captured_table('L4', 8, [['0001', '0002', 'SYNTHETIC TYPE', 'SYNTHETIC SOURCE', '1.000,00']],
                           footer, spans)
    book, reports = export(tmp_path, [section('L4', tables=[table])])
    try:
        assert book['L4']['A9'].value is None and book['L4']['E10'].value is None
        assert book['L4'].max_row == 10
        assert not any(cell.value == 'Data sumber tambahan / belum dipetakan'
                       for cells in book['L4'] for cell in cells)
        assert reports[0]['unmapped_tables'] == 1
        assert not reports[0]['complete']
    finally:
        book.close()


def test_unsafe_footer_amount_stays_literal_and_unresolved(tmp_path):
    footer = [['', 'JUMLAH', '9.999.999.999.999.999,99']]
    table = captured_table('L4', 8, [['0001', '0002', 'SYNTHETIC TYPE', 'SYNTHETIC SOURCE', '1.000,00']],
                           footer, [[1, 3, 1]])
    book, reports = export(tmp_path, [section('L4', tables=[table])])
    try:
        assert book['L4']['A9'].value == '0001'
        assert book['L4']['E10'].value == footer[0][2]
        assert book['L4']['E10'].data_type == 's'
        assert reports[0]['unmapped_rows'] == 1 and not reports[0]['complete']
    finally:
        book.close()


def test_l7_ten_year_rows_and_total_use_accounting_without_raw_appendix(tmp_path):
    footer = [['JUMLAH'] + ['9.876,50'] * 6]
    rows = [[str(index), str(2015 + index)] + ['1.234,50'] * 7 for index in range(10)]
    table = captured_table('L7', 2, rows, footer, [[3, 1, 1, 1, 1, 1, 1]])
    table['headers'][7] = 'KOMPENSASI KERUGIAN FISKAL / 2025 / TAHUN PAJAK INI - NILAI (Rp)'
    book, reports = export(tmp_path, [section('L7', tables=[table], text='Synthetic duplicate page text')])
    try:
        sheet = book['L7']
        assert sheet['B3'].value == '2015' and sheet['B12'].value == '2024'
        assert sheet['C3'].value == sheet['I12'].value == 1234.5
        assert sheet['D13'].value == sheet['I13'].value == 9876.5
        assert sheet.max_row == 13 and reports[0]['complete']
        for cell in ('C3', 'I12', 'D13', 'I13'):
            assert 'Rp' in sheet[cell].number_format and '*' in sheet[cell].number_format
    finally:
        book.close()


@pytest.mark.parametrize('unsupported', [0, False])
def test_falsy_footer_value_cannot_match_an_empty_reference_label(tmp_path, unsupported):
    footer = [['TOTAL', '100,00', unsupported, '20,00']]
    rows = [['0000000000000002', 'SYNTHETIC WITHHOLDER', '0001', 'SYNTHETIC TAX',
             '1.000,00', '2,00', '20,00']]
    table = captured_table('L4', 2, rows, footer, [[4, 1, 1, 1]])
    book, reports = export(tmp_path, [section('L4', tables=[table])])
    try:
        assert book['L4']['A3'].value is None and book['L4']['F4'].value is None
        assert reports[0]['unmapped_tables'] == 1 and not reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('value,expected', [
    ('25,1234', 0.251234), ('100,0000', 1), ('0,0000', 0), ('-2,5000', -0.025),
])
def test_l2_percentage_points_become_excel_fractions_in_both_tables_and_footer(tmp_path, value, expected):
    owner = ['0001', 'SYNTHETIC OWNER', 'SYNTHETIC ADDRESS', 'SYNTHETIC COUNTRY',
             '0000000000000002', 'SYNTHETIC DIRECTOR', '1.234,50', value, '50,00']
    affiliate = ['0002', 'SYNTHETIC AFFILIATE', 'SYNTHETIC COUNTRY', '0000000000000003',
                 '1.000,00', value, '2.000,00', '2025', '3', '4.000,00', '2025', '4']
    first = captured_table('L2', 2, [owner], [['', 'JUMLAH', '9.999,00', value, '10,00']],
                           [[1, 5, 1, 1, 1]])
    second = captured_table('L2', 12, [affiliate], [], [])
    book, reports = export(tmp_path, [section('L2', tables=[second, first])])
    try:
        sheet = book['L2']
        for coord in ('H3', 'H9', 'F13'):
            assert sheet[coord].value == expected
            assert sheet[coord].data_type == 'n'
            assert sheet[coord].number_format == '0.00%'
        assert sheet['F3'].value == 'SYNTHETIC DIRECTOR'
        assert sheet['F3'].number_format == '0.00%'
        assert sheet['F2'].value == 'Jabatan'
        assert sheet['E3'].value == '0000000000000002'
        assert sheet['D13'].value == '0000000000000003'
        assert sheet['H13'].value == '2025'
        assert sheet['G9'].value == 9999  # Captured total, not the body sum.
        assert sheet.max_row == 22 and reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('value', ['1.23,4567', '1,23456', '9999999999999999,1234',
                                 'NaN', 'Rp. | 25,0000', True])
def test_l2_unsafe_percentage_stays_literal_and_unresolved(tmp_path, value):
    row = ['0001', 'SYNTHETIC OWNER', '', '', '0000000000000002', '', '1,00', value, '0']
    table = captured_table('L2', 2, [row], [], [])
    book, reports = export(tmp_path, [section('L2', tables=[table])])
    try:
        assert book['L2']['H3'].value == value
        assert not reports[0]['complete'] and reports[0]['unmapped_rows'] == 1
        assert book['L2'].max_row == 22
    finally:
        book.close()


@pytest.mark.parametrize('amount,expected', [
    ('Rp. | 1.234,50', -1234.5), ('Rp. | -1.234,50', -1234.5), ('0', 0),
    ('9.999.999.999.999.999,99', '9.999.999.999.999.999,99'), ('1.23,45', '1.23,45'),
])
def test_l1d_verified_subtraction_account_has_negative_presentation_without_source_mutation(tmp_path, amount, expected):
    rows = [['', 'SYNTHETIC ASSET HEADING', '']]
    rows += [['0001', 'SYNTHETIC ACCOUNT', '0'] for _ in range(16)]
    rows += [['0002', 'Dikurangi: Akumulasi penyusutan aset tetap dan inventaris', amount]]
    rows[14][0] = ''  # Reference A57:B57 is an account-heading merge.
    table = captured_table('L1-D', 42, rows, [], [], source_structure={
        'form': 'L1-D', 'component': 'rsh-citr-l1d-grid2',
        'groups': ['B. LAPORAN POSISI KEUANGAN']})
    book, reports = export(tmp_path, [section('L1-D', tables=[table])])
    try:
        assert book['L1-D']['C60'].value == expected
        assert book['L1-D']['A60'].value == '0002'
        assert reports[0]['complete'] is isinstance(expected, (int, float))
    finally:
        book.close()


def test_l1d_subtraction_sign_follows_verified_caption_instead_of_coordinate(tmp_path):
    rows = [['', 'SYNTHETIC HEADING', ''],
            ['0002', 'Dikurangi: Akumulasi penyusutan aset tetap dan inventaris', '1.234,50']]
    rows += [['0001', 'SYNTHETIC ORDINARY ACCOUNT', '9,00'] for _ in range(16)]
    rows[14][0] = ''
    table = captured_table('L1-D', 42, rows, [], [], source_structure={
        'form': 'L1-D', 'component': 'rsh-citr-l1d-grid2',
        'groups': ['B. LAPORAN POSISI KEUANGAN']})
    book, reports = export(tmp_path, [section('L1-D', tables=[table])])
    try:
        assert book['L1-D']['C44'].value == -1234.5
        assert book['L1-D']['C60'].value == 9
        assert reports[0]['complete']
    finally:
        book.close()


def test_l1d_income_money_slots_gain_accounting_without_touching_account_codes(tmp_path):
    rows = [['', 'SYNTHETIC HEADING'] + [''] * 8,
            ['0001', 'SYNTHETIC ACCOUNT'] + ['1.234,50'] * 6 + ['0009', '1.234,50']]
    table = captured_table('L1-D', 5, rows, [], [], source_structure={
        'form': 'L1-D', 'component': 'rsh-citr-l1d-grid1',
        'groups': ['A. LAPORAN LABA RUGI']})
    book, reports = export(tmp_path, [section('L1-D', tables=[table])])
    try:
        for coord in ('C7', 'D7', 'E7'):
            assert book['L1-D'][coord].value == 1234.5
            assert 'Rp' in book['L1-D'][coord].number_format
            assert '*' in book['L1-D'][coord].number_format
        assert book['L1-D']['A7'].value == '0001'
        assert book['L1-D']['I7'].value == '0009'
        assert reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('amount,expected', [
    ('1.234,50', 1234.5), ('9.999.999.999.999.999,99', '9.999.999.999.999.999,99'),
])
def test_l11a_exact_entertainment_heading_binds_safe_amounts(tmp_path, amount, expected):
    row = ['01/01/2025', 'SYNTHETIC PLACE', 'SYNTHETIC ADDRESS', 'SYNTHETIC TYPE', amount,
           'SYNTHETIC PERSON', 'SYNTHETIC ROLE', 'SYNTHETIC BUSINESS', 'SYNTHETIC INDUSTRY', '']
    table = captured_table('L11-A', 7, [row], [], [])
    book, reports = export(tmp_path, [section('L11-A', tables=[table])])
    try:
        assert book['L11-A']['E8'].value == expected
        assert 'Rp' in book['L11-A']['E8'].number_format
        assert book['L11-A']['A8'].value == '01/01/2025'
        assert reports[0]['complete'] is isinstance(expected, (int, float))
        assert book['L11-A'].max_row == 305
    finally:
        book.close()


def l11b_der_tables(*, debt='9.000,00', capital='3.000,00', grow=False):
    debt_rows = [['0000000000000002', 'SYNTHETIC LENDER', 'SYNTHETIC RELATION']
                 + ['1,00'] * 13 for _ in range(3 if grow else 1)]
    capital_rows = [[f'SYNTHETIC CAPITAL {index}'] + ['2,00'] * 13
                    for index in range(7 if grow else 5)]
    return [
        captured_table('L11-B', 11, debt_rows, [['JUMLAH'] + ['4,00'] * 12 + [debt]],
                       [[3] + [1] * 13]),
        captured_table('L11-B', 17, capital_rows, [['JUMLAH'] + ['5,00'] * 12 + [capital]],
                       [[1] * 14]),
    ]


@pytest.mark.parametrize('grow', [False, True])
def test_l11b_der_uses_bound_captured_averages_and_tracks_table_growth(tmp_path, grow):
    tables = l11b_der_tables(grow=grow)
    loan = captured_table('L11-B', 30,
        [['SYNTHETIC LENDER', '1.234,50', '12,50', '10,00', '2,50']],
        [['JUMLAH', '9.000,00', '80,00', '50,00', '30,00']], [[1] * 5])
    book, reports = export(tmp_path, [section('L11-B', tables=[loan] + list(reversed(tables)))])
    try:
        sheet = book['L11-B']
        shift = 4 if grow else 0
        der_row = 26 + shift
        assert sheet.cell(der_row, 1).value == 'PERHITUNGAN DER'
        assert sheet.cell(der_row, 2).value == 'Jumlah Saldo Rata-Rata Utang'
        assert sheet.cell(der_row + 1, 2).value == 'Jumlah Saldo Rata-Rata Modal'
        assert sheet.cell(der_row, 3).value == ('=P15' if grow else '=P13')
        assert sheet.cell(der_row + 1, 3).value == ('=N27' if grow else '=N23')
        assert sheet.cell(der_row, 4).value == f'=C{der_row}/C{der_row + 1}'
        assert sheet.cell(der_row, 4).number_format == '0.00'
        assert f'A{der_row}:A{der_row + 1}' in sheet.merged_cells
        assert f'D{der_row}:D{der_row + 1}' in sheet.merged_cells
        assert sheet.cell(31 + shift, 2).value == 1234.5
        assert sheet.cell(32 + shift, 2).value == 9000
        assert sheet.cell(32 + shift, 5).value == 30
        assert sheet['A12'].value == '0000000000000002'
        for coord in (('D14', 'P15', 'B26', 'N27') if grow else ('D12', 'P13', 'B22', 'N23')):
            assert 'Rp' in sheet[coord].number_format
        assert sheet.max_row == 32 + shift and reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('capital', ['0', '', '1.23,45', '9.999.999.999.999.999,99'])
def test_l11b_der_unavailable_denominator_never_fabricates_zero_or_division(tmp_path, capital):
    book, reports = export(tmp_path, [section('L11-B', tables=l11b_der_tables(capital=capital))])
    try:
        assert book['L11-B']['D26'].value in (None, '')
        assert not reports[0]['complete']
        assert book['L11-B'].max_row == 32
        if capital not in ('0', ''):
            assert book['L11-B']['N23'].value == capital
    finally:
        book.close()


def test_l11b_missing_capital_table_does_not_emit_der_formula(tmp_path):
    book, reports = export(tmp_path, [section('L11-B', tables=l11b_der_tables()[:1])])
    try:
        assert book['L11-B']['C27'].value in (None, '')
        assert book['L11-B']['D26'].value in (None, '')
        assert not reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('content', ['borrowing_cost_only', 'scalar_only', 'empty'])
def test_observed_l11b_without_either_average_reports_missing_der_without_source_inflation(tmp_path, content):
    tables, fields = [], []
    if content == 'borrowing_cost_only':
        tables = [captured_table('L11-B', 30,
            [['SYNTHETIC LENDER', '1.234,50', '12,50', '10,00', '2,50']],
            [['JUMLAH', '9.000,00', '80,00', '50,00', '30,00']], [[1] * 5])]
    elif content == 'scalar_only':
        fields = [{'label': '1. PENGHASILAN NETO KOMERSIAL', 'value': '1.234,50'}]
    book, reports = export(tmp_path, [section('L11-B', fields=fields, tables=tables)])
    try:
        assert book.sheetnames == ['L11-B']
        for coord in ('C26', 'C27', 'D26'):
            assert book['L11-B'][coord].value in (None, '')
        if content == 'borrowing_cost_only':
            assert book['L11-B']['B31'].value == 1234.5
            assert book['L11-B']['B32'].value == 9000
        elif content == 'scalar_only':
            assert book['L11-B']['B2'].value == 1234.5
        report = reports[0]
        assert not report['complete'] and report['derived_complete'] is False
        assert report['derived_reason'] == 'missing_average'
        assert report['unmapped_fields'] == report['unmapped_tables'] == report['unmapped_rows'] == 0
        resume = load_workbook(tmp_path / 'resume.xlsx')
        try:
            audit = [row for row in resume['Audit'].iter_rows(values_only=True)
                     if row[0] == 'Pemetaan template']
            assert len(audit) == 1 and audit[0][1:5] == (0, 0, 0, 'PARTIAL')
            assert 'DER' in audit[0][5] and 'belum tersedia' in audit[0][5]
        finally:
            resume.close()
    finally:
        book.close()


def test_absent_l11b_does_not_create_a_der_form_or_der_audit(tmp_path):
    book, reports = export(tmp_path, [section('L6')])
    try:
        assert book.sheetnames == ['L6']
        assert len(reports) == 1 and reports[0]['form'] == 'L6'
        assert reports[0]['complete']
        assert 'derived_complete' not in reports[0] and 'derived_reason' not in reports[0]
    finally:
        book.close()


def test_l11b_zero_denominator_reports_der_partial_without_inventing_unmapped_source(tmp_path):
    book, reports = export(tmp_path, [section('L11-B', tables=l11b_der_tables(capital='0'))])
    try:
        report = reports[0]
        assert report['derived_complete'] is False
        assert report['derived_reason'] == 'zero_capital_average'
        assert report['unmapped_fields'] == report['unmapped_tables'] == report['unmapped_rows'] == 0
        assert not report['complete']
        assert book['L11-B']['D26'].value in (None, '')
        resume = load_workbook(tmp_path / 'resume.xlsx')
        try:
            audit = [row for row in resume['Audit'].iter_rows(values_only=True)
                     if row[0] == 'Pemetaan template']
            assert len(audit) == 1 and audit[0][1:5] == (0, 0, 0, 'PARTIAL')
            assert 'DER' in audit[0][5] and 'modal nol' in audit[0][5]
        finally:
            resume.close()
    finally:
        book.close()


@pytest.mark.parametrize('owner', [True, False])
def test_l9_inspected_kelompok_lainnya_alias_requires_source_owner(tmp_path, owner):
    tables = []
    for grid, outer in [('grid1', 'HARTA BERWUJUD'), ('grid3', 'HARTA TIDAK BERWUJUD')]:
        table = captured_table('L9', 4, [], [['TOTAL', '12,50', '']], [[7, 1, 1]])
        if owner:
            table['source_structure'] = {'form': 'L9',
                'component': 'rshshr-citr-l9-form-' + grid,
                'groups': [outer, 'KELOMPOK LAINNYA']}
        tables.append(table)
    book, reports = export(tmp_path, [section('L9', tables=list(reversed(tables)))])
    try:
        for coord in ('H212', 'H257'):
            if owner:
                assert book['L9'][coord].value == 12.5
            else:
                assert book['L9'][coord].value in (None, '')
        assert reports[0]['complete'] is owner
        assert reports[0]['unmapped_tables'] == (0 if owner else 2)
        assert not any(cell.value == 'Data sumber tambahan / belum dipetakan'
                       for cells in book['L9'] for cell in cells)
    finally:
        book.close()


@pytest.mark.parametrize('fault', ['wrong_outer', 'wrong_component', 'duplicate'])
def test_l9_lainnya_alias_does_not_relax_owner_or_duplicate_gates(tmp_path, fault):
    structure = {'form': 'L9', 'component': 'rshshr-citr-l9-form-grid1',
                 'groups': ['HARTA BERWUJUD', 'KELOMPOK LAINNYA']}
    if fault == 'wrong_outer':
        structure['groups'][0] = 'HARTA TIDAK BERWUJUD'
    elif fault == 'wrong_component':
        structure['component'] = 'rshshr-citr-l9-form-grid2'
    table = captured_table('L9', 4, [], [['TOTAL', '12,50', '']], [[7, 1, 1]],
                           source_structure=structure)
    tables = [table, deepcopy(table)] if fault == 'duplicate' else [table]
    book, reports = export(tmp_path, [section('L9', tables=tables)])
    try:
        assert book['L9']['H212'].value in (None, '')
        assert not reports[0]['complete']
        assert reports[0]['unmapped_tables'] == len(tables)
    finally:
        book.close()


@pytest.mark.parametrize('fault', ['wrong_span', 'unknown_caption', 'incomplete', 'duplicate'])
def test_l11b_borrowing_cost_binding_keeps_footer_and_completion_gates(tmp_path, fault):
    table = captured_table('L11-B', 30,
        [['SYNTHETIC LENDER', '1.234,50', '12,50', '10,00', '2,50']],
        [['JUMLAH', '9.000,00', '80,00', '50,00', '30,00']], [[1] * 5])
    if fault == 'wrong_span':
        table['footer_spans'] = [[1, 1, 1, 1, 2]]
    elif fault == 'unknown_caption':
        table['footer'][0][0] = 'SYNTHETIC UNKNOWN TOTAL'
    elif fault == 'incomplete':
        table['complete'] = False
    tables = [table, deepcopy(table)] if fault == 'duplicate' else [table]
    book, reports = export(tmp_path, [section('L11-B', tables=tables)],
                           capture_complete=fault != 'incomplete')
    try:
        assert book['L11-B']['B31'].value in (None, '')
        assert book['L11-B']['B32'].value in (None, '')
        assert reports[0]['unmapped_tables'] == len(tables)
        assert not reports[0]['complete']
        assert book['L11-B'].max_row == 32
    finally:
        book.close()


def test_l2_growth_copies_accounting_and_percentage_formats_before_later_block(tmp_path):
    owner = ['0001', 'SYNTHETIC OWNER', '', '', '0000000000000002',
             'SYNTHETIC ROLE', '1.234,50', '12,5000', '0']
    first = captured_table('L2', 2, [deepcopy(owner) for _ in range(8)],
                           [['', 'JUMLAH', '9.999,00', '100,0000', '10,00']], [[1, 5, 1, 1, 1]])
    affiliate = ['0002', 'SYNTHETIC AFFILIATE', '', '0000000000000003',
                 '1.000,00', '25,0000', '2.000,00', '2025', '3', '4.000,00', '2025', '4']
    second = captured_table('L2', 12, [affiliate], [], [])
    book, reports = export(tmp_path, [section('L2', tables=[first, second])])
    try:
        sheet = book['L2']
        assert sheet['G10'].value == 1234.5 and 'Rp' in sheet['G10'].number_format
        assert sheet['H10'].value == 0.125 and sheet['H10'].number_format == '0.00%'
        assert sheet['G11'].value == 9999 and sheet['H11'].value == 1
        assert sheet['F15'].value == 0.25 and sheet['F15'].number_format == '0.00%'
        assert sheet['G15'].value == 2000 and 'Rp' in sheet['G15'].number_format
        assert sheet['D15'].value == '0000000000000003'
        assert 'A11:F11' in sheet.merged_cells
        assert sheet.max_row == 24 and reports[0]['complete']
    finally:
        book.close()
