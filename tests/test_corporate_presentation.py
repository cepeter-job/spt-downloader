"""Explicit corporate display exclusions must not erase checkpoint/history."""
from copy import deepcopy
import hashlib
import json
import pytest
from openpyxl import Workbook, load_workbook
import coretax
import coretax_exporter as exporter
from support_templates import source_run, section


def source():
    run=source_run('PPh Badan', [
        section('L10-B', fields=[{'label':'Synthetic omitted B','value':'SYNTHETIC-B'}]),
        section('Induk', fields=[{'label':'Synthetic retained','value':'SYNTHETIC-INDUK'}]),
        section('L10-D', fields=[{'label':'Synthetic omitted D','value':'SYNTHETIC-D'}])])
    run['records'][0]['meta']['Masa Pajak']='Januari - Desember 2025'
    return run


def test_default_omits_only_approved_forms_and_links_to_first_visible_form(tmp_path):
    run=source(); before=deepcopy(run)
    paths=exporter.export_split_workbooks(run,tmp_path)
    assert run==before
    name,=[p for p in paths if p!='resume.xlsx']
    book=load_workbook(tmp_path/name)
    resume=load_workbook(tmp_path/'resume.xlsx')
    try:
        assert book.sheetnames==['Induk']
        proof=exporter.read_workbook_proof(book)
        assert proof==exporter.workbook_proof(run)
        assert proof['records'][0]['sections']==[{'name':'Induk'}]
        links=[c.hyperlink for row in resume['Daftar SPT'] for c in row if c.hyperlink]
        assert any(link.target==name and 'Induk' in (link.location or '') for link in links)
    finally:
        book.close();resume.close()
    coretax.verify_outputs(run,tmp_path)


def test_historical_proof_with_excluded_forms_remains_readable(tmp_path):
    run=source();proof=exporter.workbook_proof(run)
    proof['records'][0]['sections']=[{'name':s['name']} for s in run['records'][0]['sections']]
    book=Workbook();book.remove(book.active)
    for s in run['records'][0]['sections']:book.create_sheet(s['name'])
    book.properties.identifier=json.dumps(proof)
    path=tmp_path/'synthetic-historical.xlsx';book.save(path);book.close()
    restored=load_workbook(path)
    try:assert exporter.read_workbook_proof(restored)==proof
    finally:restored.close()


def test_noncorporate_same_literal_form_names_are_not_excluded(tmp_path):
    run=source();run['records'][0]['meta']['Jenis Pajak']='Synthetic Other Tax'
    paths=exporter.export_split_workbooks(run,tmp_path)
    name,=[p for p in paths if p!='resume.xlsx']
    book=load_workbook(tmp_path/name)
    try:assert book.sheetnames==['L10-B','Induk','L10-D']
    finally:book.close()


@pytest.mark.parametrize('verify_only', [False, True])
def test_retained_excluded_forms_keep_real_links_and_physical_proof(tmp_path, verify_only):
    current = source()
    exporter.export_split_workbooks(current, tmp_path)
    selected_name, = exporter.workbook_groups(current)
    historical = deepcopy(current)
    record = historical['records'][0]
    record.update(key='synthetic-retained', record_id='synthetic-retained')
    record['meta']['Model SPT'] = 'Pembetulan 2'
    record['sections'] = [section('L10-B'), section('L10-D')]
    historical_name, = exporter.workbook_groups(historical)
    old = Workbook()
    old.remove(old.active)
    for form in ('L10-B', 'L10-D'):
        old.create_sheet(form)
    old.properties.identifier = json.dumps(exporter.workbook_proof(historical, presentation=False))
    old.save(tmp_path / historical_name)
    old.close()
    before = hashlib.sha256((tmp_path / historical_name).read_bytes()).hexdigest()
    coretax.write_folder_resume(current, tmp_path, tmp_path / 'resume.xlsx')
    assert hashlib.sha256((tmp_path / historical_name).read_bytes()).hexdigest() == before
    if verify_only:
        coretax.verify_outputs(current, tmp_path, allow_extra=True)
    else:
        resume = load_workbook(tmp_path / 'resume.xlsx')
        try:
            links = [resume['Daftar SPT'].cell(row, 1).hyperlink for row in (5, 6)]
            assert [(link.target, link.location) for link in links] == [
                (selected_name, "'Induk'!A1"), (historical_name, "'L10-B'!A1")]
        finally:
            resume.close()
