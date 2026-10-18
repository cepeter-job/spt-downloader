# Synthetic packet-to-template contracts. No portal records or DOM tax values.
from copy import deepcopy
import json
from pathlib import Path

from openpyxl import load_workbook
from coretax_json import ORIGIN
from masa_json_fixtures import RECORD, SCOPE, packet, payload, selected

MAIN = ORIGIN + '/returnsheets/id-ID/main.cccfdb820c20596c.js'
BUILDS = {
    'PPN': {MAIN: '696ac66dc2897c937c7df7b466382dffdf25651837a20e0a4d4235787f3c9afd',
        ORIGIN + '/returnsheets/id-ID/640.36c042afd2381189.js': '4a7e33826b6d0395afb23372fb6b4a9d4be6abaab0ec25fa743c897d0e654aef'},
    'PPh Unifikasi': {MAIN: '696ac66dc2897c937c7df7b466382dffdf25651837a20e0a4d4235787f3c9afd',
        ORIGIN + '/returnsheets/id-ID/693.4de61c6eb2000c7e.js': '2694fdda5c23b7d5f0eefe0cf0a0c55caa0fe29acc590363f34684ed90188fe0'},
}
UNIF_HEADERS = ['NO.', 'NIK/NPWP', 'NAMA', 'NOMOR BUKTI POTONG', 'TANGGAL BUKTI POTONG',
    'JENIS PAJAK', 'KODE OBJEK PAJAK', 'OBJEK PAJAK', 'DASAR PENGENAAN PAJAK (Rp)',
    'TINGKAT (%)', 'PAJAK PENGHASILAN', 'FASILITAS PERPAJAKAN',
    'UANG PERSEDIAAN / PEMBAYARAN LANGSUNG (UNTUK WP INSTANSI PEMERINTAH DENGAN DANA APBN)',
    'NITKU / NOMOR IDENTITAS SUBUNIT ORGANISASI', 'STATUS', 'KAP-KJS']


def response(tax, data, endpoint, request, sequence):
    item = packet(tax, data, sequence=sequence, request=request)
    item['context']['url'] = ORIGIN + '/returnsheet/api/' + endpoint
    return item


def field(control, label, group='', column=''):
    return {'field': control, 'label': label, 'group': group, 'column_label': column, 'tag': 'INPUT'}


def native_case(tax, *, populated=True):
    outer = payload(tax)
    data = json.loads(outer['FormData'])
    data.update(Name='SYNTHETIC NATIVE NAME', Address='SYNTHETIC ADDRESS', PhoneNumber='00000000')
    forms, packets = [], []
    resource = 'ppn_format.xlsx' if tax == 'PPN' else 'unifikasi_format.xlsx'
    reference = load_workbook(Path(__file__).resolve().parents[1] / 'resources' / resource)
    try:
        def add_form(name, fields=()):
            result = {'name': name, 'fields': list(fields), 'tables': [],
                      'builds': deepcopy(BUILDS[tax]), 'locale': 'id-ID', 'timezone': 'Asia/Jakarta'}
            forms.append(result)
            return result

        def add_grid(form, name, headers, group, spans, rows, size):
            endpoint = ('loadndvat/' if tax == 'PPN' else 'loadwtr/') + name + '-grid'
            pages = []
            for first in range(0, max(1, len(rows)), size):
                sequence = len(packets) + 1
                request = {'ReturnSheetRecordId': RECORD, 'First': first, 'Rows': size, 'Filters': []}
                packets.append(response(tax, {'Data': deepcopy(rows[first:first + size]), 'TotalRecords': len(rows)},
                                        endpoint, request, sequence))
                pages.append({'request_json': request, 'offset': first, 'sequences': [sequence]})
            table = {'name': name, 'headers': headers, 'group': group, 'footer_spans': spans,
                     'endpoint': ORIGIN + '/returnsheet/api/' + endpoint,
                     'pages': pages, 'source_total': len(rows)}
            form['tables'].append(table)
            return table

        if tax == 'PPN':
            data.update(ExportTangibleTaxableGoods={'SellingPrice': 999},
                        A1FormListOfExportTaxableGoodsOrServices={'Total': 17},
                        TotalVATandSTLGCollectedByOtherParties={'VAT': 999})
            induk = add_form('INDUK', [field('ExportOfTangibleGoods', 'SYNTHETIC EXPORT'),
                field('TotalVATandSTLGCollectedByOtherPartiesVAT', 'Jumlah PPN dan PPnBM yang dipungut',
                      'VIII. PEMUNGUTAN PPN ATAU PPN DAN PPNBM OLEH PIHAK LAIN', 'PPN (Rupiah)')])
        else:
            data['IncomeTaxes'] = [{'data': {'No': '1', 'Name': 'Article 4 Section 2',
                'TotalSelfPayment': 0, 'TotalWithholding': 17, 'TotalIncomeTaxByGovernment': 19,
                'TotalIncomeTaxPaid': 0, 'TotalIncomeTaxPaidPreviousReturn': 0, 'TotalIncomeTaxAmendment': 0},
                'expanded': True, 'children': []}]
            data.update(TotalSelfPayment=0, TotalWithholding=17, TotalIncomeTaxByGoverment=23,
                        TotalIncomeTaxPaid=0, TotalIncomeTaxPaidPreviousReturn=0,
                        TotalIncomeTaxWithDirectPaymentBPUBPNR=0, TotalIncomeTaxBPUBPNR=0,
                        TotalIncomeTaxPaidWithDirectPaymentBPUBPNR=0, TotalIncomeTaxToBePaidBPUBPNR=0,
                        TotalIncomeTaxSP=0, TotalIncomeTaxToBePaidSP=0, TotalIncomeTaxToBePaidCY=0,
                        TotalIncomeTaxAttachment=0)
            induk = add_form('INDUK', [field('FullName', 'Nama Lengkap'), field('TaxIdentificationNumber', 'NPWP')])
            sh = reference['Induk']
            headers = [sh['A8'].value, sh['B8'].value, 'PAJAK PENGHASILAN / SETOR SENDIRI',
                       'PAJAK PENGHASILAN / PEMOTONGAN / PEMUNGUTAN'] + [sh.cell(8,c).value for c in range(5,9)]
            induk['tables'].append({'name': 'income-taxes', 'headers': headers, 'group': 'B. PAJAK PENGHASILAN',
                                    'footer_spans': [[2, 1, 1, 1, 1, 1, 1]]})
        outer['FormData'] = json.dumps(data)
        primary = packet(tax, outer)
        packets.append(primary)
        induk['primary'] = {'endpoint': primary['context']['url'],
            'requests': [{'request_json': primary['context']['request_json'], 'sequences': [1]}]}
        if tax == 'PPN':
            # Synthetic observed labels use the approved template to test the DTO
            # consumer; this fixture is not a capture of the portal's labels.
            specs = [('A-1','la1','A1',3,6,[[4,1,1]],''), ('A-2','la2','A2',4,10,[[5,1,1,1,1,1]],'B3'),
                     ('B-1','lb1','B1',4,9,[[5,1,1,1,1]],''), ('B-2','lb2','B2',3,10,[[5,1,1,1,1,1]],''),
                     ('B-3','lb3','B3',3,10,[[5,1,1,1,1,1]],''), ('C','lc','C',3,14,[[9,1,1,1,1,1]],'')]
            buyer = {'Name': 'SYNTHETIC BUYER', 'TIN': '0000000000000002', 'DocumentNumber': '0000123',
                     'DocumentDate': '2025-01-02T00:00:00', 'TaxBase': 1200, 'OtherTaxBase': None,
                     'VAT': 132, 'STLG': 0, 'TaxInvoiceCode': '01'}
            for name, grid, sheet, line, width, spans, group_cell in specs:
                form = add_form(name)
                headers = [reference[sheet].cell(line,c).value or f'COLUMN {c}' for c in range(1,width+1)]
                rows = [deepcopy(buyer) for _ in range(13)] if populated and name == 'A-2' else []
                table = add_grid(form, grid, headers, reference[sheet][group_cell].value if group_cell else '', spans, rows, 10)
                if name == 'C':
                    sequence = len(packets) + 1
                    request = {'ReturnSheetRecordId': RECORD, 'IsNormalVAT': True}
                    packets.append(response(tax, {'TaxBase': 1200, 'OtherTaxBase': 0, 'VAT': 35, 'STLG': 0},
                                            'loadndvat/lc-grid/footer', request, sequence))
                    table['footer_source'] = {'endpoint': packets[-1]['context']['url'],
                        'requests': [{'request_json': request, 'sequences': [sequence]}]}
        else:
            recipient = {'TaxIdentificationNumber': '0000000000000002', 'TaxpayerName': 'SYNTHETIC PAYEE',
                'WithholdingSlipsNumber': '0000123', 'WithholdingSlipsDate': '2025-01-02T00:00:00',
                'TaxArticle': 'SYNTHETIC ARTICLE', 'TaxObjectCode': 'SYNTHETIC OBJECT CODE',
                'TaxObject': 'SYNTHETIC OBJECT', 'TaxBase': 1200, 'TaxRate': 2.5, 'IncomeTax': 30,
                'TaxCertificate': None, 'PaymentMethod': None, 'BranchId': '0000000000000002000001',
                'Status': 'SYNTHETIC STATUS', 'RevenueCode': 'SYNTHETIC REVENUE'}
            specs = [('DAFTAR-I','Daftar I', [('list-i-bpu',16,'B3',[]),('list-i-bpnr',16,'B9',[[10,1,5]]*4)]),
                     ('DAFTAR-II','Daftar II', [('list-ii-sp',11,'B3',[[6,1,4]]*2),('list-ii-cy',11,'B12',[[6,1,4]])]),
                     ('LAMPIRAN-I','Lampiran I',[('attachment-i',17,'B3',[[12,1,4]])])]
            for name, sheet, grids in specs:
                form = add_form(name, [field('TaxIdentificationNumber', 'NPWP')])
                for grid, width, group_cell, spans in grids:
                    headers = list(UNIF_HEADERS) if width == 16 else [f'COLUMN {i}' for i in range(width)]
                    rows = [deepcopy(recipient) for _ in range(6)] if populated and grid == 'list-i-bpu' else []
                    add_grid(form, grid, headers, reference[sheet][group_cell].value, spans, rows, 2)
        return {'tax': tax, 'record': selected(tax), 'scope': list(SCOPE), 'packets': packets, 'forms': forms}
    finally:
        reference.close()
