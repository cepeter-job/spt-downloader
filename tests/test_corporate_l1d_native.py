"""Source-proven L1-D packet/semantic/export integration; runners only."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

import coretax
from coretax_json import JsonSafetyError, adapt_corporate_detail
import support_json_corpus as corpus
from test_corporate_l7_native import FRAME, selected, packet as primary_packet, bindings as old_bindings
from test_corporate_l11c_native import L11C_PINS
from test_corporate_l1d_projection import source, HEADERS, SPANS

ROOT = 'L1dFormData'
DETAIL = FRAME + '?view=true'


def envelope(model='Normal', migrated=False):
    form = corpus.current_form()
    form[ROOT] = source()
    row = next(n['Data'] for n in form[ROOT]['L1dTreeData'] if n['Data']['AccountCode'] == '5322')
    row['Amount'] = 77
    if migrated:
        # Migrated totals are copied, not recalculated from the zero4500/4700.
        form[ROOT]['L1dTreeTotalRow'] = {key: 13 for key in form[ROOT]['L1dTreeTotalRow']}
    return {'IsSuccessful': True, 'Payload': {'TIN': corpus.NPWP, 'PeriodYear': 2025,
        'IsDraft': True, 'ReturnSheetStatusCode': 'SUBMITTED', 'IsMigrated': migrated,
        'BusinessClassification': 'BC04',
        'ReturnSheetModel': 'NORMAL' if model == 'Normal' else 'AMENDMENT',
        'PreviousFormData': None, 'LastSubmittedFormData': None,
        'CurrentFormData': corpus.wire(form)}}


def packet(value=None):
    item = primary_packet(envelope() if value is None else value)
    item['context']['frame_url'] = DETAIL
    item['context']['component_lifetime'] = {
        'attempt': corpus.ATTEMPT, 'kind': 'corporate-submitted-spa-v1',
        'record_id': corpus.RECORD_ID, 'source_url': corpus.LIST_URL, 'detail_url': DETAIL,
        'transition_count': 1, 'same_document': True, 'source_owner': True, 'detail_owner': True}
    return item


def bindings():
    return [old_bindings()[0], {'name': 'L1-D', 'paths': [[ROOT]], 'grids': [], 'projection': {
        'kind': 'corporate-l1d-v1', 'builds': dict(L11C_PINS), 'route': DETAIL,
        'locale': 'id-ID', 'timezone': 'Asia/Jakarta',
        'tables': [{'component': f'rsh-citr-l1d-grid{index + 1}', 'headers': list(headers),
                    'footer_spans': deepcopy(SPANS[index])} for index, headers in enumerate(HEADERS)]}}]


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
@pytest.mark.parametrize('migrated', [False, True])
def test_fresh_services_l1d_packet_reaches_canonical_and_final_saved_workbook(tmp_path, model, migrated):
    record, raw, observed = selected(model), packet(envelope(model, migrated)), bindings()
    before = deepcopy((record, raw, observed))
    result = adapt_corporate_detail(iter([raw]), record, corpus.SCOPE, observed)
    assert result['eligible'] and result['coverage']['complete'], result['reason']
    assert (record, raw, observed) == before
    section, = [s for s in result['sections'] if s['name'] == 'L1-D']
    assert [t['body_row_count'] for t in section['tables']] == [35, 26, 29]
    assert all(t['complete'] and t['source_total'] == t['body_row_count'] for t in section['tables'])
    expected_total = '13' if migrated else '0'
    assert section['tables'][0]['footer'][0][2:8] == [expected_total] * 6
    assert section['tables'][0]['footer'][0][9] == expected_total
    values, = [p['values'] for p in result['provenance']['projections'] if p['form'] == 'L1-D']
    assert any(e['response_sequence'] == raw['context']['sequence']
               and e['container'] == ['Payload', 'CurrentFormData']
               and e['path'][-1] == 'AccountCode' and e['value'] == '5322' for e in values)
    record.update(capture_complete=True, sections=result['sections'])
    run = {'npwp': corpus.NPWP, 'year': 2025, 'year_basis': 'period',
           'inventory_complete': True, 'complete': True, 'errors': [], 'records': [record]}
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    name, = [n for n in run['output_files'] if n != 'resume.xlsx']
    assert model in name
    book = load_workbook(tmp_path / name)
    try:
        assert set(book.sheetnames) == {'Induk', 'L1-D'}
        sheet = book['L1-D']
        assert sheet.max_row == 98 and sheet['A20'].value == '5322' and sheet['C20'].value == 77
        assert sheet['A37'].value == '4800' and sheet['A98'].value == '3300'
        assert sheet['C37'].value == sheet['J37'].value == (13 if migrated else 0)
        assert all(f'A{r}:B{r}' in {str(m) for m in sheet.merged_cells.ranges} for r in (2, 8, 24, 30))
        assert 'Rp' in sheet['C20'].number_format and '0.00' not in sheet['C20'].number_format
    finally:
        book.close()
    report, = [r for r in run['template_mapping'] if r['form'] == 'L1-D']
    assert report['complete'] and report['unmapped_rows'] == report['unmapped_tables'] == 0
    induk, = [r for r in run['template_mapping'] if r['form'] == 'Induk']
    assert not induk['complete']  # L1-D eligibility cannot certify another form's mapping.


@pytest.mark.parametrize('defect', ['missing-lifetime', 'vendor-pin', 'main-pin', 'corporate-pin',
    'kind', 'proof-route', 'locale', 'timezone', 'owner', 'header', 'span', 'missing-table',
    'extra-table', 'reordered-table', 'induk-after', 'missing-migration', 'migration-type',
    'missing-classification', 'other-classification', 'missing-previous', 'invalid-previous',
    'nonobject-previous', 'missing-model', 'model-mismatch', 'missing-status',
    'non-submitted', 'literal-model-mismatch', 'unsubmitted-list-label',
    'missing-5322', 'unsafe-total', 'unsupported-mask'])
def test_unproved_l1d_semantic_branch_keeps_whole_record_fallback_without_mutating_source(defect):
    value, observed, record = envelope(), bindings(), selected()
    proof = observed[1]['projection']
    if defect.endswith('-pin'):
        key = next(k for k in proof['builds'] if (
            'vendor.' in k if defect == 'vendor-pin' else 'main.' in k if defect == 'main-pin' else '/525.' in k))
        proof['builds'][key] = '0' * 64
    elif defect == 'kind':
        proof['kind'] = 'SYNTHETIC-UNKNOWN'
    elif defect == 'proof-route':
        proof['route'] = FRAME
    elif defect in ('locale', 'timezone'):
        proof[defect] = 'SYNTHETIC-UNREVIEWED'
    elif defect == 'owner':
        proof['tables'][1]['component'] = 'SYNTHETIC-FOREIGN-OWNER'
    elif defect == 'header':
        proof['tables'][0]['headers'][0] = 'SYNTHETIC-UNBOUND-HEADER'
    elif defect == 'span':
        proof['tables'][1]['footer_spans'] = [[2, 1]]
    elif defect == 'missing-table':
        proof['tables'].pop()
    elif defect == 'extra-table':
        proof['tables'].append(deepcopy(proof['tables'][0]))
    elif defect == 'reordered-table':
        proof['tables'].reverse()
    elif defect == 'induk-after':
        observed.reverse()
    elif defect.startswith('missing-') and defect != 'missing-5322' and defect != 'missing-lifetime':
        field = {'missing-migration': 'IsMigrated', 'missing-classification': 'BusinessClassification',
                 'missing-previous': 'PreviousFormData', 'missing-model': 'ReturnSheetModel',
                 'missing-status': 'ReturnSheetStatusCode'}[defect]
        del value['Payload'][field]
    elif defect == 'migration-type':
        value['Payload']['IsMigrated'] = 0
    elif defect == 'other-classification':
        value['Payload']['BusinessClassification'] = 'BC01'
    elif defect == 'model-mismatch':
        value['Payload']['ReturnSheetModel'] = 'AMENDMENT'
    elif defect == 'non-submitted':
        value['Payload']['ReturnSheetStatusCode'] = 'DRAFT'
        record['source_fields']['ReturnSheetStatusCode'] = 'DRAFT'
    elif defect == 'literal-model-mismatch':
        record['meta']['Model SPT'] = 'Amendment 002'
        value['Payload']['ReturnSheetModel'] = 'AMENDMENT'
    elif defect == 'unsubmitted-list-label':
        record['source_status'] = 'SYNTHETIC OTHER LIST'
    elif defect in ('invalid-previous', 'nonobject-previous'):
        value['Payload']['PreviousFormData'] = '{' if defect == 'invalid-previous' else '[]'
    elif defect in ('missing-5322', 'unsafe-total', 'unsupported-mask'):
        form = corpus.current_form()
        form[ROOT] = source()
        if defect == 'missing-5322':
            form[ROOT]['L1dTreeData'] = [n for n in form[ROOT]['L1dTreeData'] if n['Data']['AccountCode'] != '5322']
        elif defect == 'unsafe-total':
            form[ROOT]['L1dAssets']['AC1700'] = 1000000000000000
        else:
            row = next(n['Data'] for n in form[ROOT]['L1dTreeData'] if n['Data']['AccountCode'] == '4021')
            row['BlockFields']['NonTaxableObject'] = 'SYNTHETIC-UNSAFE'
        value['Payload']['CurrentFormData'] = corpus.wire(form)
    raw = packet(value)
    if defect == 'missing-lifetime':
        del raw['context']['component_lifetime']
    before = deepcopy((record, raw, observed))
    result = adapt_corporate_detail([raw], record, corpus.SCOPE, observed)
    assert result['eligible'] is False and result['sections'] == []
    assert not result['coverage']['complete']
    assert (record, raw, observed) == before


@pytest.mark.parametrize('observed', [False, True])
def test_l1d_lifetime_contradiction_is_fatal_before_semantic_fallback(observed):
    raw = packet()
    raw['context']['component_lifetime']['detail_owner'] = False
    with pytest.raises(JsonSafetyError):
        adapt_corporate_detail([raw], selected(), corpus.SCOPE, bindings() if observed else [])
