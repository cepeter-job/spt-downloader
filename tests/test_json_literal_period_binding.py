"""Synthetic literal inventory periods must survive native identity binding."""
from copy import deepcopy
import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from coretax_json import JsonSafetyError
from coretax_masa_json import (
    _check_period_evidence,
    adapt_masa_detail,
    bind_corporate_listing_identity,
    bind_pph21_listing_identity,
    bind_ppn_listing_identity,
    bind_unifikasi_listing_identity,
)
from test_ppn_json_listing_binding import case, seal, SCOPE


FAMILIES = [
    ('PPN', 'VAT_VAT', 'VAT_VATR', bind_ppn_listing_identity),
    ('PPh Pasal 21/26', 'ICT_WIT', 'ICT_WITR', bind_pph21_listing_identity),
    ('PPh Unifikasi', 'ICT_WT', 'ICT_WTR', bind_unifikasi_listing_identity),
    ('PPh Badan', 'ICT_RCIT', 'ICT_RCITR', bind_corporate_listing_identity),
]


def listing(family, model='Normal', literal=True):
    tax, tax_code, sheet_code, _ = family
    response, record = case()
    code = '01122026' if tax == 'PPh Badan' else '02022026'
    caption = 'Januari - Desember 2026' if tax == 'PPh Badan' else 'Februari 2026'
    response['rows'][0].update(TaxTypeCode=tax_code, ReturnSheetTypeCode=sheet_code,
                               TaxPeriodCode=code, ReturnSheetModel=model)
    record['meta'].update({'Jenis Pajak': tax, 'Masa Pajak': code if literal else caption,
                           'Model SPT': model})
    seal(response)
    return response, record


@pytest.mark.parametrize('family', FAMILIES, ids=[f[0] for f in FAMILIES])
@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
@pytest.mark.parametrize('literal', [True, False], ids=['literal-code', 'legacy-caption'])
def test_bound_listing_accepts_exact_literal_period_without_rewriting_metadata(family, model, literal):
    response, record = listing(family, model, literal)
    before = deepcopy((response, record))
    result = family[3](response, record=record, scope=SCOPE)
    assert result['eligible']
    assert result['source_fields']['TaxPeriodCode'] == response['rows'][0]['TaxPeriodCode']
    assert result['source_fields']['ReturnSheetModel'] == model
    assert (response, record) == before


@pytest.mark.parametrize('family', FAMILIES, ids=[f[0] for f in FAMILIES])
@pytest.mark.parametrize('label', ['03032026', '02022025', ' 02022026', '02022026 ',
                                   '2026', 2026, None])
def test_literal_period_mismatch_or_nonliteral_alias_is_not_repaired(family, label):
    response, record = listing(family)
    record['meta']['Masa Pajak'] = label
    before = deepcopy((response, record))
    with pytest.raises(JsonSafetyError):
        family[3](response, record=record, scope=SCOPE)
    assert (response, record) == before


@pytest.mark.parametrize('tax', ['PPN', 'PPh Pasal 21/26', 'PPh Unifikasi'])
def test_primary_period_evidence_accepts_exact_literal_selected_code(tax):
    record = {'meta': {'Jenis Pajak': tax, 'Masa Pajak': '02022026'}}
    reasons = []
    assert _check_period_evidence('02022026', record, SCOPE, reasons) == '02022026'
    assert reasons == []


@pytest.mark.parametrize('label', ['03032026', '02022025'])
def test_primary_period_evidence_rejects_different_literal_code_even_without_listing(label):
    record = {'meta': {'Masa Pajak': label}}
    with pytest.raises(JsonSafetyError):
        _check_period_evidence('02022026', record, SCOPE, [])


def test_literal_pph21_period_projects_and_saves_all_five_compacted_forms(tmp_path):
    from coretax_exporter import export_split_workbooks
    from masa_json_fixtures import SCOPE as primary_scope
    from pph21_projection_fixture import case as native_case
    from support_templates import run_for

    record, bodies, forms = native_case()
    record['meta']['Masa Pajak'] = record['source_fields']['TaxPeriodCode']
    before = deepcopy((record, bodies, forms))
    result = adapt_masa_detail(bodies, record=record, scope=primary_scope, observed_forms=forms)
    assert result['eligible'] and result['coverage']['complete'], result['reason']
    assert len(result['coverage']['grids']) == 7
    assert (record, bodies, forms) == before
    run = run_for(result['sections'])
    run['records'][0]['meta']['Masa Pajak'] = record['meta']['Masa Pajak']
    reports = []
    export_split_workbooks(run, tmp_path, mapping_reports=reports)
    assert len(reports) == 5 and all(r['complete'] for r in reports)
    path = tmp_path / 'SPT Masa 21' / 'SPT_21_01012025_Normal.xlsx'
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ['Induk', 'L-IA', 'L-IB', 'L-II', 'L-III']
        assert workbook['Induk']['D4'].value == '1'
        assert workbook['Induk']['D5'].value == '2025'
        assert workbook['Induk']['D12'].value == 39
    finally:
        workbook.close()
    assert run['records'][0]['meta']['Masa Pajak'] == '01012025'


@pytest.mark.parametrize('basis', ['period', 'submission'])
@pytest.mark.parametrize('location', ['outer-month', 'inner-month', 'outer-year', 'inner-year'])
def test_unbound_pph21_primary_rejects_literal_period_conflicts_through_actual_dispatch(basis, location):
    from masa_json_fixtures import SCOPE as primary_scope
    from test_pph21_json_primary import case as primary_case, source

    record, outer, form = primary_case()
    scope = [*primary_scope[:2], basis]
    record['scope'] = list(scope)
    record['meta']['Masa Pajak'] = '01012025'
    del record['source_fields']['TaxPeriodCode']
    if location == 'outer-month':
        outer['PeriodMonth'] = 2
    elif location == 'inner-month':
        form['TaxPeriodMonth'] = 2
    elif location == 'outer-year':
        outer['PeriodYear'] = 2024
    else:
        form['TaxPeriodYear'] = 2024
    outer['FormData'] = json.dumps(form)
    item = source(outer)
    item['context']['scope'] = list(scope)
    before = deepcopy((record, item))
    with pytest.raises(JsonSafetyError):
        adapt_masa_detail([item], record=record, scope=scope, observed_forms=[])
    assert (record, item) == before


@pytest.mark.parametrize('basis', ['period', 'submission'])
def test_consistent_unbound_pph21_literal_period_does_not_grant_native_eligibility(basis):
    from masa_json_fixtures import SCOPE as primary_scope
    from test_pph21_json_primary import case as primary_case, source

    record, outer, _ = primary_case()
    scope = [*primary_scope[:2], basis]
    record['scope'] = list(scope)
    record['meta']['Masa Pajak'] = '01012025'
    del record['source_fields']['TaxPeriodCode']
    item = source(outer)
    item['context']['scope'] = list(scope)
    result = adapt_masa_detail([item], record=record, scope=scope, observed_forms=[])
    assert not result['eligible'] and not result['coverage']['complete']


def test_hybrid_cannot_hide_unbound_pph21_literal_period_conflict_with_dom_fallback(tmp_path, monkeypatch):
    import coretax_browser as driver
    import coretax_pph21_observer as observer
    from masa_json_fixtures import SCOPE as primary_scope
    from test_pph21_json_primary import case as primary_case, source

    record, outer, _ = primary_case()
    record['meta']['Masa Pajak'] = '02022025'
    del record['source_fields']['TaxPeriodCode']
    item = source(outer)
    page = SimpleNamespace(url=item['context']['frame_url'])
    capture = SimpleNamespace(request_observations=lambda: [],
        drain=lambda **kwargs: {'complete': True, 'reasons': []},
        iter_bodies=lambda **kwargs: iter([item]))
    monkeypatch.setattr(observer, 'observe_pph21_forms', lambda *args, **kwargs: [])
    fallback = []
    def dom(*args, **kwargs):
        fallback.append(True)
        return {'sections': [], 'capture_complete': False, 'errors': [], 'attachments': []}
    monkeypatch.setattr(driver, 'capture_detail', dom)
    with pytest.raises(JsonSafetyError):
        driver._hybrid_detail(page, record, primary_scope, tmp_path, capture, 1000, None)
    assert fallback == []


@pytest.mark.parametrize('basis', ['period', 'submission'])
@pytest.mark.parametrize('conflict', ['inner-code', 'selected-year'])
def test_unbound_corporate_primary_rejects_literal_period_conflict(basis, conflict):
    from coretax_json import adapt_corporate_detail
    import support_json_corpus as corpus

    record, form = corpus.selected_record(), corpus.current_form()
    scope = [*corpus.SCOPE[:2], basis]
    record['scope'] = list(scope)
    record['meta']['Masa Pajak'] = '01122025' if conflict == 'inner-code' else '01122024'
    del record['source_fields']['TaxPeriodCode']
    if conflict == 'inner-code':
        form['TaxPeriodCode'] = '02022025'
    item = corpus.primary(form)
    item['context']['scope'] = list(scope)
    with pytest.raises(JsonSafetyError):
        adapt_corporate_detail([item], record, scope, corpus.observed())


def test_consistent_unbound_corporate_literal_period_remains_ineligible():
    from coretax_json import adapt_corporate_detail
    import support_json_corpus as corpus

    record = corpus.selected_record()
    record['meta']['Masa Pajak'] = '01122025'
    del record['source_fields']['TaxPeriodCode']
    result = adapt_corporate_detail([corpus.primary()], record, corpus.SCOPE, corpus.observed())
    assert not result['eligible'] and not result['coverage']['complete']
