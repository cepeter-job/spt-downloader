"""Dated reference labels pass through full custody and approved XLSX layouts."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from coretax_masa_json import adapt_masa_detail
from masa_json_fixtures import SCOPE, RECORD
from pph21_projection_fixture import case, REFERENCE
from support_templates import run_for
from test_dated_reference_catalogs import CLOCK
from test_pph21_template import export
from pph21_observer_fixture import browser


def dated_case():
    record, bodies, forms = case()
    expected = []
    for item in bodies:
        item['context']['clock_interval'] = {**CLOCK, 'attempt': item['context']['attempt']}
        if item['context']['url'] != REFERENCE:
            continue
        envelope = json.loads(item['body'])
        for catalog in envelope['Payload']:
            revised = []
            for entry in catalog['Details']:
                chosen = {**entry, 'CodeDescription': 'DATED ' + entry['CodeDescription'],
                          'ValidTo': '2030-01-01'}
                expected.append(chosen['CodeDescription'])
                revised.extend([{**entry, 'CodeDescription': 'UNSELECTED LABEL',
                                  'ValidTo': '2029-01-01'}, chosen])
            catalog['Details'] = revised
        item['body'] = json.dumps(envelope).encode()
        item['context'].update(bytes=len(item['body']), sha256=hashlib.sha256(item['body']).hexdigest())
    return record, bodies, forms, expected


def test_dated_labels_reach_complete_clean_saved_workbook_and_keep_source_provenance(tmp_path):
    record, bodies, forms, expected = dated_case()
    before = deepcopy((record, bodies, forms))
    result = adapt_masa_detail(bodies, record=record, scope=SCOPE, observed_forms=forms)
    assert result['eligible'], result['reason']
    assert result['coverage']['complete'] and len(result['coverage']['grids']) == 7
    assert any(binding.get('projection_rule', '').startswith('reviewed-grid-active-version:')
               for binding in result['provenance']['bindings'])
    reports = []
    workbook = export(run_for(result['sections'], key=RECORD), tmp_path, reports=reports)
    try:
        assert workbook.sheetnames == ['Induk', 'L-IA', 'L-IB', 'L-II', 'L-III']
        assert len(reports) == 5 and all(report['complete'] for report in reports)
        values = {cell.value for sheet in workbook for row in sheet for cell in row
                  if isinstance(cell.value, str)}
        assert set(expected) <= values
        assert 'UNSELECTED LABEL' not in values
        assert workbook['L-IA']['B5'].value == '0000000000000002'
        assert workbook['L-IA']['G5'].value == 1200
        assert workbook['L-IA']['G5'].data_type == 'n'
    finally:
        workbook.close()
    assert (record, bodies, forms) == before


@pytest.mark.parametrize('defect', ['missing-clock', 'cutoff-crossing', 'unbound-sequence'])
def test_dated_label_ineligibility_keeps_whole_record_fallback(defect):
    record, bodies, forms, _ = dated_case()
    if defect == 'missing-clock':
        for item in bodies:
            del item['context']['clock_interval']
    elif defect == 'cutoff-crossing':
        for item in bodies:
            if item['context']['url'] != REFERENCE:
                continue
            envelope = json.loads(item['body'])
            for catalog in envelope['Payload']:
                for entry in catalog['Details'][1::2]:
                    entry['ValidTo'] = '2025-01-01T00:00:00.500Z'
            item['body'] = json.dumps(envelope).encode()
            item['context'].update(bytes=len(item['body']), sha256=hashlib.sha256(item['body']).hexdigest())
    else:
        forms[0]['references'][0]['sequences'] = [999999]
    result = adapt_masa_detail(bodies, record=record, scope=SCOPE, observed_forms=forms)
    assert not result['eligible'] and not result['coverage']['complete']
    assert result['sections'] == []


def test_real_edge_clock_and_dated_response_reach_native_workbook_without_dom_values(
        browser, tmp_path, monkeypatch):
    import coretax_browser as driver
    import coretax_pph21_projection as projection
    from pph21_observer_fixture import MAIN, ARTICLE, BUILD, observed
    from test_pph21_json_observer import feature, assert_no_values_or_actions
    monkeypatch.setattr(projection, 'BUILD_DIGESTS', {
        MAIN: hashlib.sha256(BUILD).hexdigest(), ARTICLE: hashlib.sha256(BUILD).hexdigest()})
    feature(monkeypatch)

    def forbidden(*args, **kwargs):
        pytest.fail('Dated native source must not read DOM business values')

    monkeypatch.setattr(driver, 'capture_detail', forbidden)
    with observed(browser, tmp_path, native=True) as (fixture, capture):
        envelope = json.loads(fixture.source_packets[REFERENCE])
        for catalog in envelope['Payload']:
            catalog['Details'] = [version for entry in catalog['Details'] for version in (
                {**entry, 'CodeDescription': 'UNSELECTED LABEL', 'ValidTo': '2098-01-01'},
                {**entry, 'CodeDescription': 'DATED ' + entry['CodeDescription'], 'ValidTo': '2099-01-01'})]
        fixture.source_packets[REFERENCE] = json.dumps(envelope).encode()
        result = driver._hybrid_detail(fixture.page, fixture.record, SCOPE, tmp_path, capture, 5, None)
        assert result['capture_complete'] and result['detail_source']['used'] == 'json'
        clock = capture.drain()['clock_interval']
        assert clock['timezone'] == 'Asia/Jakarta' and clock['start_ms'] <= clock['end_ms']
        assert_no_values_or_actions(fixture)
        reports = []
        workbook = export(run_for(result['sections'], key=fixture.record['record_id']),
                          tmp_path / 'export', reports=reports)
        try:
            assert len(reports) == 5 and all(report['complete'] for report in reports)
            values = {cell.value for sheet in workbook for row in sheet for cell in row
                      if isinstance(cell.value, str)}
            assert 'DATED SYNTHETIC CERTIFICATE' in values
            assert 'UNSELECTED LABEL' not in values
            assert workbook['L-IA']['B5'].value == '0000000000000002'
            assert workbook['L-IA']['G5'].value == 1200
        finally:
            workbook.close()


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
def test_dated_unifikasi_catalogs_reach_all_forms_and_preserve_period_dropdown(tmp_path, model):
    from unifikasi_projection_fixture import case as unifikasi_case, URL
    from support_masa_native import exported
    fixture = unifikasi_case(model=model)
    # The projection-only fixture uses placeholder headings for these three
    # grids. Supply approved presentation geometry for this workbook test.
    reference = load_workbook(Path(__file__).resolve().parents[1] / 'resources' / 'unifikasi_format.xlsx')
    try:
        header_rows = {'list-ii-sp': ('Daftar II', 4), 'list-ii-cy': ('Daftar II', 13),
                       'attachment-i': ('Lampiran I', 4)}
        for form in fixture['forms']:
            for table in form['tables']:
                if table['name'] in header_rows:
                    sheet, row = header_rows[table['name']]
                    table['headers'] = [reference[sheet].cell(row, col).value
                                        for col in range(1, len(table['headers']) + 1)]
                    assert all(isinstance(label, str) for label in table['headers'])
    finally:
        reference.close()
    expected = set()
    for item in fixture['packets']:
        item['context']['clock_interval'] = {**CLOCK, 'attempt': item['context']['attempt']}
        if item['context']['url'] != URL:
            continue
        envelope = json.loads(item['body'])
        for catalog in envelope['Payload']:
            if catalog['ReferenceDataType'] == 'PERIOD':
                continue
            label = 'DATED ' + catalog['ReferenceDataType']
            expected.add(label)
            catalog['Details'] = [version for entry in catalog['Details'] for version in (
                {**entry, 'CodeDescription': 'UNSELECTED LABEL', 'ValidTo': '2029-01-01'},
                {**entry, 'CodeDescription': label, 'ValidTo': '2030-01-01'})]
        item['body'] = json.dumps(envelope).encode()
        item['context'].update(bytes=len(item['body']), sha256=hashlib.sha256(item['body']).hexdigest())
    before = deepcopy(fixture)
    result = adapt_masa_detail(fixture['packets'], record=fixture['record'], scope=fixture['scope'],
                              observed_forms=fixture['forms'])
    assert result['eligible'], result['reason']
    assert len(result['coverage']['grids']) == 5 and result['coverage']['complete']
    assert any(binding.get('projection_rule', '').startswith('reviewed-grid-active-version:')
               for binding in result['provenance']['bindings'])
    assert next(field['value'] for field in result['sections'][0]['fields']
                if field['field'] == 'TaxPeriodCode') == 'Januari 2025'
    canonical_values = {value for section in result['sections'] for table in section['tables']
                        for row in table['rows'] for value in row if isinstance(value, str)}
    assert expected <= canonical_values
    workbook = exported(tmp_path, 'PPh Unifikasi', result['sections'])
    try:
        assert len(workbook.sheetnames) == 4
        values = {cell.value for sheet in workbook for row in sheet for cell in row
                  if isinstance(cell.value, str)}
        # The approved clean layout omits the redundant payment-method column;
        # it remains in the canonical source and unchanged checkpoint above.
        assert expected - {'DATED EBUPOT_PAYMENT_METHOD'} <= values
        assert 'DATED EBUPOT_PAYMENT_METHOD' not in values
        assert 'UNSELECTED LABEL' not in values
    finally:
        workbook.close()
    assert fixture == before
