"""Synthetic financial metadata contracts; no network, profiles or live records."""
from copy import deepcopy
import hashlib
import json

import pytest
from openpyxl import load_workbook

import coretax_browser as driver
from coretax_exporter import export_split_workbooks, record_sheet_titles
import test_browser_safety as protocol
from test_split_output import fixture_run

PDF = b'%PDF-1.4\n1 0 obj\n<< /Type /Catalog >>\nendobj\ntrailer\n<< /Root 1 0 R >>\n%%EOF\n'
LABEL = 'Laporan Keuangan yang Telah Diaudit*'


def metadata(record_id, *, status='COMPLETE'):
    record_key = hashlib.sha256(record_id.encode('utf-8')).hexdigest()[:32]
    filename = hashlib.sha256(b'https://coretax.synthetic.invalid/existing/statement.pdf').hexdigest() + '.pdf'
    return {'label': LABEL, 'relative_path': f'attachments/{record_key}/{filename}',
            'status': status, 'bytes': len(PDF), 'hash': hashlib.sha256(PDF).hexdigest()}


def export_fixture(*, partial=False):
    run = fixture_run()
    run['records'] = run['records'][:2]
    for record in run['records']:
        record['scope'] = [run['npwp'], run['year'], run['year_basis']]
    record = run['records'][0]
    record['sections'][0]['fields'] += [
        {'label': 'Nomor literal', 'field': 'identifier', 'value': '0000123'},
        {'label': 'Nilai literal', 'field': 'amount', 'value': '1.234,56'},
    ]
    record['sections'][0]['tables'] = [{'name': 'SYNTHETIC DUPLICATES', 'headers': ['Kode', 'Nilai'],
                                      'rows': [['0007', '1.234,56'], ['0007', '1.234,56']], 'complete': True}]
    attachment = metadata(record['record_id'])
    if partial:
        record['capture_complete'] = False
        attachment.update(status='PARTIAL', relative_path='', bytes=0, hash='')
    record['attachments'] = [attachment]
    return run


@pytest.mark.parametrize('status', ['COMPLETE', 'PARTIAL'])
def test_collect_details_passes_selected_record_folder_and_persists_pdf_metadata(tmp_path, monkeypatch, status):
    monkeypatch.setattr(driver, 'wait_idle', lambda page: None)
    monkeypatch.setattr(protocol, 'META', {**protocol.META, 'Jenis Pajak': 'PPh Badan'})
    run = protocol.run_fixture()
    record = run['records'][0]
    item = metadata(record['record_id'])
    if status == 'PARTIAL':
        item.update(status='PARTIAL', relative_path='', bytes=0, hash='')
    seen = []

    def capture(page, **kwargs):
        seen.append(kwargs)
        assert kwargs.get('record') is record, 'Capture must receive the exact selected record'
        assert kwargs.get('folder') == tmp_path, 'Capture must receive its owned private job folder'
        return {'sections': [{'name': 'Induk', 'fields': [{'label': 'Nomor literal', 'value': '0000123'}],
                              'tables': [], 'text': 'SYNTHETIC ONLY'}],
                'attachments': [deepcopy(item)], 'capture_complete': status == 'COMPLETE',
                'errors': [] if status == 'COMPLETE' else [{'object': 'Induk/' + LABEL, 'message': 'SYNTHETIC UNSUPPORTED CONTROL'}]}

    monkeypatch.setattr(driver, 'capture_detail', capture)
    page = protocol.SourcePage()
    result = driver.collect_details(page, run, tmp_path)
    assert len(seen) == 1 and page.popup.closed
    assert result['records'][0]['attachments'] == [item]
    assert result['records'][0]['scope'] == [run['npwp'], run['year'], run['year_basis']]
    assert result['complete'] is (status == 'COMPLETE')
    saved = json.loads((tmp_path / 'returns' / '0001.json').read_text(encoding='utf-8'))
    assert saved['attachments'] == [item] and saved['record_id'] == record['record_id']
    # A verified checkpoint resumes metadata without reopening a detail or redownloading.
    if status == 'COMPLETE':
        resumed_run = protocol.run_fixture()
        untouched = protocol.SourcePage()
        resumed = driver.collect_details(untouched, resumed_run, tmp_path)
        assert resumed['records'][0]['attachments'] == [item]
        assert resumed['complete'] and untouched.navigated == [] and len(seen) == 1


def test_audit_contains_literal_relative_pdf_link_and_bound_metadata_without_extra_sheets(tmp_path):
    run = export_fixture()
    original = deepcopy(run)
    files = export_split_workbooks(run, tmp_path)
    assert run == original, 'Adding PDF links must not rewrite captured source data'
    assert set(files) == {'resume.xlsx', 'SPT_Badan.xlsx'}
    item = run['records'][0]['attachments'][0]
    for name, include_details in [('resume.xlsx', False), ('SPT_Badan.xlsx', True)]:
        workbook = load_workbook(files[name])
        try:
            if include_details:
                assert workbook.sheetnames == ['001.01 Normal Induk', '002.01 Pembetulan 1 Induk']
                proof = json.loads(workbook.properties.identifier)
                assert proof['npwp'] == run['npwp'] and proof['year'] == run['year']
                assert proof['records'][0]['record_id'] == run['records'][0]['record_id']
                assert proof['records'][0]['scope'] == [run['npwp'], run['year'], run['year_basis']]
                assert proof['records'][0]['attachments'] == [item]
                source_values = [cell.value for sheet in workbook.worksheets for row in sheet for cell in row]
                assert '0000123' in source_values and '1.234,56' in source_values
                assert source_values.count('0007') == 2
            else:
                assert workbook.sheetnames == ['Ringkasan', 'Daftar SPT', 'Audit']
                linked = [cell for row in workbook['Audit'] for cell in row
                          if cell.hyperlink and cell.hyperlink.target == item['relative_path']]
                assert len(linked) == 1, 'A declared PDF needs one relative link in the resume Audit'
                values = [cell.value for cell in workbook['Audit'][linked[0].row]]
                assert all(value in values for value in [run['records'][0]['record_id'], LABEL,
                                                         'COMPLETE', len(PDF), item['hash']])
                assert linked[0].data_type == 's'
        finally:
            workbook.close()


def test_partial_pdf_is_reported_in_audit_without_a_file_hyperlink(tmp_path):
    run = export_fixture(partial=True)
    files = export_split_workbooks(run, tmp_path)
    workbook = load_workbook(files['resume.xlsx'])
    try:
        rows = [[cell.value for cell in row] for row in workbook['Audit']]
        assert any(LABEL in row and 'PARTIAL' in row and run['records'][0]['record_id'] in row for row in rows)
        assert not any(cell.hyperlink for row in workbook['Audit'] for cell in row)
    finally:
        workbook.close()


@pytest.mark.parametrize('changes', [
    {'relative_path': '../outside.pdf'},
    {'relative_path': 'https://outside.invalid/statement.pdf'},
    {'relative_path': 'C:\\outside.pdf'},
    {'relative_path': 'attachments/other-record/' + 'a' * 64 + '.pdf'},
    {'bytes': True},
    {'bytes': 0},
    {'hash': 'not-a-sha256'},
    {'status': 'UNKNOWN'},
])
def test_invalid_financial_metadata_is_rejected_before_any_workbook_is_written(tmp_path, changes):
    run = export_fixture()
    run['records'][0]['attachments'][0].update(changes)
    with pytest.raises(ValueError, match='(?i)financial|attachment|pdf'):
        export_split_workbooks(run, tmp_path)
    assert not list(tmp_path.rglob('*.xlsx'))


def test_same_pdf_path_cannot_belong_to_two_selected_return_records(tmp_path):
    run = export_fixture()
    run['records'][1]['attachments'] = deepcopy(run['records'][0]['attachments'])
    with pytest.raises(ValueError, match='(?i)financial|attachment|pdf'):
        export_split_workbooks(run, tmp_path)
    assert not list(tmp_path.rglob('*.xlsx'))


@pytest.mark.parametrize('legacy_attachments', ['missing', 'empty'])
def test_live_resume_reinspects_legacy_corporate_checkpoint_without_pdf_evidence(tmp_path, monkeypatch, legacy_attachments):
    monkeypatch.setattr(driver, 'wait_idle', lambda page: None)
    monkeypatch.setattr(protocol, 'META', {**protocol.META, 'Jenis Pajak': 'PPh Badan'})
    run = protocol.run_fixture()
    record = run['records'][0]
    changes = {} if legacy_attachments == 'missing' else {'attachments': []}
    protocol.write_checkpoint(tmp_path, record, **changes)
    item = metadata(record['record_id'])
    inspected = []

    def capture(page, **kwargs):
        inspected.append(kwargs['record']['record_id'])
        assert kwargs['record'] is record and kwargs['folder'] == tmp_path
        return {'sections': [{'name': 'Induk', 'fields': [{'label': 'Literal', 'value': '0000123'}],
                              'tables': [], 'text': 'SYNTHETIC REINSPECTION'}],
                'attachments': [deepcopy(item)], 'capture_complete': True, 'errors': []}

    monkeypatch.setattr(driver, 'capture_detail', capture)
    page = protocol.SourcePage()
    result = driver.collect_details(page, run, tmp_path)

    assert inspected == [record['record_id']], 'A legacy COMPLETE flag does not prove financial capture'
    assert page.navigated == [record['list_url']] and page.popup.closed
    assert result['complete'] and result['records'][0]['attachments'] == [item]
    saved = json.loads((tmp_path / 'returns' / '0001.json').read_text(encoding='utf-8'))
    assert saved['attachments'] == [item] and saved['scope'] == [run['npwp'], run['year'], run['year_basis']]
