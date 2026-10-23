"""Mapping counts need their own headings without moving the meaning of links."""
from copy import deepcopy
from openpyxl import load_workbook
import coretax_exporter as exporter
from support_financial_publication import financial_run


def test_resume_mapping_header_is_separate_from_inventory_and_pdf_headers(tmp_path):
    run=financial_run(tmp_path)
    run['audit']=[{'status':'SPT Dilaporkan','total':1,'scanned':1,'selected':1,'result':'RECONCILED'}]
    run['errors']=[{'object':'SYNTHETIC OBJECT','message':'SYNTHETIC LIMITATION'}]
    run['template_mapping']=[{'file':'SPT_Badan_2025_Normal.xlsx','form':'L8','complete':False,
                             'unmapped_fields':2,'unmapped_tables':1,'unmapped_rows':9}]
    before=deepcopy(run)
    exporter.export_workbook(run,tmp_path/'resume.xlsx',include_details=False)
    book=load_workbook(tmp_path/'resume.xlsx')
    try:
        sheet=book['Audit']
        assert [cell.value for cell in sheet[4]][:6]==['Status / Objek','Total sumber','Baris dibaca','Dipilih','Hasil','Keterangan']
        headers=[cell.row for row in sheet for cell in row if cell.value=='Field belum dipetakan']
        assert len(headers)==1
        row=headers[0]
        assert [sheet.cell(row,c).value for c in range(1,7)]==[
            'Status / Objek','Field belum dipetakan','Tabel belum dipetakan','Baris belum dipetakan','Hasil','Keterangan']
        assert [sheet.cell(row+1,c).value for c in range(1,6)]==['Pemetaan template',2,1,9,'PARTIAL']
        assert sheet['E5'].value=='RECONCILED'
        assert any(cell.value=='SYNTHETIC LIMITATION' for cells in sheet for cell in cells)
        links=[cell for cells in sheet for cell in cells if cell.hyperlink]
        assert len(links)==1 and links[0].hyperlink.target==run['records'][0]['attachments'][0]['relative_path']
        assert links[0].row>row+1 and run==before
    finally:
        book.close()
