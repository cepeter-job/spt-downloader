"""Synthetic PPh21 DOM shapes; no taxpayer values or captured records."""
from copy import deepcopy
from openpyxl import load_workbook
from support_templates import COMMON_HEADERS, IA_HEADERS, section
from test_pph21_template import REFERENCE, INDUK_HEADERS, row

TOOLTIP = ' This row contain the total amount of payment before amendment.'


def induk_tables():
    book = load_workbook(REFERENCE)
    tables, expected = [], {}
    try:
        for number, (caption, first, last) in enumerate([(10, 12, 17), (19, 21, 21), (24, 26, 31), (33, 35, 35)], 1):
            values = []
            for target in range(first, last + 1):
                label = book['Induk'].cell(target, 2).value
                if target in (16, 30):
                    label += TOOLTIP
                if target in (21, 35):
                    label = f'PAJAK PENGHASILAN PASAL {21 if target == 21 else 26} DITANGGUNG PEMERINTAH'
                amount = f'{target}01,25'
                values.append([book['Induk'].cell(target, 1).value, label,
                               book['Induk'].cell(target, 3).value, amount])
                expected[target] = amount
            tables.append({'name': f'Tabel {number}', 'group': book['Induk'].cell(caption, 1).value,
                           'headers': list(INDUK_HEADERS), 'rows': values, 'body_row_count': len(values),
                           'footer': [], 'footer_spans': [], 'source_total': None, 'complete': True})
    finally:
        book.close()
    return tables, expected


def appendix_table(form, group, count, *, with_totals=False):
    headers = list(IA_HEADERS if form == 'L-IA' else COMMON_HEADERS)
    if form == 'L-II':
        headers = headers[:8] + ['MASA PEEROLEHAN PENGHASILAN'] + headers[9:11] + ['STATUS']
    elif form == 'L-III':
        headers = headers[:3] + ['JENIS PAJAK'] + headers[3:6] + ['OBJEK PAJAK'] + headers[6:9] + headers[10:]
    headers = [h + ' (dd-mm-yyyy)' if h == 'TANGGAL BUKTI PEMOTONGAN' else h for h in headers]
    base = dict(zip(IA_HEADERS if form == 'L-IA' else COMMON_HEADERS, row(ia=form == 'L-IA')))
    base.update({'MASA PEEROLEHAN PENGHASILAN': 'Januari', 'JENIS PAJAK': 'PPh Pasal 21' if group == 'BP21' else 'PPh Pasal 26',
                 'OBJEK PAJAK': 'SYNTHETIC TAX OBJECT'})
    body = []
    for index in range(count):
        values = {**base, 'NO.': str(index + 1), 'NAMA': f'SYNTHETIC {group} {index}'}
        body.append([values[h.removesuffix(' (dd-mm-yyyy)')] for h in headers])
    footer, spans, expected = [], [], {}
    if with_totals:
        book = load_workbook(REFERENCE)
        try:
            first = 16 if form in ('L-IA', 'L-IB') else 15
            total_count = 1 if form == 'L-II' else 3
            footer.append([])
            spans.append([])
            for index in range(total_count):
                label = book[form].cell(first + index, 1).value
                gross, tax = f'{index + 7}.321,45', f'-{index + 2}.010,50'
                values = ['', label, gross, '', tax, ''] if form == 'L-IA' else ['', label, gross, tax, '']
                span = {'L-IA': [5, 3, 1, 1, 1, 5], 'L-IB': [5, 3, 1, 1, 5],
                        'L-II': [5, 5, 1, 1, 4], 'L-III': [5, 7, 1, 1, 4]}[form]
                footer.append(values)
                spans.append(list(span))
                expected[first + index] = (gross, tax)
        finally:
            book.close()
    return {'name': 'Tabel 1', 'group': group, 'headers': headers, 'rows': body + deepcopy(footer),
            'body_row_count': len(body), 'footer': footer, 'footer_spans': spans,
            'source_total': len(body), 'complete': True}, expected


def observed(form, tables):
    return section(form, tables=tables, text='SYNTHETIC REDUNDANT PAGE TEXT')
