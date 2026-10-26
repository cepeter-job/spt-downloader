"""Runner-only PPh21 regressions for observed DOM group/body/footer shapes."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from support_masa_binding import render_pph21_binding as render_template_form
from support_templates import assert_format, assert_layout, cells_with, run_for
from support_pph21_dom import appendix_table, induk_tables, observed
from test_pph21_template import REFERENCE
from support_pph21_numeric import money


def render(form, tables):
    book = load_workbook(REFERENCE)
    sections = [observed(form, tables)]
    original = deepcopy(sections)
    report = render_template_form(book[form], form, sections)
    assert sections == original, 'Rendering must not modify private source custody'
    assert not cells_with(book, 'Data sumber tambahan / belum dipetakan')
    assert not cells_with(book, 'SYNTHETIC REDUNDANT PAGE TEXT')
    return book, report


def test_induk_generic_names_use_groups_labels_and_exact_tooltip_alias():
    tables, expected = induk_tables()
    book, report = render('Induk', list(reversed(tables)))
    reference = load_workbook(REFERENCE)
    try:
        assert report['complete'] is True
        for row, value in expected.items():
            assert book['Induk'].cell(row, 4).value == money(value)
            assert book['Induk'].cell(row, 4).data_type == 'n'
        assert book['Induk'].max_row == 35
        assert_layout(book['Induk'], reference['Induk'])
        for row in expected:
            assert book['Induk'].cell(row, 2).value == reference['Induk'].cell(row, 2).value
    finally:
        book.close()
        reference.close()


@pytest.mark.parametrize('defect', ['conflicting-group', 'duplicate-row', 'unknown-tooltip'])
def test_induk_ambiguity_is_reported_not_guessed_or_shifted(defect):
    tables, _ = induk_tables()
    if defect == 'conflicting-group':
        tables[0]['name'] = tables[1]['group']
    elif defect == 'duplicate-row':
        tables[0]['rows'].append(deepcopy(tables[0]['rows'][0]))
        tables[0]['rows'][-1][-1] = '999,99'
        tables[0]['body_row_count'] += 1
    else:
        tables[0]['rows'][4][1] += ' UNRECOGNIZED NOTE'
    book, report = render('Induk', tables)
    try:
        assert report['complete'] is False and report['unmapped_rows'] > 0
    finally:
        book.close()


@pytest.mark.parametrize('count,first_total', [(0, 16), (2, 16), (132, 138)])
def test_lia_reported_totals_use_g_and_i_not_condensed_c_and_e(count, first_total):
    source, totals = appendix_table('L-IA', 'LIST-IA', count, with_totals=True)
    book, report = render('L-IA', [source])
    reference = load_workbook(REFERENCE)
    try:
        sheet = book['L-IA']
        assert report['complete'] is True
        for index, (gross, tax) in enumerate(totals.values()):
            target = first_total + index
            assert sheet.cell(target, 7).value == money(gross)
            assert sheet.cell(target, 9).value == money(tax)
            assert sheet.cell(target, 3).value is None and sheet.cell(target, 5).value is None
            assert sheet.cell(target, 1).value == reference['L-IA'].cell(16 + index, 1).value
            for column in range(1, 15):
                assert_format(sheet.cell(target, column), reference['L-IA'].cell(16 + index, column))
        assert sheet.max_row == first_total + 2
        assert [sheet.cell(5 + i, 3).value for i in range(count)] == [f'SYNTHETIC LIST-IA {i}' for i in range(count)]
    finally:
        book.close()
        reference.close()


@pytest.mark.parametrize('form,groups,starts,total_start', [
    ('L-IB', ('BPA1', 'BPA2'), (6, 12), 16),
    ('L-II', ('BPA1', 'BPA2'), (5, 11), 15),
    ('L-III', ('BP21', 'BP26'), (5, 11), 15),
])
@pytest.mark.parametrize('second_count', [0, 2])
def test_optional_second_group_and_reported_totals_remain_separate(form, groups, starts, total_start, second_count):
    first, _ = appendix_table(form, groups[0], 1)
    second, totals = appendix_table(form, groups[1], second_count, with_totals=True)
    book, report = render(form, [second, first])
    reference = load_workbook(REFERENCE)
    try:
        sheet = book[form]
        assert report['complete'] is True
        assert sheet.cell(starts[0], 3).value == f'SYNTHETIC {groups[0]} 0'
        assert sheet.cell(starts[0], 7).value == 1234.5
        assert sheet.cell(starts[0], 8).value == 123.45
        assert sheet.cell(starts[0], 2).value == '0000000000000002'
        assert sheet.cell(starts[0], 2).data_type == 's'
        assert sheet.cell(starts[1], 3).value == (f'SYNTHETIC {groups[1]} 0' if second_count else None)
        for index, (gross, tax) in enumerate(totals.values()):
            assert sheet.cell(total_start + index, 7).value == money(gross)
            assert sheet.cell(total_start + index, 8).value == money(tax)
        assert sheet.max_row == reference[form].max_row
        assert_layout(sheet, reference[form])
    finally:
        book.close()
        reference.close()


def test_lia_reordered_headers_do_not_reorder_physical_footer_cells():
    source, totals = appendix_table('L-IA', 'LIST-IA', 1, with_totals=True)
    source['headers'].reverse()
    source['rows'][0].reverse()
    book, report = render('L-IA', [source])
    try:
        assert report['complete'] is True
        assert book['L-IA']['C5'].value == 'SYNTHETIC LIST-IA 0'
        assert (book['L-IA']['G16'].value, book['L-IA']['I16'].value) == tuple(map(money, totals[16]))
    finally:
        book.close()


def test_first_group_growth_preserves_second_group_and_reported_totals():
    first, _ = appendix_table('L-IB', 'BPA1', 6)
    second, totals = appendix_table('L-IB', 'BPA2', 1, with_totals=True)
    book, report = render('L-IB', [second, first])
    try:
        assert report['complete'] is True
        assert book['L-IB']['C15'].value == 'SYNTHETIC BPA2 0'
        for row, values in totals.items():
            assert tuple(book['L-IB'].cell(row + 3, c).value for c in (7, 8)) == tuple(map(money, values))
    finally:
        book.close()


@pytest.mark.parametrize('defect', ['unknown-group', 'conflicting-group', 'unknown-column', 'bad-count',
                                    'footer-tail-mismatch', 'bad-spans', 'duplicate-total', 'unknown-total', 'incomplete'])
def test_invalid_appendix_evidence_stays_partial_without_displaying_raw_block(defect):
    source, _ = appendix_table('L-IB', 'BPA2', 1, with_totals=True)
    if defect == 'unknown-group':
        source['group'] = 'SYNTHETIC UNKNOWN'
    elif defect == 'conflicting-group':
        source['name'] = 'BPA1'
    elif defect == 'unknown-column':
        source['headers'][2] = 'SYNTHETIC UNKNOWN COLUMN'
    elif defect == 'bad-count':
        source['body_row_count'] = -1
    elif defect == 'footer-tail-mismatch':
        source['rows'][-1][2] = '999,00'
    elif defect == 'bad-spans':
        source['footer_spans'][1][0] = 999
    elif defect == 'duplicate-total':
        source['footer'].append(deepcopy(source['footer'][1]))
        source['footer_spans'].append(deepcopy(source['footer_spans'][1]))
        source['rows'].append(deepcopy(source['footer'][1]))
    elif defect == 'unknown-total':
        source['footer'][1][1] = source['rows'][2][1] = 'SYNTHETIC UNKNOWN TOTAL'
    else:
        source['complete'] = False
    book, report = render('L-IB', [source])
    try:
        assert report['complete'] is False and report['unmapped_tables'] > 0
    finally:
        book.close()


def test_full_five_form_export_reports_mapping_and_preserves_scope_links(tmp_path):
    import coretax_exporter as exporter
    tables, _ = induk_tables()
    sections = [observed('Induk', tables)]
    for form, group in [('L-IA', 'LIST-IA'), ('L-IB', 'BPA2'), ('L-II', 'BPA2'), ('L-III', 'BP26')]:
        source, _ = appendix_table(form, group, 1, with_totals=True)
        sections.append(observed(form, [source]))
    run = run_for(sections)
    original = deepcopy(run)
    reports = []
    files = exporter.export_split_workbooks(run, tmp_path, mapping_reports=reports)
    assert run == original
    assert len(files) == 2 and len(reports) == 5 and all(r['complete'] for r in reports)
    workbook = load_workbook(next(p for n, p in files.items() if n != 'resume.xlsx'))
    resume = load_workbook(files['resume.xlsx'])
    try:
        assert workbook.sheetnames == ['Induk', 'L-IA', 'L-IB', 'L-II', 'L-III']
        assert exporter.read_workbook_proof(workbook)['records'][0]['record_id'] == 'synthetic-21-normal'
        assert resume['Daftar SPT']['A5'].hyperlink.target == 'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx'
        assert resume['Ringkasan']['B15'].value == 'COMPLETE'
    finally:
        workbook.close()
        resume.close()
