"""Preserve previous synthetic exports when one ID has a new literal period."""
import json

from openpyxl import load_workbook
import pytest
import coretax
from spt_desktop import output
from support_runs import fixture_run
from support_financial_publication import CHANGED_PDF, declare_pdf


def one_run(period, *, captured=True):
    run = fixture_run()
    run['records'] = run['records'][:1]
    record = run['records'][0]
    record['meta']['Masa Pajak'] = period
    record['capture_complete'] = captured
    if not captured:
        record['sections'] = []
    return run


def bytes_in(folder):
    return {p.relative_to(folder).as_posix(): p.read_bytes() for p in folder.rglob('*') if p.is_file()}


@pytest.mark.parametrize('captured', [True, False])
def test_same_id_new_literal_period_publishes_separately_without_rewriting_old_files(tmp_path, captured):
    downloads = tmp_path / 'Downloads'
    previous = one_run('Januari - Desember 2025')
    coretax.save_output(previous, tmp_path / 'old-job')
    old = output.publish_workbooks(previous, tmp_path / 'old-job', downloads, 'synthetic-old')
    original = bytes_in(old)
    new = one_run('01122025', captured=captured)
    coretax.save_output(new, tmp_path / 'new-job')
    assert output.existing_outputs(new, downloads) is None, 'A retained same-ID filename must require a separate folder'
    final = output.publish_workbooks(new, tmp_path / 'new-job', downloads, 'synthetic-new')
    assert final != old and final.parent == old.parent
    assert bytes_in(old) == original
    assert set(bytes_in(final)) == set(new['output_files'])
    coretax.verify_outputs(previous, old)
    coretax.verify_outputs(new, final)
    wb = load_workbook(final / 'resume.xlsx')
    try:
        assert wb['Ringkasan']['B8'].value == 'PARTIAL'
        assert wb['Ringkasan']['B9'].value == 1
        assert wb['Ringkasan']['B10'].value == int(captured)
        assert wb['Daftar SPT']['D5'].value == '01122025'
        assert (final / wb['Daftar SPT']['A5'].hyperlink.target).is_file()
    finally:
        wb.close()


def test_period_identity_fallback_still_validates_retained_workbook_scope(tmp_path):
    previous = one_run('Januari - Desember 2025')
    coretax.save_output(previous, tmp_path / 'old-job')
    old = output.publish_workbooks(previous, tmp_path / 'old-job', tmp_path / 'Downloads', 'synthetic-old')
    form = next(old.glob('SPT_Badan_*.xlsx'))
    wb = load_workbook(form)
    proof = json.loads(wb.properties.identifier)
    proof['npwp'] = '0000000000000002'
    wb.properties.identifier = json.dumps(proof)
    wb.save(form)
    wb.close()
    before = bytes_in(old)
    new = one_run('01122025')
    coretax.save_output(new, tmp_path / 'new-job')
    with pytest.raises(ValueError, match='scope|NPWP|proof'):
        output.publish_workbooks(new, tmp_path / 'new-job', tmp_path / 'Downloads', 'synthetic-new')
    assert bytes_in(old) == before


def test_period_identity_fallback_rejects_corrupted_retained_pdf(tmp_path):
    previous = one_run('Januari - Desember 2025')
    declare_pdf(previous['records'][0], tmp_path / 'old-job')
    coretax.save_output(previous, tmp_path / 'old-job')
    old = output.publish_workbooks(previous, tmp_path / 'old-job', tmp_path / 'Downloads', 'synthetic-old')
    pdf = old / previous['records'][0]['attachments'][0]['relative_path']
    pdf.write_bytes(CHANGED_PDF)
    before = bytes_in(old)
    new = one_run('01122025')
    coretax.save_output(new, tmp_path / 'new-job')
    with pytest.raises(ValueError, match='(?i)pdf|hash|digest|financial'):
        output.publish_workbooks(new, tmp_path / 'new-job', tmp_path / 'Downloads', 'synthetic-new')
    assert bytes_in(old) == before


def test_same_literal_period_still_requires_overwrite_consent(tmp_path):
    run = one_run('Januari - Desember 2025')
    coretax.save_output(run, tmp_path / 'job')
    old = output.publish_workbooks(run, tmp_path / 'job', tmp_path / 'Downloads', 'synthetic-old')
    before = bytes_in(old)
    with pytest.raises(FileExistsError, match='overwrite consent'):
        output.publish_workbooks(run, tmp_path / 'job', tmp_path / 'Downloads', 'synthetic-new')
    assert bytes_in(old) == before
