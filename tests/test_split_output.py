"""Synthetic split-file contract fixtures; never presented as taxpayer output."""
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import coretax
import json
from openpyxl import load_workbook


def fixture_run():
    rows = [('badan-normal', 'PPh Badan', 'Januari - Desember 2025', 'Normal'),
            ('badan-amendment', 'PPh Badan', 'Januari - Desember 2025', 'Amendment 001'),
            ('21-january', 'PPh Pasal 21/26', 'Januari 2025', 'Normal'),
            ('21-february', 'PPh Pasal 21/26', 'Februari 2025', 'Normal'),
            ('unifikasi', 'PPh Unifikasi', 'Januari 2025', 'Normal')]
    return {'npwp': '0000000000000001', 'year': 2025, 'year_basis': 'period',
            'taxpayer_name': 'SYNTHETIC TEST FIXTURE ONLY', 'complete': False,
            'started_at': 'fixture', 'audit': [], 'errors': [],
            'records': [{'key': key, 'record_id': key, 'source_status': 'Fixture',
                         'meta': {'Jenis Pajak': tax, 'Masa Pajak': period, 'Model SPT': model},
                         'capture_complete': True, 'detail_url': '',
                         'sections': [{'name': 'Induk', 'text': '', 'tables': [],
                                       'fields': [{'label': 'Record marker', 'field': 'Marker', 'value': key}]}]}
                        for key, tax, period, model in rows]}


def test_save_output_separates_resume_and_period_files(tmp_path):
    run = fixture_run()
    path = coretax.save_output(run, tmp_path)
    assert path.name == 'resume.xlsx', 'Separate resume workbook not implemented'
    assert {p.relative_to(tmp_path).as_posix() for p in tmp_path.rglob('*.xlsx')} == {'resume.xlsx', 'SPT_Badan_Januari - Desember 2025_Normal.xlsx',
        'SPT_Badan_Januari - Desember 2025_Amendment 001.xlsx',
        'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx', 'SPT Masa 21/SPT_21_Februari 2025_Normal.xlsx',
        'SPT Masa Unifikasi/SPT_Unifikasi_Januari 2025_Normal.xlsx'}
    resume = load_workbook(path)
    assert resume.sheetnames == ['Ringkasan', 'Daftar SPT', 'Audit'], 'Detailed SPT must not be embedded in resume.xlsx'
    assert resume['Daftar SPT'].max_row == 4 + len(run['records'])
    assert resume['Ringkasan']['B4'].value == run['npwp']
    assert resume['Ringkasan']['B4'].data_type == 's'
    expected = {'SPT_Badan_Januari - Desember 2025_Normal.xlsx': run['records'][:1],
        'SPT_Badan_Januari - Desember 2025_Amendment 001.xlsx': run['records'][1:2],
        'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx': run['records'][2:3],
        'SPT Masa 21/SPT_21_Februari 2025_Normal.xlsx': run['records'][3:4],
        'SPT Masa Unifikasi/SPT_Unifikasi_Januari 2025_Normal.xlsx': run['records'][4:]}
    for filename, records in expected.items():
        wb = load_workbook(tmp_path / filename)
        assert len(wb.sheetnames) == len(records), 'Observed forms of this selected model were lost or combined'
        assert {'Ringkasan', 'Daftar SPT', 'Audit'}.isdisjoint(wb.sheetnames)
        proof = json.loads(wb.properties.identifier)
        assert proof['npwp'] == run['npwp'] and proof['year'] == run['year']
        assert [r['record_id'] for r in proof['records']] == [r['record_id'] for r in records]
        assert [r['meta']['Masa Pajak'] for r in proof['records']] == [r['meta']['Masa Pajak'] for r in records]
        assert [r['meta']['Model SPT'] for r in proof['records']] == [r['meta']['Model SPT'] for r in records]
        markers = [cell.value for ws in wb.worksheets for row in ws for cell in row if cell.value in {r['key'] for r in run['records']}]
        assert markers == [r['key'] for r in records], 'Another tax type leaked into this workbook'
        assert all(r['meta']['Jenis Pajak'] == records[0]['meta']['Jenis Pajak'] for r in proof['records'])
        wb.close()
    resume.close()


def test_resume_links_to_correct_external_detail_without_generic_backlink(tmp_path):
    run = fixture_run()
    coretax.save_output(run, tmp_path)
    assert (tmp_path / 'resume.xlsx').is_file(), 'Separate resume workbook missing'
    resume = load_workbook(tmp_path / 'resume.xlsx')
    links = [resume['Daftar SPT'].cell(5+i, 1).hyperlink for i in range(len(run['records']))]
    assert [link.target for link in links] == ['SPT_Badan_Januari - Desember 2025_Normal.xlsx',
        'SPT_Badan_Januari - Desember 2025_Amendment 001.xlsx',
        'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx', 'SPT Masa 21/SPT_21_Februari 2025_Normal.xlsx',
        'SPT Masa Unifikasi/SPT_Unifikasi_Januari 2025_Normal.xlsx']
    for link in links:
        target = load_workbook(tmp_path / link.target)
        sheet_name = link.location.rsplit('!', 1)[0].strip("'")
        assert sheet_name in target.sheetnames
        caption = 'Induk'
        for directory, expected in [('SPT Masa 21/', 'Induk SPT PPh Pasal 21/26'),
                                    ('SPT Masa PPN/', 'Induk SPT PPN'),
                                    ('SPT Masa Unifikasi/', 'Induk SPT PPh Unifikasi')]:
            if link.target.startswith(directory):
                caption = expected
        assert target[sheet_name]['A1'].value == caption
        assert not any(cell.hyperlink and 'resume.xlsx' in (cell.hyperlink.target or '')
                       for row in target[sheet_name] for cell in row)
        target.close()
    resume.close()


def test_empty_scope_produces_only_resume_and_no_fake_type_files(tmp_path):
    run = fixture_run()
    run['records'] = []
    path = coretax.save_output(run, tmp_path)
    assert path.name == 'resume.xlsx'
    assert [p.name for p in tmp_path.glob('*.xlsx')] == ['resume.xlsx']
    wb = load_workbook(path)
    assert wb.sheetnames == ['Ringkasan', 'Daftar SPT', 'Audit']
    assert wb['Ringkasan']['B9'].value == 0
    wb.close()


def test_split_filename_collision_is_rejected_before_writing(tmp_path):
    import pytest
    run = fixture_run()
    run['records'] = run['records'][:2]
    run['records'][0]['meta']['Jenis Pajak'] = 'A/B'
    run['records'][1]['meta']['Jenis Pajak'] = 'A:B'
    run['records'][1]['meta']['Model SPT'] = run['records'][0]['meta']['Model SPT']
    with pytest.raises(ValueError, match='filename collision'):
        coretax.save_output(run, tmp_path)
    assert not list(tmp_path.glob('*.xlsx')), 'A colliding type must not overwrite another type file'


def test_another_npwp_cannot_overwrite_existing_output_scope(tmp_path):
    import pytest
    import json
    original = fixture_run()
    coretax.save_output(original, tmp_path)
    other = fixture_run()
    other['npwp'] = '0000000000000002'
    with pytest.raises(FileExistsError, match='scope'):
        coretax.save_output(other, tmp_path)
    saved = json.loads((tmp_path / 'run.json').read_text(encoding='utf-8'))
    assert saved['npwp'] == original['npwp']
