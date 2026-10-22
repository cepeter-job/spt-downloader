"""Approved PPN/Unifikasi layout contracts; synthetic, GitHub-runner-only."""
from collections import Counter
from copy import copy, deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import zipfile

import pytest
from openpyxl import load_workbook

import coretax
import coretax_exporter as exporter
from spt_desktop import runtime
from support_templates import assert_format, assert_layout, cells_with, section

FIXTURES = Path(__file__).parent / 'fixtures'
from support_templates import (
    TYPES, export, source_run
)

FORMS = [(tax, form) for tax, config in TYPES.items() for form in config['forms']]
# Headers/groups and slots come from visible reference captions, not native API aliases.
BLOCKS = [
    ('PPN', 'A1', 3, 4, 6, ''),
    ('PPN', 'A2', 4, 5, 15, 'B3'),
    ('PPN', 'B1', 4, 5, 5, ''),
    ('PPN', 'B2', 3, 4, 7, ''),
    ('PPN', 'B3', 3, 4, 7, ''),
    ('PPN', 'C', 3, 4, 6, ''),
    ('PPh Unifikasi', 'Daftar I', 4, 5, 8, 'B3'),
    ('PPh Unifikasi', 'Daftar I', 11, 12, 15, 'B10'),
    ('PPh Unifikasi', 'Daftar II', 4, 5, 8, 'B3'),
    ('PPh Unifikasi', 'Daftar II', 13, 14, 17, 'B12'),
    ('PPh Unifikasi', 'Lampiran I', 4, 5, 5, 'B3'),
]
BLANK_REGIONS = {
    'PPN': {'A1': [(4, 6)], 'A2': [(5, 15)], 'B1': [(5, 5)],
            'B2': [(4, 7)], 'B3': [(4, 7)], 'C': [(4, 6)]},
    'PPh Unifikasi': {'Daftar I': [(5, 8), (12, 15)],
                     'Daftar II': [(5, 8), (14, 17)], 'Lampiran I': [(5, 5)]},
}


def approved_reference(tax):
    book = load_workbook(FIXTURES / TYPES[tax]['reference'])
    if tax == 'PPN':
        book['A2']['B15'].alignment = copy(book['A2']['B14'].alignment)
    if tax == 'PPh Unifikasi':
        # Only the user-approved rate/spacer deltas change the legacy layout.
        sheet = book['Daftar I']
        for coordinate in ('J5', 'J6', 'J7', 'J8', 'J11', 'J12', 'J13'):
            sheet[coordinate].number_format = '0.00%'
        sheet['J14'].border = copy(sheet['I14'].border)
        sheet.insert_rows(9)
        # The new separator inherits the approved template's blank-row style.
        for column in range(1, 17):
            sheet.cell(9, column)._style = copy(sheet.cell(2, column)._style)
        for row in sorted(list(sheet.row_dimensions), reverse=True):
            if row >= 9:
                dimension = sheet.row_dimensions.pop(row)
                dimension.index = row + 1
                sheet.row_dimensions[row + 1] = dimension
        sheet = book['Lampiran I']
        sheet.delete_rows(6, 4)
        for row in list(sheet.row_dimensions):
            if 6 <= row <= 9:
                del sheet.row_dimensions[row]
    return book


def source_table(name, headers, rows, *, complete=True):
    return {'name': name, 'headers': headers, 'rows': rows, 'complete': complete}


def headers_for(sheet, row):
    return [cell.value for cell in sheet[row] if cell.value is not None]


def synthetic_row(headers, marker):
    values = []
    for header in headers:
        normalized = header.casefold()
        if normalized.rstrip('.') == 'no':
            value = '1'
        elif any(token in normalized for token in ('npwp', 'nik', 'identitas wp', 'nitku', 'tempat kegiatan')):
            value = '0000000000000002000001' if 'nitku' in normalized or 'tempat kegiatan' in normalized else '0000000000000002'
        elif 'tanggal' in normalized:
            value = '02-01-2025'
        elif 'kap-kjs' in normalized:
            value = '411124-100'
        elif 'kode objek' in normalized:
            value = '23-100-01'
        elif '%' in normalized:
            value = '2,00'
        elif any(token in normalized for token in ('(rp)', '(rupiah)', 'dpp', 'harga jual')):
            value = '1.234,50'
        elif 'nomor' in normalized:
            value = 'SYNTHETIC-DOCUMENT-0001'
        else:
            value = marker + ' ' + header
        values.append(value)
    return values


@pytest.mark.parametrize('tax', list(TYPES))
def test_empty_observed_templates_contain_no_reference_taxpayer_or_amount_data(tmp_path, tax):
    config = TYPES[tax]
    reference_path = FIXTURES / config['reference']
    assert hashlib.sha256(reference_path.read_bytes()).hexdigest() == config['sha256']
    with zipfile.ZipFile(reference_path) as archive:
        assert archive.testzip() is None
        assert hashlib.sha256(archive.read('xl/styles.xml')).hexdigest() == config['styles']
        assert not any(name.startswith(('docProps/', 'xl/printerSettings/')) for name in archive.namelist())
    reference = approved_reference(tax)
    workbook = export(tax, source_run(tax, [section(form) for form in config['forms']]), tmp_path)
    try:
        assert workbook.sheetnames == reference.sheetnames == config['forms']
        for sheet in reference:
            assert workbook[sheet.title]['A1'].value == sheet['A1'].value
            assert_layout(workbook[sheet.title], sheet)
        for coordinate in (['D4', 'D5', 'D6', 'D7'] if tax == 'PPN' else ['D4', 'D5']):
            assert reference['Induk'][coordinate].value is None
            assert workbook['Induk'][coordinate].value is None
        if tax == 'PPN':
            assert reference['Induk']['D48'].value is None
            assert workbook['Induk']['D48'].value is None
            amount_regions = [(12, 23, 5, 8), (28, 37, 5, 8), (41, 47, 5, 5),
                              (52, 52, 5, 6), (56, 56, 5, 5), (60, 65, 5, 5),
                              (69, 71, 5, 8), (75, 78, 5, 8)]
        else:
            amount_regions = [(10, 25, 3, 8)]
        for first, last, column, end in amount_regions:
            for sheet in [reference['Induk'], workbook['Induk']]:
                assert all(c.value is None for row in sheet.iter_rows(min_row=first, max_row=last,
                           min_col=column, max_col=end) for c in row)
        for name, regions in BLANK_REGIONS[tax].items():
            for first, last in regions:
                for sheet in [reference[name], workbook[name]]:
                    assert all(c.value is None for row in sheet.iter_rows(min_row=first, max_row=last) for c in row)
    finally:
        workbook.close()
        reference.close()


@pytest.mark.parametrize('tax,form', FORMS)
def test_each_observed_form_retains_every_cell_style_and_reference_layout(tmp_path, tax, form):
    reference = approved_reference(tax)
    workbook = export(tax, source_run(tax, [section(form)]), tmp_path)
    try:
        assert workbook.sheetnames == [form], 'Absent source forms must not be invented'
        assert workbook[form]['A1'].value == reference[form]['A1'].value
        assert_layout(workbook[form], reference[form])
        assert exporter.read_workbook_proof(workbook)['records'][0]['record_id'] == 'synthetic-template-normal'
    finally:
        workbook.close()
        reference.close()


@pytest.mark.parametrize('tax', list(TYPES))
def test_induk_identity_fields_use_the_original_slots_without_reformatting(tmp_path, tax):
    fields = [{'label': label, 'value': value} for label, value in
              [('Masa Pajak', 'Januari 2025'), ('Tahun Pajak', '2025'),
               ('Status SPT', 'Pembetulan ke-2'), ('Status Penyampaian', 'Dilaporkan')]]
    workbook = export(tax, source_run(tax, [section('Induk', fields=fields)]), tmp_path)
    reference = approved_reference(tax)
    try:
        coordinates = ['D4', 'D5', 'D6', 'D7'] if tax == 'PPN' else ['D4', 'D5']
        expected_fields = fields if tax == 'PPN' else [{'value': 'Januari'}, {'value': '2025'}]
        for coordinate, field in zip(coordinates, expected_fields):
            cell = workbook['Induk'][coordinate]
            if tax == 'PPN' and coordinate == 'D7':
                assert cell.value is None
                assert workbook['Induk']['A7'].value is None
            else:
                assert cell.value == field['value'] and cell.data_type == 's'
            assert_format(cell, reference['Induk'][coordinate])
        assert_layout(workbook['Induk'], reference['Induk'])
    finally:
        workbook.close()
        reference.close()


def test_ppn_induk_keeps_repeated_row_labels_bound_to_their_visible_section(tmp_path):
    reference = load_workbook(FIXTURES / TYPES['PPN']['reference'])
    ref = reference['Induk']
    groups = [(9, 10, 12, 'C12', '120,01'), (67, 68, 69, 'B69', '690,01'),
              (73, 74, 75, 'B75', '750,01')]
    tables = []
    for caption_row, header_row, value_row, label_cell, marker in groups:
        name = str(ref[f'A{caption_row}'].value).strip() + ' ' + str(ref[f'B{caption_row}'].value).strip()
        headers = [''] + [ref.cell(header_row, column).value for column in range(5, 9)]
        tables.append(source_table(name, headers, [[ref[label_cell].value, marker, '20,02', '30,03', '40,04']]))
    workbook = export('PPN', source_run('PPN', [section('Induk', tables=tables)]), tmp_path)
    try:
        for caption_row, header_row, value_row, label_cell, marker in groups:
            assert [workbook['Induk'].cell(value_row, col).value for col in range(5, 9)] == [float(marker.replace(',', '.')), 20.02, 30.03, 40.04]
        assert_layout(workbook['Induk'], ref)
    finally:
        workbook.close()
        reference.close()


def test_ppn_compensation_choice_is_captured_data_not_a_reference_default(tmp_path):
    run = source_run('PPN', [section('Induk', fields=[
        {'label': 'Diminta Untuk', 'value': 'SYNTHETIC REQUESTED CHOICE'}])])
    workbook = export('PPN', run, tmp_path)
    reference = load_workbook(FIXTURES / TYPES['PPN']['reference'])
    try:
        assert workbook['Induk']['D48'].value == 'SYNTHETIC REQUESTED CHOICE'
        assert_format(workbook['Induk']['D48'], reference['Induk']['D48'])
        assert_layout(workbook['Induk'], reference['Induk'])
    finally:
        workbook.close()
        reference.close()


def test_unifikasi_induk_leaf_headers_fill_both_merged_income_tax_columns(tmp_path):
    reference = load_workbook(FIXTURES / TYPES['PPh Unifikasi']['reference'])
    ref = reference['Induk']
    headers = ['NO', 'DETIL', 'SETOR SENDIRI', 'PEMOTONGAN / PEMUNGUTAN'] + [ref.cell(8, col).value for col in range(5, 9)]
    values = ['1', 'Pasal 4 Ayat 2', '101,01', '102,02', '103,03', '104,04', '105,05', '106,06']
    run = source_run('PPh Unifikasi', [section('Induk', tables=[source_table('PAJAK PENGHASILAN', headers, [values])])])
    workbook = export('PPh Unifikasi', run, tmp_path)
    try:
        assert [workbook['Induk'].cell(10, col).value for col in range(3, 9)] == [101.01, 102.02, 103.03, 104.04, 105.05, 106.06]
        assert workbook['Induk']['C8'].value == 'PAJAK PENGHASILAN'
        assert workbook['Induk']['C9'].value == 'SETOR SENDIRI'
        assert workbook['Induk']['D9'].value == 'PEMOTONGAN / PEMUNGUTAN'
        assert_layout(workbook['Induk'], ref)
    finally:
        workbook.close()
        reference.close()


@pytest.mark.parametrize('tax,form,header_row,first,last,caption_cell', BLOCKS)
def test_appendix_columns_reorder_without_losing_ids_dates_or_duplicate_rows(tmp_path, tax, form, header_row, first, last, caption_cell):
    reference = approved_reference(tax)
    ref = reference[form]
    headers = headers_for(ref, header_row)
    values = synthetic_row(headers, 'SYNTHETIC MATCHED')
    name = str(ref[caption_cell].value) if caption_cell else form
    table = source_table(name, list(reversed(headers)), [list(reversed(values)), list(reversed(values))])
    workbook = export(tax, source_run(tax, [section(form, tables=[table])]), tmp_path)
    try:
        for row_number in [first, first + 1]:
            for column, value in enumerate(values, 1):
                cell = workbook[form].cell(row_number, column)
                rate = tax == 'PPh Unifikasi' and form == 'Daftar I' and column == 10
                numeric = value == '1.234,50'
                assert cell.value == (0.02 if rate else 1234.5 if numeric else value)
                assert cell.data_type == ('n' if numeric or rate else 's')
                reference_row = row_number if row_number <= last else first
                assert_format(cell, ref.cell(reference_row, column))
        if tax == 'PPN' and form == 'B1':
            # Two duplicate source rows now grow the single reserved B1 row.
            assert workbook[form].max_row == ref.max_row + 1
            for cells in ref:
                for cell in cells:
                    target_row = cell.row + 1 if cell.row > last else cell.row
                    assert_format(workbook[form].cell(target_row, cell.column), cell)
        elif tax == 'PPh Unifikasi' and form == 'Lampiran I':
            assert workbook[form]['A8'].value == ref['A6'].value
            for column in range(1, ref.max_column + 1):
                assert_format(workbook[form].cell(6, column), ref.cell(5, column))
                assert workbook[form].cell(7, column).value is None
                assert_format(workbook[form].cell(8, column), ref.cell(6, column))
        else:
            assert_layout(workbook[form], ref)
    finally:
        workbook.close()
        reference.close()


@pytest.mark.parametrize('tax,form,header_row,first,boundary,caption_cell', [
    ('PPN', 'A1', 3, 4, 7, ''),
    ('PPh Unifikasi', 'Daftar I', 4, 5, 9, 'B3'),
])
def test_row_growth_preserves_new_row_styles_and_every_shifted_header_and_total(tmp_path, tax, form, header_row, first, boundary, caption_cell):
    reference = approved_reference(tax)
    ref = reference[form]
    headers = headers_for(ref, header_row)
    rows = [synthetic_row(headers, f'SYNTHETIC ROW {i:03d}') for i in range(11)]
    name = str(ref[caption_cell].value) if caption_cell else form
    run = source_run(tax, [section(form, tables=[source_table(name, headers, rows)])])
    workbook = export(tax, run, tmp_path)
    try:
        sheet = workbook[form]
        shift = len(rows) - (boundary - first)
        for index, values in enumerate(rows):
            for column, value in enumerate(values, 1):
                cell = sheet.cell(first + index, column)
                rate = tax == 'PPh Unifikasi' and form == 'Daftar I' and column == 10
                expected_value = 0.02 if rate else 1234.5 if value == '1.234,50' else value
                assert cell.value == expected_value
                reference_row = first + index if first + index < boundary else first
                assert_format(cell, ref.cell(reference_row, column))
        for row_number in range(boundary, ref.max_row + 1):
            for column in range(1, ref.max_column + 1):
                expected = ref.cell(row_number, column)
                actual = sheet.cell(row_number + shift, column)
                assert_format(actual, expected)
                assert actual.value == expected.value
            before = ref.row_dimensions[row_number]
            after = sheet.row_dimensions[row_number + shift]
            assert (after.height, after.hidden, after.outlineLevel) == (before.height, before.hidden, before.outlineLevel)
        for coordinate in ['A1', f'A{header_row}']:
            assert_format(sheet[coordinate], ref[coordinate])
        assert {str(r) for r in sheet.merged_cells.ranges} == {str(r) for r in ref.merged_cells.ranges}
        assert sheet.column_dimensions['B'].width == ref.column_dimensions['B'].width
    finally:
        workbook.close()
        reference.close()


@pytest.mark.parametrize('tax', list(TYPES))
def test_clean_appendices_keep_partial_proof_without_duplicate_raw_block(tmp_path, tax):
    form, first = ('A1', 4) if tax == 'PPN' else ('Daftar I', 5)
    reference = approved_reference(tax)
    ref = reference[form]
    headers = headers_for(ref, 3 if tax == 'PPN' else 4)
    values = synthetic_row(headers, 'SYNTHETIC PARTIAL')
    partial_name = form if tax == 'PPN' else str(ref['B3'].value)
    tables = [source_table(partial_name, headers, [values], complete=False),
              source_table('SYNTHETIC AMBIGUOUS TABLE', ['NIK/NPWP', 'NIK/NPWP', 'Nilai'],
                           [['0000000000000002', '0000000000000002', '9.876,50']] * 2)]
    fields = [{'label': 'SYNTHETIC UNMAPPED FIELD', 'value': '=SUM(1,2)'},
              {'label': 'SYNTHETIC UNMAPPED FIELD', 'value': '=SUM(1,2)'}]
    text = 'SYNTHETIC CONTEXT MUST REMAIN VISIBLE'
    run = source_run(tax, [section(form, fields=fields, tables=tables, text=text)])
    with pytest.raises(ValueError, match='capture completion contradicts a table completion flag'):
        exporter.validate_export_run(run)
    run['complete'] = False
    run['records'][0]['capture_complete'] = False
    workbook = export(tax, run, tmp_path)
    try:
        proof = exporter.read_workbook_proof(workbook)
        assert proof['complete'] is False
        assert proof['records'][0]['capture_complete'] is False
        assert all(workbook[form].cell(first, col).value is None for col in range(1, len(headers) + 1))
        expected = Counter(value for table in tables for row in table['rows'] for value in row)
        expected.update(field['value'] for field in fields)
        actual = Counter(cell.value for sheet in workbook for row in sheet for cell in row
                         if cell.value in expected)
        assert not actual
        assert workbook[form].max_row == ref.max_row
        assert not cells_with(workbook, text)
        assert_layout(workbook[form], ref)
    finally:
        workbook.close()
        reference.close()


@pytest.mark.parametrize('tax', list(TYPES))
def test_each_amendment_uses_independent_templates_and_exact_resume_links(tmp_path, tax):
    run = source_run(tax, [section(form, fields=[{'label': 'SYNTHETIC MARKER',
           'value': 'SYNTHETIC NORMAL ' + form}]) for form in TYPES[tax]['forms']])
    amendment = deepcopy(run['records'][0])
    amendment.update(key='synthetic-template-amendment', record_id='synthetic-template-amendment')
    amendment['meta']['Model SPT'] = 'Pembetulan 2'
    for sec in amendment['sections']:
        sec['fields'][0]['value'] = 'SYNTHETIC AMENDMENT ' + sec['name']
    run['records'].append(amendment)
    normal_file = TYPES[tax]['path']
    amendment_file = normal_file.replace('_Normal.xlsx', '_Pembetulan 2.xlsx')
    files = exporter.export_split_workbooks(run, tmp_path)
    assert set(files) == {'resume.xlsx', normal_file, amendment_file}
    coretax.verify_outputs(run, tmp_path)
    workbook = load_workbook(tmp_path / normal_file)
    amended = load_workbook(tmp_path / amendment_file)
    resume = load_workbook(tmp_path / 'resume.xlsx')
    reference = approved_reference(tax)
    try:
        forms = TYPES[tax]['forms']
        assert workbook.sheetnames == amended.sheetnames == forms
        for item in [workbook, amended]:
            assert all(sheet.sheet_state == 'visible' for sheet in item)
        proofs = [exporter.read_workbook_proof(item) for item in [workbook, amended]]
        assert [[r['meta']['Model SPT'] for r in proof['records']] for proof in proofs] == [
            ['Normal'], ['Pembetulan 2']]
        assert [[r['record_id'] for r in proof['records']] for proof in proofs] == [
            [run['records'][0]['record_id']], ['synthetic-template-amendment']]
        for item, other, model in [(workbook, amended, 'NORMAL'), (amended, workbook, 'AMENDMENT')]:
            for form in forms:
                marker = 'SYNTHETIC ' + model + ' ' + form
                assert len(cells_with(item, marker)) == 0
                assert not cells_with(other, marker)
                assert_layout(item[form], reference[form])
        links = [cell.hyperlink for sheet in resume for row in sheet for cell in row if cell.hyperlink]
        targets = {(link.target, link.location) for link in links
                   if link.target in {normal_file, amendment_file}}
        assert targets == {(normal_file, "'Induk'!A1"), (amendment_file, "'Induk'!A1")}
    finally:
        workbook.close()
        amended.close()
        resume.close()
        reference.close()


@pytest.mark.parametrize('tax', list(TYPES))
def test_copied_worker_exports_using_only_explicit_packaged_template_resources(tmp_path, tax):
    resource = TYPES[tax]['resource']
    assert resource in runtime.WORKER_FILES, 'Missing approved worker resource is intentional feature RED'
    root = Path(__file__).resolve().parent.parent
    assert (root / resource).is_file(), 'Missing approved resource is intentional feature RED'
    assert (root / resource).read_bytes() == (FIXTURES / TYPES[tax]['reference']).read_bytes()
    runtime._resources(root, tmp_path / 'runtime')
    app = tmp_path / 'runtime/app'
    assert (app / resource).read_bytes() == (root / resource).read_bytes()
    run = source_run(tax, [section('Induk')])
    script = ("import json;from pathlib import Path;import coretax,coretax_exporter as e;"
              "r=json.loads(Path('synthetic-run.json').read_text());"
              "e.export_split_workbooks(r,Path('output'));coretax.verify_outputs(r,Path('output'))")
    (app / 'synthetic-run.json').write_text(json.dumps(run), encoding='utf-8')
    subprocess.run([sys.executable, '-c', script], cwd=app, check=True, capture_output=True, timeout=60)
    workbook = load_workbook(app / 'output' / TYPES[tax]['path'])
    reference = approved_reference(tax)
    try:
        assert workbook.sheetnames == ['Induk']
        assert_layout(workbook['Induk'], reference['Induk'])
    finally:
        workbook.close()
        reference.close()
