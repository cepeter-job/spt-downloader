"""Pinned final-choice invariants; synthetic sources, GitHub-runner execution."""
from copy import deepcopy
import inspect
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

import coretax
from coretax_json import JsonSafetyError, _UnsupportedMetadata, _corporate_json
from coretax_masa_bindings import project_ppn_asked_to
from masa_projection_fixture import native_case
from test_ppn_json_amendment import amend_case, decode, reseal
from test_ppn_json_surface import CHOICES
from test_ppn_json_vat_balance import raw_source as balance_source


def raw_source(choice='FAST_REFUND'):
    return {'ReturnSheetModel': 'AMENDMENT', 'AskedTo': choice,
        'VATUnderOrOverPaidDueToAmendment': {'VAT': -1},
        'ExportTangibleTaxableGoods': {'SellingPrice': 1},
        'A1FormListOfExportTaxableGoodsOrServices': {'Total': 1},
        'TaxPeriod': '01012025', 'BookkeepingPeriodEnd': 12,
        'FacilityRegisters': ['AS.09-01'],
        'DomesticDeliveriesCollectedByVATCollector': {'SellingPrice': 0},
        'DomesticDeliveriesUncollectedVAT': {'SellingPrice': 0, 'OtherTaxBase': 0,
                                           'VAT': 0, 'SalesTaxLuxuryGoods': 0},
        'DomesticTaxableGoodsWithInvoiceCode02': {'SellingPrice': 0, 'OtherTaxBase': 0,
                                               'VAT': 0, 'SalesTaxLuxuryGoods': 0},
        'DeliveryOfTaxableGoodsSubjectToVAT': {'SellingPrice': 2},
        'DomesticDeliveriesExemptedVAT': {'SellingPrice': 0},
        'TotalDeliveryOfGoodsAndServices': {'SellingPrice': 1}}


def project(raw=None, *, model='AMENDMENT', **options):
    if options:
        assert set(options) <= set(inspect.signature(project_ppn_asked_to).parameters), \
            'Missing capture-bound PPN choice clock/outer-year contract'
    return project_ppn_asked_to(_corporate_json(json.dumps(raw_source() if raw is None else raw)),
                                model=model, **options)


@pytest.mark.parametrize('choice', [None, *CHOICES])
def test_invariant_amendment_choice_needs_both_export_stages_and_facility_provenance(choice):
    source = raw_source(choice)
    before = deepcopy(source)
    result = project(source)
    assert result['value'] == choice and source == before
    paths = {tuple(e['path']) for e in result['provenance']}
    assert {('ExportTangibleTaxableGoods', 'SellingPrice'),
            ('A1FormListOfExportTaxableGoodsOrServices', 'Total'),
            ('FacilityRegisters', 0), ('TaxPeriod',), ('BookkeepingPeriodEnd',)} <= paths


def test_final_a1_export_does_not_hide_earlier_calendar_override():
    source = raw_source()
    source['ExportTangibleTaxableGoods']['SellingPrice'] = 0
    with pytest.raises(_UnsupportedMetadata):
        project(source)


@pytest.mark.parametrize('model', ['NORMAL', 'AMENDMENT'])
def test_nonnegative_forward_refund_cannot_assume_a_distinct_post_raw_clear(model):
    source = raw_source('FORWARD_REFUND')
    source.update(json.loads(json.dumps(balance_source())))
    source['ReturnSheetModel'] = model
    source.update(BasedOnType=None, InputTaxPercentageGoods=0, InputTaxPercentageServices=0,
        OutputTaxMustBeSelfCollected={'VAT': 0}, VATIsDepositedInAdvance={'VAT': 0},
        InputTaxCanBeCalculated={'VAT': 0}, VATCollectedByVATCollector={'VAT': 0},
        VATUnderOrOverPaidDueToAmendment={'VAT': 0})
    with pytest.raises(_UnsupportedMetadata):
        project(source, model=model)


@pytest.mark.parametrize('turnover,expected', [(0, None), (1, 'FAST_REFUND'),
    (4200000000, 'FAST_REFUND'), (4200000001, None)])
def test_final_facility_validation_respects_inclusive_turnover_bounds(turnover, expected):
    source = raw_source()
    source.update(TaxPeriod='12122025', BookkeepingPeriodEnd=12, FacilityRegisters=[])
    source['TotalDeliveryOfGoodsAndServices']['SellingPrice'] = turnover
    result = project(source)
    assert result['value'] == expected
    assert any(e['path'] == ['TotalDeliveryOfGoodsAndServices', 'SellingPrice']
               and e['value'] == str(turnover) for e in result['provenance'])


@pytest.mark.parametrize('code,expected', [('AS.09-01', 'FAST_REFUND'),
    ('AS.09-02', 'FAST_REFUND'), ('AS_09_01', None), ('SYNTHETIC_OTHER', None)])
def test_only_literal_source_facility_codes_prevent_the_null_write(code, expected):
    source = raw_source()
    source.update(TaxPeriod='12122025', BookkeepingPeriodEnd=12, FacilityRegisters=[code])
    source['TotalDeliveryOfGoodsAndServices']['SellingPrice'] = 0
    assert project(source)['value'] == expected


def test_invoice_and_uncollected_quartets_can_block_calendar_without_export_or_clock():
    source = raw_source()
    source['ExportTangibleTaxableGoods']['SellingPrice'] = 0
    source['A1FormListOfExportTaxableGoodsOrServices']['Total'] = 0
    source['DomesticTaxableGoodsWithInvoiceCode02']['VAT'] = 1
    result = project(source)
    assert result['value'] == 'FAST_REFUND'
    assert any(e['path'] == ['DomesticTaxableGoodsWithInvoiceCode02', 'VAT']
               for e in result['provenance'])
    source['DomesticDeliveriesUncollectedVAT'] = {k: 1 for k in
        ('SellingPrice', 'OtherTaxBase', 'VAT', 'SalesTaxLuxuryGoods')}
    with pytest.raises(_UnsupportedMetadata):
        project(source)


def test_inactive_calendar_uses_outer_year_and_same_attempt_clock_not_replay_time():
    source = raw_source()
    source['ExportTangibleTaxableGoods']['SellingPrice'] = 0
    source['A1FormListOfExportTaxableGoodsOrServices']['Total'] = 0
    clock = {'attempt': 'a' * 32, 'start_ms': 1735689600000, 'end_ms': 1735689600001,
             'timezone': 'Asia/Jakarta'}
    result = project(source, clock_interval=clock, period_year=_corporate_json('2025'))
    assert result['value'] == 'FAST_REFUND'
    assert any(e.get('clock_interval') == clock for e in result['provenance'])
    # January 1 2026 Jakarta crosses the year threshold after bookkeeping end.
    clock.update(start_ms=1767200399999, end_ms=1767200400000)
    with pytest.raises(_UnsupportedMetadata):
        project(source, clock_interval=clock, period_year=_corporate_json('2025'))
    clock['end_ms'] = clock['start_ms'] - 1
    with pytest.raises(JsonSafetyError):
        project(source, clock_interval=clock, period_year=_corporate_json('2025'))


@pytest.mark.parametrize('choice', [None, 'FAST_REFUND', 'CARRY_FORWARD'])
@pytest.mark.parametrize('balance', [0, 55])
def test_narrow_normal_nonnegative_final_balance_clears_stale_choice(choice, balance):
    source = raw_source(choice)
    source.update(json.loads(json.dumps(balance_source())))
    source['ReturnSheetModel'] = 'NORMAL'
    source.update(BasedOnType=None, InputTaxPercentageGoods=0, InputTaxPercentageServices=0,
        OutputTaxMustBeSelfCollected={'VAT': balance + 45}, VATIsDepositedInAdvance={'VAT': 12},
        InputTaxCanBeCalculated={'VAT': 30}, VATCollectedByVATCollector={'VAT': 3},
        VATUnderOrOverPaid={'VAT': -999})
    source['DomesticDeliveriesUsingOtherValue']['VAT'] = balance + 45
    result = project(source, model='NORMAL')
    assert result['value'] is None
    paths = {tuple(e['path']) for e in result['provenance']}
    assert ('OutputTaxMustBeSelfCollected', 'VAT') in paths
    assert ('VATUnderOrOverPaidDueToAmendment', 'VAT') not in paths
    assert ('VATUnderOrOverPaid', 'VAT') not in paths


@pytest.mark.parametrize('defect', ['missing-choice', 'missing-facility', 'object-facility',
    'missing-period', 'unsafe-turnover', 'string-export', 'fractional-export'])
def test_unproved_choice_dependencies_do_not_gain_native_eligibility(defect):
    source = raw_source()
    if defect == 'missing-choice':
        del source['AskedTo']
    elif defect == 'missing-facility':
        del source['FacilityRegisters']
    elif defect == 'object-facility':
        source['FacilityRegisters'] = {}
    elif defect == 'missing-period':
        del source['TaxPeriod']
    elif defect == 'unsafe-turnover':
        source.update(FacilityRegisters=[], TaxPeriod='12122025')
        source['TotalDeliveryOfGoodsAndServices']['SellingPrice'] = 1000000000000000
    elif defect == 'string-export':
        source['ExportTangibleTaxableGoods']['SellingPrice'] = '1'
    else:
        source['ExportTangibleTaxableGoods']['SellingPrice'] = 1.5
    with pytest.raises(_UnsupportedMetadata):
        project(source)


@pytest.mark.parametrize('model,choice,cleared', [('NORMAL', 'FAST_REFUND', True),
    ('AMENDMENT', 'FAST_REFUND', True), ('AMENDMENT', 'CARRY_FORWARD', False)])
def test_complete_native_packet_choice_reaches_final_compact_saved_workbook(tmp_path, model, choice, cleared):
    case = native_case('PPN') if model == 'NORMAL' else amend_case()
    outer = json.loads(case['packets'][0]['body'])
    source = json.loads(outer['Payload']['FormData'])
    additions = raw_source(choice)
    additions['TaxPeriod'] = '01012025'
    if model == 'NORMAL':
        additions.update(json.loads(json.dumps(balance_source())))
        additions.update(ReturnSheetModel='NORMAL', BasedOnType=None,
            InputTaxPercentageGoods=0, InputTaxPercentageServices=0,
            OutputTaxMustBeSelfCollected={'VAT': 1}, VATIsDepositedInAdvance={'VAT': 0},
            InputTaxCanBeCalculated={'VAT': 0}, VATCollectedByVATCollector={'VAT': 0})
        additions['DomesticDeliveriesUsingOtherValue']['VAT'] = 1
        additions['ImportOfTaxableGoods']['VAT'] = 0
    elif cleared:
        additions.update(TaxPeriod='01012025', BookkeepingPeriodEnd=1, FacilityRegisters=[])
        additions['TotalDeliveryOfGoodsAndServices']['SellingPrice'] = 0
    source.update(additions)
    outer['Payload']['FormData'] = json.dumps(source)
    reseal(case, outer)
    case['forms'][0]['fields'] += [{'field': 'AskedTo', 'label': label,
        'group': 'III. PENGHITUNGAN PPN', 'column_label': '', 'tag': 'P-RADIOBUTTON', 'option': code}
        for code, label in CHOICES.items()]
    result = decode(case)
    assert result['eligible'] and result['coverage']['complete'], result['reason']
    selected = [f for f in result['sections'][0]['fields'] if f['field'] == 'AskedTo']
    assert len(selected) == (0 if cleared else 1)
    reference = load_workbook(Path(__file__).resolve().parents[1] / 'resources/ppn_format.xlsx')
    try:
        caption = reference['Induk']['B48'].value
    finally:
        reference.close()
    record = deepcopy(case['record'])
    record.update(capture_complete=True, sections=result['sections'])
    run = {'npwp': case['scope'][0], 'year': case['scope'][1], 'year_basis': case['scope'][2],
           'inventory_complete': True, 'complete': True, 'errors': [], 'records': [record]}
    # Use actual publication/final compaction. The binding-only export helper
    # suppresses compaction and is not final saved-export evidence.
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    name, = [n for n in run['output_files'] if n != 'resume.xlsx']
    assert record['meta']['Model SPT'] in name
    book = load_workbook(tmp_path / name)
    try:
        title, = [name for name in book.sheetnames if name.casefold() == 'induk']
        sheet = book[title]
        positions = [row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 2).value == caption]
        assert len(positions) == 1
        assert sheet.cell(positions[0], 4).value == (None if cleared else CHOICES[choice][3:])
    finally:
        book.close()
