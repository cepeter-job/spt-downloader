"""Synthetic L11-C custody, observer and saved-layout contracts; CI only."""
from copy import deepcopy
import hashlib
import importlib

import pytest
from openpyxl import load_workbook

import coretax
import coretax_browser as driver
from coretax_json import JsonSafetyError, adapt_corporate_detail
import support_json_corpus as corpus
from test_corporate_l7_native import (FRAME, MAIN, PINS, L7Fixture, selected,
                                     packet as l7_packet, bindings as l7_bindings)
from test_json_hybrid_native import browser
from test_corporate_l11c_projection import ROOT, HEADERS, SPANS, raw_source


COUNTRY = corpus.ORIGIN + '/referencedata/api/currentreferencedata/id-ID/COUNTRY_CODE'
VENDOR = corpus.ORIGIN + '/returnsheets/id-ID/vendor.a03856633dae3544.js'
VENDOR_BUILD = b'/* SYNTHETIC VENDOR -- no portal code */'
L11C_PINS = {**PINS, VENDOR: '1de6335b41d5431fd302ab7b7d00b6c593dfa6640c43ab8a88fbfbdb1b0f4b68'}


def catalog():
    return [{'ReferenceDataType': 'COUNTRY_CODE', 'language': 'id-ID', 'Details': [
        {'Code': 'SYN_COUNTRY_A', 'CodeDescription': 'SYNTHETIC COUNTRY LABEL',
         'ValidFrom': None, 'ValidTo': None}]}]


def envelope(model='Normal'):
    form = corpus.current_form()
    form[ROOT] = raw_source()
    return {'IsSuccessful': True, 'Payload': {'TIN': corpus.NPWP, 'PeriodYear': 2025,
        'IsDraft': True, 'ReturnSheetStatusCode': 'SUBMITTED',
        'ReturnSheetModel': 'NORMAL' if model == 'Normal' else 'AMENDMENT',
        'PreviousFormData': None, 'LastSubmittedFormData': None,
        'CurrentFormData': corpus.wire(form)}}


def bindings():
    return [l7_bindings()[0], {'name': 'L11-C', 'paths': [[ROOT]], 'grids': [], 'projection': {
        'kind': 'corporate-l11c-v1', 'builds': dict(L11C_PINS), 'route': FRAME,
        'locale': 'id-ID', 'timezone': 'Asia/Jakarta',
        'references': [{'endpoint': COUNTRY, 'method': 'GET', 'sequences': [2]}],
        'tables': [{'component': 'rshshr-citr-l11c-form-grid1', 'headers': list(HEADERS),
                    'footer_spans': deepcopy(SPANS)}]}}]


def packets(value=None, reference=None):
    primary = l7_packet(envelope() if value is None else value)
    other = corpus.packet({'IsSuccessful': True, 'Payload': catalog() if reference is None else reference},
                          2, endpoint=COUNTRY)
    other['context'].update(method='GET', frame_url=FRAME)
    return [primary, other]


def assert_saved(result, folder, model):
    # Keep the observed Induk as well: L11-C completeness cannot erase another
    # form's legacy evidence or certify that form's approved-layout mapping.
    coretax.save_output(result, folder)
    coretax.verify_outputs(result, folder)
    name, = [n for n in result['output_files'] if n != 'resume.xlsx']
    assert model in name
    book = load_workbook(folder / name)
    try:
        assert set(book.sheetnames) == {'Induk', 'L11-C'}
        sheet = book['L11-C']
        assert sheet.max_row == 5 and sheet.max_column == 16
        assert sheet['B3'].value == sheet['B4'].value == 'SYNTHETIC CREDITOR'
        assert sheet['D3'].value == 'SYNTHETIC COUNTRY LABEL'
        assert sheet['E3'].value == '001' and sheet['F3'].value == '00015000,00'
        assert sheet['G3'].value == 1000000 and sheet['J4'].value == 1150000
        assert sheet['K3'].value == '02-01-2025' and sheet['L4'].value == '31-12-2026'
        assert sheet['A5'].value == 'JUMLAH' and sheet['G5'].value == 2000000
        assert sheet['J5'].value == 2300000 and sheet['N5'].value == 150000
        assert sheet['O5'].value == 25000
        assert 'A5:F5' in {str(area) for area in sheet.merged_cells.ranges}
        for cell in ('G3', 'J4', 'G5', 'J5', 'N5', 'O5'):
            assert 'Rp' in sheet[cell].number_format and '0.00' not in sheet[cell].number_format
    finally:
        book.close()
    report, = [r for r in result['template_mapping'] if r['form'] == 'L11-C']
    assert report['complete'] and report['unmapped_tables'] == report['unmapped_rows'] == 0
    induk, = [r for r in result['template_mapping'] if r['form'] == 'Induk']
    assert not induk['complete'] and induk['unmapped_tables'] > 0


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
def test_l11c_packet_to_final_saved_resource_grows_rows_without_source_mutation(tmp_path, model):
    record, raw, observed = selected(model), packets(envelope(model)), bindings()
    before = deepcopy((record, raw, observed))
    result = adapt_corporate_detail(iter(raw), record, corpus.SCOPE, observed)
    assert result['eligible'] and result['coverage']['complete'], result['reason']
    assert (record, raw, observed) == before
    section, = [s for s in result['sections'] if s['name'] == 'L11-C']
    assert section['tables'][0]['source_total'] == 2
    evidence = result['provenance']['projections'][0]['values']
    assert any(e['response_sequence'] == 2 and e['path'] == [0, 'Details', 0, 'CodeDescription']
               and e['value'] == 'SYNTHETIC COUNTRY LABEL' for e in evidence)
    record.update(capture_complete=True, sections=result['sections'])
    run = {'npwp': corpus.NPWP, 'year': 2025, 'year_basis': 'period', 'inventory_complete': True,
           'complete': True, 'errors': [], 'records': [record]}
    assert_saved(run, tmp_path, model)


@pytest.mark.parametrize('defect', ['vendor-pin', 'main-pin', 'timezone', 'header', 'span',
    'history', 'unknown-owner', 'missing-reference', 'wrong-reference-sequence', 'missing-previous'])
def test_unproved_l11c_binding_refuses_native_coverage(defect):
    raw, observed = packets(), bindings()
    proof = observed[1]['projection']
    if defect in ('vendor-pin', 'main-pin'):
        proof['builds'][VENDOR if defect == 'vendor-pin' else MAIN] = '0' * 64
    elif defect == 'timezone':
        proof['timezone'] = 'UTC'
    elif defect == 'header':
        proof['tables'][0]['headers'][3] = 'SYNTHETIC UNREVIEWED HEADER'
    elif defect == 'span':
        proof['tables'][0]['footer_spans'] = [[6, *([1] * 10)]]
    elif defect == 'history':
        proof['route'] = FRAME.replace('corporate-income-tax-return/', 'corporate-income-tax-return-history/')
    elif defect == 'unknown-owner':
        proof['tables'][0]['component'] = 'SYNTHETIC-FOREIGN-GRID'
    elif defect == 'missing-reference':
        proof['references'] = []
        raw.pop()
    elif defect == 'wrong-reference-sequence':
        proof['references'][0]['sequences'] = [3]
    elif defect == 'missing-previous':
        value = envelope()
        del value['Payload']['PreviousFormData']
        raw = packets(value)
    result = adapt_corporate_detail(raw, selected(), corpus.SCOPE, observed)
    assert not result['eligible'] and result['sections'] == []


@pytest.mark.parametrize('defect', ['method', 'body', 'foreign-frame', 'authorization', 'privacy',
                                   'changed-snapshot', 'clock-mismatch'])
@pytest.mark.parametrize('bound', [False, True])
def test_country_safety_is_fatal_even_before_semantic_fallback(defect, bound):
    raw = packets()
    other = raw[1]
    if defect == 'method':
        other['context']['method'] = 'POST'
    elif defect == 'body':
        other['context']['request_json'] = {'SYNTHETIC': 'unexpected'}
    elif defect == 'foreign-frame':
        other['context']['frame_url'] = FRAME.replace(corpus.RECORD_ID, corpus.FOREIGN_ID)
    elif defect == 'authorization':
        other['context']['status'] = 403
    elif defect == 'privacy':
        raw = packets(reference=[{'AccessToken': 'SYNTHETIC PRIVATE BOUNDARY'}])
    elif defect == 'changed-snapshot':
        last = deepcopy(catalog())
        last[0]['Details'][0]['CodeDescription'] = 'SYNTHETIC CHANGED'
        changed = packets(reference=last)[1]
        changed['context']['sequence'] = 3
        changed['context']['relative_path'] = changed['context']['relative_path'].replace('0002.json', '0003.json')
        raw.append(changed)
    else:
        other['context']['clock_interval'] = {'attempt': corpus.ATTEMPT, 'start_ms': 0,
            'end_ms': 1, 'timezone': 'Asia/Jakarta'}
    with pytest.raises(JsonSafetyError):
        adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, bindings() if bound else [])


@pytest.mark.parametrize('semantic', [False, True])
def test_unused_naturally_captured_country_keeps_legacy_leaf_and_l7_compatibility(semantic):
    if semantic:
        raw, observed = [l7_packet()], l7_bindings()
    else:
        raw, observed = [l7_packet()], l7_bindings()
        del observed[1]['projection']
    raw.append(packets()[1])
    result = adapt_corporate_detail(raw, selected(), corpus.SCOPE, observed)
    assert result['eligible'] and result['coverage']['complete'], result['reason']
    assert any(e['url'] == COUNTRY for e in result['provenance']['events'])


def l11c_html():
    return ('<rshshr-citr-l11c-form><rshshr-citr-l11c-form-grid1><ui-grid-offline><p-table><table>'
        '<thead><tr><th rowspan="3">NO.</th><th colspan="3">PEMBERI PINJAMAN</th>'
        '<th colspan="2">MATA UANG</th><th colspan="4">POKOK UTANG (Rp)</th>'
        '<th colspan="2">JANGKA WAKTU PINJAMAN</th><th colspan="2">BUNGA</th>'
        '<th rowspan="3">BIAYA TERKAIT PEROLEHAN PINJAMAN SELAIN BUNGA (Rp)</th>'
        '<th rowspan="3">PERUNTUKAN PINJAMAN</th></tr><tr>'
        '<th rowspan="2">NAMA</th><th rowspan="2">ALAMAT</th><th rowspan="2">NEGARA/YURISDIKSI</th>'
        '<th rowspan="2">KODE</th><th rowspan="2">KURS AKHIR TAHUN</th>'
        '<th rowspan="2">AWAL TAHUN</th><th colspan="2">MUTASI</th><th rowspan="2">AKHIR TAHUN</th>'
        '<th rowspan="2">Tanggal Mulai</th><th rowspan="2">Tanggal Jatuh Tempo</th>'
        '<th rowspan="2">TINGKAT (%)</th><th rowspan="2">JUMLAH (Rp)</th></tr><tr>'
        '<th>PENAMBAHAN</th><th>PENGURANGAN</th></tr></thead><tbody><tr>' + '<td>999999</td>' * 16 +
        '</tr></tbody><tfoot><tr><td></td><td colspan="5">JUMLAH</td>' + '<td>999999</td>' * 10 +
        '</tr></tfoot></table></p-table></ui-grid-offline></rshshr-citr-l11c-form-grid1></rshshr-citr-l11c-form>')


class L11CFixture(L7Fixture):
    def __init__(self, browser, tmp_path, monkeypatch, *, model='Normal', popup=True,
                 trap=True, country_available=True, changed_vendor=False, frame_url=FRAME):
        new_context = browser.new_context
        monkeypatch.setattr(browser, 'new_context', lambda **options:
                            new_context(timezone_id='Asia/Jakarta', **options))
        self.country_available, self.changed_vendor = country_available, changed_vendor
        super().__init__(browser, tmp_path, monkeypatch, model=model, popup=popup,
                         trap=trap, frame_url=frame_url)
        self.extra_form = ('L11-C', 'rshshr-citr-l11c-form')
        self.raw = corpus.wire(envelope(model)).encode()
        projection = importlib.import_module('coretax_corporate_projection')
        monkeypatch.setattr(projection, 'L11C_BUILD_DIGESTS', {
            **projection.BUILD_DIGESTS, VENDOR: hashlib.sha256(VENDOR_BUILD).hexdigest()}, raising=False)

    def route(self, route):
        if route.request.url == COUNTRY:
            assert route.request.method == 'GET' and route.request.post_data is None
            self.requests.append({'url': COUNTRY, 'method': 'GET', 'body': None})
            route.fulfill(content_type='application/json', body=corpus.wire({
                'IsSuccessful': True, 'Payload': catalog()}))
        elif route.request.url == VENDOR:
            route.fulfill(content_type='application/javascript', body=VENDOR_BUILD +
                          (b'/* CHANGED */' if self.changed_vendor else b''))
        else:
            super().route(route)

    def detail_html(self):
        source = super().detail_html()
        source = source.replace('<rshshr-citr-l11c-form><label>UNSUPPORTED FIXTURE FORM</label>'
                                '</rshshr-citr-l11c-form>', l11c_html())
        source = source.replace('</head>', '<script src="' + VENDOR + '"></script></head>')
        if self.country_available:
            source = source.replace('</body>', '<script>fetch(' + repr(COUNTRY) + ')</script></body>')
        return source


@pytest.mark.parametrize('popup,model', [(True, 'Normal'), (False, 'Amendment 002')])
def test_real_native_l11c_uses_no_business_dom_reads_and_saves_approved_layout(
        browser, tmp_path, monkeypatch, popup, model):
    fixture = L11CFixture(browser, tmp_path, monkeypatch, popup=popup, model=model)
    monkeypatch.setattr(driver, 'capture_detail', lambda *args, **kwargs:
                        pytest.fail('L11-C unexpectedly entered DOM fallback'))
    try:
        result = fixture.collect()
        assert result['complete'] and not result['errors'], result['errors']
        assert result['records'][0]['detail_source']['used'] == 'json'
        assert fixture.business_reads == [] and fixture.unexpected == []
        assert len([r for r in fixture.requests if r['url'] == COUNTRY]) == 1
        assert not any('/loadcit/' in r['url'] for r in fixture.requests)
        assert_saved(result, tmp_path, model)
    finally:
        fixture.close()


@pytest.mark.parametrize('defect', ['cache-only', 'changed-vendor'])
def test_l11c_unobserved_country_or_changed_vendor_falls_back(browser, tmp_path, monkeypatch, defect):
    fixture = L11CFixture(browser, tmp_path, monkeypatch, trap=False,
                         country_available=defect != 'cache-only', changed_vendor=defect == 'changed-vendor')
    try:
        result = fixture.collect()
        assert result['records'][0]['detail_source']['used'] == 'dom'
        assert fixture.unexpected == []
    finally:
        fixture.close()
