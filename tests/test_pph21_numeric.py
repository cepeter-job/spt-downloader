"""PPh21 monetary types: numeric money, literal identity, immutable source."""
from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from coretax_pph21 import render_template_form
from support_pph21_dom import appendix_table, induk_tables, observed
from support_templates import assert_format
from test_pph21_template import REFERENCE

GROUPS = [('L-IA', 'LIST-IA', 5, 16, 9), ('L-IB', 'BPA1', 6, None, 8),
          ('L-IB', 'BPA2', 12, 16, 8), ('L-II', 'BPA1', 5, None, 8),
          ('L-II', 'BPA2', 11, 15, 8), ('L-III', 'BP21', 5, None, 8),
          ('L-III', 'BP26', 11, 15, 8)]


def assert_number(cell, source):
    assert cell.data_type == 'n' and type(cell.value) in (int, float)
    expected = Decimal(source.replace('.', '').replace(',', '.'))
    assert Decimal(str(cell.value)) == expected


@pytest.mark.parametrize('amount', ['1.234,50', '0', '-123,45', '999.999.999.999.999', '0,01',
                                    '1.234.567.890.123,45'])
def test_induk_jumlah_is_native_numeric_and_kap_remains_literal(tmp_path, amount):
    tables, targets = induk_tables()
    for table in tables:
        for row in table['rows']:
            row[2], row[3] = '000007', amount
    sections = [observed('Induk', tables)]
    original = deepcopy(sections)
    book, reference = load_workbook(REFERENCE), load_workbook(REFERENCE)
    try:
        report = render_template_form(book['Induk'], 'Induk', sections)
        assert report['complete'] is True
        for row in targets:
            assert_number(book['Induk'].cell(row, 4), amount)
            assert_format(book['Induk'].cell(row, 4), reference['Induk'].cell(row, 4))
            assert book['Induk'].cell(row, 3).value == '000007'
            assert book['Induk'].cell(row, 3).data_type == 's'
        assert sections == original
        path = tmp_path / 'numeric.xlsx'
        book.save(path)
        checked = load_workbook(path)
        try:
            for row in targets:
                assert_number(checked['Induk'].cell(row, 4), amount)
        finally:
            checked.close()
    finally:
        book.close(); reference.close()


@pytest.mark.parametrize('form,group,first,total,tax_col', GROUPS)
@pytest.mark.parametrize('amount', ['1.234,50', '0', '-123,45', '1.234.567.890.123,45'])
def test_appendix_body_and_reported_totals_numeric_not_identifiers(tmp_path, form, group, first, total, tax_col, amount):
    table, _ = appendix_table(form, group, 1, with_totals=total is not None)
    for header in ('PENGHASILAN BRUTO (Rp)', 'PAJAK PENGHASILAN (Rp)'):
        table['rows'][0][table['headers'].index(header)] = amount
    for row in table['footer']:
        if row:
            row[2] = row[4 if form == 'L-IA' else 3] = amount
    table['rows'][table['body_row_count']:] = deepcopy(table['footer'])
    table['headers'].reverse()
    table['rows'][0].reverse()
    sections = [observed(form, [table])]
    original = deepcopy(sections)
    book, reference = load_workbook(REFERENCE), load_workbook(REFERENCE)
    try:
        report = render_template_form(book[form], form, sections)
        assert report['complete'] is True
        rows = [first] + ([] if total is None else list(range(total, total + (1 if form == 'L-II' else 3))))
        for row in rows:
            for col in (7, tax_col):
                assert_number(book[form].cell(row, col), amount)
                assert_format(book[form].cell(row, col), reference[form].cell(row, col))
        for col in (2, 4, 5, 6):
            assert book[form].cell(first, col).data_type == 's'
        assert book[form].cell(first, 2).value == '0000000000000002'
        assert sections == original
        path = tmp_path / 'numeric.xlsx'
        book.save(path)
        checked = load_workbook(path)
        try:
            for row in rows:
                for col in (7, tax_col):
                    assert_number(checked[form].cell(row, col), amount)
        finally:
            checked.close()
    finally:
        book.close(); reference.close()


@pytest.mark.parametrize('amount', ['1,234.50', '12.34', '00123', '1.234.567.890.123.456', '12.345.678.901.234,56', '=1+2'])
@pytest.mark.parametrize('location', ['induk', 'body', 'footer'])
def test_unsafe_money_stays_literal_and_mapping_is_partial(amount, location):
    if location == 'induk':
        tables, _ = induk_tables()
        tables[0]['rows'][0][3] = amount
        form, address = 'Induk', 'D12'
    else:
        source, _ = appendix_table('L-IA', 'LIST-IA', 1, with_totals=True)
        if location == 'body':
            source['rows'][0][source['headers'].index('PENGHASILAN BRUTO (Rp)')] = amount
            address = 'G5'
        else:
            source['footer'][1][2] = amount
            source['rows'][source['body_row_count']:] = deepcopy(source['footer'])
            address = 'G16'
        form, tables = 'L-IA', [source]
    sections = [observed(form, tables)]
    original = deepcopy(sections)
    book = load_workbook(REFERENCE)
    try:
        report = render_template_form(book[form], form, sections)
        assert book[form][address].value == amount and book[form][address].data_type == 's'
        assert report['complete'] is False
        assert report['unmapped_rows'] + report['unmapped_fields'] > 0
        assert sections == original
    finally:
        book.close()
