"""Synthetic model-separated Excel contracts; GitHub-runner execution only."""
from copy import deepcopy
import json
from pathlib import PurePosixPath

import pytest
from openpyxl import load_workbook

import coretax
import coretax_exporter as exporter


TAX_SCOPES = [
    ('PPh Badan', 'Januari - Desember 2025', None),
    ('PPh Pasal 21/26', 'Januari 2025', 'SPT Masa 21'),
    ('PPh Unifikasi', 'Januari 2025', 'SPT Masa Unifikasi'),
    ('PPN', 'Januari 2025', 'SPT Masa PPN'),
    ('PPh Pasal 22', 'Januari 2025', None),
]
from support_runs import (
    MODELS, model_run
)

def sheet_values(workbook):
    return [cell.value for sheet in workbook for row in sheet for cell in row]


@pytest.mark.parametrize('tax,period,directory', TAX_SCOPES)
def test_grouping_never_combines_normal_or_distinct_amendment_numbers(tax, period, directory):
    run = model_run(tax, period)
    original = deepcopy(run)
    groups = exporter.workbook_groups(run)
    assert len(groups) == 3
    assert [tuple(record['meta']['Model SPT'] for record in records)
            for records in groups.values()] == [('Normal',), ('Pembetulan 1',), ('Pembetulan 2',)]
    assert len({name.casefold() for name in groups}) == 3
    for (filename, records), model in zip(groups.items(), MODELS):
        path = PurePosixPath(filename)
        assert path.suffix == '.xlsx' and model in path.stem and not path.is_absolute()
        assert path.parent == PurePosixPath(directory or '.')
        assert records[0]['meta']['Jenis Pajak'] == tax
        assert records[0]['meta']['Masa Pajak'] == period
    assert run == original


@pytest.mark.parametrize('tax,period,directory', TAX_SCOPES)
def test_each_model_file_contains_only_its_own_all_form_sheets(tmp_path, tax, period, directory):
    run = model_run(tax, period)
    original = deepcopy(run)
    files = exporter.export_split_workbooks(run, tmp_path)
    assert run == original, 'Hidden presentation extras must remain intact in source custody'
    assert len(files) == 4 and 'resume.xlsx' in files
    expected = {record['record_id']: record for record in run['records']}
    observed = set()
    for filename, path in files.items():
        if filename == 'resume.xlsx':
            continue
        workbook = load_workbook(path)
        try:
            proof = exporter.read_workbook_proof(workbook)
            assert len(proof['records']) == 1
            source = proof['records'][0]
            identifier = source['record_id']
            assert identifier in expected and identifier not in observed
            observed.add(identifier)
            assert source['meta'] == expected[identifier]['meta']
            assert source['capture_complete'] is True and proof['complete'] is True
            assert workbook.sheetnames == ['Induk', 'L-Z']
            assert {'Ringkasan', 'Daftar SPT', 'Audit'}.isdisjoint(workbook.sheetnames)
            values = sheet_values(workbook)
            markers = [value for value in values
                       if isinstance(value, str) and value.startswith('SYNTHETIC synthetic-model-return-')]
            expected_forms = ['L-Z'] if tax in ('PPN', 'PPh Unifikasi') else ['Induk', 'L-Z']
            assert markers == [f'SYNTHETIC {identifier} / {form}' for form in expected_forms]
            # Unsupported Induk rows are source-only for approved clean forms;
            # the generic L-Z still displays both duplicate rows unchanged.
            copies = 2 if tax in ('PPh Badan', 'PPN', 'PPh Pasal 21/26', 'PPh Unifikasi') else 4
            assert values.count('00007') == copies and values.count('1.234,50') == copies
            if tax == 'PPN':
                # Two approved blank spacers were removed above this section.
                amount = int(expected[identifier]['sections'][0]['fields'][-1]['value'])
                assert workbook['Induk']['E43'].value == amount
                assert values.count(amount) == 1
            if tax == 'PPh Unifikasi':
                amount = int(expected[identifier]['sections'][0]['tables'][-1]['rows'][0][2])
                assert workbook['Induk']['C10'].value == amount
                assert [v for v in values if v in (101, 102, 103)] == [amount]
            cells = [cell for sheet in workbook for row in sheet for cell in row]
            assert all(cell.data_type == 's' for cell in cells if cell.value in ('00007', '1.234,50'))
            assert all(not any((getattr(sheet, item).center.text,
                                getattr(sheet, item).left.text,
                                getattr(sheet, item).right.text))
                       for sheet in workbook for item in ('oddHeader', 'oddFooter', 'evenHeader',
                                                          'evenFooter', 'firstHeader', 'firstFooter'))
        finally:
            workbook.close()
    assert observed == set(expected)


@pytest.mark.parametrize('tax,period,directory', TAX_SCOPES)
def test_resume_links_to_the_correct_model_file_and_first_form(tmp_path, tax, period, directory):
    run = model_run(tax, period)
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    resume = load_workbook(tmp_path / 'resume.xlsx')
    targets = []
    try:
        assert resume.sheetnames == ['Ringkasan', 'Daftar SPT', 'Audit']
        assert resume['Ringkasan']['B7'].value == 3
        for ordinal, expected in enumerate(run['records'], 5):
            link = resume['Daftar SPT'].cell(ordinal, 1).hyperlink
            assert link and link.location == "'Induk'!A1"
            assert resume['Daftar SPT'].cell(ordinal, 5).value == expected['meta']['Model SPT']
            targets.append(link.target)
            workbook = load_workbook(tmp_path / link.target)
            try:
                proof = exporter.read_workbook_proof(workbook)
                assert [record['record_id'] for record in proof['records']] == [expected['record_id']]
                values = [cell.value for row in workbook['Induk'] for cell in row]
                marker = f"SYNTHETIC {expected['record_id']} / Induk"
                if tax == 'PPN':
                    assert marker not in values
                    assert workbook['Induk']['E43'].value == int(expected['sections'][0]['fields'][-1]['value'])
                elif tax == 'PPh Unifikasi':
                    assert marker not in values
                    assert workbook['Induk']['C10'].value == int(expected['sections'][0]['tables'][-1]['rows'][0][2])
                else:
                    assert marker in values
            finally:
                workbook.close()
        assert len(set(targets)) == 3
    finally:
        resume.close()


@pytest.mark.parametrize('tax,period,directory', TAX_SCOPES)
def test_submission_year_does_not_merge_distinct_source_periods_of_one_model(tax, period, directory):
    run = model_run(tax, period, models=('Normal', 'Normal'))
    run['year_basis'] = 'submission'
    run['records'][0]['meta'].update({'Masa Pajak': 'Desember 2024', 'Tanggal Submit SPT': '10-01-2025'})
    run['records'][1]['meta'].update({'Masa Pajak': 'Januari 2025', 'Tanggal Submit SPT': '10-02-2025'})
    groups = exporter.workbook_groups(run)
    assert len(groups) == 2
    assert [[record['record_id'] for record in records] for records in groups.values()] == [
        ['synthetic-model-return-1'], ['synthetic-model-return-2']]
    assert [records[0]['meta']['Masa Pajak'] for records in groups.values()] == [
        'Desember 2024', 'Januari 2025']


@pytest.mark.parametrize('models', [
    ('Normal', 'normal'),
    ('Pembetulan/1', 'Pembetulan:1'),
    ('Pembetulan?2', 'Pembetulan*2'),
    ('Pembetulan 3.', 'Pembetulan 3'),
])
def test_distinct_models_that_collide_on_windows_fail_before_any_output(tmp_path, models):
    run = model_run(models=models)
    with pytest.raises(ValueError, match='collision'):
        coretax.save_output(run, tmp_path)
    assert not list(tmp_path.rglob('*.xlsx'))


def test_literal_amendment_text_and_leading_zeroes_are_not_coalesced(tmp_path):
    models = ('Amendment 001', 'Pembetulan 001', 'Pembetulan 1')
    run = model_run(models=models)
    files = exporter.export_split_workbooks(run, tmp_path)
    assert len(files) == 4
    actual = []
    for filename, path in files.items():
        if filename == 'resume.xlsx':
            continue
        workbook = load_workbook(path)
        try:
            proof = json.loads(workbook.properties.identifier)
            assert len(proof['records']) == 1
            actual.append(proof['records'][0]['meta']['Model SPT'])
        finally:
            workbook.close()
    assert actual == list(models)


def test_a_single_selected_amendment_does_not_create_normal_or_other_amendments(tmp_path):
    run = model_run(models=('Pembetulan 2',))
    files = exporter.export_split_workbooks(run, tmp_path)
    assert len(files) == 2
    filenames = set(files) - {'resume.xlsx'}
    assert len(filenames) == 1
    assert 'Pembetulan 2' in next(iter(filenames))
    coretax.verify_outputs(run, tmp_path)


def test_partial_capture_remains_partial_in_each_model_proof_and_resume(tmp_path):
    run = model_run()
    run['complete'] = False
    run['records'][1]['capture_complete'] = False
    run['records'][1]['sections'][0]['tables'][0]['complete'] = False
    run['errors'] = [{'object': 'synthetic-model-return-2', 'message': 'Synthetic incomplete capture'}]
    files = exporter.export_split_workbooks(run, tmp_path)
    assert len(files) == 4
    for filename, path in files.items():
        workbook = load_workbook(path)
        try:
            if filename == 'resume.xlsx':
                assert workbook['Ringkasan']['B6'].value == 'PARTIAL'
                assert workbook['Daftar SPT']['I6'].value == 'PARTIAL'
            else:
                proof = exporter.read_workbook_proof(workbook)
                assert len(proof['records']) == 1 and proof['complete'] is False
                record = proof['records'][0]
                assert record['capture_complete'] is (record['meta']['Model SPT'] != 'Pembetulan 1')
        finally:
            workbook.close()
