"""PPN-only completion contracts: synthetic values, GitHub-runner execution."""
from copy import copy, deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook

import coretax_exporter as exporter
from support_templates import assert_format, section

from support_templates import APPENDICES, RAW_MARKER, RESOURCE, export_ppn, source_appendix

# form, header, first body row, original total row, monetary columns
@pytest.mark.parametrize('case', APPENDICES, ids=[x[0] for x in APPENDICES])
@pytest.mark.parametrize('count', [0, 17])
def test_appendix_footer_reaches_the_styled_total_after_row_growth(tmp_path, case, count):
    form, header, first, total, columns = case
    table, amounts = source_appendix(*case, count=count)
    workbook, reports, _ = export_ppn(tmp_path, [section(form, tables=[table])])
    reference = load_workbook(RESOURCE)
    reference['A2']['B15'].alignment = copy(reference['A2']['B14'].alignment)
    try:
        sheet = workbook[form]
        target = total + max(0, count - (total - first))
        assert sheet.cell(target, 1).value == 'TOTAL'
        assert [sheet.cell(target, c).value for c in columns] == [0, -17, 1234.5, 29][:len(columns)]
        for c in columns:
            assert_format(sheet.cell(target, c), reference[form].cell(total, c))
            assert sheet.cell(target, c).data_type == 'n'
        assert [sheet.cell(first + i, 2).value for i in range(count)] == [
            f'SYNTHETIC-{i}-2' for i in range(count)]
        for row in range(first, first + count):
            # Existing body rows retain their distinct bottom borders/fonts;
            # only inserted rows copy the original first body row.
            style_row = row if row < total else first
            assert_format(sheet.cell(row, 2), reference[form].cell(style_row, 2))
        assert sheet.max_row == target
        assert RAW_MARKER not in [cell.value for row in sheet for cell in row]
        assert all(report['complete'] for report in reports)
        if form == 'C':
            assert sheet.cell(target, 9).value is None, 'Column I is an invoice ID, not money'
    finally:
        workbook.close()
        reference.close()


@pytest.mark.parametrize('bad_spans', [[True, 1, 1], [0, 1, 1], [4, 2], [3, 1, 1], [4, 1, 1, 1]])
def test_invalid_footer_spans_never_shift_amount_into_wrong_cell(tmp_path, bad_spans):
    table, _ = source_appendix(*APPENDICES[0])
    table['footer_spans'] = [bad_spans]
    workbook, reports, _ = export_ppn(tmp_path, [section('A1', tables=[table])])
    try:
        assert workbook['A1']['E7'].value is None
        assert not all(report['complete'] for report in reports)
        assert workbook['A1'].max_row == 7
    finally:
        workbook.close()


def test_footer_conflict_is_not_arbitrarily_selected_or_added(tmp_path):
    first, _ = source_appendix(*APPENDICES[0])
    second = deepcopy(first)
    second['footer'][0][1] = '999'
    workbook, reports, _ = export_ppn(tmp_path, [section('A1', tables=[first, second])])
    try:
        assert workbook['A1']['E7'].value is None
        assert not all(report['complete'] for report in reports)
        assert workbook['A1'].max_row == 7
    finally:
        workbook.close()


def test_footer_must_match_captured_row_suffix(tmp_path):
    table, _ = source_appendix(*APPENDICES[0])
    table['rows'] = [['TOTAL', '19', '']]
    workbook, reports, _ = export_ppn(tmp_path, [section('A1', tables=[table])])
    try:
        assert workbook['A1']['E7'].value is None
        assert not all(report['complete'] for report in reports)
    finally:
        workbook.close()
