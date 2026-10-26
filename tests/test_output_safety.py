"""Synthetic regressions for verified, recoverable output publication."""
import json
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import coretax
from support_runs import fixture_run


def snapshot(folder):
    return {p.relative_to(folder).as_posix(): p.read_bytes() for p in folder.rglob('*')
            if p.is_file() and not p.relative_to(folder).parts[0].startswith('.spt-export-')}


@pytest.mark.parametrize('relative_output', [
    None, 'relative-result', 'synthetic-capture/../parent-relative-result',
])
def test_offline_cli_exports_to_relative_output_folder(tmp_path, relative_output):
    run = fixture_run()
    capture = tmp_path / 'synthetic-capture' / 'capture.json'
    capture.parent.mkdir()
    capture.write_text(json.dumps(run), encoding='utf-8')
    args = [sys.executable, str(Path(coretax.__file__).resolve()),
            '--from-json', str(capture.relative_to(tmp_path))]
    folder = capture.parent
    if relative_output is not None:
        args.extend(['--output', relative_output])
        folder = tmp_path / relative_output

    result = subprocess.run(args, cwd=tmp_path, capture_output=True, text=True, timeout=30)

    assert result.returncode == 2, result.stderr
    saved_run = json.loads((folder / 'run.json').read_text(encoding='utf-8'))
    coretax.verify_outputs(saved_run, folder)
    assert len(list(folder.rglob('*.xlsx'))) == 6
    assert not list(folder.glob('.spt-export-*'))


@pytest.mark.parametrize('offline', [False, True])
@pytest.mark.parametrize('which', ['manifest', 'body'])
def test_cli_keeps_changed_selected_json_before_replacing_any_output(tmp_path, offline, which):
    from test_json_output_folder import captured_run, evidence
    source = tmp_path / 'private-json-job'
    run = captured_run(source)
    coretax.save_output(run, source)
    destination = tmp_path / 'offline-json-result' if offline else source
    options = {'json_source': source} if offline else {}
    if offline:
        coretax.save_output(run, destination, **options)
    expected = evidence(run, source)
    assert all((destination / name).exists() for name in expected)
    name = next(name for name in expected if name.endswith('manifest.json') == (which == 'manifest'))
    changed = b'{"SYNTHETIC":"KEEP MODIFIED SELECTED EVIDENCE"}'
    (destination / name).write_bytes(changed)
    before = snapshot(destination)
    with pytest.raises((ValueError, RuntimeError)):
        coretax.save_output(run, destination, **options)
    assert snapshot(destination) == before
    assert (destination / name).read_bytes() == changed


def test_resume_accepts_reused_and_new_json_attempts_in_one_selection(tmp_path):
    from copy import deepcopy
    from test_json_output_folder import captured_run, evidence
    from spt_desktop import output
    source, downloads = tmp_path / 'private-json-job', tmp_path / 'Downloads'
    first = captured_run(source)
    first['complete'] = False
    coretax.save_output(first, source)
    final = output.publish_workbooks(first, source, downloads, 'partial-json')
    old = evidence(first, source)
    second = captured_run(source, identifier='fixture-extra-record')
    run = deepcopy(first)
    run.update(complete=True, records=first['records'] + second['records'])
    coretax.save_output(run, source)
    incoming = evidence(second, source)
    assert all(not (final / name).exists() for name in incoming)
    consent = output.existing_outputs(run, downloads)
    assert consent['target'] == str(final)
    assert set(consent['json']) == set(old)
    assert all((final / name).read_bytes() == body for name, body in old.items())
    if sys.platform == 'win32':
        assert output.publish_workbooks(run, source, downloads, 'resumed-json',
            overwrite=True, expected_outputs=consent) == final
        assert all((final / name).read_bytes() == body for name, body in {**old, **incoming}.items())
        coretax.verify_outputs(run, final)


def test_relative_output_root_reparse_point_cannot_hide_during_normalization(tmp_path, monkeypatch):
    run = fixture_run()
    existing = tmp_path / 'existing-results'
    coretax.save_output(run, existing)
    before = snapshot(existing)
    alias = tmp_path / 'redirected-results'
    if sys.platform == 'win32':
        subprocess.run(['cmd', '/d', '/c', 'mklink', '/J', str(alias), str(existing)],
                       check=True, capture_output=True, text=True, timeout=10)
    else:
        alias.symlink_to(existing, target_is_directory=True)
    monkeypatch.chdir(tmp_path)
    run['taxpayer_name'] = 'SYNTHETIC REPLACEMENT'

    with pytest.raises(ValueError, match='symbolic|reparse'):
        coretax.save_output(run, Path(alias.name))

    assert snapshot(existing) == before


def test_verification_remains_enabled_under_python_optimization(tmp_path):
    run = fixture_run()
    path = coretax.save_output(run, tmp_path)
    wb = load_workbook(path)
    wb['Ringkasan']['B2'] = '0000000000000002'
    wb.save(path)
    wb.close()
    script = (
        'import json,sys,coretax; from pathlib import Path; '
        'folder=Path(sys.argv[1]); '
        'run=json.loads((folder/"run.json").read_text(encoding="utf-8")); '
        'coretax.verify_workbook(folder/"resume.xlsx",run,include_details=False)'
    )
    result = subprocess.run([sys.executable, '-O', '-c', script, str(tmp_path)],
                            cwd=Path(coretax.__file__).parent, capture_output=True, text=True)
    assert result.returncode != 0, 'Optimized Python accepted another taxpayer workbook'
    assert 'NPWP mismatch' in result.stderr


def test_verification_failure_preserves_previous_output(tmp_path, monkeypatch):
    run = fixture_run()
    coretax.save_output(run, tmp_path)
    before = snapshot(tmp_path)
    run['taxpayer_name'] = 'SYNTHETIC REPLACEMENT'
    def fail_verification(*args, **kwargs):
        raise ValueError('Synthetic invalid generated workbook')
    monkeypatch.setattr(coretax, 'verify_workbook', fail_verification)
    with pytest.raises(ValueError, match='Synthetic invalid'):
        coretax.save_output(run, tmp_path)
    assert snapshot(tmp_path) == before


def test_untracked_workbook_cannot_be_overwritten(tmp_path):
    existing = tmp_path / 'resume.xlsx'
    existing.write_bytes(b'SYNTHETIC UNTRACKED DOCUMENT')
    with pytest.raises(FileExistsError, match='scope'):
        coretax.save_output(fixture_run(), tmp_path)
    assert existing.read_bytes() == b'SYNTHETIC UNTRACKED DOCUMENT'


def test_case_insensitive_existing_filename_collision_is_rejected(tmp_path):
    run = fixture_run()
    coretax.save_output(run, tmp_path)
    existing = tmp_path / 'SPT_Badan_Januari - Desember 2025_Normal.xlsx'
    existing.rename(tmp_path / 'spt_badan_januari - desember 2025_normal.XLSX')
    before = snapshot(tmp_path)
    with pytest.raises(FileExistsError, match='collision'):
        coretax.save_output(run, tmp_path)
    assert snapshot(tmp_path) == before


def test_reduced_type_scope_does_not_leave_stale_workbooks(tmp_path):
    run = fixture_run()
    coretax.save_output(run, tmp_path)
    before = snapshot(tmp_path)
    run['records'] = run['records'][:2]
    with pytest.raises(FileExistsError, match='scope'):
        coretax.save_output(run, tmp_path)
    assert snapshot(tmp_path) == before


def test_failed_replacement_rolls_back_completed_replacements(tmp_path, monkeypatch):
    run = fixture_run()
    coretax.save_output(run, tmp_path)
    before = snapshot(tmp_path)
    run['taxpayer_name'] = 'SYNTHETIC NEW NAME'
    real_replace = Path.replace
    failed = False
    def fail_one_replace(source, target):
        nonlocal failed
        target = Path(target)
        if target == tmp_path/'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx' and not failed:
            failed = True
            raise PermissionError('Synthetic workbook lock')
        return real_replace(source, target)
    monkeypatch.setattr(Path, 'replace', fail_one_replace)
    with pytest.raises(PermissionError, match='Synthetic workbook lock'):
        coretax.save_output(run, tmp_path)
    assert failed
    assert snapshot(tmp_path) == before


@pytest.mark.parametrize('url', [
    'http://user:secret@localhost:9235', 'http://localhost:9235/path',
    'http://localhost:9235?proxy=elsewhere', 'http://localhost:9235#secret',
    'http://localhost:99999', 'http://localhost:0', 'http://localhost:',
    'http://localhost:9235\n', 'http://127.0.0.1.evil.invalid:9235',
])
def test_debugging_endpoint_rejects_ambiguous_or_credential_urls(url):
    validator = getattr(coretax, 'validate_debug_endpoint', None)
    assert callable(validator), 'Explicit debugging endpoint validation is missing'
    with pytest.raises(ValueError):
        validator(url)


@pytest.mark.parametrize('url', [
    'http://127.0.0.1:9235', 'http://localhost:9235/',
    'http://[::1]:9235', 'https://localhost:9235',
])
def test_debugging_endpoint_accepts_explicit_loopback_origins(url):
    validator = getattr(coretax, 'validate_debug_endpoint', None)
    assert callable(validator), 'Explicit debugging endpoint validation is missing'
    assert validator(url) == url


def test_checkpoint_manifest_cannot_authorize_overwriting_other_taxpayer(tmp_path):
    run = fixture_run()
    coretax.save_output(run, tmp_path)
    before = {p.name: p.read_bytes() for p in tmp_path.glob('*.xlsx')}
    other = fixture_run()
    other['npwp'] = '0000000000000002'
    # Capture checkpoints update run.json before Excel is regenerated.
    (tmp_path / 'run.json').write_text(json.dumps(other), encoding='utf-8')
    with pytest.raises(FileExistsError, match='scope'):
        coretax.save_output(other, tmp_path)
    assert {p.name: p.read_bytes() for p in tmp_path.glob('*.xlsx')} == before


def test_resume_link_must_identify_exact_amendment_not_just_existing_sheet(tmp_path):
    run = fixture_run()
    coretax.save_output(run, tmp_path)
    path = tmp_path / 'resume.xlsx'
    wb = load_workbook(path)
    wb['Daftar SPT']['A6'].hyperlink.target = wb['Daftar SPT']['A5'].hyperlink.target
    wb['Daftar SPT']['A6'].hyperlink.location = wb['Daftar SPT']['A5'].hyperlink.location
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match='type-workbook link'):
        coretax.verify_outputs(run, tmp_path)


def test_interrupt_immediately_after_replace_rolls_back_that_file(tmp_path, monkeypatch):
    run = fixture_run()
    coretax.save_output(run, tmp_path)
    before = snapshot(tmp_path)
    run['taxpayer_name'] = 'SYNTHETIC UPDATED NAME'
    real_replace = Path.replace
    interrupted = False
    def interrupt_after_replace(source, target):
        nonlocal interrupted
        result = real_replace(source, target)
        target = Path(target)
        if target == tmp_path/'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx' and not interrupted:
            interrupted = True
            raise KeyboardInterrupt()
        return result
    monkeypatch.setattr(Path, 'replace', interrupt_after_replace)
    with pytest.raises(KeyboardInterrupt):
        coretax.save_output(run, tmp_path)
    assert snapshot(tmp_path) == before


def test_blocked_rollback_keeps_recoverable_original_files(tmp_path, monkeypatch):
    run = fixture_run()
    coretax.save_output(run, tmp_path)
    before = snapshot(tmp_path)
    run['taxpayer_name'] = 'SYNTHETIC UPDATED NAME'
    real_replace = Path.replace
    def fail_publication_and_rollback(source, target):
        target = Path(target)
        if target == tmp_path/'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx' or source.parent.name == 'backup':
            raise PermissionError('Synthetic persistent file lock')
        return real_replace(source, target)
    monkeypatch.setattr(Path, 'replace', fail_publication_and_rollback)
    with pytest.raises(RuntimeError, match='recovery files retained'):
        coretax.save_output(run, tmp_path)
    recovery, = tmp_path.glob('.spt-export-*')
    assert snapshot(recovery/'backup') == before


def test_interrupt_in_rollback_probe_keeps_recoverable_original_files(tmp_path, monkeypatch):
    run = fixture_run()
    coretax.save_output(run, tmp_path)
    before = snapshot(tmp_path)
    run['taxpayer_name'] = 'SYNTHETIC UPDATED NAME'
    real_replace = Path.replace
    real_exists = Path.exists
    publication_failed = False

    def fail_later_publication(source, target):
        nonlocal publication_failed
        if Path(target) == tmp_path/'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx':
            publication_failed = True
            raise PermissionError('Synthetic workbook lock')
        return real_replace(source, target)

    def interrupt_recovery_probe(path):
        if publication_failed and path.parent.name.startswith('.spt-export-'):
            raise KeyboardInterrupt()
        return real_exists(path)

    monkeypatch.setattr(Path, 'replace', fail_later_publication)
    monkeypatch.setattr(Path, 'exists', interrupt_recovery_probe)
    with pytest.raises(RuntimeError, match='recovery files retained'):
        coretax.save_output(run, tmp_path)
    recovery, = tmp_path.glob('.spt-export-*')
    assert snapshot(recovery/'backup') == before
