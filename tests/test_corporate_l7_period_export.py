"""Literal-period L7 exports through saved files; synthetic data only."""
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook

import coretax
import coretax_exporter as exporter
from support_templates import section, source_run


RESOURCE = Path(__file__).resolve().parents[1] / 'resources/corporate_format.xlsx'


def l7_run(period, *, year=2025, header_year=None, captured=True):
    reference = load_workbook(RESOURCE)
    try:
        headers = [reference['L7'].cell(2, column).value for column in range(1, 10)]
    finally:
        reference.close()
    headers[7] = (f'KOMPENSASI KERUGIAN FISKAL / {header_year or year} / '
                  'TAHUN PAJAK INI - NILAI (Rp)')
    rows = [[f'{i:04d}', str(year - 10 + i), '1.000', '0', '0', '0', '0', '0', '0']
            for i in range(1, 11)]
    footer = [['JUMLAH', '0', '0', '0', '0', '0', '0']]
    table = {'name': 'SYNTHETIC L7', 'group': 'PENGHITUNGAN KOMPENSASI KERUGIAN FISKAL',
             'headers': headers, 'rows': rows + footer, 'body_row_count': 10,
             'source_total': 10, 'complete': True, 'footer': footer,
             'footer_spans': [[3, 1, 1, 1, 1, 1, 1]]}
    run = source_run('PPh Badan', [section('L7', tables=[table])])
    run['year'] = year
    run['records'][0]['meta']['Masa Pajak'] = period
    run['records'][0]['capture_complete'] = captured
    return run


@pytest.mark.parametrize('year,model,captured', [
    (2024, 'Normal', True), (2025, 'Amendment 001', True),
    (2025, 'Normal', False),
])
def test_literal_annual_l7_exports_body_total_and_preserves_source(tmp_path, year, model, captured):
    period = f'0112{year}'
    run = l7_run(period, year=year, captured=captured)
    run['records'][0]['meta']['Model SPT'] = model
    before = deepcopy(run)
    reports = []
    files = exporter.export_split_workbooks(run, tmp_path, mapping_reports=reports)
    name = f'SPT_Badan_{period}_{model}.xlsx'
    assert name in files
    book = load_workbook(tmp_path / name)
    resume = load_workbook(tmp_path / 'resume.xlsx')
    try:
        sheet = book['L7']
        assert sheet['H2'].value == run['records'][0]['sections'][0]['tables'][0]['headers'][7]
        assert sheet.max_row == 13 and sheet.max_column == 9
        assert [sheet.cell(row, 1).value for row in range(3, 13)] == [f'{i:04d}' for i in range(1, 11)]
        assert [sheet.cell(row, 2).value for row in range(3, 13)] == [str(year - 10 + i) for i in range(1, 11)]
        assert sheet['C3'].value == 1000 and sheet['H12'].value == 0
        assert sheet['A13'].value == 'JUMLAH'
        assert [sheet.cell(13, c).value for c in range(4, 10)] == [0] * 6
        assert 'A13:C13' in sheet.merged_cells
        assert 'Rp' in sheet['H12'].number_format and '0.00' not in sheet['H12'].number_format
        assert reports == [{'file': name, 'form': 'L7', 'complete': True,
                            'unmapped_fields': 0, 'unmapped_tables': 0, 'unmapped_rows': 0}]
        assert resume['Daftar SPT']['D5'].value == period
        assert resume['Daftar SPT']['E5'].value == model
        assert resume['Daftar SPT']['I5'].value == ('OK' if captured else 'PARTIAL')
        assert period in book.properties.identifier
        assert run == before
        coretax.verify_outputs(run, tmp_path)
    finally:
        book.close()
        resume.close()


@pytest.mark.parametrize('period,header_year', [
    ('01122025', 2024), ('01012025', 2025), ('02122025', 2025),
    ('01112025', 2025), ('01122025 ', 2025), (' 01122025', 2025),
    ('0112202', 2025), ('011220250', 2025), ('０１１２２０２５', 2025),
])
def test_unproved_l7_period_or_header_stays_unmapped(tmp_path, period, header_year):
    run = l7_run(period, header_year=header_year)
    before = deepcopy(run)
    reports = []
    files = exporter.export_split_workbooks(run, tmp_path, mapping_reports=reports)
    name, = [name for name in files if name != 'resume.xlsx']
    book = load_workbook(tmp_path / name)
    try:
        assert book['L7']['H2'].value is None
        assert book['L7']['C3'].value is None
        assert reports[0]['unmapped_tables'] == 1
        assert reports[0]['unmapped_rows'] == 11
        assert reports[0]['complete'] is False
        assert run == before
    finally:
        book.close()
