"""Synthetic coverage for ordered overflow and final saved L7 row growth."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

import coretax
from coretax_json import _UnsupportedMetadata, _corporate_json, adapt_corporate_detail
import support_json_corpus as corpus
from test_corporate_l7_projection import BODY, ROOT, SUFFIXES, raw_source, project
from test_corporate_l7_native import selected, envelope, packet, bindings


def test_l7_rejects_intermediate_overflow_even_when_final_reported_total_is_safe():
    raw = raw_source()
    raw[BODY].append(deepcopy(raw[BODY][0]))
    raw[BODY][-1]['TaxYear'] = '2026'
    for row, value in zip(raw[BODY], ('999999999999999', '1', '-1')):
        row['FiscalLossCompensationYearMinus4'] = corpus.Number(value)
    for suffix in SUFFIXES:
        raw['TotalFiscalLossCompensation' + suffix] = corpus.Number(str(sum(
            int(row['FiscalLossCompensation' + suffix]) for row in raw[BODY])))
    source = _corporate_json(corpus.wire(raw))
    before = deepcopy(source)
    with pytest.raises(_UnsupportedMetadata, match='unsafe_corporate_l7_ordered_total'):
        project(source)
    assert source == before


def test_l7_twelve_packet_rows_grow_saved_layout_and_preserve_captured_footer(tmp_path):
    outer = envelope()
    current = corpus.current_form()
    current['MainFormData']['NetIncomeBeforeTax'] = corpus.Number('1000')
    source = raw_source()
    first, last = source[BODY]
    source[BODY] = [dict(deepcopy(first), TaxYear=str(year)) for year in range(2000, 2011)]
    source[BODY].append(deepcopy(last))
    for suffix in SUFFIXES:
        source['TotalFiscalLossCompensation' + suffix] = corpus.Number(str(sum(
            int(row['FiscalLossCompensation' + suffix]) for row in source[BODY])))
    current[ROOT] = source
    outer['Payload']['CurrentFormData'] = corpus.wire(current)
    record, response, observed = selected(), packet(outer), bindings()
    before = deepcopy((record, response, observed))
    adapted = adapt_corporate_detail(iter([response]), record, corpus.SCOPE, observed)
    assert adapted['eligible'] and adapted['coverage']['complete'], adapted['reason']
    section, = [s for s in adapted['sections'] if s['name'] == 'L7']
    assert section['tables'][0]['body_row_count'] == section['tables'][0]['source_total'] == 12
    assert (record, response, observed) == before
    record.update(capture_complete=True, sections=[section])
    run = {'npwp': corpus.NPWP, 'year': 2025, 'year_basis': 'period',
           'inventory_complete': True, 'complete': True, 'errors': [], 'records': [record]}
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    path, = [tmp_path / name for name in run['output_files'] if name != 'resume.xlsx']
    workbook = load_workbook(path)
    try:
        sheet = workbook['L7']
        assert sheet.max_row == 15
        assert sheet['B3'].value == '2000' and sheet['B14'].value == '2025'
        assert sheet['C14'].value == 1000
        assert sheet['A15'].value == 'JUMLAH'
        assert sheet['D15'].value == 1200 and sheet['I15'].value == 7200
        assert 'A15:C15' in {str(area) for area in sheet.merged_cells.ranges}
    finally:
        workbook.close()
    report, = [r for r in run['template_mapping'] if r['form'] == 'L7']
    assert report['complete'] and report['unmapped_rows'] == report['unmapped_tables'] == 0
