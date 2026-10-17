"""Synthetic templates helpers; no collected tests or shared mutable test state."""
from copy import copy
from openpyxl import load_workbook
import coretax
import coretax_exporter as exporter
from copy import deepcopy
from pathlib import Path


COMMON_HEADERS = ['NO.', 'NIK/NPWP', 'NAMA', 'NOMOR BUKTI POTONG',
                  'TANGGAL BUKTI PEMOTONGAN', 'KODE OBJEK PAJAK',
                  'PENGHASILAN BRUTO (Rp)', 'PAJAK PENGHASILAN (Rp)',
                  'FASILITAS PERPAJAKAN', 'Negara', 'ID TEMPAT KEGIATAN USAHA',
                  'KAP-KJS', 'STATUS']


IA_HEADERS = COMMON_HEADERS[:7] + ['TARIF PAJAK (%)'] + COMMON_HEADERS[7:]


def table(name, values, *, ia=False, headers=None, complete=True):
    return {'name': name, 'headers': list(headers or (IA_HEADERS if ia else COMMON_HEADERS)),
            'rows': values, 'complete': complete}


def section(name, *, tables=None, fields=None, text=''):
    return {'name': name, 'fields': fields or [], 'tables': tables or [], 'text': text}


def run_for(sections, *, model='Normal', key='synthetic-21-normal'):
    return {'npwp': '0000000000000001', 'year': 2025, 'year_basis': 'period',
            'taxpayer_name': 'SYNTHETIC FIXTURE ONLY', 'complete': False,
            'started_at': 'fixture', 'audit': [], 'errors': [],
            'records': [{'key': key, 'record_id': key, 'source_status': 'Fixture',
                         'meta': {'Jenis Pajak': 'PPh Pasal 21/26',
                                  'Masa Pajak': 'Januari 2025', 'Model SPT': model},
                         'capture_complete': True, 'detail_url': '', 'sections': sections}]}


def cells_with(workbook, value):
    return [c for ws in workbook for cells in ws for c in cells if c.value == value]


def assert_format(actual, reference):
    # Style IDs are workbook-local; compare rendered style components instead.
    for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
        assert copy(getattr(actual, attribute)) == copy(getattr(reference, attribute))
    assert actual.number_format == reference.number_format


def assert_layout(actual, reference):
    assert set(str(r) for r in actual.merged_cells.ranges) == set(str(r) for r in reference.merged_cells.ranges)
    for key, dimension in reference.column_dimensions.items():
        got = actual.column_dimensions[key]
        for attr in ('width', 'hidden', 'min', 'max', 'bestFit', 'outlineLevel', 'collapsed'):
            assert getattr(got, attr) == getattr(dimension, attr)
    for key, dimension in reference.row_dimensions.items():
        got = actual.row_dimensions[key]
        for attr in ('height', 'hidden', 'outlineLevel', 'collapsed'):
            assert getattr(got, attr) == getattr(dimension, attr)
    for attr in ('sheet_format', 'page_margins', 'print_options'):
        assert copy(getattr(actual, attr)) == copy(getattr(reference, attr))
    for attr in ('orientation', 'paperSize', 'fitToWidth', 'fitToHeight'):
        assert getattr(actual.page_setup, attr) == getattr(reference.page_setup, attr)
    for cells in reference:
        for c in cells:
            assert_format(actual[c.coordinate], c)


TYPES = {
    'PPN': {'reference': 'ppn_format.xlsx',
            'sha256': '99d63e34602b38414e3987ee72c49a8fb770fcd80cc99448266cbf587489d193',
            'styles': 'c26a67d76129ee20220826fb9070d6e90bdeffd73d35747b1d19b1ce4280ec9d',
            'forms': ['Induk', 'A1', 'A2', 'B1', 'B2', 'B3', 'C'],
            'path': 'SPT Masa PPN/SPT_PPN_Januari 2025_Normal.xlsx',
            'resource': 'resources/ppn_format.xlsx'},
    'PPh Unifikasi': {'reference': 'unifikasi_format.xlsx',
                     'sha256': '2ad49e997b6404a687c273dc7d895245e6d03593b8d76678b0928f8573efc7a2',
                     'styles': 'e4fb3f3684fc70914742af99a321fb17d950cc75b44b8b553df49ca071f37f0c',
                     'forms': ['Induk', 'Daftar I', 'Daftar II', 'Lampiran I'],
                     'path': 'SPT Masa Unifikasi/SPT_Unifikasi_Januari 2025_Normal.xlsx',
                     'resource': 'resources/unifikasi_format.xlsx'},
}


def source_run(tax, sections):
    run = run_for(sections, key='synthetic-template-normal')
    run['records'][0]['meta']['Jenis Pajak'] = tax
    return run


def export(tax, run, folder):
    files = exporter.export_split_workbooks(run, folder)
    coretax.verify_outputs(run, folder)
    assert set(files) == {'resume.xlsx', TYPES[tax]['path']}
    return load_workbook(folder / TYPES[tax]['path'])


RESOURCE = Path(__file__).resolve().parents[1] / 'resources' / 'ppn_format.xlsx'


APPENDICES = [
    ('A1', 3, 4, 7, (5,)),
    ('A2', 4, 5, 16, (6, 7, 8, 9)),
    ('B1', 4, 5, 6, (6, 7, 8)),
    ('B2', 3, 4, 8, (6, 7, 8, 9)),
    ('B3', 3, 4, 8, (6, 7, 8, 9)),
    ('C', 3, 4, 7, (10, 11, 12, 13)),
]


RAW_MARKER = 'Data sumber tambahan / belum dipetakan'


GROUP_III = 'III. PENGHITUNGAN PPN KURANG BAYAR / LEBIH BAYAR'


def export_ppn(tmp_path, sections):
    run = run_for(sections)
    run['records'][0]['meta']['Jenis Pajak'] = 'PPN'
    before = deepcopy(run)
    reports = []
    paths = exporter.export_split_workbooks(run, tmp_path, mapping_reports=reports)
    assert run == before, 'Display projection must not rewrite captured values or identity'
    path = next(path for path in paths if path != 'resume.xlsx')
    return load_workbook(tmp_path / path), reports, before


def source_appendix(form, header, first, total, columns, count=0):
    reference = load_workbook(RESOURCE)
    try:
        headers = [reference[form].cell(header, c).value
                   for c in range(1, reference[form].max_column + 1)]
        while headers and headers[-1] is None:
            headers.pop()
    finally:
        reference.close()
    if form == 'B3':
        headers[1] = 'Nama Penjual BKP/BKP Tidak Berwujud/Pemberi JKP'
    rows = [[str(i + 1)] + [f'SYNTHETIC-{i}-{c}' for c in range(2, len(headers) + 1)]
            for i in range(count)]
    amounts = ['0', '-17,00', '1.234,50', '29,00'][:len(columns)]
    footer = ['TOTAL'] + amounts + [''] * (len(headers) - columns[-1])
    spans = [columns[0] - 1] + [1] * (len(footer) - 1)
    return {'name': 'Tabel 1', 'headers': headers, 'rows': rows + [footer],
            'body_row_count': count, 'source_total': count, 'complete': True,
            'footer': [footer], 'footer_spans': [spans]}, amounts
