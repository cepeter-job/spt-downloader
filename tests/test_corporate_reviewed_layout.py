"""Operator-reviewed corporate layout; all inputs and values are synthetic."""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

import coretax_exporter as exporter
from support_templates import section, source_run
from test_badan_shouldbe_presentation import fields
from test_corporate_compact_export import export, table


INCOME_GROUP = ('C. PENGHASILAN YANG DIKENAKAN PPh YANG BERSIFAT FINAL DAN '
                'YANG TIDAK TERMASUK OBJEK PAJAK')
INCOME_CAPTIONS = [
    '2. Apakah Wajib Pajak menerima atau memperoleh penghasilan yang dikenakan PPh yang bersifat final? * / Tidak / Ya',
    '3. Apakah Wajib Pajak menerima atau memperoleh penghasilan yang tidak termasuk objek pajak? * / Tidak / Ya',
]


def income_fields():
    return [{'group': INCOME_GROUP, 'field': '', 'label': caption, 'value': value}
            for caption, value in zip(INCOME_CAPTIONS, ['1.234,50', '0'])]


def test_reviewed_induk_two_amounts_and_section_geometry_survive_saved_export(tmp_path):
    raw = fields() + income_fields()
    # Selected radio fields have different semantics from the bound amounts.
    raw += [{'group': INCOME_GROUP, 'field': 'IsIncomeSubject', 'label': 'Ya', 'value': '1'},
            {'group': INCOME_GROUP, 'field': 'IsIncomeExcluded', 'label': 'Tidak', 'value': '0'}]
    before = deepcopy(raw)
    book, reports = export(tmp_path, [section('Induk', fields=raw)])
    try:
        sheet = book['Induk']
        assert sheet.max_row == 33 and sheet.max_column == 2
        assert all(c.value is None for c in sheet[12])
        assert sheet['A13'].value == INCOME_GROUP[3:] + ' '
        assert sheet['A14'].value == 'Apakah Wajib Pajak menerima atau memperoleh penghasilan yang dikenakan PPh yang bersifat final?'
        assert sheet['A15'].value == ' Apakah Wajib Pajak menerima atau memperoleh penghasilan yang tidak termasuk objek pajak?'
        assert [sheet[c].value for c in ('B14', 'B15')] == [1234.5, 0]
        for address in ('B14', 'B15'):
            assert sheet[address].data_type == 'n'
            assert '#,##0' in sheet[address].number_format and '#,##0.' not in sheet[address].number_format
        for row, caption in [(16, 'PENGHITUNGAN PPh'), (21, 'PENGURANG PPh TERUTANG '),
                             (26, 'PPh KURANG/LEBIH BAYAR ')]:
            assert sheet.cell(row, 1).value == caption
            for cell in sheet[row][:2]:
                assert cell.font.bold
                assert cell.border.top.style == cell.border.bottom.style == 'thin'
                assert cell.border.left.style is None and cell.border.right.style is None
        assert sheet['A13'].font.bold and sheet['A13'].border.top.style is None
        assert sheet['B6'].value == '0000000000000002' and sheet['B6'].data_type == 's'
        assert sheet['B33'].value == 'Tarif Fasilitas sebagaimana Pasal 31E ayat (1) UU PPh'
        assert sheet.column_dimensions['A'].width == 110.7109375
        assert sheet.column_dimensions['B'].width == 49.7109375
        assert reports[0]['unmapped_fields'] == 2 and not reports[0]['complete']
        assert raw == before
    finally:
        book.close()


@pytest.mark.parametrize('mutation', ['owner', 'field', 'caption', 'duplicate', 'unsafe'])
def test_new_induk_amount_requires_one_exact_owned_safe_binding(tmp_path, mutation):
    raw = income_fields()[:1]
    if mutation == 'owner': raw[0]['group'] = 'SYNTHETIC FOREIGN GROUP'
    elif mutation == 'field': raw[0]['field'] = 'IsIncomeSubject'
    elif mutation == 'caption': raw[0]['label'] += ' changed'
    elif mutation == 'duplicate': raw += deepcopy(raw)
    else: raw[0]['value'] = '1234567890123456'
    book, reports = export(tmp_path, [section('Induk', fields=raw)])
    try:
        report = reports[0]
        assert not report['complete'] and report['unmapped_fields'] == len(raw)
        assert not any(c.value == 1234.5 for row in book['Induk'] for c in row)
        if mutation == 'unsafe':
            assert book['Induk']['B3'].value == '1234567890123456'
            assert book['Induk']['B3'].data_type == 's'
    finally:
        book.close()


def test_empty_new_induk_amounts_omit_heading_and_body_while_zero_is_content(tmp_path):
    raw = income_fields()
    for field in raw: field['value'] = ''
    book, reports = export(tmp_path, [section('Induk', fields=raw)])
    try:
        assert book['Induk'].max_row == 1
        assert reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('form,year_field,year_label,tin_label', [
    ('L2', 'TaxPeriodYear', 'Tahun Pajak/Bagian Tahun Pajak', 'Nomor Identitas WP'),
    ('L3', 'TaxPeriodYear', 'TAHUN PAJAK', 'Nomor Identitas WP'),
    ('L6', 'TaxPeriodYear', 'TAHUN PAJAK', 'Nomor Identitas WP'),
    ('L8', 'TaxPeriodYear', 'TAHUN PAJAK', 'NPWP'),
    ('L9', 'PeriodYear', 'TAHUN PAJAK', 'NPWP'),
    ('L10-A', 'PeriodYear', 'Tahun Pajak', 'Nomor Identitas WP'),
])
def test_reviewed_header_omission_keeps_unmapped_counts_and_checkpoint(tmp_path, form, year_field, year_label, tin_label):
    raw = [{'group': 'HEADER', 'field': year_field, 'label': year_label, 'value': '2025'},
           {'group': 'HEADER', 'field': 'CollectorTin', 'label': tin_label, 'value': '0000000000000002'},
           {'group': 'SYNTHETIC FOREIGN GROUP', 'field': 'CollectorTin', 'label': tin_label, 'value': '0000000000000002'}]
    book, reports = export(tmp_path, [section(form, fields=raw)])
    try:
        assert reports[0]['reviewed_redundant_fields'] == 2
        assert reports[0]['unmapped_fields'] == 3 and not reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('form,group', [
    ('L2', 'A. DAFTAR PEMEGANG SAHAM/PEMILIK MODAL DAN JUMLAH DIVIDEN/PEMBAGIAN LABA YANG DIBAGIKAN SERTA DAFTAR SUSUNAN PENGURUS DAN KOMISARIS'),
    ('L2', 'B. DAFTAR PENYERTAAN MODAL, UTANG, DAN/ATAU PIUTANG PADA PERUSAHAAN AFILIASI'),
    ('L3', 'A. PENGHASILAN DARI LUAR NEGERI'),
    ('L3', 'B. PPh YANG DIPOTONG/DIPUNGUT PIHAK LAIN'),
    ('L9', 'KELOMPOK 1'), ('L9', 'KELOMPOK 2'), ('L9', 'KELOMPOK 3'),
    ('L9', 'KELOMPOK 4'), ('L9', 'KELOMPOK LAINNYA'),
    ('L9', 'PERMANEN'), ('L9', 'TIDAK PERMANEN'),
    ('L10-A', 'Daftar Transaksi Yang Dipengaruhi Hubungan Istimewa'),
])
def test_reviewed_unnamed_remnant_is_annotated_without_certifying_its_origin(tmp_path, form, group):
    raw = [{'group': group, 'field': '', 'label': 'Unnamed field', 'value': '10'},
           {'group': group, 'field': 'SYNTHETIC UNKNOWN', 'label': 'Unnamed field', 'value': '10'}]
    book, reports = export(tmp_path, [section(form, fields=raw)])
    try:
        assert reports[0]['reviewed_redundant_fields'] == 1
        assert reports[0]['unmapped_fields'] == 2 and not reports[0]['complete']
    finally:
        book.close()


def test_l11b_review_rows_77_to_79_remain_unqualified(tmp_path):
    raw = [{'group': group, 'field': '', 'label': 'Unnamed field', 'value': '10'} for group in (
        'A. PENGHITUNGAN RATA-RATA SALDO UTANG', 'B. PENGHITUNGAN SALDO RATA-RATA MODAL',
        'III. PENGHITUNGAN BIAYA PINJAMAN')]
    book, reports = export(tmp_path, [section('L11-B', fields=raw)])
    try:
        assert reports[0]['unmapped_fields'] == 3 and not reports[0]['complete']
        assert reports[0].get('reviewed_redundant_fields', 0) == 0
        assert reports[0]['derived_reason'] == 'missing_average'
    finally:
        book.close()


@pytest.mark.parametrize('changed', ['none', 'nonempty', 'incomplete', 'foreign'])
def test_empty_l3_remnant_annotation_does_not_relax_table_completion(tmp_path, changed):
    source = {'name': 'Tabel 2', 'group': 'A. PENGHASILAN DARI LUAR NEGERI',
              'headers': [], 'rows': [], 'body_row_count': 0, 'complete': True, 'footer': []}
    if changed == 'nonempty': source.update(rows=[['SYNTHETIC VALUE']], body_row_count=1)
    elif changed == 'incomplete': source['complete'] = False
    elif changed == 'foreign': source['group'] = 'SYNTHETIC FOREIGN GROUP'
    book, reports = export(tmp_path, [section('L3', tables=[source])], capture_complete=changed != 'incomplete')
    try:
        assert reports[0].get('reviewed_redundant_tables', 0) == (1 if changed == 'none' else 0)
        assert reports[0]['unmapped_tables'] == 1 and not reports[0]['complete']
    finally:
        book.close()


def test_saved_resume_explains_reviewed_omission_without_changing_mapping_status(tmp_path):
    run = source_run('PPh Badan', [section('L6', fields=[
        {'group': 'HEADER', 'field': 'TaxPeriodYear', 'label': 'TAHUN PAJAK', 'value': '2025'}])])
    before = deepcopy(run)
    exporter.export_split_workbooks(run, tmp_path)
    book = load_workbook(tmp_path / 'resume.xlsx')
    try:
        row, = [row for row in book['Audit'].values if row[0] == 'Pemetaan template']
        assert row[1:5] == (1, 0, 0, 'PARTIAL')
        assert '1 field dan 0 tabel disetujui sebagai tampilan redundan' in row[5]
        assert run == before
    finally:
        book.close()


def test_captured_l9_first_total_is_bold_without_recalculation(tmp_path):
    source = table('L9', 4, [['0001', 'SYNTHETIC ASSET', '01/2025', '100', '80', 'GL', 'GL', '20', '']],
                   footer=[['TOTAL', '999', '']], spans=[[7, 1, 1]],
                   source_structure={'form': 'L9', 'component': 'rshshr-citr-l9-form-grid1',
                                     'groups': ['HARTA BERWUJUD', 'KELOMPOK 1']})
    book, _ = export(tmp_path, [section('L9', tables=[source])])
    try:
        assert book['L9']['H6'].value == 999 and book['L9']['H6'].font.bold
        assert '0.00' not in book['L9']['H6'].number_format
    finally:
        book.close()


def test_empty_l10c_monetary_slot_keeps_accounting_for_later_rows(tmp_path):
    book, reports = export(tmp_path, [section('L10-C')])
    try:
        assert book['L10-C']['E3'].value is None
        assert 'Rp' in book['L10-C']['E3'].number_format
        assert '#,##0.' not in book['L10-C']['E3'].number_format
        assert reports[0]['complete']
    finally:
        book.close()
