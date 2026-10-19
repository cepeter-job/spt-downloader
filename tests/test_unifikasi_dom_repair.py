"""Synthetic regressions for user-approved Unifikasi DOM repairs; runner-only."""
from copy import copy, deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook

import coretax_exporter as exporter
from support_templates import section, source_run

RESOURCE = Path(__file__).resolve().parents[1] / 'resources/unifikasi_format.xlsx'
TAX = 'PPh Unifikasi'
ACCOUNTING_RP = '_-"Rp"* #,##0_-;\\-"Rp"* #,##0_-;_-"Rp"* "-"??_-;_-@_-'
HEADERS = ['NO', 'DETIL', 'PAJAK PENGHASILAN / SETOR SENDIRI',
           'PAJAK PENGHASILAN / PEMOTONGAN / PEMUNGUTAN', 'PPh DITANGGUNG PEMERINTAH',
           'JUMLAH PPh DIBAYAR', 'JUMLAH PPh DIBAYAR DARI SPT YANG DIBETULKAN ' +
           'This row contain the total amount of payment deducted by the refunded amount prior to the amendment.',
           'PPh KURANG BAYAR / LEBIH BAYAR KARENA PEMBETULAN']


def export(tmp_path, sections, *, captured=True):
    run = source_run(TAX, sections)
    run['records'][0]['capture_complete'] = captured
    original = deepcopy(run)
    reports = []
    files = exporter.export_split_workbooks(run, tmp_path, mapping_reports=reports)
    assert run == original
    path, = [p for p in files if p != 'resume.xlsx']
    return load_workbook(tmp_path / path), reports


def source_table(headers, body, group, footers=(), spans=()):
    footers = list(footers)
    return {'name': 'Tabel 1', 'group': group, 'headers': headers,
            'rows': body + footers, 'body_row_count': len(body), 'source_total': len(body),
            'complete': True, 'footer': footers, 'footer_spans': list(spans)}


@pytest.mark.parametrize('form,reference', [('DAFTAR-I', 'Daftar I'), ('DAFTAR-II', 'Daftar II'), ('LAMPIRAN-I', 'Lampiran I')])
def test_hyphenated_dom_forms_use_template(tmp_path, form, reference):
    template = load_workbook(RESOURCE)
    book, _ = export(tmp_path, [section(form)])
    try:
        actual, expected = book[form], template[reference]
        assert actual['A1'].value == expected['A1'].value
        assert {str(m) for m in actual.merged_cells.ranges} == {str(m) for m in expected.merged_cells.ranges}
        assert copy(actual['A1'].font) == copy(expected['A1'].font)
    finally:
        book.close()
        template.close()


def test_induk_period_split_and_removed_status_rows(tmp_path):
    fields = [{'label': 'Masa Pajak', 'value': 'Januari 2025'},
              {'label': 'Status SPT', 'value': 'Normal'}, {'label': 'Status Penyampaian', 'value': 'Dilaporkan'}]
    book, _ = export(tmp_path, [section('INDUK', fields=fields)])
    try:
        sheet = book['INDUK']
        assert (sheet['D4'].value, sheet['D5'].value) == ('Januari', '2025')
        assert sheet['D4'].data_type == sheet['D5'].data_type == 's'
        assert sheet['B10'].value == 'Pasal 4 Ayat 2'
        assert sheet['A25'].value == 'TOTAL OF INCOME TAX'
        assert not any(c.value in ('Status SPT', 'Status Penyampaian') for row in sheet for c in row)
    finally:
        book.close()


@pytest.mark.parametrize('value,expected', [('1.234,50', 1234.5), ('0', 0), ('-17', -17),
                                         ('1234567890123456', '1234567890123456'), ('1e3', '1e3')])
def test_induk_rincian_money_and_reported_total(tmp_path, value, expected):
    body = [['1', 'Pasal 4 Ayat 2'] + [value] * 6]
    footer = [['TOTAL OF INCOME TAX'] + ['9.876,50'] * 6]
    table = source_table(HEADERS, body, 'B. PAJAK PENGHASILAN', footer, [[2, 1, 1, 1, 1, 1, 1]])
    book, reports = export(tmp_path, [section('INDUK', tables=[table])])
    try:
        for col in range(3, 9):
            assert book['INDUK'].cell(10, col).value == expected
            assert book['INDUK'].cell(10, col).data_type == ('s' if isinstance(expected, str) else 'n')
            assert book['INDUK'].cell(25, col).value == 9876.5
            assert book['INDUK'].cell(25, col).data_type == 'n'
            assert book['INDUK'].cell(10, col).number_format == ACCOUNTING_RP
            assert book['INDUK'].cell(25, col).number_format == ACCOUNTING_RP
        if isinstance(expected, str):
            assert reports and not reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('count', [0, 7, 41])
def test_daftar_i_dpp_numeric_and_body_boundary_style(tmp_path, count):
    headers = ['NO.', 'NIK/NPWP', 'NAMA', 'NOMOR BUKTI POTONG', 'TANGGAL BUKTI POTONG',
               'JENIS PAJAK', 'KODE OBJEK PAJAK', 'OBJEK PAJAK', 'DASAR PENGENAAN PAJAK (Rp)',
               'TINGKAT (%)', 'PAJAK PENGHASILAN', 'FASILITAS PERPAJAKAN',
               'UANG PERSEDIAAN / PEMBAYARAN LANGSUNG (UNTUK WP INSTANSI PEMERINTAH DENGAN DANA APBN)',
               'NITKU / NOMOR IDENTITAS SUBUNIT ORGANISASI', 'STATUS', 'KAP-KJS']
    body = [[str(i + 1), '0000000000000002', 'SYNTHETIC NAME', f'SYNTHETIC-{i}', '02-01-2025',
             'PPh Pasal 23', '23-100-01', 'SYNTHETIC OBJECT', '1.234,50', '2,00', '24,69', '',
             'SYNTHETIC PAYMENT METHOD', '0000000000000002000001', 'Normal', '411124-100'] for i in range(count)]
    table = source_table(headers, body, 'TABEL I. BPPU')
    book, _ = export(tmp_path, [section('DAFTAR-I', tables=[table])])
    try:
        sheet = book['DAFTAR-I']
        assert sheet['I4'].value == 'DASAR PENGENAAN PAJAK (Rp)'
        assert 'A1:O1' in {str(m) for m in sheet.merged_cells.ranges}
        for row in range(5, 5 + count):
            assert sheet.cell(row, 9).value == 1234.5 and sheet.cell(row, 9).data_type == 'n'
            assert sheet.cell(row, 10).value == 24.69 and sheet.cell(row, 10).data_type == 'n'
            assert sheet.cell(row, 2).value == '0000000000000002' and sheet.cell(row, 2).data_type == 's'
            assert sheet.cell(row, 13).value == '0000000000000002000001'
            assert sheet.cell(row, 12).value is None
            assert sheet.cell(row, 9).number_format == sheet.cell(row, 10).number_format == ACCOUNTING_RP
        if count >= 7:
            for col in range(1, 16):
                for attr in ('font', 'fill', 'border', 'alignment'):
                    assert copy(getattr(sheet.cell(8, col), attr)) == copy(getattr(sheet.cell(7, col), attr))
                    assert copy(getattr(sheet.cell(8, col), attr)) == copy(getattr(sheet.cell(9, col), attr))
        if count:
            assert any(c.value == 'SYNTHETIC PAYMENT METHOD' for row in sheet for c in row)
    finally:
        book.close()


@pytest.mark.parametrize('form,reference,group,footer_row,amount_col', [
    ('DAFTAR-II', 'Daftar II', 'TABEL I. DAFTAR PPh YANG DIBAYAR SENDIRI DAN/ATAU DISETOR SENDIRI', 10, 7),
    ('DAFTAR-II', 'Daftar II', 'TABEL II. DAFTAR PAJAK PENGHASILAN - PEMBAYARAN KUMULATIF', 18, 7),
    ('LAMPIRAN-I', 'Lampiran I', 'TABEL I. ATC', 10, 13),
])
def test_empty_appendix_preserves_reported_zero_without_fake_body(tmp_path, form, reference, group, footer_row, amount_col):
    ref = load_workbook(RESOURCE)
    try:
        header_row = 13 if footer_row == 18 else 4
        headers = [c.value for c in ref[reference][header_row]]
        label = ref[reference].cell(footer_row, 1).value
    finally:
        ref.close()
    footer = [[], [label, '0', '']]
    table = source_table(headers, [], group, footer, [[], [amount_col - 1, 1, len(headers) - amount_col]])
    book, reports = export(tmp_path, [section(form, tables=[table])])
    try:
        sheet = book[form]
        assert sheet.cell(footer_row, amount_col).value == 0
        assert sheet.cell(footer_row, amount_col).data_type == 'n'
        assert sheet.cell(footer_row, amount_col).number_format == ACCOUNTING_RP
        assert sheet['A5'].value is None and sheet['B5'].value is None
        assert reports and reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('defect', ['duplicate', 'unknown-header', 'foreign-group', 'incomplete', 'count'])
def test_empty_source_contradictions_never_claim_complete(tmp_path, defect):
    ref = load_workbook(RESOURCE)
    try:
        headers = [c.value for c in ref['Daftar I'][4]]
    finally:
        ref.close()
    table = source_table(headers, [], 'TABEL I. BPPU')
    if defect == 'unknown-header':
        table['headers'] = ['UNSUPPORTED SOURCE FIELD']
    elif defect == 'foreign-group':
        table['group'] = 'SYNTHETIC UNRELATED GROUP'
    elif defect == 'incomplete':
        table['complete'] = False
    elif defect == 'count':
        table['source_total'] = 1
    tables = [table, deepcopy(table)] if defect == 'duplicate' else [table]
    book, reports = export(tmp_path, [section('DAFTAR-I', tables=tables)], captured=defect != 'incomplete')
    try:
        assert reports and not reports[0]['complete']
        assert book['DAFTAR-I']['A5'].value is None
    finally:
        book.close()


def test_monetary_template_formats_match_ppn_and_pph21():
    resources = RESOURCE.parent
    for filename, address in [('ppn_format.xlsx', 'E12'), ('pph21_format.xlsx', 'D12')]:
        reference = load_workbook(resources / filename)
        try:
            assert reference['Induk'][address].number_format == ACCOUNTING_RP
        finally:
            reference.close()
    targets = {
        'Induk': [(range(10, 26), range(3, 9))],
        'Daftar I': [(range(5, 9), (9, 10)), (range(11, 15), (9, 10)), (range(15, 19), (10,))],
        'Daftar II': [(range(5, 9), (5, 7)), (range(14, 18), (5, 7)), ((9, 10, 18), (7,))],
        'Lampiran I': [(range(5, 10), (11, 13)), ((10,), (13,))],
    }
    workbook = load_workbook(RESOURCE)
    try:
        mismatches = [(name, workbook[name].cell(row, col).coordinate)
                      for name, regions in targets.items() for rows, columns in regions
                      for row in rows for col in columns
                      if workbook[name].cell(row, col).number_format != ACCOUNTING_RP]
        assert mismatches == []
        # Periods, identifiers, rates and dates must not inherit the currency format.
        for name, cells in [('Induk', ('D4', 'D5')), ('Daftar I', ('B5', 'D5', 'E5', 'M5')),
                            ('Daftar II', ('C5', 'F5', 'I5')), ('Lampiran I', ('B5', 'I5', 'J5', 'L5', 'O5'))]:
            assert all(workbook[name][cell].number_format != ACCOUNTING_RP for cell in cells)
    finally:
        workbook.close()


def test_conflicting_period_year_preserves_raw_evidence(tmp_path):
    fields = [{'label': 'Masa Pajak', 'value': 'Januari 2025'}, {'label': 'Tahun Pajak', 'value': '2024'}]
    book, reports = export(tmp_path, [section('INDUK', fields=fields)])
    try:
        assert reports and not reports[0]['complete']
        assert book['INDUK']['D4'].value is None and book['INDUK']['D5'].value is None
        assert any(c.value == '2024' for row in book['INDUK'] for c in row)
    finally:
        book.close()
