"""Complete PPh21 JSON sections and native numeric workbooks, with no DOM values."""
from copy import deepcopy
import pytest
from openpyxl import load_workbook

from coretax_json import JsonSafetyError
from coretax_masa_json import adapt_masa_detail
from masa_json_fixtures import SCOPE, RECORD
from pph21_projection_fixture import case
from support_templates import assert_format, assert_layout, run_for
from test_pph21_template import REFERENCE, export


def adapt(record, bodies, forms):
    return adapt_masa_detail(bodies, record=record, scope=SCOPE, observed_forms=forms)


def test_complete_attempt_projects_all_forms_rows_footers_and_preserves_original_sources():
    record, bodies, forms = case()
    before = deepcopy((record, bodies, forms))
    result = adapt(record, bodies, forms)
    assert result['eligible'], result['reason']
    assert result['coverage']['complete']
    assert [s['name'] for s in result['sections']] == ['Induk', 'L-IA', 'L-IB', 'L-II', 'L-III']
    assert len(result['coverage']['grids']) == 7
    assert all(grid['source_total'] == 2 and grid['complete'] for grid in result['coverage']['grids'])
    assert result['provenance']['bindings']
    assert (record, bodies, forms) == before


def test_json_sections_export_all_five_clean_sheets_with_money_numeric_and_identifiers_text(tmp_path):
    record, bodies, forms = case()
    result = adapt(record, bodies, forms)
    assert result['eligible'], result['reason']
    reports = []
    workbook = export(run_for(result['sections'], key=RECORD), tmp_path, reports=reports)
    reference = load_workbook(REFERENCE)
    try:
        assert workbook.sheetnames == ['Induk', 'L-IA', 'L-IB', 'L-II', 'L-III']
        assert len(reports) == 5 and all(report['complete'] for report in reports)
        for name in workbook.sheetnames:
            assert_layout(workbook[name], reference[name])
            for row in reference[name]:
                for cell in row:
                    assert_format(workbook[name][cell.coordinate], cell)
        for coordinate, value in [('D4', '1'), ('D5', '2025'), ('D6', 'SYNTHETIC MODEL'), ('D7', 'Submitted')]:
            assert workbook['Induk'][coordinate].value == value
            assert workbook['Induk'][coordinate].data_type == 's'
        assert workbook['Induk']['D12'].value == sum([12, 13, 14])
        assert workbook['Induk']['D12'].data_type == 'n'
        for name, row in [('L-IA', 5), ('L-IB', 6), ('L-IB', 12), ('L-II', 5), ('L-II', 11), ('L-III', 5), ('L-III', 11)]:
            assert workbook[name].cell(row, 2).value == '0000000000000002'
            assert workbook[name].cell(row, 2).data_type == 's'
            assert workbook[name].cell(row, 7).value == 1200
            assert workbook[name].cell(row, 7).data_type == 'n'
        assert workbook['L-IA']['G16'].value == 800
        assert workbook['L-IA']['G16'].data_type == 'n'
        assert all(table['body_row_count'] == 2 for s in result['sections'][1:] for table in s['tables'])
    finally:
        workbook.close()
        reference.close()


@pytest.mark.parametrize('defect', ['missing-form', 'extra-form', 'missing-grid', 'unknown-field',
    'dom-value', 'wrong-build', 'missing-page', 'bad-footer', 'missing-reference', 'unbound-lookup', 'duplicate-amount'])
def test_incomplete_or_unreviewed_projection_stays_ineligible(defect):
    record, bodies, forms = case()
    if defect == 'missing-form':
        forms.pop()
    elif defect == 'extra-form':
        forms.append(deepcopy(forms[-1]))
    elif defect == 'missing-grid':
        forms[-1]['tables'].pop()
    elif defect == 'unknown-field':
        forms[0]['fields'].append({'field': 'Unreviewed', 'tag': 'INPUT', 'label': 'Unreviewed', 'group': '', 'column_label': ''})
    elif defect == 'dom-value':
        forms[0]['fields'][0]['value'] = 'DOM VALUE MUST NOT BE USED'
    elif defect == 'wrong-build':
        forms[0]['builds'] = {}
    elif defect == 'missing-page':
        forms[1]['tables'][0]['pages'].pop()
    elif defect == 'bad-footer':
        forms[1]['tables'][0]['footer_spans'] = []
    elif defect == 'missing-reference':
        forms[0]['references'].pop()
    elif defect == 'unbound-lookup':
        forms[0]['lookups'] = {'ReturnSheetModel': {'NORMAL': 'UNBOUND'}}
        forms[0]['references'] = []
    else:
        forms[0]['tables'][0]['row_fields'][2] = forms[0]['tables'][0]['row_fields'][1][:]
    result = adapt(record, bodies, forms)
    assert not result['eligible'] and not result['coverage']['complete']


@pytest.mark.parametrize('name', ['L-IB', 'L-II'])
def test_empty_first_group_remains_separate_from_second_group_and_reported_footer(name):
    import hashlib
    import json
    record, bodies, forms = case()
    table = next(f for f in forms if f['name'] == name)['tables'][0]
    removed = table['pages'][1]['sequences'][0]
    table['pages'] = table['pages'][:1]
    table['source_total'] = 0
    bodies = [p for p in bodies if p['context']['sequence'] != removed]
    item = next(p for p in bodies if p['context']['sequence'] == table['pages'][0]['sequences'][0])
    item['body'] = json.dumps({'IsSuccessful': True, 'Payload': {'Data': [], 'TotalRecords': 0}}).encode()
    item['context'].update(bytes=len(item['body']), sha256=hashlib.sha256(item['body']).hexdigest())
    result = adapt(record, bodies, forms)
    assert result['eligible'], result['reason']
    projected = next(s for s in result['sections'] if s['name'] == name)['tables']
    assert projected[0]['body_row_count'] == 0 and projected[0]['group'] == 'BPA1'
    assert projected[1]['body_row_count'] == 2 and projected[1]['group'] == 'BPA2'
    assert projected[1]['footer']


def test_foreign_grid_packet_remains_fatal_before_missing_form_fallback():
    record, bodies, forms = case()
    bodies[-1]['context']['request_json']['ReturnSheetRecordId'] = '00000000-0000-4000-8000-000000000099'
    with pytest.raises(JsonSafetyError):
        adapt(record, bodies, [])


def test_reference_sequence_cannot_point_to_a_grid_response():
    record, bodies, forms = case()
    forms[0]['references'][0]['sequences'] = [4]
    with pytest.raises(JsonSafetyError):
        adapt(record, bodies, forms)
