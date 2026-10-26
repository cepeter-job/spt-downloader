"""Real worker/publication flow; only the authorized portal is replaced."""
import inspect
from pathlib import Path
from types import SimpleNamespace

import pytest

from spt_desktop import worker, output
from support_runs import fixture_run
from support_selected_capture import synthetic_browser


from support_selected_capture import (
    selected_job
)

def test_selected_capture_never_opens_unconfirmed_returns_and_can_complete_subset(tmp_path, monkeypatch):
    status, events, state, inventory = selected_job(tmp_path, monkeypatch, ['21-january'])
    assert state['captured'] == ['21-january']
    assert state['requests'] == ['selection']
    assert status == 0 and events[-1]['status'] == 'COMPLETE'
    assert state['browser_closed'] and state['driver_closed']
    folder = Path(events[-1]['output'])
    assert {path.relative_to(folder).as_posix() for path in folder.rglob('*.xlsx')} == {
        'resume.xlsx', 'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx'}
    assert len(inventory['records']) == 5 and not inventory['complete']
    from openpyxl import load_workbook
    workbook = load_workbook(folder / 'resume.xlsx')
    try:
        assert workbook['Ringkasan']['B7'].value == 1
        assert any('Pilihan pengguna: 1 dari 5' in str(cell.value)
                   for row in workbook['Ringkasan'] for cell in row)
    finally:
        workbook.close()


@pytest.mark.parametrize('selected_ids', [[], ['unknown'], ['21-january', '21-january']])
def test_invalid_selection_never_captures_or_publishes_any_return(tmp_path, monkeypatch, selected_ids):
    status, events, state, _ = selected_job(tmp_path, monkeypatch, selected_ids)
    assert status != 0 and events[-1]['status'] != 'COMPLETE'
    assert state['captured'] == []
    assert state['browser_closed'] and state['driver_closed']
    assert not (tmp_path / 'Downloads').exists()


def test_stop_after_selection_reply_wins_before_any_detail_is_opened(tmp_path, monkeypatch):
    status, events, state, _ = selected_job(tmp_path, monkeypatch, ['21-january'], stop_on_choice=True)
    assert status == 4 and events[-1]['status'] == 'CANCELLED'
    assert state['captured'] == []
    assert state['browser_closed'] and state['driver_closed']
    assert not (tmp_path / 'Downloads').exists()


def test_worker_declining_overwrite_keeps_previous_output_byte_for_byte(tmp_path, monkeypatch):
    run = fixture_run()
    worker.coretax.save_output(run, tmp_path / 'old-source')
    final = output.publish_workbooks(run, tmp_path / 'old-source', tmp_path / 'Downloads', 'old-job')
    old = {path.relative_to(final).as_posix(): path.read_bytes() for path in final.rglob('*.xlsx')}
    status, events, state, _ = selected_job(tmp_path, monkeypatch, ['21-january'], overwrite=False)
    assert state['requests'] == ['selection', 'overwrite']
    assert status != 0 and events[-1]['status'] != 'COMPLETE' and events[-1]['output'] is None
    assert {path.relative_to(final).as_posix(): path.read_bytes() for path in final.rglob('*.xlsx')} == old


def test_worker_approved_overwrite_preserves_unselected_workbooks(tmp_path, monkeypatch):
    run = fixture_run()
    worker.coretax.save_output(run, tmp_path / 'old-source')
    final = output.publish_workbooks(run, tmp_path / 'old-source', tmp_path / 'Downloads', 'old-job')
    february = final / 'SPT Masa 21' / 'SPT_21_Februari 2025_Normal.xlsx'
    before = february.read_bytes()
    marker = 'SYNTHETIC APPROVED JANUARY UPDATE'
    status, events, state, _ = selected_job(tmp_path, monkeypatch, ['21-january'],
                                           overwrite=True, marker=marker)
    assert state['requests'] == ['selection', 'overwrite']
    assert state['captured'] == ['21-january']
    assert status == 0 and events[-1]['status'] == 'COMPLETE'
    assert Path(events[-1]['output']) == final
    assert february.read_bytes() == before
    from openpyxl import load_workbook
    workbook = load_workbook(final / 'SPT Masa 21' / 'SPT_21_Januari 2025_Normal.xlsx')
    try:
        assert any(cell.value == marker for sheet in workbook for row in sheet for cell in row)
    finally:
        workbook.close()


def test_worker_without_a_consent_channel_never_captures_all_records(tmp_path, monkeypatch):
    state = {'captured': []}
    synthetic_browser(worker, monkeypatch, state)
    run = fixture_run()
    run['inventory_complete'] = True
    monkeypatch.setattr(worker, 'collect_inventory', lambda *args, **kwargs: run)

    def capture(page, data, folder, on_progress, check_cancel):
        state['captured'] = [record['record_id'] for record in data['records']]
        data['complete'] = True
        return data

    monkeypatch.setattr(worker, 'collect_details', capture)
    monkeypatch.setattr(worker, 'downloads_folder', lambda: tmp_path / 'Downloads')
    events = []
    status = worker.run_job('0000000000000001', '2025', tmp_path / 'workspace',
                            'unconfirmed-job', 'msedge', True, events.append, lambda: False)
    assert state['captured'] == []
    assert status != 0 and events[-1]['status'] != 'COMPLETE'
    assert not state.get('browser_started') or state['browser_closed']
    assert not state.get('driver_started') or state['driver_closed']
    assert not (tmp_path / 'Downloads').exists()


def test_worker_announces_zero_completed_details_before_starting_capture(tmp_path, monkeypatch):
    observed = []

    def at_capture_start(data, events):
        details = tuple(dict(event) for event in events if event.get('type') == 'progress'
                        and event.get('stage') == 'details')
        observed.append((len(data['records']), details))

    status, events, state, _ = selected_job(tmp_path, monkeypatch, ['21-january'],
                                           before_capture=at_capture_start)
    assert status == 0 and events[-1]['status'] == 'COMPLETE'
    assert len(observed) == 1 and state['captured'] == ['21-january']
    record_count, details = observed[0]
    assert len(details) == 1, 'Capture started before its initial progress event'
    event = details[0]
    assert event['processed'] == event['verified'] == 0
    assert event['total'] == record_count == 1
    assert event['message'] == 'SPT diproses 0/1; terverifikasi 0.'
    assert event['percent'] < 100
