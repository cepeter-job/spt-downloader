"""Pinned L8 current-source pure projection; all execution runs on runners."""
from copy import deepcopy
from fractions import Fraction

import pytest
from openpyxl import load_workbook

import coretax
import coretax_corporate_projection as projection
from coretax_json import _UnsupportedMetadata, _corporate_json
from support_json_corpus import Number, wire
from support_templates import section, source_run


ROOT = 'CalculationFacilitiesIncomeTaxRateReduction'
KEYS = [
    'GrossTurnOver',
    'TaxableIncomeShareOfGrossCirculationGrantedFacility',
    'TaxableIncomeShareOfGrossCirculationNotGrantedFacility',
    'IncomeTaxPayableFromTaxableIncomeShareOfGrossCirculationGrantedFacility',
    'IncomeTaxPayableFromTaxableIncomeShareOfGrossCirculationNotGrantedFacility',
    'TotalIncomeTax',
]
MAIN = ['TaxableIncome', 'TotalGrossIncomeResearch']
GROUP = ('PERHITUNGAN FASILITAS PENGURANGAN TARIF PPh BAGI WAJIB PAJAK BADAN DALAM NEGERI '
         'BERDASARKAN PASAL 31E AYAT (1) UNDANG-UNDANG PPh')
HEADERS = ['NO.', 'DESKRIPSI', 'AMOUNT (Rupiah)']
CAPTIONS = [
    'Jumlah Peredaran Bruto',
    'Penghasilan Kena Pajak',
    'Penghasilan Kena Pajak dari bagian peredaran bruto yang memperoleh fasilitas',
    'Penghasilan Kena Pajak dari bagian peredaran bruto yang tidak memperoleh fasilitas',
    'Pajak Terutang',
    'PPh Terutang atas Penghasilan Kena Pajak dari bagian peredararan bruto yang memperoleh fasilitas',
    'PPh Terutang atas Penghasilan Kena Pajak dari bagian peredararan bruto yang tidak memperoleh fasilitas',
]
SPANS = [[1, 2], [1, 1, 1], [1, 2], [1, 1, 1], [1, 1, 1], [1, 2], [1, 1, 1], [1, 1, 1]]
DESTINATIONS = [
    {'table': 'L8', 'row': 1, 'column': 2},
    {'table': 'L8', 'row': 3, 'column': 2},
    {'table': 'L8', 'row': 4, 'column': 2},
    {'table': 'L8', 'row': 6, 'column': 2},
    {'table': 'L8', 'row': 7, 'column': 2},
    {'table': 'L8', 'footer': 0, 'column': 2},
]


def inputs(*, turnover=100000, taxable=12345, year=2025):
    masked = taxable - taxable % 1000
    tax = (masked // 100) * 11
    return (dict(zip(KEYS, [turnover, masked, 0, tax, 0, tax])),
            {'TaxYear': year, 'TaxableIncome': taxable, 'TotalGrossIncomeResearch': 0})


def project(form, main_form, *, period_year=2025, **overrides):
    function = getattr(projection, 'project_l8', None)
    assert callable(function), 'L8 pure source projector is not implemented'
    options = {'headers': deepcopy(HEADERS), 'footer_spans': [[1, 1, 1]],
               'tax_type': 'ICT_RCIT', 'locale': 'id-ID'}
    options.update(overrides)
    current, main, year = [_corporate_json(wire(value)) for value in (form, main_form, period_year)]
    before = deepcopy((current, main, year, options))
    try:
        return function(current, main_form=main, period_year=year, **options)
    finally:
        assert (current, main, year, options) == before


def evidence_for(result, path):
    entry, = [entry for entry in result['provenance'] if entry['path'] == path]
    return entry


def test_l8_exact_eight_physical_rows_spans_footer_and_source_destinations():
    result = project(*inputs())
    table = result['table']
    body = [
        ['1.', CAPTIONS[0]], ['', CAPTIONS[0], '100.000'],
        ['2.', CAPTIONS[1]], ['', CAPTIONS[2], '12.000'], ['', CAPTIONS[3], '0'],
        ['3.', CAPTIONS[4]], ['', CAPTIONS[5], '1.320'], ['', CAPTIONS[6], '0'],
    ]
    footer = [['', 'Jumlah PPh Terutang', '1.320']]
    assert table['headers'] == HEADERS and table['group'] == GROUP
    assert table['rows'] == body + footer and table['footer'] == footer
    assert table['body_spans'] == SPANS
    assert table['body_rowspans'] == [[1] * len(row) for row in body]
    assert table['footer_spans'] == [[1, 1, 1]]
    assert table['source_structure'] == {
        'form': 'L8', 'component': 'rsh-citr-l8-form-grid1', 'groups': [GROUP]}
    assert table['pages'] == 1 and table['body_row_count'] == table['source_total'] == 8
    assert table['complete'] is True
    form, parent = inputs()
    for key, destination in zip(KEYS, DESTINATIONS):
        entry = evidence_for(result, [ROOT, key])
        assert entry['source'] == 'current' and entry['type'] == 'number'
        assert entry['value'] == str(form[key]) and entry['destination'] == destination
        assert entry['rule'] == 'reconciled-captured/input-number-integer-id'
    for key in KEYS[3:5]:
        assert evidence_for(result, [ROOT, key])['also_supports'] == [DESTINATIONS[-1]]
    for key in MAIN:
        entry = evidence_for(result, ['MainFormData', key])
        assert entry['source'] == 'current' and entry['type'] == 'number'
        assert entry['value'] == str(parent[key])
        assert entry['destinations'] == DESTINATIONS[1:]
    for path, source, control in [(['MainFormData', 'TaxYear'], 'current', 'rate-year'),
                                  (['PeriodYear'], 'primary', 'header-year')]:
        entry = evidence_for(result, path)
        assert (entry['source'], entry['type'], entry['value']) == (source, 'number', '2025')
        assert entry['destination'] == {'table': 'L8', 'control': control}
    assert len(result['provenance']) == 10


@pytest.mark.parametrize('turnover', [0, 1, 4800000000])
@pytest.mark.parametrize('taxable,masked,tax', [
    (0, 0, 0), (999, 0, 0), (1000, 1000, 110), (12345, 12000, 1320),
    (999999999999999, 999999999999000, 109999999999890),
])
def test_l8_reviewed_turnover_boundary_and_main_thousand_mask(turnover, taxable, masked, tax):
    form, parent = inputs(turnover=turnover, taxable=taxable)
    assert form[KEYS[1]] == masked and form[KEYS[3]] == tax
    table = project(form, parent)['table']
    rendered = lambda value: format(value, ',d').replace(',', '.')
    assert [table['rows'][row][2] for row in (1, 3, 4, 6, 7)] == [
        rendered(turnover), rendered(masked), '0', rendered(tax), '0']
    assert table['footer'][0][2] == rendered(tax)


def test_l8_binary64_rate_bound_covers_the_supported_maximum_without_rounding_tax():
    result = project(*inputs(taxable=999999999999999))
    rate = Fraction.from_float(0.11)
    error = rate - Fraction(11, 100)
    assert rate == Fraction(7926335344172073, 72057594037927936)
    assert error == Fraction(1, 1801439850948198400)
    assert error / Fraction(11, 100) == Fraction(1, 198158383604301824) < Fraction(1, 2 ** 54)
    assert 109999999999890 < 2 ** 47
    assert result['table']['footer'][0][2] == '109.999.999.999.890'


@pytest.mark.parametrize('location,key', [('form', key) for key in KEYS] + [('main', key) for key in MAIN])
@pytest.mark.parametrize('suffix', ['.0', 'e0'])
def test_l8_original_numeric_tokens_survive_rendering_and_reconciliation(location, key, suffix):
    form, parent = inputs()
    source = form if location == 'form' else parent
    token = str(source[key]) + suffix
    source[key] = Number(token)
    result = project(form, parent)
    entry = evidence_for(result, [ROOT if location == 'form' else 'MainFormData', key])
    assert entry['type'] == 'number' and entry['value'] == token
    assert result['table']['footer'][0][2] == '1.320'


@pytest.mark.parametrize('location,key', [('form', key) for key in KEYS] + [('main', key) for key in MAIN])
@pytest.mark.parametrize('value', [None, '', '0', True, [], {}, -1, 0.5, 1000000000000000])
def test_l8_unsupported_numeric_operands_are_never_coerced_or_defaulted(location, key, value):
    form, parent = inputs()
    (form if location == 'form' else parent)[key] = value
    with pytest.raises(_UnsupportedMetadata):
        project(form, parent)


@pytest.mark.parametrize('location,key', [('form', key) for key in KEYS] + [('main', key) for key in MAIN])
def test_l8_missing_operand_is_not_an_implicit_zero(location, key):
    form, parent = inputs()
    del (form if location == 'form' else parent)[key]
    with pytest.raises(_UnsupportedMetadata):
        project(form, parent)


@pytest.mark.parametrize('location,key', [('form', key) for key in KEYS] + [('main', key) for key in MAIN])
@pytest.mark.parametrize('token', ['-0', '-0.0', '-0e0'])
def test_l8_unproved_signed_zero_control_stays_unsupported(location, key, token):
    form, parent = inputs(turnover=0, taxable=0)
    (form if location == 'form' else parent)[key] = Number(token)
    with pytest.raises(_UnsupportedMetadata):
        project(form, parent)


@pytest.mark.parametrize('key', KEYS[1:])
def test_l8_every_reported_share_tax_and_total_must_reconcile(key):
    form, parent = inputs()
    form[key] += 1
    with pytest.raises(_UnsupportedMetadata):
        project(form, parent)


def test_l8_captured_taxable_share_cannot_be_silently_masked_into_eligibility():
    form, parent = inputs()
    form[KEYS[1]] = parent['TaxableIncome']
    with pytest.raises(_UnsupportedMetadata):
        project(form, parent)


@pytest.mark.parametrize('defect', ['above-threshold', 'research-nonzero', 'extra-root',
    'list-root', 'null-root', 'list-main', 'null-main', 'header-usd', 'header-alias',
    'header-order', 'footer-span', 'tax-usd', 'tax-oil', 'locale'])
def test_l8_other_source_shapes_and_unproved_branches_remain_unsupported(defect):
    form, parent = inputs()
    options = {}
    if defect == 'above-threshold': form[KEYS[0]] = 4800000001
    elif defect == 'research-nonzero': parent['TotalGrossIncomeResearch'] = 1
    elif defect == 'extra-root': form['SYNTHETIC_UNREVIEWED'] = 0
    elif defect == 'list-root': form = []
    elif defect == 'null-root': form = None
    elif defect == 'list-main': parent = []
    elif defect == 'null-main': parent = None
    elif defect == 'header-usd': options['headers'] = ['NO.', 'DESKRIPSI', 'AMOUNT (USD)']
    elif defect == 'header-alias': options['headers'] = ['NO.', 'DESCRIPTION', HEADERS[2]]
    elif defect == 'header-order': options['headers'] = list(reversed(HEADERS))
    elif defect == 'footer-span': options['footer_spans'] = [[2, 1]]
    elif defect in ('tax-usd', 'tax-oil'):
        options['tax_type'] = 'ICT_DCIT' if defect == 'tax-usd' else 'ICT_ROGCIT'
    elif defect == 'locale': options['locale'] = 'en-US'
    with pytest.raises(_UnsupportedMetadata):
        project(form, parent, **options)


@pytest.mark.parametrize('year', [2020, 2025, 2100])
def test_l8_original_matching_outer_and_main_years_select_the_reviewed_rate(year):
    result = project(*inputs(year=year), period_year=year)
    for path in (['PeriodYear'], ['MainFormData', 'TaxYear']):
        assert evidence_for(result, path)['value'] == str(year)


@pytest.mark.parametrize('location', ['outer', 'main'])
@pytest.mark.parametrize('year', [None, '2025', True, 2019, 2101, 2025.5, Number('2025.0'), Number('2.025e3')])
def test_l8_year_encoding_or_unreviewed_rate_year_is_not_inferred(location, year):
    form, parent = inputs()
    options = {}
    if location == 'main': parent['TaxYear'] = year
    else: options['period_year'] = year
    with pytest.raises(_UnsupportedMetadata):
        project(form, parent, **options)


@pytest.mark.parametrize('main_year', ['missing', 2019, 2024])
def test_l8_original_main_rate_year_cannot_be_replaced_with_outer_header_year(main_year):
    form, parent = inputs()
    if main_year == 'missing': del parent['TaxYear']
    else: parent['TaxYear'] = main_year
    with pytest.raises(_UnsupportedMetadata):
        project(form, parent)


@pytest.mark.parametrize('year', [2019, 2101])
def test_l8_matching_original_years_still_require_the_reviewed_rate_year_range(year):
    form, parent = inputs(year=year)
    with pytest.raises(_UnsupportedMetadata):
        project(form, parent, period_year=year)


@pytest.mark.parametrize('value', [None, 0, Number('-0.0'), Number('1.2345e3')])
def test_l8_optional_exchange_control_is_retained_without_affecting_rupiah_amounts(value):
    form, parent = inputs()
    expected = project(form, parent)
    form['ExchangeRateValue'] = value
    result = project(form, parent)
    assert result['table'] == expected['table']
    entry = evidence_for(result, [ROOT, 'ExchangeRateValue'])
    assert (entry['source'], entry['type'], entry['value']) == (
        'current', 'null' if value is None else 'number', None if value is None else str(value))
    assert entry['destination'] == {'table': 'L8', 'control': 'unused-rupiah-exchange-rate'}
    assert entry['rule'] == 'retained-unused-rupiah-source-control'
    assert len(result['provenance']) == len(expected['provenance']) + 1
    assert not result.get('fields')


@pytest.mark.parametrize('value', ['1', True, [], {}])
def test_l8_unproved_optional_exchange_control_shape_is_unsupported(value):
    form, parent = inputs()
    form['ExchangeRateValue'] = value
    with pytest.raises(_UnsupportedMetadata):
        project(form, parent)


@pytest.mark.parametrize('taxable', [0, 12345, 999999999999999])
def test_l8_final_saved_export_keeps_both_literal_models_counts_merges_and_links(tmp_path, taxable):
    result = project(*inputs(taxable=taxable))
    run = source_run('PPh Badan', [section('L8', tables=[deepcopy(result['table'])])])
    normal = run['records'][0]
    normal['meta']['Masa Pajak'] = '01012025'
    amendment = deepcopy(normal)
    amendment['key'] = amendment['record_id'] = 'synthetic-l8-amendment-002'
    amendment['meta']['Model SPT'] = 'Amendment 002'
    run['records'].append(amendment)
    source_before = deepcopy([record['sections'] for record in run['records']])
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    assert [record['sections'] for record in run['records']] == source_before
    assert {record['meta']['Masa Pajak'] for record in run['records']} == {'01012025'}
    files = [name for name in run['output_files'] if name != 'resume.xlsx']
    assert len(files) == 2 and len(set(files)) == 2
    assert all(name.startswith('SPT_Badan_01012025_') for name in files)
    assert any('Normal' in name for name in files) and any('Amendment 002' in name for name in files)
    expected_amounts = [100000, taxable - taxable % 1000, 0, ((taxable - taxable % 1000) // 100) * 11, 0]
    for name in files:
        book = load_workbook(tmp_path / name)
        try:
            assert book.sheetnames == ['L8']
            sheet = book['L8']
            assert (sheet.max_row, sheet.max_column) == (10, 3)
            assert [cell.value for cell in sheet[1]] == HEADERS
            assert {str(area) for area in sheet.merged_cells.ranges} == {'B2:C2', 'B4:C4', 'B7:C7'}
            assert [sheet.cell(row, 2).value for row in range(2, 11)] == [
                CAPTIONS[0], CAPTIONS[0], CAPTIONS[1], CAPTIONS[2], CAPTIONS[3],
                CAPTIONS[4], CAPTIONS[5], CAPTIONS[6], 'Jumlah PPh Terutang']
            assert [sheet.cell(row, 3).value for row in (3, 5, 6, 8, 9)] == expected_amounts
            assert sheet['C10'].value == expected_amounts[3]
            for row in (3, 5, 6, 8, 9, 10):
                cell = sheet.cell(row, 3)
                assert cell.data_type == 'n' and 'Rp' in cell.number_format and '0.00' not in cell.number_format
                assert cell.border.top.style == cell.border.bottom.style == 'thin'
            assert [sheet.column_dimensions[column].width for column in 'ABC'] == [4.42578125, 96.28515625, 36]
            assert sheet.column_dimensions['A'].bestFit and sheet.column_dimensions['B'].bestFit
            assert sheet.page_setup.orientation == 'landscape'
            assert not any(cell.value == 'Data sumber tambahan / belum dipetakan' for row in sheet for cell in row)
        finally:
            book.close()
    assert len(run['template_mapping']) == 2
    assert all(report['form'] == 'L8' and report['complete'] and report['unmapped_fields'] ==
               report['unmapped_rows'] == report['unmapped_tables'] == 0 for report in run['template_mapping'])
    summary = load_workbook(tmp_path / 'resume.xlsx')
    try:
        assert 'Audit' in summary.sheetnames
        audit = [row for row in summary['Audit'].values if row[0] == 'Pemetaan template']
        assert len(audit) == 2 and all(row[1:5] == (0, 0, 0, 'COMPLETE') for row in audit)
        targets = [cell.hyperlink.target for sheet in summary for row in sheet for cell in row if cell.hyperlink]
        assert all(any(name in target for target in targets) for name in files)
        assert all(not target.startswith(('http:', 'https:', 'file:', '/', '\\')) for target in targets)
        assert all(record['capture_complete'] is True for record in run['records'])
    finally:
        summary.close()
