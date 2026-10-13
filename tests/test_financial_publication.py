"""Synthetic declared-PDF custody contracts, executed only on GitHub runners."""
from copy import deepcopy
import hashlib
import json
import shutil
import subprocess
import sys

import pytest
from openpyxl import load_workbook

import coretax
from spt_desktop import output
from test_split_output import fixture_run

PDF = b'%PDF-1.7\n1 0 obj\n(ORIGINAL SYNTHETIC FIXTURE)\nendobj\n%%EOF\n'
CHANGED_PDF = PDF.replace(b'ORIGINAL', b'MUTATION')


def declare_pdf(record, folder, data=PDF):
    record['scope'] = ['0000000000000001', 2025, 'period']
    key = hashlib.sha256(record['record_id'].encode()).hexdigest()[:32]
    url = 'https://coretax.synthetic.invalid/pdf/' + record['record_id']
    relative = 'attachments/' + key + '/' + hashlib.sha256(url.encode()).hexdigest() + '.pdf'
    record['attachments'] = [{'label': 'Laporan Keuangan', 'relative_path': relative,
                              'status': 'COMPLETE', 'bytes': len(data),
                              'hash': hashlib.sha256(data).hexdigest()}]
    path = folder / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path


def financial_run(folder, count=1, data=PDF):
    run = fixture_run()
    run['records'] = run['records'][:count]
    run['inventory_complete'] = True
    run['complete'] = True
    for record in run['records']:
        declare_pdf(record, folder, data)
    return run


def declarations(run):
    return [{**item, 'record_id': record['record_id'], 'scope': record['scope']}
            for record in run['records'] for item in record.get('attachments', [])
            if item['status'] == 'COMPLETE']


def assert_audit_targets(folder, name, expected):
    workbook = load_workbook(folder / name)
    try:
        if name == 'resume.xlsx':
            links = [cell.hyperlink.target for row in workbook['Audit'] for cell in row
                     if cell.hyperlink and cell.hyperlink.target.endswith('.pdf')]
        else:
            assert {'Ringkasan', 'Daftar SPT', 'Audit'}.isdisjoint(workbook.sheetnames)
            proof = json.loads(workbook.properties.identifier)
            assert proof['npwp'] == '0000000000000001'
            assert proof['year'] == 2025 and proof['year_basis'] == 'period'
            links = [item['relative_path'] for record in proof['records']
                     for item in record.get('attachments', []) if item['status'] == 'COMPLETE']
        assert set(links) == set(expected) and len(links) == len(expected)
        for relative in links:
            assert (folder / relative).read_bytes() == expected[relative]
    finally:
        workbook.close()


def test_cli_declares_bound_pdfs_separately_and_preserves_private_bytes(tmp_path):
    source = tmp_path / 'private-job'
    run = financial_run(source, count=2)
    before = {item['relative_path']: (source / item['relative_path']).read_bytes()
              for item in declarations(run)}
    assert coretax.save_output(run, source) == source / 'resume.xlsx'
    assert run.get('output_attachments') == declarations(run), 'Bound PDF declarations are missing'
    saved = json.loads((source / 'run.json').read_text(encoding='utf-8'))
    assert saved['output_attachments'] == declarations(run)
    assert set(run['output_files']) == {'resume.xlsx', 'SPT_Badan.xlsx'}
    assert {relative: (source / relative).read_bytes() for relative in before} == before
    for name in run['output_files']:
        assert_audit_targets(source, name, before)
    coretax.verify_outputs(run, source)


@pytest.mark.parametrize('invalid', ['missing', 'digest', 'size', 'not-pdf', 'directory'])
def test_invalid_pdf_is_rejected_before_any_workbook_write(tmp_path, invalid):
    source = tmp_path / 'private-job'
    run = financial_run(source, data=b'SYNTHETIC NOT A PDF' if invalid == 'not-pdf' else PDF)
    item = run['records'][0]['attachments'][0]
    path = source / item['relative_path']
    if invalid == 'missing':
        path.unlink()
    elif invalid == 'digest':
        item['hash'] = '0' * 64
    elif invalid == 'size':
        item['bytes'] += 1
    elif invalid == 'directory':
        path.unlink()
        path.mkdir()
    sentinel = source / 'keep-private.txt'
    sentinel.write_bytes(b'SYNTHETIC PRIVATE MATERIAL; KEEP EXACTLY')
    with pytest.raises((ValueError, OSError), match='(?i)pdf|attachment|financial|file'):
        coretax.save_output(run, source)
    assert not list(source.rglob('*.xlsx'))
    assert not (source / 'run.json').exists()
    assert sentinel.read_bytes() == b'SYNTHETIC PRIVATE MATERIAL; KEEP EXACTLY'


def test_verify_outputs_rechecks_declared_pdf_bytes_after_export(tmp_path):
    source = tmp_path / 'private-job'
    run = financial_run(source)
    coretax.save_output(run, source)
    item = run['records'][0]['attachments'][0]
    (source / item['relative_path']).write_bytes(CHANGED_PDF)
    assert len(CHANGED_PDF) == item['bytes']
    with pytest.raises(ValueError, match='(?i)pdf|attachment|financial|hash|digest'):
        coretax.verify_outputs(run, source)


def test_desktop_publishes_only_selected_declared_pdfs_with_relative_links(tmp_path):
    source = tmp_path / 'private-job'
    run = financial_run(source)
    foreign = deepcopy(fixture_run()['records'][1])
    unselected = declare_pdf(foreign, source, CHANGED_PDF)
    undeclared = source / 'attachments' / 'unrelated-private.pdf'
    undeclared.write_bytes(b'SYNTHETIC PRIVATE FILE; NEVER PUBLISH')
    checkpoint = source / 'private-unselected-checkpoint.json'
    checkpoint.write_text('{"synthetic": true}', encoding='utf-8')
    coretax.save_output(run, source)
    downloads = tmp_path / 'Downloads'
    final = output.publish_workbooks(run, source, downloads, 'synthetic-financial')
    expected = {item['relative_path']: PDF for item in declarations(run)}
    assert {path.relative_to(final).as_posix(): path.read_bytes()
            for path in final.rglob('*.pdf')} == expected, 'Selected PDF bytes were not published'
    assert not list(downloads.rglob('*.json'))
    assert unselected.read_bytes() == CHANGED_PDF
    assert undeclared.read_bytes() == b'SYNTHETIC PRIVATE FILE; NEVER PUBLISH'
    assert checkpoint.is_file()
    for name in run['output_files']:
        assert_audit_targets(final, name, expected)
    coretax.verify_outputs(run, final)


def test_partial_metadata_and_audit_have_no_publishable_file(tmp_path):
    source = tmp_path / 'private-job'
    run = financial_run(source)
    record = run['records'][0]
    private_pdf = source / record['attachments'][0]['relative_path']
    record['attachments'] = [{'label': 'Laporan Keuangan', 'relative_path': '',
                              'status': 'PARTIAL', 'bytes': 0, 'hash': '',
                              'reason': 'SYNTHETIC unsupported read control'}]
    record['capture_complete'] = False
    run['complete'] = False
    coretax.save_output(run, source)
    assert run.get('output_attachments') == [], 'PARTIAL must not declare a publishable file'
    saved = json.loads((source / 'run.json').read_text(encoding='utf-8'))
    assert saved['records'][0]['attachments'] == record['attachments']
    final = output.publish_workbooks(run, source, tmp_path / 'Downloads', 'synthetic-partial')
    assert not list(final.rglob('*.pdf'))
    assert private_pdf.read_bytes() == PDF
    workbook = load_workbook(final / 'resume.xlsx')
    try:
        rows = list(workbook['Audit'].iter_rows(values_only=True))
        assert any(record['record_id'] in row and 'Laporan Keuangan' in row
                   and 'PARTIAL' in row for row in rows)
        assert not [cell for row in workbook['Audit'] for cell in row
                    if cell.hyperlink and cell.hyperlink.target.endswith('.pdf')]
    finally:
        workbook.close()


def test_extra_unbound_pdf_declaration_blocks_downloads_creation(tmp_path):
    source = tmp_path / 'private-job'
    run = financial_run(source)
    foreign = deepcopy(fixture_run()['records'][1])
    declare_pdf(foreign, source, CHANGED_PDF)
    coretax.save_output(run, source)
    run['output_attachments'] = declarations(run) + [
        {**foreign['attachments'][0], 'record_id': foreign['record_id'], 'scope': foreign['scope']}]
    downloads = tmp_path / 'Downloads'
    with pytest.raises(ValueError, match='(?i)pdf|attachment|financial|scope|declaration'):
        output.publish_workbooks(run, source, downloads, 'synthetic-unbound')
    assert not downloads.exists()
    assert (source / 'resume.xlsx').is_file()


def test_pdf_consent_names_selected_replacements_and_preserves_unselected_bytes(tmp_path):
    old_source = tmp_path / 'first-private-job'
    old_run = financial_run(old_source, count=2)
    coretax.save_output(old_run, old_source)
    downloads = tmp_path / 'Downloads'
    final = output.publish_workbooks(old_run, old_source, downloads, 'synthetic-first')
    first_pdf, untouched_pdf = [final / item['relative_path'] for item in declarations(old_run)]
    assert first_pdf.read_bytes() == untouched_pdf.read_bytes() == PDF
    old_bytes = {path.relative_to(final).as_posix(): path.read_bytes()
                 for path in final.rglob('*') if path.is_file()}
    new_source = tmp_path / 'second-private-job'
    selected = deepcopy(old_run)
    selected['records'] = selected['records'][:1]
    selected.pop('output_attachments', None)
    declare_pdf(selected['records'][0], new_source, CHANGED_PDF)
    coretax.save_output(selected, new_source)
    snapshot = output.existing_outputs(selected, downloads)
    selected_pdf = selected['records'][0]['attachments'][0]['relative_path']
    assert selected_pdf in snapshot['selected_files']
    assert selected_pdf in snapshot['files'], 'Consent must show the PDF that will be replaced'
    with pytest.raises(FileExistsError, match='overwrite|consent'):
        output.publish_workbooks(selected, new_source, downloads, 'synthetic-decline',
                                 overwrite=False, expected_outputs=snapshot)
    assert {path.relative_to(final).as_posix(): path.read_bytes()
            for path in final.rglob('*') if path.is_file()} == old_bytes
    result = output.publish_workbooks(selected, new_source, downloads, 'synthetic-replace',
                                      overwrite=True, expected_outputs=snapshot)
    assert result == final
    assert first_pdf.read_bytes() == CHANGED_PDF
    assert untouched_pdf.read_bytes() == PDF
    assert_audit_targets(final, 'resume.xlsx', {selected_pdf: CHANGED_PDF})
    coretax.verify_outputs(selected, final, allow_extra=True)


def test_attachment_directory_cannot_redirect_to_another_scope(tmp_path):
    source = tmp_path / 'private-job'
    run = financial_run(source)
    coretax.save_output(run, source)
    outside = tmp_path / 'outside-private-scope'
    outside.mkdir()
    sentinel = outside / 'keep.txt'
    sentinel.write_bytes(b'SYNTHETIC OUTSIDE DATA; KEEP EXACTLY')
    redirected = source / 'attachments'
    shutil.rmtree(redirected)
    if sys.platform == 'win32':
        subprocess.run(['cmd', '/d', '/c', 'mklink', '/J', str(redirected), str(outside)],
                       check=True, capture_output=True, text=True, timeout=10)
    else:
        redirected.symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError, match='(?i)pdf|attachment|path|reparse|symbolic|redirect'):
        output.publish_workbooks(run, source, tmp_path / 'Downloads', 'synthetic-junction')
    assert {path.name: path.read_bytes() for path in outside.iterdir()} == {
        'keep.txt': b'SYNTHETIC OUTSIDE DATA; KEEP EXACTLY'}
    assert not (tmp_path / 'Downloads').exists()


@pytest.mark.skipif(sys.platform != 'win32', reason='Real Windows sharing/native publication contract')
def test_staged_pdf_mutation_cannot_advertise_complete_export(tmp_path):
    source = tmp_path / 'private-job'
    run = financial_run(source)
    coretax.save_output(run, source)
    downloads = tmp_path / 'Downloads'
    attempted = []

    def corrupt_staged_pdf():
        if attempted:
            return
        for path in downloads.rglob('*.pdf'):
            if any(parent.name.startswith('.spt-') for parent in path.parents):
                attempted.append(path)
                # Native sharing must deny this write, or publication must
                # refuse the changed bytes before advertising a final index.
                path.write_bytes(CHANGED_PDF)
                return

    with pytest.raises((OSError, ValueError), match='(?i)pdf|attachment|hash|digest|changed|denied|sharing|access'):
        output.publish_workbooks(run, source, downloads, 'synthetic-stage-race',
                                 check_cancel=corrupt_staged_pdf)
    assert attempted, 'The real staged-PDF publication barrier was not exercised'
    assert (source / run['records'][0]['attachments'][0]['relative_path']).read_bytes() == PDF
    assert not [path for path in downloads.rglob('resume.xlsx')
                if not any(parent.name.startswith('.spt-') for parent in path.parents)]


def test_pdf_changed_while_consent_is_open_is_never_authorized(tmp_path):
    source = tmp_path / 'old-private-job'
    run = financial_run(source)
    coretax.save_output(run, source)
    downloads = tmp_path / 'Downloads'
    final = output.publish_workbooks(run, source, downloads, 'synthetic-old')
    replacement = tmp_path / 'new-private-job'
    selected = financial_run(replacement)
    coretax.save_output(selected, replacement)
    snapshot = output.existing_outputs(selected, downloads)
    relative = declarations(run)[0]['relative_path']
    assert (final / relative).is_file(), 'The old selected PDF was not published'
    (final / relative).write_bytes(CHANGED_PDF)
    before = {path.relative_to(final).as_posix(): path.read_bytes()
              for path in final.rglob('*') if path.is_file()}
    with pytest.raises(ValueError, match='(?i)pdf|attachment|hash|digest|changed|consent'):
        output.publish_workbooks(selected, replacement, downloads, 'synthetic-stale-pdf',
                                 overwrite=True, expected_outputs=snapshot)
    assert {path.relative_to(final).as_posix(): path.read_bytes()
            for path in final.rglob('*') if path.is_file()} == before
    assert (replacement / relative).read_bytes() == PDF


@pytest.mark.parametrize('format_name', ['legacy', 'forms'])
@pytest.mark.parametrize('tamper', ['record_id', 'digest'])
def test_old_pdf_overwrite_requires_independent_physical_declaration_binding(tmp_path, tamper, format_name):
    source = tmp_path / 'old-private-job'
    run = financial_run(source)
    coretax.save_output(run, source)
    downloads = tmp_path / 'Downloads'
    final = output.publish_workbooks(run, source, downloads, 'synthetic-old-audit')
    assert all((final / item['relative_path']).is_file() for item in declarations(run)), (
        'The physical old PDF required by its Audit was not published')
    path = final / 'SPT_Badan.xlsx'
    if format_name == 'legacy':
        from coretax_exporter import export_workbook
        export_workbook(run, path)
    workbook = load_workbook(path)
    try:
        if format_name == 'legacy':
            for row in workbook['Audit'].iter_rows(min_row=5, max_col=6):
                if row[5].hyperlink and row[5].hyperlink.target.endswith('.pdf'):
                    row[0 if tamper == 'record_id' else 4].value = (
                        'SYNTHETIC FOREIGN RETURN' if tamper == 'record_id' else '0' * 64)
                    break
            else:
                pytest.fail('The real old physical PDF Audit row was not exercised')
        else:
            assert isinstance(workbook.properties.identifier, str), 'Nonvisual PDF proof is missing'
            proof = json.loads(workbook.properties.identifier)
            record = proof['records'][0]
            assert record['record_id'] == run['records'][0]['record_id']
            assert record['attachments'] == run['records'][0]['attachments']
            if tamper == 'record_id':
                record['record_id'] = 'SYNTHETIC FOREIGN RETURN'
            else:
                record['attachments'][0]['hash'] = '0' * 64
            workbook.properties.identifier = json.dumps(proof, ensure_ascii=False)
        workbook.save(path)
    finally:
        workbook.close()
    before = {path.relative_to(final).as_posix(): path.read_bytes()
              for path in final.rglob('*') if path.is_file()}
    replacement = tmp_path / 'replacement-private-job'
    selected = financial_run(replacement)
    coretax.save_output(selected, replacement)
    with pytest.raises(ValueError, match='(?i)pdf|attachment|audit|scope|hash|digest'):
        output.existing_outputs(selected, downloads)
    assert {path.relative_to(final).as_posix(): path.read_bytes()
            for path in final.rglob('*') if path.is_file()} == before


@pytest.mark.skipif(sys.platform != 'win32', reason='Native directory rename byte-custody boundary')
def test_pdf_mutated_after_actual_native_directory_move_is_quarantined(tmp_path, monkeypatch):
    source = tmp_path / 'private-job'
    run = financial_run(source)
    coretax.save_output(run, source)
    downloads = tmp_path / 'Downloads'
    relative = declarations(run)[0]['relative_path']
    real_move = output._move_no_replace
    attempted = []

    def mutate_after_move(original, destination):
        # Intercept the real move, not a mock pretending publication succeeded.
        is_staging_tree = original.is_dir() and original.name.startswith('.spt-')
        result = real_move(original, destination)
        if is_staging_tree and not attempted:
            path = destination / relative
            if path.is_file():
                attempted.append(path)
                path.write_bytes(CHANGED_PDF)
        return result

    monkeypatch.setattr(output, '_move_no_replace', mutate_after_move)
    with pytest.raises((OSError, ValueError, RuntimeError),
                       match='(?i)pdf|attachment|hash|digest|changed|recovery|sharing|access'):
        output.publish_workbooks(run, source, downloads, 'synthetic-rename-race')
    assert attempted, 'The real native directory-move PDF boundary was not exercised'
    assert (source / relative).read_bytes() == PDF
    assert not [path for path in downloads.rglob('resume.xlsx')
                if not any(parent.name.startswith('.spt-') for parent in path.parents)]
    assert any(path.read_bytes() == CHANGED_PDF for path in downloads.rglob('*.pdf')), (
        'The actual changed bytes must remain recoverable, not be silently deleted')
