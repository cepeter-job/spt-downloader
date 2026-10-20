"""Synthetic clean-form, percentage scale and shifted spacer contracts."""
from copy import copy, deepcopy
from decimal import Decimal
import pytest
from openpyxl import load_workbook
import coretax_exporter as exporter
from support_templates import section, source_run
from test_unifikasi_rate_layout import HEADERS


def export(tmp_path, sections):
    run = source_run('PPh Unifikasi', sections)
    run['records'][0]['capture_complete'] = all(
        table.get('complete', True) for item in sections for table in item.get('tables', []))
    before = deepcopy(run)
    reports = []
    paths = exporter.export_split_workbooks(run, tmp_path, mapping_reports=reports)
    assert run == before
    path, = [name for name in paths if name != 'resume.xlsx']
    return load_workbook(tmp_path / path), reports


def table(group='BPPU', rate='2,50', count=1):
    rows = [[str(i+1), '0000000000000002', 'SYNTHETIC NAME', 'SYNTHETIC-DOC',
             '02-01-2025', 'PPh Pasal 23', 'SYNTHETIC-CODE', 'SYNTHETIC OBJECT',
             '1.234,50', rate, '30,86', '', 'SYNTHETIC EXTRA',
             '0000000000000002000001', 'Normal', 'SYNTHETIC-KAP'] for i in range(count)]
    return {'name':'Tabel 1', 'group':'TABEL I. '+group, 'headers':HEADERS,
            'rows':rows, 'body_row_count':count, 'source_total':count,
            'complete':True, 'footer':[], 'footer_spans':[]}


@pytest.mark.parametrize('group,first', [('BPPU',5),('BPNR',11)])
@pytest.mark.parametrize('literal,expected', [('2,00','0.02'),('2,50','0.025'),('0','0'),('100','1'),
                                             ('12,3456789012345','0.123456789012345')])
def test_numeric_percent_scale_and_unchanged_adjacent_values(tmp_path,group,first,literal,expected):
    book,reports=export(tmp_path,[section('DAFTAR-I',tables=[table(group,literal)])])
    try:
        sheet=book['DAFTAR-I']; cell=sheet.cell(first,10)
        assert cell.data_type=='n'
        assert Decimal(str(cell.value))==Decimal(expected)
        assert cell.number_format=='0.00%'
        assert sheet.cell(first,9).value==1234.5
        assert sheet.cell(first,11).value==30.86
        assert sheet.cell(first,2).value=='0000000000000002'
        assert reports[0]['complete']
    finally:book.close()


@pytest.mark.parametrize('literal',['abc','2.5','2,50%','-1','101','0,1234567890123456',
                                            '99,999999999999999999999999999999'])
def test_unsupported_rate_preserves_literal_and_partial(tmp_path,literal):
    book,reports=export(tmp_path,[section('DAFTAR-I',tables=[table(rate=literal)])])
    try:
        assert book['DAFTAR-I']['J5'].value==literal
        assert book['DAFTAR-I']['J5'].data_type=='s'
        assert reports[0]['complete'] is False
        assert reports[0]['unmapped_rows']>=1
    finally:book.close()


def test_blank_rate_not_zero(tmp_path):
    book,_=export(tmp_path,[section('DAFTAR-I',tables=[table(rate='')])])
    try:assert book['DAFTAR-I']['J5'].value is None
    finally:book.close()


@pytest.mark.parametrize('count',[0,43])
def test_grown_rates_and_blank_spacer_border(tmp_path,count):
    book,reports=export(tmp_path,[section('DAFTAR-I',tables=[table(count=count),table('BPNR',count=0)])])
    try:
        sheet=book['DAFTAR-I']; offset=max(0,count-4)
        for row in list(range(5,9+offset))+list(range(11+offset,14+offset)):
            assert sheet.cell(row,10).number_format=='0.00%'
        for row in range(5,count+5):
            assert sheet.cell(row,10).value==0.025
            assert sheet.cell(row,11).value==30.86
        spacer=sheet.cell(14+offset,10)
        assert spacer.value is None
        assert copy(spacer.border)==copy(sheet.cell(14+offset,9).border)
        assert all(getattr(spacer.border,e).style is None for e in ('left','right','top','bottom'))
        assert sheet.max_row==18+offset
        assert reports[0]['complete']
    finally:book.close()


@pytest.mark.parametrize('name,last',[('DAFTAR-I',18),('DAFTAR-II',18),('LAMPIRAN-I',10)])
@pytest.mark.parametrize('populated',[False,True])
def test_clean_appendices_preserve_unresolved_audit(tmp_path,name,last,populated):
    source=section(name,fields=[{'field':'SyntheticUnknown','label':'Unbound','value':'SYNTHETIC FIELD'}],
                   tables=[{'name':'UNKNOWN','headers':['Synthetic'],
                            'rows':[['SYNTHETIC ROW']] if populated else [],'complete':False}],
                   text='SYNTHETIC CONTEXT')
    book,reports=export(tmp_path,[source])
    try:
        sheet=book[name]
        assert sheet.max_row==last
        values=[c.value for row in sheet for c in row]
        for marker in ('Data sumber tambahan / belum dipetakan','SYNTHETIC FIELD','SYNTHETIC ROW','SYNTHETIC CONTEXT'):
            assert marker not in values
        assert reports[0]['complete'] is False
        assert reports[0]['unmapped_fields']==1
        assert reports[0]['unmapped_tables']==1
        assert reports[0]['unmapped_rows']==int(populated)
    finally:book.close()
