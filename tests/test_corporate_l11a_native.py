"""Original-byte current-only L11-A integration contracts; runners only."""
from copy import deepcopy
import json

import pytest
from openpyxl import load_workbook

import coretax
import coretax_browser as driver
from coretax_json import JsonSafetyError, adapt_corporate_detail
import support_json_corpus as corpus
from test_corporate_l7_native import bindings as old_bindings, selected
from test_corporate_l11c_native import L11C_PINS
from test_corporate_l1d_native import DETAIL, packet as fresh_packet
from test_corporate_l3_native import packet
from test_corporate_l11a_projection import (ROOT, HEADERS, TYPES, LABELS, RECAP, ENDPOINTS,
                                          CAPTIONS, GROUP as RECAP_GROUP, inputs)


BASE = corpus.ORIGIN + '/referencedata/api/currentreferencedata/id-ID/'
REFERENCES = {kind: BASE + kind for kind in TYPES}
PAIR = BASE + 'DEDUCTION_METHOD,TYPE_OF_FULFILLMENT'
GRIDS = {name: corpus.ORIGIN + '/returnsheet/api/loadcit/' + name for name in ENDPOINTS}
OWNERS = ['rshshr-citr-l11a-grid1', 'rshshr-citr-l11a-grid2', 'rshshr-citr-l11a-grid3',
          'rsh-citr-l11a-grid4', 'rshshr-citr-l11a-grid5']
GROUPS = [
    ['I. DAFTAR NOMINATIF BIAYA PROMOSI DAN PENJUALAN, SERTA PENGGANTIAN ATAU IMBALAN DALAM BENTUK NATURA DAN/ATAU KENIKMATAN'],
    ['II. DAFTAR NOMINATIF BIAYA ENTERTAINMENT'],
    ['III. DAFTAR PIUTANG YANG NYATA-NYATA TIDAK DAPAT DITAGIH'],
    ['IV. RINCIAN BAGI WAJIB PAJAK PEMBERI NATURA DAN/ATAU KENIKMATAN',
     'IV.A DAFTAR SARANA DAN FASILITAS SERTA PENYUSUTANNYA'],
    ['V. DAFTAR DEBITUR NON-PERFORMING LOAN']]


def request(first):
    return {'ReturnSheetRecordId': corpus.RECORD_ID, 'First': first, 'Rows': 10,
            'SortField': '', 'SortOrder': 1, 'Filters': [], 'LanguageId': 'id-ID'}


def envelope(model='Normal', form=None, *, previous=None):
    root = corpus.current_form()
    root['MainFormData']['IsDeclareEntertainmentExpense'] = 1
    root[ROOT] = inputs()[0] if form is None else form
    return {'IsSuccessful': True, 'Payload': {'TIN': corpus.NPWP, 'PeriodYear': 2025,
        'IsDraft': True, 'ReturnSheetStatusCode': 'SUBMITTED', 'IsDeclareEntertainmentExpense': 1,
        'ReturnSheetModel': 'NORMAL' if model == 'Normal' else 'AMENDMENT',
        'PreviousFormData': previous, 'LastSubmittedFormData': None,
        'CurrentFormData': corpus.wire(root)}}


def catalog(kind):
    return {'ReferenceDataType': kind, 'language': 'id-ID', 'Details': [
        {'Code': code, 'CodeDescription': text, 'ValidFrom': None, 'ValidTo': None}
        for code, text in LABELS[kind].items()]}


def transports(transport='singletons'):
    if transport == 'singletons':
        return [(url, [kind]) for kind, url in REFERENCES.items()]
    return [(REFERENCES[kind], [kind]) for kind in TYPES[:3]] + [
        (PAIR, TYPES[3:] if transport == 'pair' else list(reversed(TYPES[3:])))]


def packets(model='Normal', transport='singletons', *, form=None, pages=None, previous=None):
    pages = inputs()[1] if pages is None else pages
    raw = [fresh_packet(envelope(model, form, previous=previous))]
    for endpoint, values in pages.items():
        for page in values:
            raw.append(packet({'IsSuccessful': True, 'Payload': {
                'Data': page['data'], 'TotalRecords': page['total']}}, len(raw) + 1,
                GRIDS[endpoint], request(page['offset'])))
    for url, kinds in transports(transport):
        raw.append(packet({'IsSuccessful': True, 'Payload': [catalog(kind) for kind in kinds]},
                          len(raw) + 1, url, None, method='GET'))
    return raw


def bindings(transport='singletons', *, pages=None):
    pages = inputs()[1] if pages is None else pages
    sequence = 2 + sum(len(values) for values in pages.values())
    refs = [{'endpoint': url, 'method': 'GET', 'sequences': [sequence + index]}
            for index, (url, _) in enumerate(transports(transport))]
    return [old_bindings()[0], {'name': 'L11-A', 'paths': [[ROOT]], 'grids': [],
        'responses': [{'name': 'L11-A/' + endpoint, 'endpoint': GRIDS[endpoint],
                       'requests': [request(page['offset']) for page in values]}
                      for endpoint, values in pages.items()],
        'projection': {'kind': 'corporate-l11a-v1', 'builds': dict(L11C_PINS), 'route': DETAIL,
            'locale': 'id-ID', 'timezone': 'Asia/Jakarta', 'references': refs,
            'tables': [{'component': owner, 'headers': list(HEADERS[index]), 'footer_spans': [],
                'source_structure': {'form': 'L11-A', 'component': owner, 'groups': list(GROUPS[index])}}
                       for index, owner in enumerate(OWNERS)]}}]


def replace_payload(raw, index, payload, *, status=200):
    context = raw[index]['context']
    raw[index] = packet({'IsSuccessful': status == 200, 'Payload': payload}, context['sequence'],
        context['url'], context.get('request_json'), method=context['method'], status=status)


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
@pytest.mark.parametrize('transport', ['singletons', 'pair', 'pair-reordered'])
@pytest.mark.parametrize('order', ['source', 'completion'])
def test_l11a_original_packets_keep_five_tables_recap_and_compact_saved_form(tmp_path, model, transport, order):
    raw, observed, record = packets(model, transport), bindings(transport), selected(model)
    if order == 'completion': raw = [raw[0], *reversed(raw[1:])]
    before = deepcopy((raw, observed, record))
    result = adapt_corporate_detail(iter(raw), record, corpus.SCOPE, observed)
    assert result['eligible'] and result['coverage']['complete'], result['reason']
    assert (raw, observed, record) == before
    section, = [s for s in result['sections'] if s['name'] == 'L11-A']
    assert [t['body_row_count'] for t in section['tables']] == [12, 2, 2, 2, 2]
    assert all(t['footer'] == t['footer_spans'] == [] for t in section['tables'])
    assert [t['source_structure'] for t in section['tables']] == [
        table['source_structure'] for table in observed[1]['projection']['tables']]
    assert section['tables'][0]['rows'][0][0] == section['tables'][0]['rows'][10][0] == '0007'
    assert section['tables'][1]['rows'][0][5] == 'SYNTHETIC PERSON'
    displays = ['SYNTHETIC LOCATION', '0001', '02-12-2025', '0002', '02-12-2025',
                '1', '2', '3', '4', '5', '6', '21']
    assert section['fields'] == [{'group': RECAP_GROUP, 'field': key, 'label': caption, 'value': value}
                                 for key, caption, value in zip(RECAP, CAPTIONS, displays)]
    evidence, = [p['values'] for p in result['provenance']['projections'] if p['form'] == 'L11-A']
    assert any(e['response_sequence'] == 3 and e['container'] == ['Payload']
               and e['path'] == ['Data', 0, 'Index'] and e['type'] == 'string' and e['value'] == '0007'
               for e in evidence)
    assert any(e['response_sequence'] == 1 and e['container'] == ['Payload', 'CurrentFormData']
               and e['path'] == [ROOT, 'Recapitulation', 'TotalAmountofCosts']
               and e['type'] == 'number' and e['value'] == '21' for e in evidence)
    raw_values = ['SYNTHETIC LOCATION', '0001', '2025-12-01T23:30:00.000Z', '0002', '2025-12-02',
                  '1', '2', '3', '4', '5', '6', '21']
    for index, (key, value) in enumerate(zip(RECAP, raw_values)):
        assert any(e['response_sequence'] == 1 and e['container'] == ['Payload', 'CurrentFormData']
                   and e['path'] == [ROOT, 'Recapitulation', key] and e['value'] == value
                   and e['type'] == ('string' if index < 5 else 'number')
                   and e.get('destination') == {'field': key} for e in evidence)
    for ref in raw:
        if ref['context']['url'] not in {url for url, _ in transports(transport)}: continue
        catalogs = json.loads(ref['body'])['Payload']
        for index, source in enumerate(catalogs):
            assert any(e['response_sequence'] == ref['context']['sequence'] and e['container'] == ['Payload']
                       and e['path'] == [index, 'Details', 0, 'CodeDescription']
                       and e['value'] == source['Details'][0]['CodeDescription'] for e in evidence)
    record.update(capture_complete=True, sections=result['sections'])
    run = {'npwp': corpus.NPWP, 'year': 2025, 'year_basis': 'period', 'inventory_complete': True,
           'complete': True, 'errors': [], 'records': [record]}
    fields_before = deepcopy(section['fields'])
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    assert section['fields'] == fields_before
    assert next(s for s in run['records'][0]['sections'] if s['name'] == 'L11-A')['fields'] == fields_before
    name, = [name for name in run['output_files'] if name != 'resume.xlsx']
    assert model in name
    book = load_workbook(tmp_path / name)
    try:
        sheet = book['L11-A']
        values = [cell.value for row in sheet for cell in row]
        assert values.count('0007') == 14 and values.count('0005') == 2
        assert values.count('00000002') == 16 and values.count('00001') == 12
        assert values.count('0025') == values.count('SYNTHETIC PERSON') == 2
        assert 'JUMLAH' not in values and 'Data sumber tambahan / belum dipetakan' not in values
        assert 'A1:D1' in {str(area) for area in sheet.merged_cells.ranges}
        assert sheet.max_row < 60
        money = [row[4] for row in sheet if len(row) >= 6 and row[5].value == 'SYNTHETIC PERSON']
        assert len(money) == 2 and all(c.value == 7 and 'Rp' in c.number_format and '0.00' not in c.number_format for c in money)
    finally: book.close()
    mapping, = [r for r in run['template_mapping'] if r['form'] == 'L11-A']
    assert not mapping['complete'] and mapping['unmapped_fields'] == 12
    assert mapping['unmapped_tables'] == mapping['unmapped_rows'] == 0
    assert not next(r for r in run['template_mapping'] if r['form'] == 'Induk')['complete']


@pytest.mark.parametrize('transport', ['singletons', 'pair', 'pair-reordered'])
@pytest.mark.parametrize('previous', [None, 'null', '{}', '{"L11AFormData":null}'])
def test_l11a_explicit_null_or_absent_previous_form_and_empty_inventories_remain_observed(transport, previous):
    form, pages = inputs()
    form['Facilities'] = []
    pages = {endpoint: [{'offset': 0, 'rows': 10, 'total': 0, 'data': []}] for endpoint in pages}
    raw = packets(transport=transport, form=form, pages=pages, previous=previous)
    observed = bindings(transport, pages=pages)
    before = deepcopy((raw, observed))
    result = adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, observed)
    assert result['eligible'], result['reason']
    section = next(s for s in result['sections'] if s['name'] == 'L11-A')
    assert all(t['rows'] == t['footer'] == [] and t['body_row_count'] == 0 for t in section['tables'])
    assert len(section['fields']) == 12 and (raw, observed) == before


@pytest.mark.parametrize('defect', ['pin', 'route', 'locale', 'timezone', 'owner', 'group', 'nested-group',
    'header', 'footer', 'missing-table', 'induk-after', 'missing-lifetime', 'missing-previous',
    'previous-form', 'previous-empty-object', 'previous-list', 'missing-flag', 'missing-current-flag',
    'disabled-flag', 'boolean-flag', 'model', 'status', 'missing-page', 'missing-grid', 'result-only',
    'index-schema', 'failed-page', 'missing-reference', 'reference-sequence', 'pair-extra-catalog',
    'pair-duplicate-type', 'pair-missing-type', 'reference-failed', 'identical-overlap',
    'sort', 'sort-order-bool', 'filters', 'language-id', 'extra-language', 'rows', 'missing-request-key'])
def test_unproved_l11a_lifecycle_structure_or_dependency_keeps_whole_record_fallback(defect):
    transport = 'pair'
    raw, observed = packets(transport=transport), bindings(transport)
    proof = observed[1]['projection']
    if defect == 'pin': proof['builds'][next(iter(proof['builds']))] = '0' * 64
    elif defect == 'route': proof['route'] = DETAIL.replace('?view=true', '?history=true')
    elif defect in ('locale', 'timezone'): proof[defect] = 'SYNTHETIC UNPROVED'
    elif defect == 'owner': proof['tables'][0]['component'] = 'synthetic-foreign-owner'
    elif defect == 'group': proof['tables'][0]['source_structure']['groups'] = GROUPS[1]
    elif defect == 'nested-group': proof['tables'][3]['source_structure']['groups'].pop(0)
    elif defect == 'header': proof['tables'][0]['headers'][0] = 'SYNTHETIC UNPROVED'
    elif defect == 'footer': proof['tables'][4]['footer_spans'] = [[8]]
    elif defect == 'missing-table': proof['tables'].pop()
    elif defect == 'induk-after': observed.reverse()
    elif defect == 'missing-lifetime':
        for p in raw: del p['context']['component_lifetime']
    elif defect in ('missing-previous', 'previous-form', 'previous-empty-object', 'previous-list',
                    'missing-flag', 'missing-current-flag', 'disabled-flag', 'boolean-flag', 'model', 'status'):
        value = envelope()
        payload = value['Payload']
        root = json.loads(payload['CurrentFormData'])
        if defect == 'missing-previous': del payload['PreviousFormData']
        elif defect.startswith('previous-'):
            payload['PreviousFormData'] = corpus.wire({'L11AFormData': {
                'previous-form': inputs()[0], 'previous-empty-object': {}, 'previous-list': []}[defect]})
        elif defect == 'missing-flag': del payload['IsDeclareEntertainmentExpense']
        elif defect == 'missing-current-flag': del root['MainFormData']['IsDeclareEntertainmentExpense']
        elif defect in ('disabled-flag', 'boolean-flag'):
            payload['IsDeclareEntertainmentExpense'] = root['MainFormData']['IsDeclareEntertainmentExpense'] = (0 if defect == 'disabled-flag' else True)
        elif defect == 'model': payload['ReturnSheetModel'] = 'AMENDMENT'
        else: payload['ReturnSheetStatusCode'] = 'DRAFT'
        payload['CurrentFormData'] = corpus.wire(root)
        raw[0] = fresh_packet(value)
    elif defect == 'missing-page': raw.pop(2)
    elif defect == 'missing-grid':
        raw = [p for p in raw if p['context']['url'] != GRIDS[ENDPOINTS[1]]]
    elif defect in ('result-only', 'index-schema', 'failed-page'):
        value = json.loads(raw[1]['body'])['Payload']
        if defect == 'result-only': value['Result'] = {'Data': value.pop('Data')}
        elif defect == 'index-schema': value['Index'] = None
        replace_payload(raw, 1, value, status=500 if defect == 'failed-page' else 200)
    elif defect == 'missing-reference': raw.pop()
    elif defect == 'reference-sequence': proof['references'][0]['sequences'] = [999]
    elif defect in ('pair-extra-catalog', 'pair-duplicate-type', 'pair-missing-type', 'reference-failed'):
        value = json.loads(raw[-1]['body'])['Payload']
        if defect == 'pair-extra-catalog': value.append(catalog(TYPES[0]))
        elif defect == 'pair-duplicate-type': value[1]['ReferenceDataType'] = value[0]['ReferenceDataType']
        elif defect == 'pair-missing-type': value.pop()
        replace_payload(raw, len(raw) - 1, value, status=500 if defect == 'reference-failed' else 200)
    elif defect == 'identical-overlap':
        raw.append(packet({'IsSuccessful': True, 'Payload': [catalog(TYPES[3])]}, 99, REFERENCES[TYPES[3]], None, method='GET'))
        proof['references'].append({'endpoint': REFERENCES[TYPES[3]], 'method': 'GET', 'sequences': [99]})
    else:
        requests = [p['context']['request_json'] for p in raw if p['context']['url'] in GRIDS.values()]
        requests += [r for response in observed[1]['responses'] for r in response['requests']]
        for r in requests:
            if defect == 'sort': r['SortField'] = 'Index'
            elif defect == 'sort-order-bool': r['SortOrder'] = True
            elif defect == 'filters': r['Filters'] = [{'PropertyName': 'Name', 'Value': 'SYNTHETIC'}]
            elif defect == 'language-id': r['LanguageId'] = 'SYNTHETIC UNPROVED'
            elif defect == 'extra-language': r['Language'] = 'id-ID'
            elif defect == 'rows': r['Rows'] = 25
            else: del r['Rows']
    before = deepcopy((raw, observed))
    # A contradictory supplied status is an existing fatal identity contract.
    if defect == 'status':
        with pytest.raises(JsonSafetyError, match='status'):
            adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, observed)
    else:
        result = adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, observed)
        assert not result['eligible'] and result['sections'] == []
    assert (raw, observed) == before


@pytest.mark.parametrize('endpoint', ENDPOINTS)
@pytest.mark.parametrize('bound', [False, True])
@pytest.mark.parametrize('defect', ['identity', 'snapshot', 'count'])
def test_l11a_each_known_grid_contradiction_is_fatal_before_fallback(endpoint, bound, defect):
    raw, observed = packets(), bindings()
    index = next(i for i, p in enumerate(raw) if p['context']['url'] == GRIDS[endpoint])
    value = json.loads(raw[index]['body'])['Payload']
    if defect == 'identity': raw[index]['context']['request_json']['ReturnSheetRecordId'] = corpus.FOREIGN_ID
    elif defect == 'snapshot':
        value['Data'][0]['RecordId'] = 'SYNTHETIC CHANGED ROW'
        raw.append(packet({'IsSuccessful': True, 'Payload': value}, 99, GRIDS[endpoint], request(0)))
    else:
        row = deepcopy(value['Data'][0])
        replace_payload(raw, index, {'TotalRecords': 12, 'Data': [deepcopy(row) for _ in range(10)]})
        raw = [p for p in raw if not (p['context']['url'] == GRIDS[endpoint] and p['context']['request_json']['First'] != 0)]
        raw.append(packet({'IsSuccessful': True, 'Payload': {'TotalRecords': 13, 'Data': [deepcopy(row) for _ in range(3)]}},
                          99, GRIDS[endpoint], request(10)))
        next(r for r in observed[1]['responses'] if r['endpoint'] == GRIDS[endpoint])['requests'] = [request(0), request(10)]
    before = deepcopy(raw)
    with pytest.raises(JsonSafetyError, match='contradicts selected return' if defect == 'identity' else 'snapshot|count'):
        adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, observed if bound else [])
    assert raw == before


@pytest.mark.parametrize('endpoint', ENDPOINTS)
def test_l11a_known_grid_identity_is_fatal_before_observer_or_build_eligibility(endpoint):
    with pytest.raises(JsonSafetyError, match='contradicts selected return'):
        driver._check_corporate_request_identities([{'url': GRIDS[endpoint], 'method': 'POST',
            'request_json': {'ReturnSheetRecordId': corpus.FOREIGN_ID}}], selected(), corpus.SCOPE)


@pytest.mark.parametrize('transport', ['pair', 'pair-reordered'])
@pytest.mark.parametrize('bound', [False, True])
@pytest.mark.parametrize('defect', ['description', 'metadata', 'duplicate-order', 'numeric-token', 'numeric-type'])
def test_l11a_cross_endpoint_same_type_catalog_changes_are_fatal_before_overlap_fallback(transport, bound, defect):
    raw, observed = packets(transport=transport), bindings(transport)
    source = catalog(TYPES[3])
    if defect in ('metadata', 'numeric-token', 'numeric-type'): source['SyntheticMetadata'] = 3
    if defect == 'duplicate-order': source['Details'].append({**source['Details'][0], 'CodeDescription': 'SYNTHETIC OTHER'})
    pair = [source, catalog(TYPES[4])]
    if transport == 'pair-reordered': pair.reverse()
    replace_payload(raw, len(raw) - 1, pair)
    other = deepcopy(source)
    if defect == 'description': other['Details'][0]['CodeDescription'] = 'SYNTHETIC CHANGED'
    elif defect == 'metadata': other['SyntheticMetadata'] = 4
    elif defect == 'duplicate-order': other['Details'].reverse()
    elif defect == 'numeric-token': other['SyntheticMetadata'] = 3.0
    else: other['SyntheticMetadata'] = '3'
    raw.append(packet({'IsSuccessful': True, 'Payload': [other]}, 99, REFERENCES[TYPES[3]], None, method='GET'))
    observed[1]['projection']['references'].append({'endpoint': REFERENCES[TYPES[3]], 'method': 'GET', 'sequences': [99]})
    before = deepcopy((raw, observed))
    with pytest.raises(JsonSafetyError, match='catalog|snapshot'):
        adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, observed if bound else [])
    assert (raw, observed) == before


@pytest.mark.parametrize('url', [BASE + 'TYPE_OF_FULFILLMENT,DEDUCTION_METHOD',
    BASE + 'DEDUCTION_METHOD,TYPE_OF_FULFILLMENT,COUNTRY_CODE', PAIR + '/checksum', PAIR + '?unproved=true'])
def test_l11a_reference_admission_does_not_expand_to_uninspected_transports(url):
    raw = packets(transport='pair')
    raw[-1]['context']['url'] = url
    with pytest.raises(JsonSafetyError, match='endpoint|unsafe'):
        adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, [])


@pytest.mark.parametrize('transport', ['singletons', 'pair', 'pair-reordered'])
@pytest.mark.parametrize('kind', TYPES)
@pytest.mark.parametrize('bound', [False, True])
@pytest.mark.parametrize('defect', ['snapshot', 'locale'])
def test_l11a_every_catalog_transport_keeps_fatal_snapshot_and_locale_before_fallback(transport, kind, bound, defect):
    raw = packets(transport=transport)
    for index, source in enumerate(raw):
        value = json.loads(source['body'])['Payload']
        if not isinstance(value, list): continue
        matching = [c for c in value if c.get('ReferenceDataType') == kind]
        if not matching: continue
        if defect == 'snapshot':
            matching[0]['Details'][0]['CodeDescription'] = 'SYNTHETIC CHANGED'
            raw.append(packet({'IsSuccessful': True, 'Payload': value}, 99, source['context']['url'], None, method='GET'))
        else:
            matching[0]['language'] = 'en-US'
            replace_payload(raw, index, value)
        break
    else: raise AssertionError('Synthetic required catalog was not present')
    before = deepcopy(raw)
    with pytest.raises(JsonSafetyError, match=defect):
        adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, bindings(transport) if bound else [])
    assert raw == before


@pytest.mark.parametrize('endpoint', [*REFERENCES.values(), PAIR])
@pytest.mark.parametrize('bound', [False, True])
@pytest.mark.parametrize('defect', ['authorization', 'privacy', 'catalog-locale', 'detail-locale',
                                    'mixed-clock', 'invalid-clock', 'mixed-lifetime', 'invalid-lifetime'])
def test_l11a_catalog_admission_preserves_fatal_access_privacy_environment_categories(endpoint, bound, defect):
    transport = 'pair' if endpoint == PAIR else 'singletons'
    raw = packets(transport=transport)
    index = next(i for i, p in enumerate(raw) if p['context']['url'] == endpoint)
    value = json.loads(raw[index]['body'])['Payload']
    if defect == 'authorization': raw[index]['context']['status'] = 403
    elif defect in ('privacy', 'catalog-locale', 'detail-locale'):
        if defect == 'privacy': value[0]['AccessToken'] = 'SYNTHETIC SECRET BOUNDARY'
        elif defect == 'catalog-locale': value[0]['language'] = 'en-US'
        else: value[0]['Details'][0]['language'] = 'en-US'
        replace_payload(raw, index, value)
    elif defect in ('mixed-clock', 'invalid-clock'):
        raw[index]['context']['clock_interval'] = {'attempt': corpus.ATTEMPT, 'start_ms': 0,
            'end_ms': 1 if defect == 'mixed-clock' else False, 'timezone': 'Asia/Jakarta'}
    elif defect == 'mixed-lifetime': del raw[index]['context']['component_lifetime']
    else: raw[index]['context']['component_lifetime']['same_document'] = False
    category = ('authorization' if defect == 'authorization' else
        'Authentication material|privacy|forbidden|prohibited secret' if defect == 'privacy' else
        'locale' if defect.endswith('locale') else 'clock' if defect.endswith('clock') else 'lifetime')
    before = deepcopy(raw)
    with pytest.raises(JsonSafetyError, match=category):
        adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, bindings(transport) if bound else [])
    assert raw == before


@pytest.mark.parametrize('transport', ['singletons', 'pair', 'pair-reordered'])
@pytest.mark.parametrize('kind', TYPES)
def test_l11a_each_catalog_preserves_sorted_label_and_original_catalog_detail_index(transport, kind):
    raw, observed = packets(transport=transport), bindings(transport)
    for index, source in enumerate(raw):
        value = json.loads(source['body'])['Payload']
        if not isinstance(value, list): continue
        matching = [(j, c) for j, c in enumerate(value) if c.get('ReferenceDataType') == kind]
        if not matching: continue
        catalog_index, current = matching[0]
        current['Details'][0]['CodeDescription'] = 'SYNTHETIC ZULU'
        current['Details'].append({**current['Details'][0], 'CodeDescription': 'SYNTHETIC ALPHA'})
        replace_payload(raw, index, value)
        sequence = source['context']['sequence']
        break
    else: raise AssertionError('Synthetic required catalog was not present')
    before = deepcopy((raw, observed))
    result = adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, observed)
    assert result['eligible'], result['reason']
    section = next(s for s in result['sections'] if s['name'] == 'L11-A')
    table, column = {TYPES[0]: (0, 5), TYPES[1]: (3, 1), TYPES[2]: (4, 7),
                     TYPES[3]: (2, 6), TYPES[4]: (2, 7)}[kind]
    assert section['tables'][table]['rows'][0][column] == 'SYNTHETIC ALPHA'
    evidence = next(p['values'] for p in result['provenance']['projections'] if p['form'] == 'L11-A')
    assert any(e['response_sequence'] == sequence and e['container'] == ['Payload']
               and e['path'] == [catalog_index, 'Details', 1, 'CodeDescription']
               and e.get('destination') == {'table': f'L11-A/grid{table + 1}', 'row': 0, 'column': column}
               and e['value'] == 'SYNTHETIC ALPHA' for e in evidence)
    assert (raw, observed) == before


def test_l11a_and_l3_share_one_primary_without_mixing_auxiliary_sources():
    from test_corporate_l3_native import packets as l3_packets, bindings as l3_bindings
    raw = l3_packets()
    primary = json.loads(raw[0]['body'])
    root = json.loads(primary['Payload']['CurrentFormData'])
    root[ROOT] = inputs()[0]
    root['MainFormData']['IsDeclareEntertainmentExpense'] = 1
    primary['Payload']['IsDeclareEntertainmentExpense'] = 1
    primary['Payload']['CurrentFormData'] = corpus.wire(root)
    raw[0] = fresh_packet(primary)
    offset = len(raw) - 1
    extra = packets()
    for source in extra[1:]: source['context']['sequence'] += offset
    raw.extend(extra[1:])
    observed = l3_bindings() + bindings()[1:]
    for ref in observed[-1]['projection']['references']:
        ref['sequences'] = [sequence + offset for sequence in ref['sequences']]
    before = deepcopy((raw, observed))
    result = adapt_corporate_detail(iter(raw), selected(), corpus.SCOPE, observed)
    assert result['eligible'] and result['coverage']['complete'], result['reason']
    sections = {s['name']: s for s in result['sections']}
    assert sections['L3']['tables'][1]['footer'][-1] == ['JUMLAH KREDIT PAJAK', '', '41', '']
    assert sections['L11-A']['tables'][0]['body_row_count'] == 12
    evidence = {p['form']: p['values'] for p in result['provenance']['projections']}
    assert any(e['response_sequence'] == 2 and e['path'] == ['IncomeTax'] and e['value'] == '40' for e in evidence['L3'])
    assert any(e['response_sequence'] == 3 + offset and e['path'] == ['Data', 0, 'Index'] and e['value'] == '0007'
               for e in evidence['L11-A'])
    assert (raw, observed) == before
