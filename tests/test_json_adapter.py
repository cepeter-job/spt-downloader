"""Synthetic corporate-detail adapter contracts; GitHub-runner-only execution.

Native field names/shapes are schema evidence, not taxpayer sample values. UI
paths and request/offset bindings below are explicit trusted caller inputs; the
synthetic request keys do not claim to describe a native API request schema.
"""
import builtins
from copy import deepcopy
import gc
import hashlib
import json
from pathlib import Path
import re
import weakref

import pytest
from openpyxl import load_workbook

import coretax
import coretax_exporter as exporter
import coretax_json as native

ORIGIN = 'https://coretax.intranet.pajak.go.id'
NPWP = '0000000000000001'
RECORD_ID = '00000000-0000-4000-8000-000000000001'
AGGREGATE_ID = '00000000-0000-4000-8000-000000000002'
TAXPAYER_ID = '00000000-0000-4000-8000-000000000003'
FOREIGN_ID = '00000000-0000-4000-8000-000000000099'
ATTEMPT = 'a' * 32
SCOPE = [NPWP, 2025, 'period']
LIST_URL = ORIGIN + '/returnsheets/id-ID/submitted-returnsheets/' + TAXPAYER_ID
FRAME_URL = ORIGIN + '/returnsheets/id-ID/fixture-return/' + TAXPAYER_ID + '/' + RECORD_ID
PRIMARY_URL = ORIGIN + '/returnsheet/api/prefillreturnsheet/corporate-income'
GRID_URL = ORIGIN + '/returnsheet/api/loadcit/l3-table-b-grid'
FOOTER_URL = GRID_URL + '/footer'
LOOKUP_URL = ORIGIN + '/returnsheet/api/prefillreturnsheet/getinfotin'
HEADERS = ['JSON path', 'JSON type', 'Source value']


class Number(str):
    """Fixture-only JSON number token; avoids creating source bytes via float."""


def wire(value):
    if isinstance(value, Number):
        assert re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?', value)
        return str(value)
    if isinstance(value, dict):
        return '{' + ','.join(json.dumps(k) + ':' + wire(v) for k, v in value.items()) + '}'
    if isinstance(value, list):
        return '[' + ','.join(wire(v) for v in value) + ']'
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def selected_record():
    return {'key': 'synthetic-corporate-normal', 'record_id': RECORD_ID,
            'scope': list(SCOPE), 'list_url': LIST_URL, 'source_status': 'SPT Dilaporkan',
            'capture_complete': False,
            'meta': {'Jenis Pajak': 'PPh Badan', 'Masa Pajak': 'Tahun 2025',
                     'Model SPT': 'Normal', 'Tanggal Submit SPT': '02-01-2025'},
            'source_fields': {'RecordId': RECORD_ID, 'AggregateIdentifier': AGGREGATE_ID,
                              'TaxpayerAggregateIdentifier': TAXPAYER_ID,
                              'TaxTypeCode': 'ICT_RCIT', 'TaxPeriodCode': '01122025'},
            'source_identity': {'record_id': RECORD_ID, 'aggregate_id': AGGREGATE_ID,
                                'taxpayer_id': TAXPAYER_ID, 'source_url': LIST_URL}}


def current_form():
    return {'TaxType': 'ICT_RCIT', 'IsDraft': False, 'TaxPeriodCode': '01122025',
            'ReturnSheetRecordIdentifier': RECORD_ID,
            'ReturnSheetAggregateIdentifier': AGGREGATE_ID,
            'MainFormData': {'Tin': NPWP, 'TaxYear': 2025, 'Name': 'SYNTHETIC COMPANY',
                             'NetIncomeBeforeTax': Number('123.4500'),
                             'IsIncomeTaxOverseasWithholding': False,
                             'TaxRate': 'SYNTHETIC RATE'}}


def context(raw, sequence=1, *, endpoint=PRIMARY_URL, request=None, status=200):
    return {'relative_path': f'returns/json/0001/{ATTEMPT}/response-{sequence:04d}.json',
            'bytes': len(raw), 'sha256': hashlib.sha256(raw).hexdigest(),
            'url': endpoint, 'method': 'POST', 'status': status, 'frame_url': FRAME_URL,
            'sequence': sequence, 'attempt': ATTEMPT, 'record_id': RECORD_ID,
            'scope': list(SCOPE), 'request_json': request, 'identity_validated': False}


def packet(envelope, sequence=1, *, endpoint=PRIMARY_URL, request=None, status=200):
    raw = envelope if isinstance(envelope, bytes) else wire(envelope).encode('utf-8')
    # The collector's parsed payload may have rounded floats; raw bytes win.
    parsed = json.loads(raw) if isinstance(envelope, dict) else None
    return {'body': raw, 'payload': parsed,
            'context': context(raw, sequence, endpoint=endpoint, request=request, status=status)}


def primary(form=None, *, outer_draft=True, sequence=1):
    form = current_form() if form is None else form
    payload = {'TIN': NPWP, 'PeriodYear': 2025, 'IsDraft': outer_draft,
               'ReturnSheetStatusCode': 'SUBMITTED', 'ReturnSheetModel': 'Normal',
               'PreviousFormData': None, 'LastSubmittedFormData': None,
               'CurrentFormData': wire(form)}
    return packet({'IsSuccessful': True, 'ErrorCode': 0, 'Payload': payload}, sequence)


def observed(*extra_paths):
    return [{'name': 'Induk', 'paths': [['MainFormData'], *extra_paths],
             'primary': {'endpoint': PRIMARY_URL, 'request_json': None},
             'grids': [], 'responses': [], 'dependencies': []}]


def request(group, page=0):
    return {'synthetic_group': group, 'synthetic_page': page}


def declare_grid(forms, group='A', *, pages=None, name=None):
    grid = {'name': name or 'SYNTHETIC GRID ' + group, 'endpoint': GRID_URL,
            'pages': pages or [{'request_json': request(group), 'offset': 0}]}
    forms[0]['grids'].append(grid)
    return grid


def grid_packet(rows, total=None, *, group='A', page=0, sequence=2, index=None):
    payload = {'Data': rows, 'TotalRecords': len(rows) if total is None else total, 'Index': index}
    return packet({'IsSuccessful': True, 'ErrorCode': 0, 'Payload': payload}, sequence,
                  endpoint=GRID_URL, request=request(group, page))


def source_row(marker='SYNTHETIC SLIP', amount=None):
    return {'Name': marker, 'TIN': '0000000000000002', 'Index': 0,
            'TaxBase': Number('12.3400') if amount is None else amount,
            'IncomeTax': Number('0'), 'IsManually': False, 'Action': None}


def adapter():
    target = getattr(native, 'adapt_corporate_detail', None)
    assert callable(target), 'Corporate detail adapter is missing: intentional feature RED'
    return target


def adapt(bodies=None, *, record=None, scope=None, forms=None):
    target = adapter()
    result = target([primary()] if bodies is None else bodies,
                    selected_record() if record is None else record,
                    list(SCOPE) if scope is None else scope,
                    observed() if forms is None else forms)
    assert set(result) == {'eligible', 'reason', 'sections', 'coverage', 'provenance'}
    assert type(result['eligible']) is bool
    assert isinstance(result['reason'], str) and result['reason']
    assert isinstance(result['sections'], list)
    assert isinstance(result['coverage'], dict) and isinstance(result['provenance'], dict)
    assert type(result['coverage']['complete']) is bool
    return result


def eligible(result):
    assert result['eligible'] is True and result['coverage']['complete'] is True
    assert result['sections']


def ineligible(result):
    assert result['eligible'] is False and result['coverage']['complete'] is False
    assert result['sections'] == [], 'Do not expose partial JSON for mixed-record export'


def source_rows(result, table_name=None):
    rows = []
    for section in result['sections']:
        for table in section['tables']:
            assert table['headers'] == HEADERS and table['complete'] is True
            if table_name is None or table['name'] == table_name:
                for row in table['rows']:
                    assert len(row) == 3
                    assert all(value is None or type(value) in (str, bool, int, float) for value in row)
                    rows.append(row)
    return rows


def values_at(result, pointer):
    return [(kind, value) for path, kind, value in source_rows(result) if path == pointer]


def test_verified_current_form_is_eligible_without_claiming_record_or_pdf_complete():
    record = selected_record()
    result = adapt(record=record)
    eligible(result)
    assert [s['name'] for s in result['sections']] == ['Induk']
    assert values_at(result, '/MainFormData/Tin') == [('string', NPWP)]
    assert record['capture_complete'] is False
    assert 'complete' not in result and 'attachments' not in result


@pytest.mark.parametrize('token', ['9007199254740993', '123.4500', '1.2300e+03', '-0',
                                   '-0.000', '0', '1e+400', '12345678901234567890.00000000000001'])
def test_embedded_current_form_preserves_numeric_lexemes(token):
    form = current_form()
    form['MainFormData']['NetIncomeBeforeTax'] = Number(token)
    result = adapt([primary(form)])
    eligible(result)
    assert values_at(result, '/MainFormData/NetIncomeBeforeTax') == [('number', token)]


@pytest.mark.parametrize(('value', 'kind', 'expected'), [(None, 'null', None), (False, 'boolean', False),
                         (True, 'boolean', True), (Number('0'), 'number', '0'),
                         ('0', 'string', '0'), ('', 'string', ''),
                         ('0000000000000002', 'string', '0000000000000002'),
                         ('02-01-2025', 'string', '02-01-2025'), ('=SUM(1,2)', 'string', '=SUM(1,2)')])
def test_null_boolean_zero_string_identifier_date_and_formula_remain_distinct(value, kind, expected):
    form = current_form()
    form['MainFormData']['TaxRate'] = value
    result = adapt([primary(form)])
    eligible(result)
    assert values_at(result, '/MainFormData/TaxRate') == [(kind, expected)]
    assert type(values_at(result, '/MainFormData/TaxRate')[0][1]) is type(expected)


def test_raw_bytes_not_rounded_collector_payload_are_authority():
    item = primary()
    item['payload']['Payload']['CurrentFormData'] = 'SYNTHETIC ROUNDED PARSED COPY'
    result = adapt([item])
    eligible(result)
    assert values_at(result, '/MainFormData/NetIncomeBeforeTax') == [('number', '123.4500')]


def test_outer_prefill_and_inner_current_form_isdraft_are_independent_observations():
    result = adapt([primary(outer_draft=True)])
    eligible(result)
    assert result['provenance']['draft_flags'] == {'outer': True, 'inner': False}
    assert selected_record()['source_status'] == 'SPT Dilaporkan'


def test_hierarchy_parent_children_duplicate_values_empty_arrays_and_paths_are_lossless():
    form = current_form()
    child = {'Data': {'AccountCode': 'SYNTHETIC CHILD', 'Description': 'SYNTHETIC LINE',
                      'Amount': Number('50.000')}, 'Children': [], 'Expanded': False}
    parent = {'Data': {'AccountCode': 'SYNTHETIC PARENT', 'Description': 'SYNTHETIC TOTAL',
                       'Amount': Number('100.000')}, 'Children': [deepcopy(child), deepcopy(child)],
              'Expanded': True}
    form['L1dFormData'] = {'L1dTreeData': [parent], 'L1dTreeTotalRow': {'Amount': Number('100.000')}}
    form['FacilityRegisters'] = []
    form['OtherDocuments'] = []
    result = adapt([primary(form)], forms=observed(['L1dFormData'], ['FacilityRegisters'], ['OtherDocuments']))
    eligible(result)
    assert values_at(result, '/L1dFormData/L1dTreeData/0/Data/Amount') == [('number', '100.000')]
    for index in [0, 1]:
        assert values_at(result, f'/L1dFormData/L1dTreeData/0/Children/{index}/Data/Amount') == [('number', '50.000')]
    assert values_at(result, '/L1dFormData/L1dTreeTotalRow/Amount') == [('number', '100.000')]
    assert values_at(result, '/FacilityRegisters') == [('array', '[]')]
    assert values_at(result, '/OtherDocuments') == [('array', '[]')]
    amounts = [v for p, k, v in source_rows(result) if p.endswith('/Amount')]
    assert amounts == ['100.000', '50.000', '50.000', '100.000'], 'Never compute or merge ancestor totals'


def test_json_pointer_escapes_field_tokens_without_normalizing_them():
    form = current_form()
    form['MainFormData']['SYNTHETIC/a~b'] = 'SYNTHETIC POINTER VALUE'
    result = adapt([primary(form)])
    eligible(result)
    assert values_at(result, '/MainFormData/SYNTHETIC~1a~0b') == [('string', 'SYNTHETIC POINTER VALUE')]


def test_independently_observed_ui_paths_partition_the_actual_source_groups():
    form = current_form()
    form['L11BFormData'] = {'Ebitda': Number('25.0000')}
    forms = observed()
    forms.append({'name': 'Lampiran / L11-B', 'paths': [['L11BFormData']],
                  'grids': [], 'responses': [], 'dependencies': []})
    result = adapt([primary(form)], forms=forms)
    eligible(result)
    assert [s['name'] for s in result['sections']] == ['Induk', 'Lampiran / L11-B']
    assert result['sections'][1]['tables'][0]['rows'] == [['/L11BFormData/Ebitda', 'number', '25.0000']]


def test_complete_bound_grid_preserves_duplicate_rows_and_zero_row_indices():
    forms = observed()
    grid = declare_grid(forms)
    line = source_row()
    result = adapt([primary(), grid_packet([deepcopy(line), deepcopy(line)], 2)], forms=forms)
    eligible(result)
    rows = source_rows(result, grid['name'])
    assert [v for p, k, v in rows if p.endswith('/TaxBase')] == ['12.3400', '12.3400']
    assert [v for p, k, v in rows if p.endswith('/Index')] == ['0', '0']
    assert next(g for g in result['coverage']['grids'] if g['name'] == grid['name'])['received'] == 2


def test_ten_of_sixty_is_not_complete_even_when_every_row_index_is_zero():
    forms = observed()
    declare_grid(forms)
    result = adapt([primary(), grid_packet([source_row()] * 10, 60)], forms=forms)
    ineligible(result)
    grid = result['coverage']['grids'][0]
    assert grid['received'] == 10 and grid['total'] == 60


def test_pages_cover_the_exact_source_span_using_observed_offsets_not_row_index():
    forms = observed()
    grid = declare_grid(forms, pages=[{'request_json': request('A', 0), 'offset': 0},
                                     {'request_json': request('A', 1), 'offset': 2}])
    result = adapt([primary(), grid_packet([source_row()] * 2, 3, sequence=2),
                    grid_packet([source_row('SYNTHETIC LAST')], 3, page=1, sequence=3)], forms=forms)
    eligible(result)
    assert [v for p, k, v in source_rows(result, grid['name']) if p.endswith('/TaxBase')] == ['12.3400'] * 3
    assert values_at(result, '/Payload/Data/2/Name') == [('string', 'SYNTHETIC LAST')]
    coverage = result['coverage']['grids'][0]
    assert coverage['total'] == coverage['received'] == 3


@pytest.mark.parametrize('failure', ['missing-page', 'gap', 'unknown-request', 'no-request-context', 'nonnull-index'])
def test_unproven_grid_pages_are_ineligible_not_counted_as_complete(failure):
    forms = observed()
    declare_grid(forms, pages=[{'request_json': request('A', 0), 'offset': 0},
                               {'request_json': request('A', 1), 'offset': 2}])
    pages = [grid_packet([source_row()] * 2, 3), grid_packet([source_row()], 3, page=1, sequence=3)]
    if failure == 'missing-page':
        pages.pop()
    elif failure == 'gap':
        pages[0] = grid_packet([source_row()], 3)
    elif failure == 'unknown-request':
        pages[1]['context']['request_json'] = request('UNOBSERVED', 1)
    elif failure == 'no-request-context':
        pages[1]['context']['request_json'] = None
    else:
        pages[1] = grid_packet([source_row()], 3, page=1, sequence=3, index=1)
    ineligible(adapt([primary(), *pages], forms=forms))


@pytest.mark.parametrize('failure', ['overlap', 'inconsistent-total', 'conflicting-repeated-page', 'rows-exceed-total'])
def test_contradictory_grid_page_count_or_snapshot_fails_closed(failure):
    forms = observed()
    declare_grid(forms, pages=[{'request_json': request('A', 0), 'offset': 0},
                               {'request_json': request('A', 1), 'offset': 1 if failure == 'overlap' else 2}])
    pages = [grid_packet([source_row()] * 2, 3), grid_packet([source_row()], 3, page=1, sequence=3)]
    if failure == 'inconsistent-total':
        pages[1] = grid_packet([source_row()], 4, page=1, sequence=3)
    elif failure == 'conflicting-repeated-page':
        pages[1] = grid_packet([source_row('SYNTHETIC CHANGED')], 3, sequence=3)
    elif failure == 'rows-exceed-total':
        pages[0] = grid_packet([source_row()] * 2, 1)
    with pytest.raises(native.JsonSafetyError):
        adapt([primary(), *pages], forms=forms)


def test_same_body_in_different_request_groups_is_not_deduplicated():
    forms = observed()
    first = declare_grid(forms, 'A')
    second = declare_grid(forms, 'B')
    a = grid_packet([source_row()], 1, group='A', sequence=2)
    b = grid_packet([source_row()], 1, group='B', sequence=3)
    assert a['body'] == b['body'] and a['context']['sha256'] == b['context']['sha256']
    result = adapt([primary(), a, b], forms=forms)
    eligible(result)
    assert len([v for p, k, v in source_rows(result, first['name']) if p.endswith('/TaxBase')]) == 1
    assert len([v for p, k, v in source_rows(result, second['name']) if p.endswith('/TaxBase')]) == 1
    assert len(result['provenance']['events']) == 3


def test_repeated_identical_page_event_is_evidence_not_extra_source_rows():
    forms = observed()
    grid = declare_grid(forms)
    a = grid_packet([source_row()], 1, sequence=2)
    b = grid_packet([source_row()], 1, sequence=3)
    result = adapt([primary(), a, b], forms=forms)
    eligible(result)
    assert len([v for p, k, v in source_rows(result, grid['name']) if p.endswith('/TaxBase')]) == 1
    assert len(result['provenance']['events']) == 3
    assert result['coverage']['grids'][0]['received'] == 1


@pytest.mark.parametrize('empty', [[], None])
def test_empty_grid_is_complete_only_with_explicit_bound_zero_total(empty):
    forms = observed()
    grid = declare_grid(forms)
    result = adapt([primary(), grid_packet(empty, 0)], forms=forms)
    eligible(result)
    assert source_rows(result, grid['name']) == []
    assert result['coverage']['grids'][0]['total'] == result['coverage']['grids'][0]['received'] == 0


def test_declared_grid_without_response_does_not_claim_empty_complete():
    forms = observed()
    declare_grid(forms)
    ineligible(adapt([primary()], forms=forms))


def test_known_footer_is_separate_source_evidence_not_a_grid_row_or_computed_total():
    forms = observed()
    req = request('FOOTER')
    forms[0]['responses'].append({'name': 'SYNTHETIC FOOTER', 'endpoint': FOOTER_URL, 'requests': [req]})
    footer = packet({'IsSuccessful': True, 'Payload': {'IncomeTax': Number('1.2500'), 'TaxBase': Number('9.0000')}},
                    2, endpoint=FOOTER_URL, request=req)
    result = adapt([primary(), footer], forms=forms)
    eligible(result)
    assert source_rows(result, 'SYNTHETIC FOOTER') == [['/Payload/IncomeTax', 'number', '1.2500'],
                                                    ['/Payload/TaxBase', 'number', '9.0000']]
    assert result['coverage']['grids'] == []


@pytest.mark.parametrize('required', [False, True])
def test_ancillary_http400_errorcode_zero_only_blocks_a_declared_dependency(required):
    forms = observed()
    req = request('LOOKUP')
    if required:
        forms[0]['dependencies'].append({'endpoint': LOOKUP_URL, 'request_json': req})
    failure = packet({'IsSuccessful': False, 'ErrorCode': 0, 'Payload': None}, 2,
                     endpoint=LOOKUP_URL, request=req, status=400)
    result = adapt([primary(), failure], forms=forms)
    if required:
        ineligible(result)
    else:
        eligible(result)
        assert len(result['provenance']['events']) == 2


@pytest.mark.parametrize('failure', ['http400', 'false-success', 'truthy-success', 'missing-success'])
def test_primary_requires_both_successful_http_and_explicit_boolean_success(failure):
    item = primary()
    env = json.loads(item['body'])
    status = 400 if failure == 'http400' else 200
    if failure == 'false-success':
        env['IsSuccessful'] = False
    elif failure == 'truthy-success':
        env['IsSuccessful'] = 1
    elif failure == 'missing-success':
        del env['IsSuccessful']
    ineligible(adapt([packet(env, status=status)]))


@pytest.mark.parametrize(('path', 'foreign'), [(('ReturnSheetRecordIdentifier',), FOREIGN_ID),
                          (('ReturnSheetAggregateIdentifier',), FOREIGN_ID),
                          (('MainFormData', 'Tin'), '0000000000000002'),
                          (('MainFormData', 'TaxYear'), 2024),
                          (('TaxType',), 'ICT_VAT'), (('TaxPeriodCode',), '01122024')])
def test_primary_identity_taxpayer_tax_kind_and_native_period_conflicts_fail_closed(path, foreign):
    form = current_form()
    target = form
    for key in path[:-1]:
        target = target[key]
    target[path[-1]] = foreign
    with pytest.raises(native.JsonSafetyError):
        adapt([primary(form)])


def test_envelope_tin_conflict_cannot_be_hidden_by_matching_inner_form():
    env = json.loads(primary()['body'])
    env['Payload']['TIN'] = '0000000000000002'
    with pytest.raises(native.JsonSafetyError):
        adapt([packet(env)])


def test_submission_year_scope_does_not_relabel_or_reject_verified_previous_tax_year():
    scope = [NPWP, 2025, 'submission']
    record = selected_record()
    record['scope'] = scope
    record['meta']['Masa Pajak'] = 'Tahun 2024'
    record['source_fields']['TaxPeriodCode'] = '01122024'
    form = current_form()
    form['TaxPeriodCode'] = '01122024'
    form['MainFormData']['TaxYear'] = 2024
    env = json.loads(primary(form)['body'])
    env['Payload']['PeriodYear'] = 2024
    item = packet(env)
    item['context']['scope'] = scope
    result = adapt([item], record=record, scope=scope)
    eligible(result)
    assert values_at(result, '/MainFormData/TaxYear') == [('number', '2024')]
    assert record['meta']['Masa Pajak'] == 'Tahun 2024'


@pytest.mark.parametrize('missing', ['aggregate', 'taxpayer', 'primary', 'main', 'raw-form', 'unmapped-group', 'missing-path'])
def test_missing_verified_source_or_form_binding_is_safe_ineligible(missing):
    record = selected_record()
    form = current_form()
    forms = observed()
    bodies = [primary(form)]
    if missing == 'aggregate':
        record['source_identity'].pop('aggregate_id')
        record['source_fields'].pop('AggregateIdentifier')
    elif missing == 'taxpayer':
        record['source_identity'].pop('taxpayer_id')
        record['source_fields'].pop('TaxpayerAggregateIdentifier')
    elif missing == 'primary':
        bodies = []
    elif missing == 'main':
        form.pop('MainFormData')
        bodies = [primary(form)]
    elif missing == 'raw-form':
        env = json.loads(primary()['body'])
        env['Payload']['CurrentFormData'] = None
        bodies = [packet(env)]
    elif missing == 'unmapped-group':
        form['L10bFormData'] = {'SYNTHETIC UNMAPPED VALUE': 'SYNTHETIC EVIDENCE'}
        bodies = [primary(form)]
    else:
        forms[0]['paths'].append(['L10cFormData'])
    ineligible(adapt(bodies, record=record, forms=forms))


@pytest.mark.parametrize('kind', ['empty-caller-bindings', 'unknown-grid', 'unknown-footer', 'unsupported-form-object', 'unexportable-text'])
def test_unsupported_observed_coverage_never_exposes_a_partial_json_section(kind):
    forms = observed()
    bodies = [primary()]
    if kind == 'empty-caller-bindings':
        forms = []
    elif kind in {'unknown-grid', 'unknown-footer'}:
        extra = grid_packet([source_row()]) if kind == 'unknown-grid' else packet(
            {'IsSuccessful': True, 'Payload': {'IncomeTax': Number('2.0000')}}, 2,
            endpoint=FOOTER_URL, request=request('UNOBSERVED'))
        bodies.append(extra)
    elif kind == 'unsupported-form-object':
        env = json.loads(primary()['body'])
        env['Payload']['CurrentFormData'] = current_form()
        bodies = [packet(env)]
    else:
        form = current_form()
        form['MainFormData']['Name'] = 'SYNTHETIC\u0001UNEXPORTABLE'
        bodies = [primary(form)]
    ineligible(adapt(bodies, forms=forms))


@pytest.mark.parametrize('scope', [['000000000000001', 2025, 'period'], [NPWP + ' ', 2025, 'period'],
                                  [NPWP, True, 'period'], [NPWP, 2025.0, 'period'],
                                  [NPWP, 2025, 'UNKNOWN']])
def test_scope_tokens_are_not_normalized_or_accepted_by_truthiness(scope):
    with pytest.raises(native.JsonSafetyError):
        adapt(scope=scope)


@pytest.mark.parametrize('mutation', ['wrong-record', 'wrong-scope', 'wrong-frame-record', 'foreign-frame-origin',
                         'foreign-endpoint-origin', 'credentials-in-url', 'secret-query', 'wrong-method',
                         'bad-byte-count', 'bad-hash', 'missing-attempt', 'bad-sequence', 'http401', 'http403'])
def test_unsafe_captured_descriptor_is_not_a_dom_fallback_reason(mutation):
    item = primary()
    ctx = item['context']
    if mutation == 'wrong-record':
        ctx['record_id'] = FOREIGN_ID
    elif mutation == 'wrong-scope':
        ctx['scope'] = [NPWP, 2024, 'period']
    elif mutation == 'wrong-frame-record':
        ctx['frame_url'] = FRAME_URL.replace(RECORD_ID, FOREIGN_ID)
    elif mutation == 'foreign-frame-origin':
        ctx['frame_url'] = FRAME_URL.replace(ORIGIN, 'https://example.invalid')
    elif mutation == 'foreign-endpoint-origin':
        ctx['url'] = PRIMARY_URL.replace(ORIGIN, 'https://example.invalid')
    elif mutation == 'credentials-in-url':
        ctx['url'] = PRIMARY_URL.replace('https://', 'https://fixture:fixture@')
    elif mutation == 'secret-query':
        ctx['url'] = PRIMARY_URL + '?access_token=SYNTHETIC_NOT_A_CREDENTIAL'
    elif mutation == 'wrong-method':
        ctx['method'] = 'DELETE'
    elif mutation == 'bad-byte-count':
        ctx['bytes'] += 1
    elif mutation == 'bad-hash':
        ctx['sha256'] = '0' * 64
    elif mutation == 'missing-attempt':
        ctx['attempt'] = ''
    elif mutation == 'bad-sequence':
        ctx['sequence'] = True
    else:
        ctx['status'] = 401 if mutation == 'http401' else 403
    with pytest.raises(native.JsonSafetyError):
        adapt([item])


@pytest.mark.parametrize('conflict', ['attempt', 'sequence', 'primary-data', 'binding-overlap'])
def test_mixed_attempt_duplicate_event_conflicting_primary_or_overlapping_mapping_fails_closed(conflict):
    forms = observed()
    a = primary(sequence=1)
    b = primary(sequence=2)
    if conflict == 'attempt':
        b['context']['attempt'] = 'b' * 32
    elif conflict == 'sequence':
        b['context']['sequence'] = 1
    elif conflict == 'primary-data':
        form = current_form()
        form['MainFormData']['NetIncomeBeforeTax'] = Number('999.0000')
        b = primary(form, sequence=2)
    else:
        forms[0]['paths'].append(['MainFormData', 'NetIncomeBeforeTax'])
    with pytest.raises(native.JsonSafetyError):
        adapt([a, b], forms=forms)


@pytest.mark.parametrize('bad', ['duplicate-outer-key', 'duplicate-inner-key', 'nan-inner', 'infinity-inner',
                                 'invalid-utf8', 'embedded-secret', 'double-encoded-secret', 'secret-text'])
def test_ambiguous_nonfinite_or_secret_source_is_rejected_without_echoing_source(bad):
    if bad == 'duplicate-outer-key':
        raw = b'{"IsSuccessful":true,"IsSuccessful":false,"Payload":{}}'
    elif bad == 'invalid-utf8':
        raw = b'\xff'
    else:
        env = json.loads(primary()['body'])
        if bad == 'duplicate-inner-key':
            env['Payload']['CurrentFormData'] = '{"TaxType":"ICT_RCIT","TaxType":"ICT_VAT"}'
        elif bad in {'nan-inner', 'infinity-inner'}:
            env['Payload']['CurrentFormData'] = '{"MainFormData":{"NetIncomeBeforeTax":' + ('NaN' if bad == 'nan-inner' else 'Infinity') + '}}'
        else:
            form = current_form()
            if bad == 'embedded-secret':
                form['MainFormData']['Authorization'] = 'SYNTHETIC_DO_NOT_ECHO'
            elif bad == 'double-encoded-secret':
                form['MainFormData']['Name'] = json.dumps(json.dumps({'access_token': 'SYNTHETIC_DO_NOT_ECHO'}))
            else:
                form['MainFormData']['Name'] = 'Bearer SYNTHETIC_DO_NOT_ECHO'
            env['Payload']['CurrentFormData'] = wire(form)
        raw = wire(env).encode('utf-8')
    with pytest.raises(native.JsonSafetyError) as error:
        adapt([packet(raw)])
    assert 'SYNTHETIC_DO_NOT_ECHO' not in str(error.value)


def test_selected_record_conflicting_native_and_canonical_identity_fails_closed():
    record = selected_record()
    record['source_fields']['AggregateIdentifier'] = FOREIGN_ID
    with pytest.raises(native.JsonSafetyError):
        adapt(record=record)


def test_declared_request_equality_is_type_sensitive_not_python_true_equals_one():
    forms = observed()
    grid = declare_grid(forms)
    grid['pages'][0]['request_json'] = {'synthetic_group': True}
    item = grid_packet([source_row()], 1)
    item['context']['request_json'] = {'synthetic_group': 1}
    ineligible(adapt([primary(), item], forms=forms))


def test_decoder_does_not_open_manifest_relative_paths_or_any_source_files(monkeypatch):
    target = adapter()
    item = primary()
    item['context']['relative_path'] = 'unresolved-private-reference.json'
    def forbidden(*args, **kwargs):
        raise AssertionError('Pure adapter attempted source filesystem I/O')
    monkeypatch.setattr(builtins, 'open', forbidden)
    monkeypatch.setattr(Path, 'open', forbidden)
    monkeypatch.setattr(Path, 'read_bytes', forbidden)
    monkeypatch.setattr(Path, 'read_text', forbidden)
    result = target(iter([item]), selected_record(), list(SCOPE), observed())
    eligible(result)


def test_lazy_body_stream_is_consumed_once_without_retaining_prior_packets():
    target = adapter()
    forms = observed()
    declare_grid(forms)
    class ProbePacket(dict):
        pass
    seen = []
    def stream():
        first = ProbePacket(primary())
        first_ref = weakref.ref(first)
        seen.append('primary')
        yield first
        del first
        second = ProbePacket(grid_packet([source_row()], 1))
        seen.append('grid')
        yield second
        del second
        gc.collect()
        assert first_ref() is None, 'Decoder retained earlier raw/parsed packet instead of derived evidence'
        seen.append('end')
    result = target(stream(), selected_record(), list(SCOPE), forms)
    eligible(result)
    assert seen == ['primary', 'grid', 'end']


def test_cancellation_or_source_stream_failure_propagates_without_fallback_or_partial_result():
    target = adapter()
    class SyntheticStop(RuntimeError):
        pass
    def stream():
        yield primary()
        raise SyntheticStop('SYNTHETIC STOP')
    with pytest.raises(SyntheticStop):
        target(stream(), selected_record(), list(SCOPE), observed())


def test_decoder_does_not_mutate_record_scope_bindings_or_collector_items():
    bodies, record, scope, forms = [primary()], selected_record(), list(SCOPE), observed()
    before = deepcopy((bodies, record, scope, forms))
    eligible(adapt(bodies, record=record, scope=scope, forms=forms))
    assert (bodies, record, scope, forms) == before


def test_unknown_masa_has_no_invented_corporate_decoder():
    record = selected_record()
    record['meta']['Jenis Pajak'] = 'PPN'
    record['meta']['Masa Pajak'] = 'Juli 2025'
    record['source_fields']['TaxTypeCode'] = 'SYNTHETIC UNKNOWN MASA CODE'
    record['source_fields']['TaxPeriodCode'] = '07072025'
    opaque = packet({'IsSuccessful': True, 'Payload': {'SYNTHETIC OPAQUE MASA': None}},
                    endpoint=ORIGIN + '/returnsheet/api/prefillreturnsheet/fixture-masa')
    ineligible(adapt([opaque], record=record))


def test_reopened_real_export_preserves_raw_values_types_paths_duplicates_and_resume_links(tmp_path):
    form = current_form()
    form['MainFormData']['NetIncomeBeforeTax'] = Number('9007199254740993.0000')
    form['MainFormData']['TaxRate'] = '=SUM(1,2)'
    forms = observed()
    declare_grid(forms)
    result = adapt([primary(form), grid_packet([source_row()] * 2, 2)], forms=forms)
    eligible(result)
    record = selected_record()
    record['meta']['Model SPT'] = 'Pembetulan 001'
    record['sections'] = result['sections']
    run = {'npwp': NPWP, 'year': 2025, 'year_basis': 'period', 'complete': False,
           'inventory_complete': False, 'taxpayer_name': 'SYNTHETIC COMPANY',
           'started_at': 'fixture', 'audit': [], 'errors': [], 'records': [record]}
    files = exporter.export_split_workbooks(run, tmp_path)
    assert set(files) == {'resume.xlsx', 'SPT_Badan_Tahun 2025_Pembetulan 001.xlsx'}
    coretax.verify_outputs(run, tmp_path)
    workbook = load_workbook(tmp_path / 'SPT_Badan_Tahun 2025_Pembetulan 001.xlsx')
    resume = load_workbook(tmp_path / 'resume.xlsx')
    try:
        assert workbook.sheetnames == ['Induk']
        proof = exporter.read_workbook_proof(workbook)
        assert proof['complete'] is False and proof['records'][0]['capture_complete'] is False
        assert proof['records'][0]['meta']['Model SPT'] == 'Pembetulan 001'
        rows = [list(row) for row in workbook['Induk'].iter_rows(values_only=True)]
        assert any(row[:3] == ['/MainFormData/NetIncomeBeforeTax', 'number', '9007199254740993.0000'] for row in rows)
        assert any(row[:3] == ['/MainFormData/Tin', 'string', NPWP] for row in rows)
        for index in [0, 1]:
            assert any(row[:3] == [f'/Payload/Data/{index}/TaxBase', 'number', '12.3400'] for row in rows)
            assert any(row[:3] == [f'/Payload/Data/{index}/IsManually', 'boolean', False] for row in rows)
            assert any(row[:3] == [f'/Payload/Data/{index}/Action', 'null', None] for row in rows)
        literals = [cell for sheet in workbook for row in sheet for cell in row
                    if cell.value in {'9007199254740993.0000', '12.3400', NPWP, '=SUM(1,2)'}]
        assert literals and all(cell.data_type == 's' for cell in literals)
        links = [c.hyperlink for sheet in resume for row in sheet for c in row if c.hyperlink]
        assert any(link.target == 'SPT_Badan_Tahun 2025_Pembetulan 001.xlsx' and link.location == "'Induk'!A1" for link in links)
        assert not list(tmp_path.rglob('*.pdf')), 'Decoder metadata is not downloaded PDF custody'
    finally:
        workbook.close()
        resume.close()


def test_financial_statement_metadata_remains_source_evidence_not_complete_pdf_custody():
    form = current_form()
    form['FinancialStatement'] = [{'DocumentId': 'SYNTHETIC DOCUMENT',
                                 'DocumentAggregateIdentifier': FOREIGN_ID,
                                 'DocumentType': 'FinancialStatement',
                                 'FileName': 'SYNTHETIC-metadata-only.pdf',
                                 'CreationDate': '02-01-2025'}]
    result = adapt([primary(form)], forms=observed(['FinancialStatement']))
    eligible(result)
    assert values_at(result, '/FinancialStatement/0/FileName') == [('string', 'SYNTHETIC-metadata-only.pdf')]
    assert 'attachments' not in result and selected_record()['capture_complete'] is False


@pytest.mark.parametrize('failure', ['undeclared-endpoint', 'wrong-primary-request', 'missing-primary-binding'])
def test_primary_shape_never_replaces_trusted_endpoint_and_request_binding(failure):
    forms = observed()
    item = primary()
    if failure == 'undeclared-endpoint':
        item['context']['url'] = ORIGIN + '/returnsheet/api/fixture-unbound'
    elif failure == 'wrong-primary-request':
        item['context']['request_json'] = request('UNOBSERVED PRIMARY')
    else:
        forms[0].pop('primary')
    ineligible(adapt([item], forms=forms))


@pytest.mark.parametrize('numeric_location', ['inner-tin', 'outer-tin'])
def test_taxpayer_numbers_cannot_bind_literal_npwp_via_numeric_string_subclass(numeric_location):
    literal = '1000000000000001'
    scope = [literal, 2025, 'period']
    record = selected_record()
    record['scope'] = list(scope)
    form = current_form()
    form['MainFormData']['Tin'] = Number(literal) if numeric_location == 'inner-tin' else literal
    envelope = json.loads(primary(form)['body'])
    envelope['Payload']['TIN'] = Number(literal) if numeric_location == 'outer-tin' else literal
    item = packet(envelope)
    item['context']['scope'] = list(scope)
    with pytest.raises(native.JsonSafetyError):
        adapt([item], record=record, scope=scope)


@pytest.mark.parametrize('unsupported', ['missing-aggregate-binding', 'missing-ui-paths', 'failed-http', 'missing-main'])
def test_known_primary_identity_contradiction_aborts_even_when_source_is_already_ineligible(unsupported):
    record, forms, form = selected_record(), observed(), current_form()
    form['ReturnSheetRecordIdentifier'] = FOREIGN_ID
    if unsupported == 'missing-aggregate-binding':
        record['source_identity'].pop('aggregate_id')
        record['source_fields'].pop('AggregateIdentifier')
    elif unsupported == 'missing-ui-paths':
        forms[0]['paths'] = []
    elif unsupported == 'missing-main':
        form.pop('MainFormData')
    item = primary(form)
    if unsupported == 'failed-http':
        item['context']['status'] = 400
    with pytest.raises(native.JsonSafetyError):
        adapt([item], record=record, forms=forms)


def test_changed_bound_page_is_safety_failure_even_when_both_indices_are_unsupported():
    forms = observed()
    declare_grid(forms)
    a = grid_packet([source_row()], 1, sequence=2, index=1)
    b = grid_packet([source_row('SYNTHETIC CHANGED')], 1, sequence=3, index=1)
    with pytest.raises(native.JsonSafetyError):
        adapt([primary(), a, b], forms=forms)


def test_empty_object_is_visible_as_source_type_and_literal_not_discarded():
    form = current_form()
    form['L11BFormData'] = {}
    result = adapt([primary(form)], forms=observed(['L11BFormData']))
    eligible(result)
    assert values_at(result, '/L11BFormData') == [('object', '{}')]


def test_source_root_count_is_distinct_from_expanded_evidence_row_count():
    form = current_form()
    form['MainFormData']['AccountingMethod'] = 'SYNTHETIC ACCOUNTING METHOD'
    result = adapt([primary(form)])
    eligible(result)
    assert result['coverage']['source_root_count'] == len(form)
    assert result['coverage']['evidence_row_count'] == len(source_rows(result))
    assert result['coverage']['evidence_row_count'] != result['coverage']['source_root_count']
    assert ['/IsDraft', 'boolean', False] in result['provenance']['controls']


def test_additional_grid_payload_source_is_not_silently_discarded():
    forms = observed()
    declare_grid(forms)
    item = grid_packet([source_row()], 1)
    envelope = json.loads(item['body'])
    envelope['Payload']['Footer'] = {'TaxBase': Number('9.0000')}
    extra = packet(envelope, 2, endpoint=GRID_URL, request=request('A'))
    ineligible(adapt([primary(), extra], forms=forms))


@pytest.mark.parametrize(('kind', 'explicit_null'), [('footer', False), ('footer', True),
                                                    ('dependency', False), ('dependency', True)])
def test_missing_payload_is_not_manufactured_as_explicit_null_source(kind, explicit_null):
    forms, req = observed(), request('SYNTHETIC NULL SOURCE')
    endpoint = FOOTER_URL if kind == 'footer' else LOOKUP_URL
    if kind == 'footer':
        forms[0]['responses'].append({'name': 'SYNTHETIC NULL FOOTER', 'endpoint': endpoint, 'requests': [req]})
    else:
        forms[0]['dependencies'].append({'endpoint': endpoint, 'request_json': req})
    envelope = {'IsSuccessful': True}
    if explicit_null:
        envelope['Payload'] = None
    result = adapt([primary(), packet(envelope, 2, endpoint=endpoint, request=req)], forms=forms)
    if explicit_null:
        eligible(result)
        if kind == 'footer':
            assert source_rows(result, 'SYNTHETIC NULL FOOTER') == [['/Payload', 'null', None]]
    else:
        ineligible(result)


@pytest.mark.parametrize('field', ['TIN', 'PeriodYear'])
def test_provable_outer_identity_conflict_is_not_hidden_by_unsupported_current_form(field):
    envelope = json.loads(primary()['body'])
    if field == 'TIN':
        envelope['Payload']['CurrentFormData'] = None
        envelope['Payload']['TIN'] = '0000000000000002'
    else:
        form = current_form()
        form.pop('MainFormData')
        envelope['Payload']['CurrentFormData'] = wire(form)
        envelope['Payload']['PeriodYear'] = 2024
    with pytest.raises(native.JsonSafetyError):
        adapt([packet(envelope)])
