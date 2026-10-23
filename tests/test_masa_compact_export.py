"""Final saved layouts use synthetic rows and hand-checked reference positions."""
from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

import coretax
import coretax_exporter as exporter
from support_templates import section, run_for, source_appendix, APPENDICES
from support_pph21_dom import appendix_table


def saved(tmp_path, tax, sections):
    run = run_for(sections, model='Amendment 001')
    run['records'][0]['meta']['Jenis Pajak'] = tax
    before = deepcopy(run)
    reports = []
    paths = exporter.export_split_workbooks(run, tmp_path, mapping_reports=reports)
    coretax.verify_outputs(run, tmp_path)
    assert run == before
    path, = [p for p in paths if p != 'resume.xlsx']
    book = load_workbook(tmp_path/path)
    proof = exporter.read_workbook_proof(book)
    assert proof['records'][0]['meta']['Model SPT'] == 'Amendment 001'
    return book, reports


@pytest.mark.parametrize('form,header,first,total,columns', APPENDICES)
@pytest.mark.parametrize('count', [0, 2, 18])
def test_ppn_keeps_only_used_body_or_one_empty_slot_and_captured_total(tmp_path, form, header, first, total, columns, count):
    table, amounts = source_appendix(form, header, first, total, columns, count=count)
    # Every body value is a real synthetic row, including a literal zero.
    for values in table['rows'][:count]:
        for col in columns:
            values[col-1] = '0'
    book, reports = saved(tmp_path, 'PPN', [section(form, tables=[table])])
    try:
        ws = book[form]
        target = first + max(1, count)
        assert ws.cell(target, 1).value == 'TOTAL'
        assert ws.max_row == target
        assert [ws.cell(target, col).value for col in columns] == [float(Decimal(v.replace('.', '').replace(',', '.'))) for v in amounts]
        assert len([r for r in range(first,target) if ws.cell(r,1).value is not None]) == count
        assert all(r['complete'] for r in reports)
        if count:
            assert ws.cell(first, columns[0]).value == 0
    finally:
        book.close()


def test_ppn_induk_removes_only_two_redundant_spacers_and_moves_merges(tmp_path):
    book, _ = saved(tmp_path, 'PPN', [section('Induk')])
    try:
        ws=book['Induk']
        assert ws.max_row == 76
        assert ws['A8'].value == 'I.'
        assert 'A9:D9' in {str(x) for x in ws.merged_cells.ranges}
        assert all(ws.cell(7,c).value is None for c in range(1,9))
    finally:
        book.close()


@pytest.mark.parametrize('form,groups,counts,total,headers', [
    ('L-IA', ['LIST-IA'], [7], 12, [3]),
    ('L-IB', ['BPA1','BPA2'], [0,0], 11, [4,8]),
    ('L-II', ['BPA1','BPA2'], [0,0], 11, [4,8]),
    ('L-III', ['BP21','BP26'], [3,0], 13, [4,10]),
])
def test_pph21_matches_reference_body_capacity_headers_spacers_and_totals(tmp_path,form,groups,counts,total,headers):
    tables=[]
    for index,(group,count) in enumerate(zip(groups,counts)):
        table,_=appendix_table(form,group,count,with_totals=index==len(groups)-1)
        tables.append(table)
    book,reports=saved(tmp_path,'PPh Pasal 21/26',[section(form,tables=tables)])
    try:
        ws=book[form]
        assert ws.cell(total,1).value.startswith('TOTAL' if form=='L-II' else 'JUMLAH')
        assert ws.max_row == total + (0 if form=='L-II' else 2)
        assert [ws.cell(r,1).value for r in headers] == ['NO.']*len(headers)
        assert all(r['complete'] for r in reports)
        if form=='L-IA':
            assert ws['B4'].value == '0000000000000002'
            assert ws['H4'].value == 0.1 and ws['H4'].data_type == 'n'
            assert ws['H4'].number_format == '0.00%'
            assert all(ws.cell(11,c).value is None for c in range(1,15))
    finally:
        book.close()


@pytest.mark.parametrize('rate,valid,expected', [('0',True,0),('2,50',True,0.025),('100,00',True,1),
    (0,True,0),(2.5,True,0.025),(100,True,1),(None,True,None),('',True,None),
    ('12,3456789012345',True,0.123456789012345),
    ('1.5',False,'1.5'),('1,234',True,0.01234),
    ('1,000000000000000000000000000001',False,'1,000000000000000000000000000001'),
    pytest.param('0,' + '0' * 400 + '1', False, '0,' + '0' * 400 + '1', id='underflow'),
    pytest.param('0,' + '0' * 305 + '1', False, '0,' + '0' * 305 + '1', id='subnormal'),
    pytest.param('100,0000000000000', False, '100,0000000000000', id='source-precision-before-division'),
    ('999.999.999.999.999,99',False,'999.999.999.999.999,99')])
def test_lia_rate_fraction_is_strict_and_unsafe_source_stays_partial(tmp_path,rate,valid,expected):
    table,_=appendix_table('L-IA','LIST-IA',1)
    table['rows'][0][table['headers'].index('TARIF PAJAK (%)')]=rate
    book,reports=saved(tmp_path,'PPh Pasal 21/26',[section('L-IA',tables=[table])])
    try:
        assert book['L-IA']['H4'].value == expected
        if expected is not None:
            assert book['L-IA']['H4'].data_type == ('n' if valid else 's')
        assert reports[0]['complete'] is valid
        assert reports[0]['unmapped_tables'] == (0 if valid else 1)
        assert reports[0]['unmapped_rows'] == (0 if valid else 1)
    finally:
        book.close()


def test_pph21_accounting_uses_whole_display_without_rounding_saved_amounts(tmp_path):
    table,_=appendix_table('L-III','BP26',1,with_totals=True)
    book,_=saved(tmp_path,'PPh Pasal 21/26',[section('L-III',tables=[table])])
    try:
        money=[c for ws in book for row in ws for c in row if '"Rp"' in c.number_format]
        assert money
        assert all('.00' not in c.number_format.split(';')[0] for c in money)
        assert any(c.value==1234.5 for c in money)
    finally:
        book.close()


def test_lia_filters_empty_display_rows_but_keeps_zero_false_and_duplicates(tmp_path):
    table,_=appendix_table('L-IA','LIST-IA',1)
    row=table['rows'][0]
    row[0],row[2]=0,False
    table['rows']=[deepcopy(row),[None]*len(row),deepcopy(row)]
    table['body_row_count']=table['source_total']=3
    book,reports=saved(tmp_path,'PPh Pasal 21/26',[section('L-IA',tables=[table])])
    try:
        ws=book['L-IA']
        assert [ws.cell(r,1).value for r in (4,5)] == [0,0]
        assert [ws.cell(r,3).value for r in (4,5)] == [False,False]
        assert ws['A6'].value.startswith('JUMLAH')
        assert reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('form,groups,counts,total,starts', [
    ('L-IA',['LIST-IA'],[18],23,[4]),
    ('L-IB',['BPA1','BPA2'],[6,9],24,[5,14]),
    ('L-II',['BPA1','BPA2'],[6,9],24,[5,14]),
    ('L-III',['BP21','BP26'],[6,9],24,[5,14]),
])
def test_pph21_growing_groups_keep_all_rows_styles_and_footer_custody(tmp_path,form,groups,counts,total,starts):
    from copy import copy
    from pathlib import Path
    reference=load_workbook(Path(__file__).resolve().parents[1]/'resources/pph21_format.xlsx')
    tables=[appendix_table(form,group,count,with_totals=i==len(groups)-1)[0]
            for i,(group,count) in enumerate(zip(groups,counts))]
    book,reports=saved(tmp_path,'PPh Pasal 21/26',[section(form,tables=list(reversed(tables)))])
    try:
        ws=book[form]
        assert ws.max_row == total+(0 if form=='L-II' else 2)
        for group,count,start in zip(groups,counts,starts):
            assert [ws.cell(r,3).value for r in range(start,start+count)] == [f'SYNTHETIC {group} {i}' for i in range(count)]
            assert copy(ws.cell(start,2).border) == copy(reference[form].cell(5 if form!='L-IB' else 6,2).border)
        assert ws.cell(total,7).value == 7321.45
        assert set(map(str,ws.merged_cells.ranges)) == {'A1:N1'}
        assert all(r['complete'] for r in reports)
    finally:
        book.close()
        reference.close()
