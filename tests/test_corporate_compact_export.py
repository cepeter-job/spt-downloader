"""Clean corporate presentation through the public exporter; synthetic data only."""
from copy import copy, deepcopy
from pathlib import Path

from openpyxl import load_workbook
import pytest

import coretax
import coretax_exporter as exporter
from support_templates import section, source_run


RESOURCE = Path(__file__).resolve().parents[1] / 'resources/corporate_format.xlsx'


def table(form, header, rows, *, footer=None, spans=None, **extra):
    reference = load_workbook(RESOURCE)
    try:
        headers = [cell.value for cell in reference[form][header]]
        while headers and headers[-1] is None:
            headers.pop()
    finally:
        reference.close()
    return {'name': 'SYNTHETIC TABLE', 'headers': headers,
            'rows': rows + (footer or []), 'body_row_count': len(rows),
            'source_total': len(rows), 'complete': True,
            'footer': footer or [], 'footer_spans': spans or [], **extra}


def export(tmp_path, sections):
    run = source_run('PPh Badan', sections)
    before = deepcopy(run)
    reports = []
    paths = exporter.export_split_workbooks(run, tmp_path, mapping_reports=reports)
    assert run == before
    coretax.verify_outputs(run, tmp_path)
    name, = [name for name in paths if name != 'resume.xlsx']
    return load_workbook(tmp_path / name), reports


@pytest.mark.parametrize('count', [1, 2, 32])
def test_only_populated_rows_extend_related_party_table(tmp_path, count):
    rows = [['SYNTHETIC PARTY', '0000000000000002', 'COUNTRY', 'RELATION',
             'ACTIVITY', 'TRANSACTION', '0', 'METHOD', 'REASON'] for _ in range(count)]
    rows.insert(1, [''] * 9)
    book, reports = export(tmp_path, [section('L10-A', tables=[table('L10-A', 2, rows)])])
    try:
        sheet = book['L10-A']
        assert sheet.max_row == count + 2
        assert [sheet.cell(row, 2).value for row in range(3, count + 3)] == ['0000000000000002'] * count
        assert [sheet.cell(row, 7).value for row in range(3, count + 3)] == [0] * count
        assert reports[0]['complete'] and reports[0]['unmapped_rows'] == 0
    finally:
        book.close()


def test_empty_tables_keep_one_body_slot_and_intentional_separator(tmp_path):
    book, reports = export(tmp_path, [section('L2')])
    try:
        sheet = book['L2']
        assert sheet.max_row == 9
        assert sheet['A2'].value is not None and sheet['A7'].value is not None
        assert all(cell.value is None for cell in sheet[5])
        assert reports[0]['complete']
    finally:
        book.close()


def test_induk_uses_bound_fields_in_clean_two_column_layout(tmp_path):
    fields = [
        {'label': 'Tahun Pajak/Bagian Tahun Pajak', 'field': 'TaxYear', 'group': 'HEADER', 'value': '2025'},
        {'label': '1. NPWP', 'field': 'Tin', 'group': 'A. IDENTITAS WAJIB PAJAK', 'value': '0000000000000002'},
        {'label': '14. Angsuran PPh Pasal 25', 'field': '', 'group': 'E. PENGURANG PPh TERUTANG', 'value': '0'},
        {'label': 'SYNTHETIC UNBOUND', 'field': 'Unknown', 'group': 'HEADER', 'value': 'UNMAPPED SOURCE'},
    ]
    book, reports = export(tmp_path, [section('Induk', fields=fields, text='SYNTHETIC RAW TEXT')])
    try:
        sheet = book['Induk']
        assert sheet.max_column == 2
        rows = list(sheet.values)
        assert rows[:3] == [('INDUK', None), ('Tahun Pajak/Bagian Tahun Pajak', '2025'), ('1. NPWP', '0000000000000002')]
        assert rows[-1] == ('Angsuran PPh Pasal 25', '0')
        assert rows[3:6] == [(None, None)] * 3
        assert sheet.max_row == 7
        assert reports[0]['unmapped_fields'] == 1 and not reports[0]['complete']
        assert all('Data sumber tambahan' not in str(cell.value) for row in sheet for cell in row)
    finally:
        book.close()


def test_induk_duplicate_owner_is_unresolved_instead_of_choosing_value(tmp_path):
    fields = [{'label': '1. NPWP', 'field': 'Tin', 'group': 'A. IDENTITAS WAJIB PAJAK', 'value': value}
              for value in ('0000000000000002', '0000000000000003')]
    book, reports = export(tmp_path, [section('Induk', fields=fields)])
    try:
        assert book['Induk'].max_column == 2
        assert all(cell.value not in {field['value'] for field in fields}
                   for row in book['Induk'] for cell in row)
        assert reports[0]['unmapped_fields'] == 2 and not reports[0]['complete']
    finally:
        book.close()


def test_l8_clean_table_retains_footer_codes_and_safe_amounts(tmp_path):
    source = {'name': 'SYNTHETIC L8', 'group': 'PERHITUNGAN FASILITAS PENGURANGAN TARIF PPh BAGI WAJIB PAJAK BADAN DALAM NEGERI BERDASARKAN PASAL 31E AYAT (1) UNDANG-UNDANG PPh',
              'headers': ['NO.', 'DESKRIPSI', 'AMOUNT (Rupiah)'],
              'rows': [['01', 'SYNTHETIC DESCRIPTION', '1.234'], ['02', 'SYNTHETIC ZERO', '0'],
                       ['JUMLAH', '', '1.234']],
              'body_row_count': 2, 'source_total': 2, 'complete': True,
              'footer': [['JUMLAH', '', '1.234']], 'footer_spans': [[1, 1, 1]]}
    book, reports = export(tmp_path, [section('L8', tables=[source])])
    try:
        sheet = book['L8']
        assert sheet.max_row == 4 and sheet.max_column == 3
        assert sheet['A2'].value == '01' and sheet['C2'].value == 1234
        assert sheet['C3'].value == 0 and sheet['A4'].value == 'JUMLAH'
        assert sheet['C4'].value == 1234
        assert 'Rp' in sheet['C4'].number_format and '0.00' not in sheet['C4'].number_format
        assert reports[0]['complete']
    finally:
        book.close()


def test_l3_short_body_moves_all_three_captured_totals_and_merges(tmp_path):
    rows = [[str(i), 'SYNTHETIC WITHHOLDER', '0000000000000002', 'SYNTHETIC TAX',
             '1.000,00', '20,00', f'000{i}', '01-01-2025'] for i in range(60)]
    footer = [['JUMLAH', '8.000,00', '900,00', ''],
              ['KREDIT PAJAK LUAR NEGERI', '', '100,00', ''],
              ['JUMLAH KREDIT PAJAK', '', '1.000,00', '']]
    source = table('L3', 7, rows, footer=footer, spans=[[4, 1, 1, 2]] * 3)
    book, reports = export(tmp_path, [section('L3', tables=[source])])
    try:
        sheet = book['L3']
        assert sheet.max_row == 70
        assert sheet['C67'].value == '0000000000000002'
        assert sheet['E68'].value == 8000 and sheet['F68'].value == 900
        assert sheet['A69'].value == 'KREDIT PAJAK LUAR NEGERI'
        assert sheet['F69'].value == 100 and sheet['F70'].value == 1000
        assert 'A70:D70' in sheet.merged_cells and reports[0]['complete']
        for address in ('F4', 'G4', 'J4'):
            assert 'Rp' in sheet[address].number_format
        reference = load_workbook(RESOURCE)
        try:
            for target, original in [(8, 8), (68, 4662), (69, 4663), (70, 4664)]:
                for attribute in ('height', 'hidden', 'outlineLevel', 'collapsed'):
                    assert getattr(sheet.row_dimensions[target], attribute) == getattr(
                        reference['L3'].row_dimensions[original], attribute)
                for column in range(1, 9):
                    for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
                        assert copy(getattr(sheet.cell(target, column), attribute)) == copy(
                            getattr(reference['L3'].cell(original, column), attribute))
        finally:
            reference.close()
    finally:
        book.close()


@pytest.mark.parametrize('count', [1, 7])
def test_der_references_follow_shrinking_and_growing_capital_body(tmp_path, count):
    debt = table('L11-B', 11, [['0000000000000002', 'SYNTHETIC LENDER', 'RELATION'] + ['1,00'] * 13],
                 footer=[['JUMLAH'] + ['4,00'] * 12 + ['9.000,00']], spans=[[3] + [1] * 13])
    capital = table('L11-B', 17, [['SYNTHETIC CAPITAL'] + ['2,00'] * 13 for _ in range(count)],
                    footer=[['JUMLAH'] + ['5,00'] * 12 + ['3.000,00']], spans=[[1] * 14])
    book, reports = export(tmp_path, [section('L11-B', tables=[capital, debt])])
    try:
        sheet = book['L11-B']
        row = 21 + count
        assert sheet.max_row == 27 + count
        assert sheet.cell(row, 1).value == 'PERHITUNGAN DER'
        assert sheet.cell(row, 3).value == '=P13'
        assert sheet.cell(row + 1, 3).value == f'=N{18 + count}'
        assert sheet.cell(row, 4).value == f'=C{row}/C{row + 1}'
        assert sheet.cell(row, 4).number_format == '0.00'
        assert f'A{row}:A{row + 1}' in sheet.merged_cells
        assert f'D{row}:D{row + 1}' in sheet.merged_cells
        assert reports[0]['complete']
    finally:
        book.close()


def test_l1d_full_accounts_keep_all_three_totals_after_removing_identity_banner(tmp_path):
    reference = load_workbook(RESOURCE)
    tables = []
    try:
        for grid, header, first, count, label in [
            (1, 5, 6, 34, 'Laba (Rugi) Sebelum Pajak'),
            (2, 42, 43, 26, 'Jumlah Aset'),
            (3, 71, 72, 29, 'Jumlah Liabilitas dan Ekuitas'),
        ]:
            width = 10 if grid == 1 else 3
            rows = [['' if f'A{first+i}:B{first+i}' in reference['L1-D'].merged_cells else '0001',
                     f'SYNTHETIC ACCOUNT {i}'] + ['0'] * (width - 2) for i in range(count)]
            footer = [['0099', label] + ['1.234,50'] * (width - 2)]
            if grid == 1:
                for row in rows + footer:
                    row[8] = ''
            tables.append(table('L1-D', header, rows, footer=footer, spans=[[1] * width],
                                source_structure={'form': 'L1-D', 'component': f'rsh-citr-l1d-grid{grid}',
                                'groups': ['A. LAPORAN LABA RUGI' if grid == 1 else 'B. LAPORAN POSISI KEUANGAN']}))
    finally:
        reference.close()
    book, reports = export(tmp_path, [section('L1-D', tables=tables, text='SYNTHETIC RAW TEXT')])
    try:
        sheet = book['L1-D']
        assert sheet.max_row == 97 and sheet['A1'].value == 'KODE AKUN'
        for row in (36, 65, 97):
            assert sheet.cell(row, 1).value == '0099'
            assert sheet.cell(row, 3).value == 1234.5
        assert 'A2:B2' in sheet.merged_cells and reports[0]['complete']
    finally:
        book.close()


@pytest.mark.parametrize('form,last', [
    ('L1-D', 11), ('L2', 9), ('L3', 11), ('L4', 10), ('L6', 8),
    ('L7', 4), ('L9', 79), ('L10-A', 3), ('L10-C', 3), ('L11-A', 21),
    ('L11-B', 28), ('L11-C', 3),
])
def test_unused_body_capacity_is_removed_without_adding_unobserved_forms(tmp_path, form, last):
    book, reports = export(tmp_path, [section(form)])
    try:
        assert book.sheetnames == [form]
        assert book[form].max_row == last
        assert reports[0]['unmapped_fields'] == reports[0]['unmapped_tables'] == 0
    finally:
        book.close()


def test_l9_captured_totals_and_scalar_merges_follow_compaction(tmp_path):
    rows = [['0001', 'SYNTHETIC ASSET', '01/2025', '1.234,50', '1.000,00', 'GL', 'GL', '34,50', '']
            for _ in range(4)]
    first = table('L9', 4, rows, footer=[['TOTAL', '99,00', '']], spans=[[7, 1, 1]],
                  source_structure={'form': 'L9', 'component': 'rshshr-citr-l9-form-grid1',
                                    'groups': ['HARTA BERWUJUD', 'KELOMPOK 1']})
    other = table('L9', 210, [], footer=[['TOTAL', '12,50', '']], spans=[[7, 1, 1]],
                  source_structure={'form': 'L9', 'component': 'rshshr-citr-l9-form-grid1',
                                    'groups': ['HARTA BERWUJUD', 'KELOMPOK LAINNYA']})
    reference = load_workbook(RESOURCE)
    try:
        fields = [{'label': reference['L9'].cell(row, 1).value, 'value': '1.234,50'}
                  for row in (226, 259)]
    finally:
        reference.close()
    book, reports = export(tmp_path, [section('L9', tables=[other, first], fields=fields,
                                             text='SYNTHETIC RAW TEXT')])
    try:
        sheet = book['L9']
        assert sheet.max_row == 82 and sheet['A5'].value == '0001'
        assert sheet['H9'].value == 99 and '0.00' not in sheet['H9'].number_format
        assert sheet['H33'].value == 12.5 and '0.00' not in sheet['H33'].number_format
        assert sheet['H47'].value == 1234.5 and '0.00' not in sheet['H47'].number_format
        assert sheet['H80'].value == 1234.5 and '0.00' not in sheet['H80'].number_format
        assert 'A47:G47' in sheet.merged_cells and 'A80:G80' in sheet.merged_cells
        assert reports[0]['complete']
    finally:
        book.close()


def test_unsafe_l8_amount_remains_literal_and_mapping_partial(tmp_path):
    source = {'name': 'SYNTHETIC L8', 'group': 'PERHITUNGAN FASILITAS PENGURANGAN TARIF PPh BAGI WAJIB PAJAK BADAN DALAM NEGERI BERDASARKAN PASAL 31E AYAT (1) UNDANG-UNDANG PPh',
              'headers': ['NO.', 'DESKRIPSI', 'AMOUNT (Rupiah)'],
              'rows': [['01', 'SYNTHETIC AMOUNT', '9.999.999.999.999.999,99']],
              'body_row_count': 1, 'source_total': 1, 'footer': [], 'footer_spans': [], 'complete': True}
    book, reports = export(tmp_path, [section('L8', tables=[source])])
    try:
        assert book['L8']['C2'].value == source['rows'][0][2]
        assert book['L8']['C2'].data_type == 's'
        assert reports[0]['unmapped_rows'] == 1 and not reports[0]['complete']
    finally:
        book.close()


def test_accounting_defaults_to_zero_decimals_without_rounding_source(tmp_path):
    source = table('L10-A', 2, [['SYNTHETIC PARTY', '0000000000000002', 'COUNTRY', 'RELATION',
                               'ACTIVITY', 'TRANSACTION', '1.234,50', 'METHOD', 'REASON']])
    book, reports = export(tmp_path, [section('L10-A', tables=[source]), section('L3'), section('L9')])
    try:
        assert book['L10-A']['G3'].value == 1234.5
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if 'Rp' in cell.number_format and '*' in cell.number_format:
                        assert '0.00' not in cell.number_format
        assert all(report['complete'] for report in reports)
    finally:
        book.close()
