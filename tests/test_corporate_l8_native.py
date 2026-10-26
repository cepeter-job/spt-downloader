"""Synthetic L8 source/fixed-point admission and saved exports; runners only."""
from copy import deepcopy
import json

import pytest
from openpyxl import load_workbook

import coretax
from coretax_json import JsonSafetyError, adapt_corporate_detail
import support_json_corpus as corpus
from test_corporate_l7_native import FRAME, REQUEST, selected
from test_corporate_l11c_native import L11C_PINS
from test_corporate_l8_projection import ROOT, HEADERS, inputs

DETAIL = FRAME + '?view=true'
FLAGS = ('IsCarriedForward', 'IsNetIncomeReduction', 'IsInvestmentOtherIncome',
         'IsGrossIncomeVocational', 'IsGrossIncomeResearch')
ZEROS = ('TotalNetIncomeReduction', 'TotalGrossIncomeVocational',
         'TotalCarriedForward', 'TotalGrossIncomeResearch')
DEPENDENTS = ('CalculationOfFiscalLossCompensation', 'L13AListOfTaxFacilityInInvestment',
              'L13BListOfAdditionalGrossIncomeDeduction')
GROUP = ('PERHITUNGAN FASILITAS PENGURANGAN TARIF PPh BAGI WAJIB PAJAK BADAN DALAM NEGERI '
         'BERDASARKAN PASAL 31E AYAT (1) UNDANG-UNDANG PPh')


def envelope(model='Normal', *, net=12345, turnover=100000):
    form = corpus.current_form()
    amounts, parent = inputs(turnover=turnover, taxable=net-net % 1000)
    form[ROOT] = amounts
    form['MainFormData'].update(parent, TaxRate='TX03', NetIncomeBeforeTax=net,
        FiscalNetIncomeAfterTax=net, AccountingPeriodStart=1, AccountingPeriodEnd=12,
        **dict.fromkeys(FLAGS, 0), **dict.fromkeys(ZEROS, 0))
    return {'IsSuccessful': True, 'Payload': {'TIN': corpus.NPWP, 'PeriodYear': 2025, 'IsDraft': True,
        'PeriodMonthStart': 1, 'PeriodMonthEnd': 12, 'TaxRate': 'TX03',
        **dict.fromkeys(FLAGS, 0), 'ReturnSheetStatusCode': 'SUBMITTED',
        'ReturnSheetModel': 'NORMAL' if model == 'Normal' else 'AMENDMENT',
        'PreviousFormData': None, 'LastSubmittedFormData': None,
        'CurrentFormData': corpus.wire(form)}}


def lifetime(*, attempt=corpus.ATTEMPT):
    return {'kind': 'corporate-l8-spa-v1', 'attempt': attempt, 'record_id': corpus.RECORD_ID,
        'source_url': corpus.LIST_URL, 'detail_url': DETAIL, 'fresh_root': True,
        'stable_root': True, 'no_dependent_forms': True}


def packet(value=None, *, sequence=1):
    result = corpus.packet(envelope() if value is None else value, sequence,
                           request=dict(REQUEST))
    result['context'].update(frame_url=DETAIL, component_lifetime={
        'attempt': corpus.ATTEMPT, 'kind': 'corporate-submitted-spa-v1',
        'record_id': corpus.RECORD_ID, 'source_url': corpus.LIST_URL, 'detail_url': DETAIL,
        'transition_count': 1, 'same_document': True, 'source_owner': True, 'detail_owner': True},
        corporate_l8_lifetime=lifetime())
    return result


def bindings():
    return [{'name': 'Induk', 'paths': [['MainFormData']], 'grids': [],
        'primary': {'endpoint': corpus.PRIMARY_URL, 'request_json': dict(REQUEST)}},
        {'name': 'L8', 'paths': [[ROOT]], 'grids': [], 'projection': {
            'kind': 'corporate-l8-v1', 'builds': dict(L11C_PINS), 'route': DETAIL,
            'locale': 'id-ID', 'timezone': 'Asia/Jakarta', 'tables': [{
                'component': 'rsh-citr-l8-form-grid1', 'headers': list(HEADERS),
                'footer_spans': [[1, 1, 1]], 'source_structure': {
                    'form': 'L8', 'component': 'rsh-citr-l8-form-grid1', 'groups': [GROUP]}}]}}]


@pytest.mark.parametrize('model', ['Normal', 'Amendment 002'])
@pytest.mark.parametrize('net', [0, 12345, 999999999999999])
def test_l8_packet_fixed_point_reaches_original_provenance_and_saved_workbook(tmp_path, model, net):
    record, raw, observed = selected(model), packet(envelope(model, net=net)), bindings()
    record['meta']['Masa Pajak'] = '01122025'
    before = deepcopy((record, raw, observed))
    result = adapt_corporate_detail(iter([raw]), record, corpus.SCOPE, observed)
    assert result['eligible'] and result['coverage']['complete'], result['reason']
    assert (record, raw, observed) == before
    section, = [s for s in result['sections'] if s['name'] == 'L8']
    table, = section['tables']
    assert table['body_row_count'] == table['source_total'] == 8
    assert len(table['rows']) == 9 and table['complete'] is True
    assert table['rows'][1][2] == '100.000'
    assert table['rows'][3][2] == f'{net-net % 1000:,}'.replace(',', '.')
    values, = [p['values'] for p in result['provenance']['projections'] if p['form'] == 'L8']
    assert all(v['response_sequence'] == 1 for v in values)
    assert any(v['path'] == ['PeriodYear'] and v['container'] == ['Payload'] for v in values)
    for key in ('NetIncomeBeforeTax', 'FiscalNetIncomeAfterTax', *ZEROS, *FLAGS):
        assert any(v['path'] == ['MainFormData', key]
                   and v['container'] == ['Payload', 'CurrentFormData'] for v in values)
    assert any(v['path'] == ['IsInvestmentOtherIncome'] and v['container'] == ['Payload'] for v in values)
    record.update(capture_complete=True, sections=result['sections'])
    run = {'npwp': corpus.NPWP, 'year': 2025, 'year_basis': 'period',
        'inventory_complete': True, 'complete': True, 'errors': [], 'records': [record]}
    checkpoint = deepcopy(record['sections'])
    coretax.save_output(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    filename, = [n for n in run['output_files'] if n != 'resume.xlsx']
    assert model in filename and '01122025' in filename
    book = load_workbook(tmp_path / filename)
    try:
        sheet = book['L8']
        assert sheet.max_row == 10 and sheet['C3'].value == 100000
        assert sheet['C5'].value == net-net % 1000
        assert sheet['C10'].value == ((net-net % 1000)//100)*11
        assert {str(m) for m in sheet.merged_cells.ranges} == {'B2:C2', 'B4:C4', 'B7:C7'}
        assert sheet.column_dimensions['B'].width == 96.28515625
        assert sheet.page_setup.orientation == 'landscape'
        assert 'Rp' in sheet['C10'].number_format and '0.00' not in sheet['C10'].number_format
    finally:
        book.close()
    assert record['sections'] == checkpoint
    report, = [r for r in run['template_mapping'] if r['form'] == 'L8']
    assert report['complete'] and report['unmapped_rows'] == report['unmapped_tables'] == 0
    summary = load_workbook(tmp_path / 'resume.xlsx')
    try:
        targets = [c.hyperlink.target for s in summary for row in s for c in row if c.hyperlink]
        assert any(filename in target for target in targets)
        assert all(not t.startswith(('http:', 'https:', 'file:', '/', '\\')) for t in targets)
    finally:
        summary.close()


@pytest.mark.parametrize('owner', ['outer', 'current'])
@pytest.mark.parametrize('flag', FLAGS)
@pytest.mark.parametrize('value', [None, '0', False, 1, corpus.Number('0.0')])
def test_l8_visibility_flags_require_original_explicit_numeric_zero(owner, flag, value):
    source = envelope()
    root = json.loads(source['Payload']['CurrentFormData'])
    (source['Payload'] if owner == 'outer' else root['MainFormData'])[flag] = value
    source['Payload']['CurrentFormData'] = corpus.wire(root)
    result = adapt_corporate_detail([packet(source)], selected(), corpus.SCOPE, bindings())
    assert not result['eligible'] and result['sections'] == []


@pytest.mark.parametrize('owner', ['outer', 'current'])
@pytest.mark.parametrize('flag', FLAGS)
def test_l8_visibility_flags_cannot_be_inferred_when_missing(owner, flag):
    source = envelope()
    root = json.loads(source['Payload']['CurrentFormData'])
    del (source['Payload'] if owner == 'outer' else root['MainFormData'])[flag]
    source['Payload']['CurrentFormData'] = corpus.wire(root)
    result = adapt_corporate_detail([packet(source)], selected(), corpus.SCOPE, bindings())
    assert not result['eligible'] and result['sections'] == []


@pytest.mark.parametrize('key', ('NetIncomeBeforeTax', 'FiscalNetIncomeAfterTax', 'TaxableIncome', *ZEROS))
@pytest.mark.parametrize('value', [None, '0', False, -1, corpus.Number('1e1000000'), corpus.Number('-0'),
                                 corpus.Number('0.5'), corpus.Number('1000000000000000')])
def test_l8_original_main_unsafe_operands_keep_whole_record_fallback(key, value):
    source = envelope()
    root = json.loads(source['Payload']['CurrentFormData'])
    root['MainFormData'][key] = value
    source['Payload']['CurrentFormData'] = corpus.wire(root)
    result = adapt_corporate_detail([packet(source)], selected(), corpus.SCOPE, bindings())
    assert not result['eligible'] and result['sections'] == []


@pytest.mark.parametrize('key', ('NetIncomeBeforeTax', 'FiscalNetIncomeAfterTax', 'TaxableIncome', *ZEROS))
def test_l8_missing_zero_main_operands_cannot_be_defaulted(key):
    source = envelope(net=0)
    root = json.loads(source['Payload']['CurrentFormData'])
    del root['MainFormData'][key]
    source['Payload']['CurrentFormData'] = corpus.wire(root)
    result = adapt_corporate_detail([packet(source)], selected(), corpus.SCOPE, bindings())
    assert not result['eligible'] and result['sections'] == []


@pytest.mark.parametrize('key', DEPENDENTS)
@pytest.mark.parametrize('value', [None, {}, [], {'Total': 0}])
def test_l8_null_or_populated_dependency_roots_keep_strict_source_coverage(key, value):
    source = envelope()
    root = json.loads(source['Payload']['CurrentFormData'])
    root[key] = value
    source['Payload']['CurrentFormData'] = corpus.wire(root)
    result = adapt_corporate_detail([packet(source)], selected(), corpus.SCOPE, bindings())
    assert not result['eligible'] and result['sections'] == []


@pytest.mark.parametrize('defect', ['missing-strong', 'missing-both',
    'start', 'end', 'main-start', 'main-end', 'rate', 'main-rate', 'net-mismatch',
    'fiscal-mismatch', 'taxable-mismatch', 'previous', 'last-submitted', 'missing-parent',
    'turnover', 'captured-tax', 'no-induk', 'pin', 'timezone', 'header', 'owner', 'group', 'spans'])
def test_l8_unproved_branch_or_structure_is_explicit_whole_record_fallback(defect):
    source, observed = envelope(), bindings()
    root = json.loads(source['Payload']['CurrentFormData'])
    payload, main = source['Payload'], root['MainFormData']
    if defect == 'start': payload['PeriodMonthStart'] = 2
    elif defect == 'end': payload['PeriodMonthEnd'] = 11
    elif defect == 'main-start': main['AccountingPeriodStart'] = '01'
    elif defect == 'main-end': main['AccountingPeriodEnd'] = '12'
    elif defect == 'rate': payload['TaxRate'] = 'TX02'
    elif defect == 'main-rate': main['TaxRate'] = 'TX01'
    elif defect == 'net-mismatch': main['NetIncomeBeforeTax'] += 1000
    elif defect == 'fiscal-mismatch': main['FiscalNetIncomeAfterTax'] += 1000
    elif defect == 'taxable-mismatch': main['TaxableIncome'] = 12345
    elif defect == 'previous': payload['PreviousFormData'] = '{}'
    elif defect == 'last-submitted': payload['LastSubmittedFormData'] = '{}'
    elif defect == 'missing-parent': del payload['PreviousFormData']
    elif defect == 'turnover': root[ROOT]['GrossTurnOver'] = 4800000001
    elif defect == 'captured-tax': root[ROOT]['TotalIncomeTax'] += 1
    elif defect == 'no-induk': observed = observed[1:]
    elif defect == 'pin': observed[1]['projection']['builds'] = {}
    elif defect == 'timezone': observed[1]['projection']['timezone'] = 'UTC'
    elif defect == 'header': observed[1]['projection']['tables'][0]['headers'][1] = 'UNKNOWN'
    elif defect == 'owner': observed[1]['projection']['tables'][0]['component'] = 'foreign'
    elif defect == 'group': observed[1]['projection']['tables'][0]['source_structure']['groups'] = ['foreign']
    elif defect == 'spans': observed[1]['projection']['tables'][0]['footer_spans'] = [[3]]
    payload['CurrentFormData'] = corpus.wire(root)
    raw = packet(source)
    if defect.startswith('missing-') and defect != 'missing-parent':
        raw['context'].pop('corporate_l8_lifetime')
        if defect == 'missing-both': raw['context'].pop('component_lifetime')
    result = adapt_corporate_detail([raw], selected(), corpus.SCOPE, observed)
    assert not result['eligible'] and result['sections'] == []


@pytest.mark.parametrize('owner', ['outer', 'main'])
@pytest.mark.parametrize('with_binding', [False, True])
def test_l8_known_year_contradictions_stay_fatal_before_unsupported_fallback(owner, with_binding):
    source = envelope()
    if owner == 'outer':
        source['Payload']['PeriodYear'] = 2024
    else:
        root = json.loads(source['Payload']['CurrentFormData'])
        root['MainFormData']['TaxYear'] = 2024
        source['Payload']['CurrentFormData'] = corpus.wire(root)
    with pytest.raises(JsonSafetyError):
        adapt_corporate_detail([packet(source)], selected(), corpus.SCOPE,
                              bindings() if with_binding else bindings()[:1])


@pytest.mark.parametrize('defect', ['mixed', 'forged-record', 'false-fresh', 'missing-base'])
def test_l8_present_unbound_or_mixed_history_is_fatal_without_semantic_binding(defect):
    first, second = packet(), packet(sequence=2)
    if defect == 'mixed': del second['context']['corporate_l8_lifetime']
    elif defect == 'forged-record': second['context']['corporate_l8_lifetime']['record_id'] = 'foreign'
    elif defect == 'false-fresh': second['context']['corporate_l8_lifetime']['fresh_root'] = False
    else: del second['context']['component_lifetime']
    with pytest.raises(JsonSafetyError):
        adapt_corporate_detail([first, second], selected(), corpus.SCOPE, [])


@pytest.mark.parametrize('value', [None, corpus.Number('1e1000000'), corpus.Number('-0'), corpus.Number('1e-1000000')])
def test_l8_parent_exchange_control_remains_original_unused_evidence(value):
    source = envelope()
    root = json.loads(source['Payload']['CurrentFormData'])
    root[ROOT]['ExchangeRateValue'] = value
    source['Payload']['CurrentFormData'] = corpus.wire(root)
    result = adapt_corporate_detail([packet(source)], selected(), corpus.SCOPE, bindings())
    assert result['eligible'], result['reason']
    values, = [p['values'] for p in result['provenance']['projections'] if p['form'] == 'L8']
    item, = [v for v in values if v['path'] == [ROOT, 'ExchangeRateValue']]
    assert item['value'] == (None if value is None else str(value))
    assert item['rule'] == 'retained-unused-rupiah-source-control'
    section, = [s for s in result['sections'] if s['name'] == 'L8']
    assert section['tables'][0]['footer'][0][2] == '1.320'
