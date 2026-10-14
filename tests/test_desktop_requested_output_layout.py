"""Synthetic requested layout/publication regressions; GitHub-runner execution only."""
from collections import Counter
from copy import deepcopy
import posixpath
from pathlib import PurePosixPath

import pytest
from openpyxl import load_workbook

import coretax
import coretax_exporter as exporter
from spt_desktop import output
from test_desktop_selected_capture import selected_job
from test_financial_publication import PDF, financial_run
from test_split_output import fixture_run


ADMIN = {'Ringkasan', 'Daftar SPT', 'Audit'}
JANUARY = 'SPT Masa 21/SPT_21_Januari 2025.xlsx'
FEBRUARY = 'SPT Masa 21/SPT_21_Februari 2025.xlsx'


def values(workbook):
    return [cell.value for sheet in workbook for row in sheet for cell in row]


def selected_run(index):
    run = fixture_run()
    run['records'] = run['records'][index:index + 1]
    return run


def publish_first_period(tmp_path):
    run = selected_run(2)
    source = tmp_path / 'first-private-job'
    coretax.save_output(run, source)
    return output.publish_workbooks(run, source, tmp_path / 'Downloads', 'first-layout-job')


def second_period_source(tmp_path):
    run = selected_run(3)
    source = tmp_path / 'second-private-job'
    coretax.save_output(run, source)
    return run, source


def test_split_spt_workbooks_have_only_forms_and_administrative_sheets_stay_in_resume(tmp_path):
    files = exporter.export_split_workbooks(fixture_run(), tmp_path)
    for name, path in files.items():
        workbook = load_workbook(path)
        try:
            if name == 'resume.xlsx':
                assert workbook.sheetnames == ['Ringkasan', 'Daftar SPT', 'Audit']
            else:
                assert ADMIN.isdisjoint(workbook.sheetnames), 'Administrative sheets still leak into SPT files'
                assert all(sheet.sheet_state == 'visible' for sheet in workbook)
                caption = 'Induk'
                for directory, expected in [('SPT Masa 21/', 'Induk SPT PPh Pasal 21/26'),
                                            ('SPT Masa PPN/', 'Induk SPT PPN'),
                                            ('SPT Masa Unifikasi/', 'Induk SPT PPh Unifikasi')]:
                    if name.startswith(directory):
                        caption = expected
                assert all(sheet['A1'].value == caption for sheet in workbook)
        finally:
            workbook.close()


@pytest.mark.parametrize('tax,directory,prefix', [
    ('PPh Pasal 21/26', 'SPT Masa 21', 'SPT_21'),
    ('PPh Unifikasi', 'SPT Masa Unifikasi', 'SPT_Unifikasi'),
    ('PPN', 'SPT Masa PPN', 'SPT_PPN'),
])
def test_every_observed_periodic_tab_and_amendment_has_its_own_sheet(tmp_path, tax, directory, prefix):
    run = selected_run(2)
    normal = run['records'][0]
    normal['meta']['Jenis Pajak'] = tax
    names = ['Induk', 'Lampiran / Rincian']
    normal['sections'] = [
        {'name': name, 'fields': [{'label': 'Marker', 'value': 'SYNTHETIC NORMAL ' + name}],
         'tables': [], 'text': 'SYNTHETIC SOURCE CONTEXT ' + name}
        for name in names
    ]
    # A repeated capture of the same tab stays on that tab, with every raw value retained.
    normal['sections'].append({'name': names[1], 'fields': [
        {'label': 'Marker', 'value': 'SYNTHETIC NORMAL DUPLICATE'}], 'tables': [], 'text': ''})
    amendment = deepcopy(normal)
    amendment.update(record_id='synthetic-period-amendment', key='synthetic-period-amendment')
    amendment['meta']['Model SPT'] = 'Amendment 002'
    for section in amendment['sections']:
        for field in section['fields']:
            field['value'] = field['value'].replace('NORMAL', 'AMENDMENT')
    run['records'].append(amendment)
    files = exporter.export_split_workbooks(run, tmp_path)
    filename = f'{directory}/{prefix}_Januari 2025.xlsx'
    assert set(files) == {'resume.xlsx', filename}
    workbook = load_workbook(files[filename])
    try:
        forms = [sheet for sheet in workbook if sheet.title not in ADMIN]
        assert len(forms) == 4, 'Masa tabs are still flattened into one sheet per return'
        assert len({sheet.title.casefold() for sheet in forms}) == 4
        induk = {'PPh Pasal 21/26': 'Induk SPT PPh Pasal 21/26',
                 'PPN': 'Induk SPT PPN',
                 'PPh Unifikasi': 'Induk SPT PPh Unifikasi'}[tax]
        captions = [induk, names[1]]
        assert [sheet['A1'].value for sheet in forms] == captions + captions
        assert [sheet.title for sheet in forms] == ['Induk', 'Lampiran - Rincian',
                                                   'Induk (2)', 'Lampiran - Rincian (2)']
        proof = exporter.read_workbook_proof(workbook)
        assert [record['meta']['Model SPT'] for record in proof['records']] == ['Normal', 'Amendment 002']
        expected = Counter(field['value'] for record in run['records']
                           for section in record['sections'] for field in section['fields'])
        observed = Counter(cell.value for sheet in forms for row in sheet for cell in row
                           if isinstance(cell.value, str) and cell.value.startswith('SYNTHETIC ')
                           and 'SOURCE CONTEXT' not in cell.value)
        assert observed == expected
    finally:
        workbook.close()


def test_spreadsheet_banners_and_print_headers_are_removed_without_losing_source_data(tmp_path):
    run = selected_run(0)
    record = run['records'][0]
    record['detail_url'] = 'https://coretax.synthetic.invalid/returns/badan-normal'
    record['sections'][0].update(
        fields=[{'label': 'NPWP', 'value': run['npwp']},
                {'label': 'Pendapatan', 'value': '1.234,50'}],
        tables=[{'name': 'SYNTHETIC SOURCE TABLE', 'headers': ['Kode', 'Kode', 'Nilai'],
                 'rows': [['001', '001', '1.000,00'], ['001', '001', '1.000,00']], 'complete': True}],
        text='SYNTHETIC SOURCE CONTEXT — retained exactly')
    exporter.export_split_workbooks(run, tmp_path)
    workbook = load_workbook(tmp_path / 'SPT_Badan.xlsx')
    try:
        detail = next(sheet for sheet in workbook if sheet.title not in ADMIN)
        assert detail['A1'].value == 'Induk', 'Generic header still precedes the actual source form'
        cells = [cell for row in detail for cell in row]
        actual = [cell.value for cell in cells]
        assert all(value not in actual for value in exporter.detail_labels(run, record))
        assert 'Resume' not in actual and 'Daftar SPT' not in actual
        assert record['detail_url'] not in actual
        assert 'NPWP' in actual and run['npwp'] in actual and '1.234,50' in actual
        assert 'SYNTHETIC SOURCE CONTEXT — retained exactly' in actual
        assert any(list(row[:3]) == ['Kode', 'Kode', 'Nilai']
                   for row in detail.iter_rows(values_only=True))
        assert sum(value == '001' for value in actual) == 4
        assert sum(value == '1.000,00' for value in actual) == 2
        assert all(cell.data_type == 's' for cell in cells if cell.value in {run['npwp'], '001'})
        assert detail.tables and detail.freeze_panes and detail.print_area
        assert not detail.print_title_rows
        for kind in ('oddHeader', 'evenHeader', 'firstHeader', 'oddFooter', 'evenFooter', 'firstFooter'):
            for position in ('left', 'center', 'right'):
                assert not getattr(getattr(detail, kind), position).text
    finally:
        workbook.close()


def test_forms_only_file_keeps_nonvisual_scope_and_declared_pdf_proof(tmp_path):
    source = tmp_path / 'private-financial-job'
    run = financial_run(source)
    coretax.save_output(run, source)
    workbook = load_workbook(source / 'SPT_Badan.xlsx')
    try:
        assert ADMIN.isdisjoint(workbook.sheetnames)
        proof = workbook.properties.identifier
        assert isinstance(proof, str), 'Scope/PDF proof has not migrated out of visible Audit sheets'
        attachment = run['records'][0]['attachments'][0]
        assert run['npwp'] in proof and run['records'][0]['record_id'] in proof
        assert attachment['hash'] in proof and attachment['relative_path'] in proof
    finally:
        workbook.close()
    coretax.verify_outputs(run, source)
    final = output.publish_workbooks(run, source, tmp_path / 'Downloads', 'financial-layout')
    snapshot = output.existing_outputs(run, tmp_path / 'Downloads')
    assert snapshot['target'] == str(final)
    item = run['records'][0]['attachments'][0]
    assert (final / item['relative_path']).read_bytes() == PDF
    coretax.verify_outputs(run, final)


@pytest.mark.parametrize('mutation', ['missing', 'npwp', 'year', 'basis', 'record-id', 'digest'])
def test_forms_only_verification_rejects_changed_nonvisual_proof(tmp_path, mutation):
    source = tmp_path / 'private-financial-job'
    run = financial_run(source)
    coretax.save_output(run, source)
    path = source / 'SPT_Badan.xlsx'
    workbook = load_workbook(path)
    try:
        proof = workbook.properties.identifier
        assert isinstance(proof, str), 'Nonvisual proof is required before this tampering regression'
        replacements = {
            'npwp': (run['npwp'], '0000000000000002'),
            'year': ('2025', '2024'),
            'basis': ('period', 'submit'),
            'record-id': ('badan-normal', 'badan-foreign'),
            'digest': (run['records'][0]['attachments'][0]['hash'], '0' * 64),
        }
        if mutation == 'missing':
            workbook.properties.identifier = None
        else:
            original, changed = replacements[mutation]
            assert original in proof
            workbook.properties.identifier = proof.replace(original, changed)
        workbook.save(path)
    finally:
        workbook.close()
    with pytest.raises(ValueError):
        coretax.verify_outputs(run, source)


def test_new_period_automatically_replaces_only_resume_without_spt_overwrite_consent(tmp_path):
    final = publish_first_period(tmp_path)
    original = (final / JANUARY).read_bytes()
    old_resume = (final / 'resume.xlsx').read_bytes()
    run, source = second_period_source(tmp_path)
    snapshot = output.existing_outputs(run, tmp_path / 'Downloads')
    assert snapshot['files'] == ['resume.xlsx']
    second = output.publish_workbooks(run, source, tmp_path / 'Downloads', 'auto-resume',
                                      overwrite=False, expected_outputs=snapshot)
    assert second == final
    assert (final / JANUARY).read_bytes() == original
    assert (final / 'resume.xlsx').read_bytes() != old_resume
    assert (final / FEBRUARY).is_file()
    coretax.verify_outputs(run, final, allow_extra=True)


def test_worker_does_not_prompt_when_only_resume_already_exists(tmp_path, monkeypatch):
    old = selected_run(3)
    source = tmp_path / 'old-private-job'
    coretax.save_output(old, source)
    final = output.publish_workbooks(old, source, tmp_path / 'Downloads', 'old-february')
    untouched = (final / FEBRUARY).read_bytes()
    status, events, state, _ = selected_job(tmp_path, monkeypatch, ['21-january'])
    assert state['requests'] == ['selection'], 'resume.xlsx still requires an overwrite prompt'
    assert status == 0 and events[-1]['status'] == 'COMPLETE'
    assert events[-1]['output'] == str(final)
    assert (final / FEBRUARY).read_bytes() == untouched and (final / JANUARY).is_file()


def test_resume_auto_overwrite_does_not_authorize_replacing_selected_spt_files(tmp_path):
    final = publish_first_period(tmp_path)
    before = {path.relative_to(final).as_posix(): path.read_bytes() for path in final.rglob('*.xlsx')}
    run = selected_run(2)
    source = tmp_path / 'replacement-private-job'
    run['records'][0]['sections'][0]['fields'][0]['value'] = 'SYNTHETIC UNAPPROVED REPLACEMENT'
    coretax.save_output(run, source)
    with pytest.raises(FileExistsError, match='overwrite|consent'):
        output.publish_workbooks(run, source, tmp_path / 'Downloads', 'denied-spt', overwrite=False)
    assert {path.relative_to(final).as_posix(): path.read_bytes() for path in final.rglob('*.xlsx')} == before


def test_resume_auto_overwrite_keeps_snapshot_race_guards(tmp_path):
    final = publish_first_period(tmp_path)
    run, source = second_period_source(tmp_path)
    snapshot = output.existing_outputs(run, tmp_path / 'Downloads')
    resume_path = final / 'resume.xlsx'
    workbook = load_workbook(resume_path)
    workbook['Ringkasan']['B3'] = 'SYNTHETIC CONCURRENT CHANGE'
    workbook.save(resume_path)
    workbook.close()
    concurrent = resume_path.read_bytes()
    with pytest.raises(ValueError, match='changed|snapshot|consent'):
        output.publish_workbooks(run, source, tmp_path / 'Downloads', 'auto-race',
                                 overwrite=False, expected_outputs=snapshot)
    assert resume_path.read_bytes() == concurrent and not (final / FEBRUARY).exists()


def test_resume_auto_overwrite_refuses_foreign_scope_without_new_period_writes(tmp_path):
    final = publish_first_period(tmp_path)
    run, source = second_period_source(tmp_path)
    resume_path = final / 'resume.xlsx'
    workbook = load_workbook(resume_path)
    workbook['Ringkasan']['B4'] = '0000000000000002'
    workbook.save(resume_path)
    workbook.close()
    foreign = resume_path.read_bytes()
    with pytest.raises(ValueError, match='scope|NPWP'):
        output.publish_workbooks(run, source, tmp_path / 'Downloads', 'foreign-auto-resume', overwrite=False)
    assert resume_path.read_bytes() == foreign and not (final / FEBRUARY).exists()


def test_legacy_administrative_spt_workbooks_remain_scope_checked_during_upgrade(tmp_path):
    final = publish_first_period(tmp_path)
    legacy = selected_run(2)
    filename = JANUARY
    back = posixpath.relpath('resume.xlsx', str(PurePosixPath(filename).parent))
    exporter.export_workbook(legacy, final / filename, resume_file=back)
    workbook = load_workbook(final / filename)
    assert ADMIN.issubset(workbook.sheetnames), 'Explicit legacy fixture has lost its administrative proof'
    workbook.close()
    original = (final / filename).read_bytes()
    run, source = second_period_source(tmp_path)
    snapshot = output.existing_outputs(run, tmp_path / 'Downloads')
    assert snapshot['files'] == ['resume.xlsx']
    second = output.publish_workbooks(run, source, tmp_path / 'Downloads', 'legacy-new-period',
                                      overwrite=False, expected_outputs=snapshot)
    assert second == final and (final / filename).read_bytes() == original
    coretax.verify_outputs(run, final, allow_extra=True)
