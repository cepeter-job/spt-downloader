"""A selected refresh must not orphan the retained workbook index."""
from copy import deepcopy
from openpyxl import load_workbook
import coretax
from spt_desktop import output
from support_runs import fixture_run


def test_resume_indexes_retained_workbooks_after_selected_refresh(tmp_path):
    run = fixture_run()
    coretax.save_output(run, tmp_path/'first')
    final = output.publish_workbooks(run, tmp_path/'first', tmp_path/'Downloads', 'first')
    kept = {p.relative_to(final).as_posix():p.read_bytes() for p in final.rglob('*.xlsx')
            if p.name not in ('resume.xlsx','SPT_21_Januari 2025_Normal.xlsx')}
    selected = deepcopy(run)
    selected['records'] = selected['records'][2:3]
    coretax.save_output(selected, tmp_path/'second')
    published = output.publish_workbooks(selected, tmp_path/'second', tmp_path/'Downloads', 'second', overwrite=True)
    assert published == final
    assert all((final/name).read_bytes() == data for name,data in kept.items())
    wb = load_workbook(final/'resume.xlsx')
    try:
        assert wb['Ringkasan']['B7'].value == len(run['records'])
        assert wb['Daftar SPT'].max_row == 4 + len(run['records'])
        links = [wb['Daftar SPT'].cell(r,1).hyperlink for r in range(5,wb['Daftar SPT'].max_row+1)]
        assert {link.target for link in links} == {p.relative_to(final).as_posix() for p in final.rglob('*.xlsx') if p.name != 'resume.xlsx'}
    finally:
        wb.close()
    coretax.verify_outputs(selected, final, allow_extra=True)
