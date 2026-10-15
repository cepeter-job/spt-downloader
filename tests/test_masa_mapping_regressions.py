"""DOM-shaped export regressions, with synthetic data only."""
from copy import deepcopy
from pathlib import Path
import pytest
from openpyxl import load_workbook
import coretax_exporter as exporter
from test_pph21_template import run_for, section, table, row, IA_HEADERS, INDUK_HEADERS, assert_format

RESOURCES = Path(__file__).resolve().parents[1] / 'resources'


def exported(tmp_path, tax, sections):
    run = run_for(sections)
    run['records'][0]['meta']['Jenis Pajak'] = tax
    files = exporter.export_split_workbooks(run, tmp_path)
    return load_workbook(tmp_path / next(f for f in files if f != 'resume.xlsx'))


@pytest.mark.parametrize('tax,observed,reference,resource', [
    *[('PPN', v[0]+'-'+v[1], v, 'ppn_format.xlsx') for v in ['A1','A2','B1','B2','B3']],
    *[('PPh Unifikasi', a, b, 'unifikasi_format.xlsx') for a,b in
      [('DAFTAR-I','Daftar I'),('DAFTAR-II','Daftar II'),('LAMPIRAN-I','Lampiran I')]],
])
def test_observed_alias_selects_template(tmp_path, tax, observed, reference, resource):
    wb = exported(tmp_path, tax, [section(observed)])
    ref = load_workbook(RESOURCES / resource)
    try:
        assert wb.sheetnames == [observed]
        assert wb[observed]['A1'].value == ref[reference]['A1'].value
        assert_format(wb[observed]['A1'], ref[reference]['A1'])
        assert set(map(str, wb[observed].merged_cells.ranges)) == set(map(str, ref[reference].merged_cells.ranges))
    finally:
        wb.close(); ref.close()


def test_pph21_ordinal_table_bound_by_entire_row_signature(tmp_path):
    ref = load_workbook(RESOURCES / 'pph21_format.xlsx')
    values = [[str(ref['Induk'].cell(r,c).value) for c in range(1,4)] + ['17,00'] for r in range(12,18)]
    wb = exported(tmp_path, 'PPh Pasal 21/26', [section('Induk', tables=[table('Tabel 1', values, headers=INDUK_HEADERS)])])
    try:
        assert [wb['Induk'][f'D{r}'].value for r in range(12,18)] == ['17,00'] * 6
        assert wb['Induk']['D21'].value is None
        assert_format(wb['Induk']['D12'], ref['Induk']['D12'])
    finally:
        wb.close(); ref.close()


def test_ordinal_single_ambiguous_pph21_row_not_guessed(tmp_path):
    data = table('Tabel 1', [['1','PPH PASAL 21 YANG DIPOTONG','411121-100','17,00']], headers=INDUK_HEADERS)
    wb = exported(tmp_path, 'PPh Pasal 21/26', [section('Induk', tables=[data])])
    try:
        assert wb['Induk']['D12'].value is None
        assert wb['Induk']['D21'].value is None
        assert any(c.value == '17,00' for cells in wb['Induk'] for c in cells)
    finally:
        wb.close()


def test_lia_date_header_maps_body_but_not_short_footer(tmp_path):
    headers = list(IA_HEADERS)
    headers[4] += ' (dd-mm-yyyy)'
    source = row('SYNTHETIC CAPTURE ROW', ia=True)
    footer = ['JUMLAH SINTETIS','17,00']
    wb = exported(tmp_path, 'PPh Pasal 21/26', [section('L-IA', tables=[table('Tabel 1', [source, [], footer], headers=headers)])])
    try:
        assert [wb['L-IA'].cell(5,c).value for c in range(1,15)] == source
        assert wb['L-IA']['A6'].value is None
        assert any(c.value == footer[0] for cells in wb['L-IA'] for c in cells)
    finally:
        wb.close()


def test_pph21_group_evidence_not_ordinal(tmp_path):
    data = table('Tabel 1', [row('SYNTHETIC BPA2')])
    data['group'] = 'Daftar Bukti Pemotongan BPA2'
    wb = exported(tmp_path, 'PPh Pasal 21/26', [section('L-IB', tables=[data])])
    try:
        assert wb['L-IB']['C12'].value == 'SYNTHETIC BPA2'
        assert wb['L-IB']['C6'].value is None
    finally:
        wb.close()


def test_unifikasi_qualified_headers_preserve_raw(tmp_path):
    ref = load_workbook(RESOURCES / 'unifikasi_format.xlsx')
    sh = ref['Induk']
    headers = [sh['A10'].value,sh['B10'].value,'PAJAK PENGHASILAN / SETOR SENDIRI',
               'PAJAK PENGHASILAN / PEMOTONGAN / PEMUNGUTAN'] + [sh.cell(10,c).value for c in range(5,9)]
    headers[4] += ' Tooltip : The amount of income tax that has been paid'
    data = {'name':'Tabel 1','headers':headers,'rows':[['1','Pasal 4 Ayat 2','0','17,00','0','0','0','0']],'complete':True}
    before = deepcopy(data)
    wb = exported(tmp_path, 'PPh Unifikasi', [section('INDUK',tables=[data])])
    try:
        assert wb['INDUK']['D12'].value == '17,00'
        assert wb['INDUK']['C12'].value == '0'
        assert data == before
    finally:
        wb.close(); ref.close()


PPN_LABEL = 'PPN kurang atau (lebih) bayar (III.A - III.B - III.C - III.D)'


def test_ppn_scalar_field_maps_to_amount_slot(tmp_path):
    wb = exported(tmp_path, 'PPN', [section('INDUK', fields=[{'label':PPN_LABEL,'field':'','value':'-17,00'}])])
    try:
        assert wb['INDUK']['E45'].value == '-17,00'
        assert wb['INDUK']['E45'].data_type == 's'
    finally:
        wb.close()


def test_ppn_duplicate_scalar_not_arbitrarily_selected(tmp_path):
    fields = [{'label':PPN_LABEL,'field':'','value':v} for v in ['17,00','19,00']]
    wb = exported(tmp_path, 'PPN', [section('Induk', fields=fields)])
    try:
        assert wb['Induk']['E45'].value is None
        assert {c.value for cells in wb['Induk'] for c in cells} >= {'17,00','19,00'}
    finally:
        wb.close()


def test_ppn_repeated_label_uses_group_column_not_order(tmp_path):
    label = 'Jumlah PPN dan PPnBM yang dipungut'
    fields = [{'label':label,'field':'','value':'19,00','group':'VIII. PEMUNGUTAN PPN ATAU PPN DAN PPNBM OLEH PIHAK LAIN','column_label':'PPN (Rupiah)'},
              {'label':label,'field':'','value':'17,00','group':'VII. PEMUNGUTAN PPN ATAU PPN DAN PPNBM OLEH PEMUNGUT PPN','column_label':'PPN (Rupiah)'}]
    wb = exported(tmp_path, 'PPN', [section('Induk', fields=fields)])
    try:
        assert wb['Induk']['G69'].value == '17,00'
        assert wb['Induk']['G75'].value == '19,00'
        assert wb['Induk']['H69'].value is None
    finally:
        wb.close()


def test_ppn_appendix_dom_headers_and_extra_source_column(tmp_path):
    ref = load_workbook(RESOURCES/'ppn_format.xlsx')
    headers = [ref['A2'].cell(4,c).value for c in range(1,11)]
    ref.close()
    headers[4] += ' (dd-MM-yyyy)'
    headers.append('Masa Pajak')
    values = ['1','SYNTHETIC BUYER','0000000000000001','SYN-INVOICE','01-01-2025','17','19','23','0','SYN-REPLACED-INVOICE','EXTRA SYNTHETIC PROOF']
    wb = exported(tmp_path, 'PPN', [section('A-2',tables=[table('Tabel 1',[values],headers=headers)])])
    try:
        assert wb['A-2']['B5'].value == 'SYNTHETIC BUYER'
        assert wb['A-2']['E5'].value == '01-01-2025'
        assert wb['A-2']['H5'].value == '23'
        assert any(c.value == 'EXTRA SYNTHETIC PROOF' for cells in wb['A-2'] for c in cells)
    finally:
        wb.close()


def test_unifikasi_observed_group_and_extra_headers_do_not_invent_gross_income(tmp_path):
    ref = load_workbook(RESOURCES/'unifikasi_format.xlsx')
    headers = ['NO.','NIK/NPWP','NAMA','NOMOR BUKTI POTONG','TANGGAL BUKTI POTONG',
               'JENIS PAJAK','KODE OBJEK PAJAK','OBJEK PAJAK','DASAR PENGENAAN PAJAK (Rp)',
               'TINGKAT (%)','PAJAK PENGHASILAN','FASILITAS PERPAJAKAN',
               'UANG PERSEDIAAN / PEMBAYARAN LANGSUNG (UNTUK WP INSTANSI PEMERINTAH DENGAN DANA APBN)',
               'NITKU / NOMOR IDENTITAS SUBUNIT ORGANISASI','STATUS','KAP-KJS']
    values = ['1','0000000000000001','SYNTHETIC PAYEE','SYN-ID','01-01-2025','SYN-TAX',
              'SYN-CODE','SYN-OBJECT','17','1','19','SYN-FACILITY','SYN-PAYMENT','SYN-TKU','Normal','SYN-KAP']
    data = table('Tabel 1',[values],headers=headers)
    data['group'] = ref['Daftar I']['B3'].value
    wb = exported(tmp_path,'PPh Unifikasi',[section('DAFTAR-I',tables=[data])])
    try:
        assert wb['DAFTAR-I']['C5'].value == 'SYNTHETIC PAYEE'
        assert wb['DAFTAR-I']['J5'].value == '19'
        assert wb['DAFTAR-I']['M5'].value == 'SYN-TKU'
        assert wb['DAFTAR-I']['I5'].value is None, 'DPP is not evidence of gross income'
        assert any(c.value == '17' for cells in wb['DAFTAR-I'] for c in cells)
    finally:
        wb.close(); ref.close()


def test_resume_reports_unmapped_template_evidence(tmp_path):
    run = run_for([section('Induk', fields=[{'label':'UNSUPPORTED SYNTHETIC AMOUNT','field':'','value':'17'}])])
    exporter.export_split_workbooks(run,tmp_path)
    wb = load_workbook(tmp_path/'resume.xlsx')
    try:
        assert wb['Ringkasan']['B23'].value == 'PARTIAL'
        assert any(c.value == 'Pemetaan template' for cells in wb['Audit'] for c in cells)
    finally:
        wb.close()
