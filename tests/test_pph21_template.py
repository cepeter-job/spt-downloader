"""PPh21 format contracts use a sanitized reference and synthetic source only."""
from copy import copy, deepcopy
import hashlib
from pathlib import Path
import sys
import zipfile

import pytest
from openpyxl import load_workbook

import coretax
import coretax_exporter as exporter
from spt_desktop import runtime
from support_pph21_numeric import money_row

REFERENCE = Path(__file__).parent / 'fixtures' / 'pph21_format.xlsx'
JANUARY = 'SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx'
FORMS = ['Induk', 'L-IA', 'L-IB', 'L-II', 'L-III']
from support_templates import IA_HEADERS, assert_format, assert_layout, cells_with, run_for, section, table

INDUK_HEADERS = ['NO.', 'URAIAN', 'KAP-KJS', 'JUMLAH (Rp)']


def row(marker='SYNTHETIC EMPLOYEE', *, ia=False):
    values = ['1', '0000000000000002', marker, 'SYNTHETIC-WITHHOLDING-001',
              '02-01-2025', '21-100-01', '1.234,50', '123,45',
              'SYNTHETIC FACILITY', 'SYNTHETIC COUNTRY', '0000000000000002000001',
              '411121-100', 'SYNTHETIC STATUS']
    return values[:7] + ['10,00'] + values[7:] if ia else values


def export(run, folder, *, expected_file=JANUARY, reports=None):
    exporter.export_split_workbooks(run, folder, mapping_reports=reports)
    coretax.verify_outputs(run, folder)
    return load_workbook(folder / expected_file)


def test_empty_observed_forms_export_clean_template_without_sample_data(tmp_path):
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == 'a0bac29353d75f4af644ad72169e46a8b882dac7f9874debbbb086f32bb16c5e'
    with zipfile.ZipFile(REFERENCE) as archive:
        assert hashlib.sha256(archive.read('xl/styles.xml')).hexdigest() == 'f824574321958fe2e858a66a1b9d1e60f5b897dc5a73f1573db0f9f9b7b842d0'
        assert not any(p.startswith(('docProps/', 'xl/printerSettings/')) for p in archive.namelist())
    with_ref = load_workbook(REFERENCE)
    exported = export(run_for([section(form) for form in FORMS]), tmp_path)
    try:
        assert with_ref.sheetnames == FORMS
        assert exported.sheetnames == FORMS
        assert [ws['A1'].value for ws in exported] == [ws['A1'].value for ws in with_ref]
        for coordinate in ['D4', 'D5', 'D6', 'D7', 'D12', 'D21', 'D26', 'D35']:
            assert with_ref['Induk'][coordinate].value is None
            assert exported['Induk'][coordinate].value is None
        for name, regions in [('L-IA', [(5, 15)]), ('L-IB', [(6, 8), (12, 14)]),
                              ('L-II', [(5, 7), (11, 14)]), ('L-III', [(5, 7), (11, 13)])]:
            for first, last in regions:
                assert all(c.value is None for cells in with_ref[name].iter_rows(min_row=first, max_row=last) for c in cells)
                assert all(c.value is None for cells in exported[name].iter_rows(min_row=first, max_row=last) for c in cells)
    finally:
        with_ref.close()
        exported.close()


@pytest.mark.parametrize('form', FORMS)
def test_observed_form_preserves_every_reference_cell_format_and_layout(tmp_path, form):
    wb = export(run_for([section(form)]), tmp_path)
    reference = load_workbook(REFERENCE)
    try:
        assert wb.sheetnames == [form], 'Do not invent unobserved forms'
        assert wb[form]['A1'].value == reference[form]['A1'].value
        assert_layout(wb[form], reference[form])
        assert exporter.read_workbook_proof(wb)['records'][0]['record_id'] == 'synthetic-21-normal'
    finally:
        wb.close()
        reference.close()


def test_induk_header_fields_and_amounts_fill_only_their_reference_slots(tmp_path):
    fields = [{'label': label, 'field': label, 'value': value} for label, value in
              [('Masa Pajak', 'Januari 2025'), ('Tahun Pajak', '2025'),
               ('Status SPT', 'Pembetulan ke-2'), ('Status Penyampaian', 'Dilaporkan')]]
    tables = [table('I. PAJAK PENGHASILAN PASAL 21 YANG DIPOTONG',
                    [['1', 'PPH PASAL 21 YANG DIPOTONG', '411121-100', '1.234,50']], headers=INDUK_HEADERS),
              table('II. PAJAK PENGHASILAN PASAL 21 YANG DITANGGUNG OLEH PEMERINTAH',
                    [['1', 'PPH PASAL 21 YANG DIPOTONG', '411121-100', '21,00']], headers=INDUK_HEADERS),
              table('I.PAJAK PENGHASILAN PASAL 26 YANG DIPOTONG',
                    [['1', 'PAJAK PENGHASILAN PASAL 26 YANG DIPOTONG', '411127-100', '26,00']], headers=INDUK_HEADERS),
              table('II.PAJAK PENGHASILAN PASAL 26 YANG DITANGGUNG OLEH PEMERINTAH',
                    [['1', 'PAJAK PENGHASILAN PASAL 26 YANG DIPOTONG', '411127-100', '126,00']], headers=INDUK_HEADERS)]
    wb = export(run_for([section('Induk', fields=fields, tables=tables)], model='Pembetulan ke-2'),
                tmp_path, expected_file='SPT Masa 21/SPT_21_Januari 2025_Pembetulan ke-2.xlsx')
    reference = load_workbook(REFERENCE)
    try:
        sheet = wb['Induk']
        assert [sheet[f'D{r}'].value for r in [4, 5, 6, 7, 12, 21, 26, 35]] == [
            'Januari 2025', '2025', 'Pembetulan ke-2', 'Dilaporkan', 1234.5, 21, 26, 126]
        assert sheet['D12'].data_type == 'n'
        assert sheet['D5'].data_type == 's'
        assert_layout(sheet, reference['Induk'])
        assert sheet['C12'].value == '411121-100'
        assert sheet['C26'].value == '411127-100'
    finally:
        wb.close()
        reference.close()


def test_lia_rows_keep_column_order_duplicates_identifiers_and_numeric_amounts(tmp_path):
    source = row(ia=True)
    wb = export(run_for([section('L-IA', tables=[table('L-IA', [source, source], ia=True)])]), tmp_path)
    reference = load_workbook(REFERENCE)
    try:
        sheet = wb['L-IA']
        assert [sheet.cell(4, i).value for i in range(1, 15)] == IA_HEADERS
        assert [[sheet.cell(r, c).value for c in range(1, 15)] for r in [5, 6]] == [money_row(source, ia=True)] * 2
        assert sheet['B5'].data_type == 's' and sheet['L5'].data_type == 's'
        assert sheet['G5'].value == 1234.5 and sheet['G5'].data_type == 'n'
        assert sheet['E5'].value == '02-01-2025'
        assert_layout(sheet, reference['L-IA'])
    finally:
        wb.close()
        reference.close()


def test_reordered_visible_headers_are_mapped_without_changing_template_columns(tmp_path):
    source = row(ia=True)
    headers = list(reversed(IA_HEADERS))
    wb = export(run_for([section('L-IA', tables=[table('L-IA', [list(reversed(source))], headers=headers)])]), tmp_path)
    try:
        assert [wb['L-IA'].cell(5, c).value for c in range(1, 15)] == money_row(source, ia=True)
        assert [wb['L-IA'].cell(4, c).value for c in range(1, 15)] == IA_HEADERS
    finally:
        wb.close()


@pytest.mark.parametrize('form,first,second,first_row,second_row', [
    ('L-IB', 'BPA1', 'BPA2', 6, 12), ('L-II', 'BPA1', 'BPA2', 5, 11),
    ('L-III', 'BP21', 'BP26', 5, 11),
])
def test_two_groups_are_bound_to_captions_not_capture_order(tmp_path, form, first, second, first_row, second_row):
    first_values, second_values = row('SYNTHETIC FIRST'), row('SYNTHETIC SECOND')
    wb = export(run_for([section(form, tables=[table(second, [second_values]), table(first, [first_values])])]), tmp_path)
    reference = load_workbook(REFERENCE)
    try:
        sheet = wb[form]
        assert [sheet.cell(first_row, c).value for c in range(1, 14)] == money_row(first_values)
        assert [sheet.cell(second_row, c).value for c in range(1, 14)] == money_row(second_values)
        assert_layout(sheet, reference[form])
    finally:
        wb.close()
        reference.close()


def test_long_lia_table_grows_with_original_row_styles_without_losing_totals(tmp_path):
    rows = [row(f'SYNTHETIC EMPLOYEE {i}', ia=True) for i in range(18)]
    wb = export(run_for([section('L-IA', tables=[table('L-IA', rows, ia=True)])]), tmp_path)
    reference = load_workbook(REFERENCE)
    try:
        sheet = wb['L-IA']
        assert [[sheet.cell(r, c).value for c in range(1, 15)] for r in range(5, 23)] == [money_row(v, ia=True) for v in rows]
        assert sheet['A23'].value == 'JUMLAH PENGHASILAN BRUTO DAN PAJAK PENGHASILAN YANG DITANGGUNG OLEH PEMERINTAH'
        for r in range(16, 23):
            for c in range(1, 15):
                assert_format(sheet.cell(r, c), reference['L-IA'].cell(5, c))
            assert sheet.row_dimensions[r].height == reference['L-IA'].row_dimensions[5].height
        assert sheet.row_dimensions[23].height == reference['L-IA'].row_dimensions[16].height
        assert set(str(m) for m in sheet.merged_cells.ranges) == {'A1:N1'}
    finally:
        wb.close()
        reference.close()


def test_first_group_growth_moves_second_group_and_preserves_its_format(tmp_path):
    first = [row(f'SYNTHETIC BPA1 {i}') for i in range(6)]
    second = [row('SYNTHETIC BPA2')]
    wb = export(run_for([section('L-IB', tables=[table('BPA2', second), table('BPA1', first)])]), tmp_path)
    reference = load_workbook(REFERENCE)
    try:
        sheet = wb['L-IB']
        assert [[sheet.cell(r, c).value for c in range(1, 14)] for r in range(6, 12)] == [money_row(v) for v in first]
        assert sheet['B13'].value == 'BPA2'
        assert [sheet.cell(15, c).value for c in range(1, 14)] == money_row(second[0])
        assert sheet['A19'].value == reference['L-IB']['A16'].value
        for c in range(1, 15):
            assert_format(sheet.cell(15, c), reference['L-IB'].cell(12, c))
        assert sheet.row_dimensions[19].height == reference['L-IB'].row_dimensions[16].height
    finally:
        wb.close()
        reference.close()


def test_multiple_tables_for_same_group_append_rows_in_source_order(tmp_path):
    one, two = row('SYNTHETIC ONE'), row('SYNTHETIC TWO')
    wb = export(run_for([section('L-IB', tables=[table('BPA1', [one]), table('BPA1', [two])])]), tmp_path)
    try:
        assert [wb['L-IB'].cell(6, c).value for c in range(1, 14)] == money_row(one)
        assert [wb['L-IB'].cell(7, c).value for c in range(1, 14)] == money_row(two)
    finally:
        wb.close()


def test_unknown_fields_are_kept_in_checkpoint_and_reported_without_raw_display(tmp_path):
    value = '=HYPERLINK("https://invalid.example", "SYNTHETIC RAW")'
    data = section('Induk', fields=[{'label': 'SYNTHETIC UNKNOWN FIELD', 'field': 'unknown', 'value': value}],
                   tables=[table('SYNTHETIC UNKNOWN TABLE', [['0000000000000003', '987,65'],
                                                            ['0000000000000003', '987,65']], headers=['SYNTHETIC ID', 'SYNTHETIC VALUE'])],
                   text='SYNTHETIC PAGE CONTEXT')
    run = run_for([data])
    original = deepcopy(run)
    reports = []
    wb = export(run, tmp_path, reports=reports)
    reference = load_workbook(REFERENCE)
    try:
        assert run == original
        assert len(reports) == 1 and reports[0]['complete'] is False
        assert reports[0]['unmapped_fields'] == 1 and reports[0]['unmapped_rows'] == 2
        assert not cells_with(wb, value)
        assert not cells_with(wb, '0000000000000003')
        assert not cells_with(wb, '987,65')
        assert not cells_with(wb, 'SYNTHETIC PAGE CONTEXT')
        assert wb['Induk'].max_row == 35
        assert_layout(wb['Induk'], reference['Induk'])
    finally:
        wb.close()
        reference.close()


def test_ambiguous_group_is_reported_not_guessed_as_bpa1(tmp_path):
    run = run_for([section('L-IB', tables=[table('SYNTHETIC UNCLASSIFIED', [row()])])])
    original = deepcopy(run)
    reports = []
    wb = export(run, tmp_path, reports=reports)
    try:
        assert wb['L-IB']['B6'].value is None and wb['L-IB']['B12'].value is None
        assert run == original
        assert len(reports) == 1 and reports[0]['complete'] is False
        assert reports[0]['unmapped_tables'] == 1
        assert not cells_with(wb, 'SYNTHETIC EMPLOYEE')
        assert wb['L-IB'].max_row == 18
    finally:
        wb.close()


def test_duplicate_header_fields_preserve_both_checkpoint_values_and_report_ambiguity(tmp_path):
    fields = [{'label': 'Status SPT', 'field': 'status', 'value': 'SYNTHETIC FIRST STATUS'},
              {'label': 'Status SPT', 'field': 'status', 'value': 'SYNTHETIC SECOND STATUS'}]
    run = run_for([section('Induk', fields=fields)])
    original = deepcopy(run)
    reports = []
    wb = export(run, tmp_path, reports=reports)
    try:
        assert len(cells_with(wb, 'SYNTHETIC FIRST STATUS')) == 1
        assert run == original
        assert len(reports) == 1 and reports[0]['complete'] is False
        assert reports[0]['unmapped_fields'] == 1
        assert not cells_with(wb, 'SYNTHETIC SECOND STATUS')
    finally:
        wb.close()


def test_normal_amendment_and_next_period_keep_separate_layouts_proof_and_resume_links(tmp_path):
    normal = run_for([section('Induk'), section('L-IA', tables=[table('L-IA', [row('SYNTHETIC NORMAL', ia=True)], ia=True)])])
    amendment = deepcopy(normal['records'][0])
    amendment['key'] = amendment['record_id'] = 'synthetic-21-amendment'
    amendment['meta']['Model SPT'] = 'Pembetulan ke-2'
    amendment['sections'][1]['tables'][0]['rows'][0][2] = 'SYNTHETIC AMENDMENT'
    february = deepcopy(amendment)
    february['key'] = february['record_id'] = 'synthetic-21-february'
    february['meta']['Masa Pajak'] = 'Februari 2025'
    normal['records'] += [amendment, february]
    wb = export(normal, tmp_path)
    reference = load_workbook(REFERENCE)
    resume = load_workbook(tmp_path / 'resume.xlsx')
    amendment_file = 'SPT Masa 21/SPT_21_Januari 2025_Pembetulan ke-2.xlsx'
    february_file = 'SPT Masa 21/SPT_21_Februari 2025_Pembetulan ke-2.xlsx'
    amended = load_workbook(tmp_path / amendment_file)
    feb = load_workbook(tmp_path / february_file)
    try:
        assert wb.sheetnames == amended.sheetnames == feb.sheetnames == ['Induk', 'L-IA']
        assert wb['L-IA']['C5'].value == 'SYNTHETIC NORMAL'
        assert amended['L-IA']['C5'].value == feb['L-IA']['C5'].value == 'SYNTHETIC AMENDMENT'
        assert_layout(wb['L-IA'], reference['L-IA'])
        assert_layout(amended['L-IA'], reference['L-IA'])
        assert_layout(feb['L-IA'], reference['L-IA'])
        proofs = [exporter.read_workbook_proof(item) for item in [wb, amended, feb]]
        assert [[r['record_id'] for r in proof['records']] for proof in proofs] == [
            ['synthetic-21-normal'], ['synthetic-21-amendment'], ['synthetic-21-february']]
        assert [[r['meta']['Model SPT'] for r in proof['records']] for proof in proofs] == [
            ['Normal'], ['Pembetulan ke-2'], ['Pembetulan ke-2']]
        assert [(resume['Daftar SPT'].cell(r, 1).hyperlink.target,
                 resume['Daftar SPT'].cell(r, 1).hyperlink.location) for r in [5, 6, 7]] == [
            (JANUARY, "'Induk'!A1"), (amendment_file, "'Induk'!A1"),
            (february_file, "'Induk'!A1")]
    finally:
        wb.close()
        amended.close()
        reference.close()
        resume.close()
        feb.close()


def test_template_is_shipped_and_used_from_the_copied_worker_resource_tree(tmp_path):
    resources = {'coretax_pph21.py', 'resources/pph21_format.xlsx'}
    assert resources <= set(runtime.WORKER_FILES), 'Template renderer/resource missing from isolated worker'
    # Use the real copying boundary, then test public export in its app tree.
    # Provisioning another Python/network environment is outside this layout test.
    worker_root = tmp_path / 'worker'
    runtime._resources(Path(__file__).resolve().parents[1], worker_root)
    copied = worker_root / 'app'
    assert hashlib.sha256((copied / 'resources/pph21_format.xlsx').read_bytes()).hexdigest() == hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
    # Exercise the real separate worker imports, not the already-loaded repo modules.
    import subprocess
    program = ('from pathlib import Path; from coretax_exporter import export_split_workbooks; '
               'import json; from openpyxl import load_workbook; '
               'r=json.loads(Path("fixture-input.json").read_text()); '
               'export_split_workbooks(r,Path("out")); '
               'w=load_workbook(Path("out")/"SPT Masa 21/SPT_21_Januari 2025_Normal.xlsx"); '
               'assert w["Induk"]["A1"].value=="Induk SPT PPh Pasal 21/26"; w.close()')
    import json
    (copied / 'fixture-input.json').write_text(json.dumps(run_for([section('Induk')])), encoding='utf-8')
    result = subprocess.run([sys.executable, '-c', program], cwd=copied, capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr


def test_other_tax_types_keep_the_existing_generic_form_output(tmp_path):
    run = run_for([section('Induk', fields=[{'label': 'Record marker', 'field': 'marker', 'value': 'SYNTHETIC CORPORATE'}])])
    run['records'][0]['meta']['Jenis Pajak'] = 'PPh Pasal 22'
    paths = exporter.export_split_workbooks(run, tmp_path)
    coretax.verify_outputs(run, tmp_path)
    detail, = [path for name, path in paths.items() if name != 'resume.xlsx']
    wb = load_workbook(detail)
    try:
        assert wb['Induk']['A1'].value == 'Induk'
        assert wb['Induk'].freeze_panes == 'A3'
        assert wb['Induk'].page_setup.orientation == 'landscape'
        assert cells_with(wb, 'SYNTHETIC CORPORATE')
    finally:
        wb.close()
