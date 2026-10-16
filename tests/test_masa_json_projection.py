# Exercise the real safety adapter, native projections and workbook consumer.
from copy import deepcopy
import json
from pathlib import Path
import shutil
import subprocess
import sys

import pytest

from coretax_json import JsonSafetyError
from coretax_masa_json import adapt_masa_detail
from masa_projection_fixture import MAIN, native_case, response
from test_masa_mapping_regressions import exported
from test_pph21_template import assert_format
from openpyxl import load_workbook


def decode(case):
    return adapt_masa_detail(case['packets'], record=case['record'], scope=case['scope'], observed_forms=case['forms'])


@pytest.mark.parametrize('tax', ['PPN', 'PPh Unifikasi'])
@pytest.mark.parametrize('populated', [False, True])
def test_verified_packets_become_business_sections_not_corporate_leaf_tables(tax, populated):
    case = native_case(tax, populated=populated)
    before = deepcopy(case)
    result = decode(case)
    assert result['eligible'] is True, result['reason']
    assert result['coverage']['complete'] is True
    assert [s['name'] for s in result['sections']] == [f['name'] for f in case['forms']]
    assert len(result['coverage']['grids']) == (6 if tax == 'PPN' else 5)
    assert all(t['complete'] for s in result['sections'] for t in s['tables'])
    assert case == before
    entries = result['provenance']['bindings']
    assert entries and all(entry['sequence'] in {p['context']['sequence'] for p in case['packets']} for entry in entries)
    if tax == 'PPN':
        fields = {f['field']: f['value'] for f in result['sections'][0]['fields']}
        assert fields['ExportOfTangibleGoods'] == '17'
        assert fields['TotalVATandSTLGCollectedByOtherPartiesVAT'] == '35'
        assert any(e['path'] == ['A1FormListOfExportTaxableGoodsOrServices', 'Total']
                   and e['container_path'] == ['Payload', 'FormData'] and e['value'] == '17' for e in entries)
        footer_sequence = case['packets'][-1]['context']['sequence']
        assert any(e['sequence'] == footer_sequence and e['path'] == ['VAT'] and e['value'] == '35' for e in entries)
        table = result['sections'][2]['tables'][0]
        assert table['body_row_count'] == (13 if populated else 0)
        if populated:
            assert [row[1] for row in table['rows'][:13]] == ['SYNTHETIC BUYER'] * 13
            assert table['rows'][10][0] == '1', 'PPN row numbers restart on each source page'
    else:
        tree = result['sections'][0]['tables'][0]
        assert tree['rows'][0][3:5] == ['17', '19']
        assert tree['footer'][0][3] == '23'
        assert tree['rows'][tree['body_row_count']:] == tree['footer']
        grid = result['sections'][1]['tables'][0]
        assert grid['body_row_count'] == (6 if populated else 0)
        if populated:
            assert [r[0] for r in grid['rows']] == ['1.', '2.', '3.', '4.', '5.', '6.']
            assert [r[2] for r in grid['rows']] == ['SYNTHETIC PAYEE'] * 6


@pytest.mark.parametrize('tax', ['PPN', 'PPh Unifikasi'])
def test_native_business_dto_reaches_approved_template_slots_and_preserves_growth(tmp_path, tax):
    case = native_case(tax)
    result = decode(case)
    assert result['eligible'] is True, result['reason']
    workbook = exported(tmp_path, tax, result['sections'])
    resource = 'ppn_format.xlsx' if tax == 'PPN' else 'unifikasi_format.xlsx'
    reference = load_workbook(Path(__file__).resolve().parents[1] / 'resources' / resource)
    try:
        if tax == 'PPN':
            assert workbook['INDUK']['G75'].value == '35'
            assert [workbook['A-2'].cell(r,2).value for r in range(5,18)] == ['SYNTHETIC BUYER'] * 13
            assert workbook['A-2']['H5'].value == '132'
            assert_format(workbook['A-2']['B17'], reference['A2']['B5'])
        else:
            assert workbook['INDUK']['D12'].value == '17'
            assert workbook['INDUK']['E12'].value == '19'
            assert workbook['INDUK']['E27'].value == '23'
            assert [workbook['DAFTAR-I'].cell(r,3).value for r in range(5,11)] == ['SYNTHETIC PAYEE'] * 6
            assert workbook['DAFTAR-I']['J5'].value == '30'
            assert workbook['DAFTAR-I']['I5'].value is None, 'TaxBase is not gross income'
            assert_format(workbook['DAFTAR-I']['C10'], reference['Daftar I']['C5'])
    finally:
        workbook.close()
        reference.close()


@pytest.mark.parametrize('tax', ['PPN', 'PPh Unifikasi'])
@pytest.mark.parametrize('defect', ['missing_form', 'duplicate_form', 'changed_build', 'missing_packet',
                                    'unknown_control', 'dom_value', 'footer_spans'])
def test_incomplete_or_unverified_binding_never_grants_native_coverage(tax, defect):
    case = native_case(tax)
    if defect == 'missing_form': case['forms'].pop()
    elif defect == 'duplicate_form': case['forms'].append(deepcopy(case['forms'][-1]))
    elif defect == 'changed_build': case['forms'][0]['builds'][MAIN] = '0' * 64
    elif defect == 'missing_packet': case['packets'].pop(1)
    elif defect == 'unknown_control': case['forms'][0]['fields'][0]['field'] = 'UNREVIEWED_CONTROL'
    elif defect == 'dom_value': case['forms'][0]['fields'][0]['value'] = 'DOM-POISONED-AMOUNT'
    else: case['forms'][1]['tables'][0]['footer_spans'] = [[999]]
    result = decode(case)
    assert result['eligible'] is False and result['coverage']['complete'] is False
    assert result['sections'] == []


@pytest.mark.parametrize('tax', ['PPN', 'PPh Unifikasi'])
def test_bound_sequence_cannot_be_replaced_with_an_unrelated_packet(tax):
    case = native_case(tax)
    case['forms'][1]['tables'][0]['pages'][0]['sequences'] = [1]
    with pytest.raises(JsonSafetyError):
        decode(case)


@pytest.mark.parametrize('tax', ['PPN', 'PPh Unifikasi'])
def test_conflicting_grid_snapshots_are_fatal_even_without_observer_eligibility(tax):
    case = native_case(tax)
    first = case['packets'][1]
    changed = json.loads(first['body'])['Payload']
    changed['TotalRecords'] += 1
    endpoint = first['context']['url'].split('/returnsheet/api/')[1]
    case['packets'].append(response(tax, changed, endpoint, first['context']['request_json'], len(case['packets']) + 1))
    case['forms'] = []
    with pytest.raises(JsonSafetyError):
        decode(case)


@pytest.mark.parametrize('tax', ['PPN', 'PPh Unifikasi'])
def test_observed_grid_count_must_reconcile_with_response_count(tax):
    case = native_case(tax)
    case['forms'][1]['tables'][0]['source_total'] += 1
    with pytest.raises(JsonSafetyError):
        decode(case)


def test_allowlisted_worker_can_execute_native_masa_projection_in_isolation(tmp_path):
    from spt_desktop.runtime import WORKER_FILES
    root = Path(__file__).resolve().parents[1]
    worker = tmp_path / 'worker'
    for name in WORKER_FILES:
        target = worker / name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(root / name, target)
    case = native_case('PPN')
    for item in case['packets']:
        item['body'] = item['body'].decode()
    script = """import json,sys
sys.path.insert(0,sys.argv[1])
from coretax_masa_json import adapt_masa_detail
case=json.load(sys.stdin)
for item in case['packets']: item['body']=item['body'].encode()
r=adapt_masa_detail(case['packets'],record=case['record'],scope=case['scope'],observed_forms=case['forms'])
assert r['eligible'],r['reason']
assert len(r['sections'])==7
print('native Masa business DTO verified')
"""
    completed = subprocess.run([sys.executable, '-I', '-c', script, str(worker)], cwd=tmp_path,
        input=json.dumps(case), text=True, capture_output=True, timeout=60)
    assert completed.returncode == 0, completed.stderr or completed.stdout
    assert completed.stdout.strip() == 'native Masa business DTO verified'
