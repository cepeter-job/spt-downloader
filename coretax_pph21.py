"""Render observed PPh21/26 forms without restyling their reference cells."""
from copy import copy
from pathlib import Path
import re

from openpyxl import load_workbook


RESOURCE = Path(__file__).resolve().parent / 'resources' / 'pph21_format.xlsx'
APPENDIX_BLOCKS = {
    'l-ia': [(None, 4, 5, 15)],
    'l-ib': [('BPA1', 5, 6, 8), ('BPA2', 11, 12, 14)],
    'l-ii': [('BPA1', 4, 5, 7), ('BPA2', 10, 11, 14)],
    'l-iii': [('BP21', 4, 5, 7), ('BP26', 10, 11, 13)],
}
INDUK_BLOCKS = [(10, 11, 12, 17), (19, 20, 21, 21),
                (24, 25, 26, 31), (33, 34, 35, 35)]


def template_workbook(reserved_titles):
    workbook = load_workbook(RESOURCE)
    expected = ['Induk', 'L-IA', 'L-IB', 'L-II', 'L-III']
    if workbook.sheetnames != expected:
        workbook.close()
        raise ValueError('Resource format PPh21/26 tidak sesuai.')
    templates = {}
    reserved = {title.casefold() for title in reserved_titles}
    counter = 0
    for sheet in workbook.worksheets:
        templates[sheet.title.casefold()] = sheet
        # Internal copies must not collide with any literal observed form name.
        while f'_reference_{counter}'.casefold() in reserved:
            counter += 1
        sheet.title = f'_reference_{counter}'
        reserved.add(sheet.title.casefold())
        counter += 1
    return workbook, templates


def _key(value):
    return re.sub(r'[\W_]+', '', str(value or '').casefold())


def _write(sheet, row, column, value):
    cell = sheet.cell(row, column)
    cell.value = value if value is not None else ''
    if isinstance(cell.value, str):
        # Excel-injection protection must not replace the user's number format.
        cell.data_type = 's'


def _mapping(table, labels):
    headers = table.get('headers', [])
    keys = [_key(value) for value in headers]
    target = {_key(value): index + 1 for index, value in enumerate(labels) if value is not None}
    if (not keys or any(not value for value in keys) or len(set(keys)) != len(keys)
            or any(value not in target for value in keys)
            or any(len(row) > len(headers) for row in table.get('rows', []))):
        return None
    return [target[value] for value in keys]


def _copy_row_style(sheet, source_row, target_row, width):
    for column in range(1, width + 1):
        sheet.cell(target_row, column)._style = copy(sheet.cell(source_row, column)._style)
    if source_row in sheet.row_dimensions:
        dimension = copy(sheet.row_dimensions[source_row])
        dimension.index = target_row
        sheet.row_dimensions[target_row] = dimension


def _grow(sheet, boundary, count, style_row):
    if not count:
        return
    # openpyxl shifts cells, not row dimensions or merged ranges.
    merged = [copy(region) for region in sheet.merged_cells.ranges]
    if any(region.min_row < boundary <= region.max_row for region in merged):
        raise ValueError('Baris data PPh21/26 memotong merge template.')
    moving_merges = [region for region in merged if region.min_row >= boundary]
    for region in moving_merges:
        sheet.unmerge_cells(str(region))
    sheet.insert_rows(boundary, count)
    for row in sorted(list(sheet.row_dimensions), reverse=True):
        if row >= boundary:
            dimension = sheet.row_dimensions.pop(row)
            dimension.index = row + count
            sheet.row_dimensions[row + count] = dimension
    for region in moving_merges:
        region.shift(row_shift=count)
        sheet.merge_cells(str(region))
    for row in range(boundary, boundary + count):
        _copy_row_style(sheet, style_row, row, sheet.max_column)


def _induk(sheet, sections):
    fields_left, tables_left = [], []
    used_fields, used_rows = set(), set()
    slots = {_key(sheet.cell(row, 1).value): row for row in range(4, 8)}
    for section in sections:
        for field in section.get('fields', []):
            label = field.get('label') or field.get('field', '')
            row = slots.get(_key(label))
            if row is not None and row not in used_fields:
                _write(sheet, row, 4, field.get('value', ''))
                used_fields.add(row)
            else:
                fields_left.append(field)
        for table in section.get('tables', []):
            name = _key(table.get('name', ''))
            block = next((b for b in INDUK_BLOCKS if _key(sheet.cell(b[0], 1).value) == name), None)
            labels = [sheet.cell(block[1], c).value for c in range(1, 5)] if block else []
            mapping = _mapping(table, labels) if block else None
            if mapping is None:
                tables_left.append(table)
                continue
            pending = []
            # Fixed Induk rows are distinguished by section and visible labels,
            # never by the capture order or by recalculating taxpayer amounts.
            for values in table.get('rows', []):
                projected = {column: value for column, value in zip(mapping, values)}
                matches = [r for r in range(block[2], block[3] + 1)
                           if _key(projected.get(1)) == _key(sheet.cell(r, 1).value)
                           and _key(projected.get(2)) == _key(sheet.cell(r, 2).value)]
                if len(matches) != 1 or matches[0] in used_rows:
                    pending.append(values)
                    continue
                row = matches[0]
                for column, value in projected.items():
                    _write(sheet, row, column, value)
                used_rows.add(row)
            if pending:
                tables_left.append({**table, 'rows': pending})
            elif not table.get('complete', True):
                tables_left.append({**table, 'rows': []})
    return fields_left, tables_left


def _appendix(sheet, form, sections):
    blocks = APPENDIX_BLOCKS[form]
    grouped = [[] for _ in blocks]
    fields_left, tables_left = [], []
    for section in sections:
        fields_left.extend(section.get('fields', []))
        for table in section.get('tables', []):
            matches = set(re.findall(r'(?<![A-Za-z0-9])(BPA1|BPA2|BP21|BP26)(?![A-Za-z0-9])',
                                     table.get('name', ''), re.IGNORECASE))
            matches = {match.upper() for match in matches}
            candidates = [index for index, (group, _, _, _) in enumerate(blocks)
                          if group is None or matches == {group}]
            if len(candidates) != 1:
                tables_left.append(table)
                continue
            index = candidates[0]
            header = blocks[index][1]
            width = 14 if form == 'l-ia' else 13
            labels = [sheet.cell(header, column).value for column in range(1, width + 1)]
            mapping = _mapping(table, labels)
            if mapping is None:
                tables_left.append(table)
                continue
            grouped[index].extend((mapping, values) for values in table.get('rows', []))
            if not table.get('complete', True):
                tables_left.append({**table, 'rows': []})
    offset = 0
    for (_, _, first, last), rows in zip(blocks, grouped):
        first, last = first + offset, last + offset
        extra = max(0, len(rows) - (last - first + 1))
        _grow(sheet, last + 1, extra, first)
        for row, (mapping, values) in enumerate(rows, first):
            for column, value in zip(mapping, values):
                _write(sheet, row, column, value)
        offset += extra
    return fields_left, tables_left


def _raw_evidence(sheet, sections, fields, tables, style_row):
    texts = [section.get('text', '') for section in sections if section.get('text')]
    if not fields and not tables and not texts:
        return
    row = sheet.max_row + 2
    original_width = sheet.max_column
    def append(values):
        nonlocal row
        width = max(original_width, len(values))
        for column in range(1, width + 1):
            sheet.cell(row, column)._style = copy(sheet.cell(style_row, min(column, original_width))._style)
        if style_row in sheet.row_dimensions:
            dimension = copy(sheet.row_dimensions[style_row])
            dimension.index = row
            sheet.row_dimensions[row] = dimension
        for column, value in enumerate(values, 1):
            _write(sheet, row, column, value)
        row += 1
    # Unsupported or ambiguous content is still visible source evidence. It must
    # not be guessed into the wrong BPA/BP group, lost, or duplicated as amounts.
    append(['Data sumber tambahan / belum dipetakan'])
    for field in fields:
        append([field.get('label') or field.get('field', ''), field.get('value', '')])
    for table in tables:
        append([table.get('name', 'Tabel')])
        if table.get('headers'):
            append(table['headers'])
        for values in table.get('rows', []):
            append(values)
        if not table.get('complete', True):
            append(['PARTIAL: paginasi tidak selesai'])
    for text in texts:
        append(['Teks halaman / konteks sumber'])
        for start in range(0, len(text), 30000):
            append([text[start:start + 30000]])


def render_template_form(sheet, form, sections):
    key = form.casefold()
    if key == 'induk':
        fields, tables = _induk(sheet, sections)
        style_row = 4
    else:
        fields, tables = _appendix(sheet, key, sections)
        style_row = APPENDIX_BLOCKS[key][0][2]
    _raw_evidence(sheet, sections, fields, tables, style_row)
