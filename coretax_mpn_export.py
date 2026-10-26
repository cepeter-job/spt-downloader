"""Separate MPN workbook and physical scope/row proof; no raw Accounting archive."""
from copy import deepcopy
from decimal import Decimal
import json
import math
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.hyperlink import Hyperlink

from coretax_mpn import validate_mpn, validate_provenance, mpn_filename, MAX_BODY, MAX_EVIDENCE, MAX_RESPONSES
from coretax_resume_layout import summary_scope

COLUMNS = (
    ('period_label', 'Masa Pajak'), ('period_code', 'Kode Masa'),
    ('transaction_date', 'Tanggal Transaksi'), ('posting_date', 'Tanggal Posting'),
    ('transaction_number', 'Nomor Transaksi'), ('accounting_type', 'Jenis Pencatatan'),
    ('accounting_detail', 'Detail Jenis Pencatatan'), ('currency', 'Mata Uang'),
    ('currency_amount', 'Nilai Mata Uang'), ('exchange_rate', 'Nilai Tukar Transaksi Induk'),
    ('amount', 'Jumlah'), ('reference', 'Referensi Detail'),
    ('parent_reference', 'Referensi Induk'), ('revenue_code', 'KAP'), ('source', 'Sumber'),
)
MONEY = frozenset(('currency_amount', 'exchange_rate', 'amount'))


def _excel_number(lexeme):
    value = Decimal(lexeme)
    try:
        number = float(value)
    except (OverflowError, ValueError):
        return lexeme, False
    if (not math.isfinite(number) or 0 < abs(number) < sys.float_info.min
            or len(value.as_tuple().digits) > 15 or Decimal(str(number)) != value
            or value.is_zero() and value.is_signed()):
        return lexeme, False
    return int(value) if value == value.to_integral_value() else number, True


def _row_values(row):
    return [_excel_number(row[key])[0] if key in MONEY else row[key] for key, _ in COLUMNS]


def _proof(run):
    validate_mpn(run)
    value = run['mpn']
    return {'format': 'spt-downloader-mpn', 'version': 1,
            **{key: deepcopy(value[key]) for key in (
                'scope', 'taxpayer_id', 'detail_source', 'status', 'inventory_complete',
                'details_complete', 'headers', 'details', 'rows', 'audit')},
            'evidence': [{key: item[key] for key in ('sha256', 'bytes')} for item in value['evidence']],
            'provenance': deepcopy(value.get('provenance')), 'coverage': deepcopy(value.get('coverage'))}


def _audit_entries(proof):
    counts = proof.get('coverage') or {}
    return [('NPWP', proof['scope'][0]), ('Tahun', proof['scope'][1]),
            ('Dasar Tahun', proof['scope'][2]), ('Status MPN', proof['status']),
            ('Induk dibaca', proof['headers']), ('Detail dibaca', proof['details']),
            ('Mode diminta', proof['detail_source']), ('Pembayaran dipilih', len(proof['rows'])),
            ('Sumber DOM', sum(row['source'] == 'dom' for row in proof['rows'])),
            ('Sumber JSON', sum(row['source'] == 'json' for row in proof['rows'])),
            ('Detail dikecualikan', counts.get('excluded')), ('Detail unresolved', counts.get('unresolved')),
            ('Detail duplikat', counts.get('duplicates'))]


def export_mpn_workbook(run, path):
    from coretax_exporter import put, header, setup_sheet, band
    proof = _proof(run)
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'MPN'
    setup_sheet(sheet, [26, 16, 28, 28, 28, 28, 32, 14, 24, 24, 24, 36, 36, 16, 14], 'A5')
    band(sheet, 1, f"MPN | Masa/Tahun Pajak {run['year']}", len(COLUMNS))
    put(sheet, 2, 1, f"NPWP {run['npwp']} | {proof['status']} | non-cancelled; Amount < 0")
    header(sheet, 4, [label for _, label in COLUMNS])
    audit = wb.create_sheet('Audit')
    entries = _audit_entries(proof)
    audit_header = len(entries) + 3
    setup_sheet(audit, [28, 32, 28, 80], f'A{audit_header + 1}')
    band(audit, 1, 'MPN | Audit dan lingkup', 4)
    for row, (label, value) in enumerate(entries, 2):
        put(audit, row, 1, label)
        put(audit, row, 2, value)
    header(audit, audit_header, ['No. / Objek', 'Hasil', 'Field', 'Keterangan'])
    ar = audit_header + 1
    for ordinal, row in enumerate(proof['rows'], 5):
        for col, ((key, _), value) in enumerate(zip(COLUMNS, _row_values(row)), 1):
            cell = put(sheet, ordinal, col, value)
            if key in MONEY:
                _, safe = _excel_number(row[key])
                if safe:
                    places = min(15, max(0, -Decimal(row[key]).as_tuple().exponent))
                    cell.number_format = '#,##0' + ('.' + '0' * places if places else '')
                else:
                    for ac, av in enumerate((ordinal - 4, 'LITERAL / PRECISION', key,
                        'Presisi sumber dipertahankan sebagai teks; tidak dibulatkan.'), 1):
                        put(audit, ar, ac, av)
                    ar += 1
    for entry in proof['audit']:
        for col, value in enumerate(('MPN', 'PARTIAL', '', entry['message']), 1):
            put(audit, ar, col, value)
        ar += 1
    back = put(audit, ar + 1, 1, 'Resume')
    back.hyperlink = Hyperlink(ref=back.coordinate, target='resume.xlsx', location="'Ringkasan'!A1")
    back.font = Font(color='2463A0', underline='single')
    for ws in wb:
        ws.print_area = f'A1:{"O" if ws.title == "MPN" else "D"}{ws.max_row}'
    sheet.print_title_rows = '1:4'
    raw = json.dumps(proof, ensure_ascii=False, separators=(',', ':'))
    if len(raw.encode('utf-8')) > MAX_BODY:
        raise ValueError('MPN workbook proof exceeds its byte budget')
    wb.properties.identifier = raw
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def _unique(pairs):
    value = {}
    for key, item in pairs:
        if key in value:
            raise ValueError('MPN proof has duplicate keys')
        value[key] = item
    return value


def read_mpn_proof(workbook):
    raw = workbook.properties.identifier
    if raw is None:
        return None
    try:
        if not isinstance(raw, str) or len(raw.encode('utf-8')) > MAX_BODY:
            raise ValueError('MPN proof byte budget is invalid')
        proof = json.loads(raw, object_pairs_hook=_unique)
        if isinstance(proof, dict) and proof.get('format') != 'spt-downloader-mpn':
            return None
        required = {'format', 'version', 'scope', 'taxpayer_id', 'detail_source', 'status',
                    'inventory_complete', 'details_complete', 'headers', 'details', 'rows', 'audit', 'evidence',
                    'provenance', 'coverage'}
        if (not isinstance(proof, dict) or set(proof) != required or type(proof['version']) is not int
                or proof['version'] != 1 or not isinstance(proof['scope'], list) or len(proof['scope']) != 3):
            raise ValueError('MPN proof schema is unsupported')
        from coretax_exporter import validate_scope, validate_excel_value
        npwp, year, basis = proof['scope']
        validate_scope(npwp, year)
        value = {key: deepcopy(proof[key]) for key in required - {'format', 'provenance'}}
        value.update(required=True, evidence=[])
        validate_mpn({'npwp': npwp, 'year': year, 'year_basis': basis, 'complete': False, 'mpn': value})
        import re
        evidence = proof['evidence']
        if (not isinstance(evidence, list) or len(evidence) > MAX_RESPONSES
                or any(not isinstance(item, dict) or set(item) != {'sha256', 'bytes'}
                       or type(item['sha256']) is not str or not re.fullmatch(r'[0-9a-f]{64}', item['sha256'])
                       or type(item['bytes']) is not int or not 0 < item['bytes'] <= MAX_BODY for item in evidence)
                or sum(item['bytes'] for item in evidence) > MAX_EVIDENCE):
            raise ValueError('MPN proof original evidence declaration is invalid')
        if proof['provenance'] is not None:
            validate_provenance(proof['rows'], proof['provenance'], evidence)
        if basis not in ('period', 'submission') or workbook.sheetnames != ['MPN', 'Audit']:
            raise ValueError('MPN scope/sheets are unsupported')
        if any(ws.sheet_state != 'visible' for ws in workbook):
            raise ValueError('MPN has hidden sheets')
        sheet, audit = workbook['MPN'], workbook['Audit']
        if (sheet.max_row != 4 + len(proof['rows']) or sheet.max_column != len(COLUMNS)
                or [sheet.cell(4, col).value for col in range(1, len(COLUMNS) + 1)] != [label for _, label in COLUMNS]):
            raise ValueError('MPN physical headers/count differ from proof')
        for ordinal, row in enumerate(proof['rows'], 5):
            for col, value in enumerate(_row_values(row), 1):
                validate_excel_value(value, 'MPN proof cell')
                actual = sheet.cell(ordinal, col)
                # openpyxl reloads a saved empty literal as None.
                if (actual.value != (None if value == '' else value)
                        or isinstance(value, str) and value != '' and actual.data_type != 's'):
                    raise ValueError('MPN physical row differs from its original lexical proof')
        expected = _audit_entries(proof)
        if (any(audit.cell(row, 1).value != label or audit.cell(row, 2).value != value
                for row, (label, value) in enumerate(expected, 2)) or audit['B2'].data_type != 's'):
            raise ValueError('MPN physical audit scope differs from proof')
        return proof
    except (ValueError, TypeError, KeyError, AttributeError) as error:
        raise ValueError('MPN workbook scope/row proof cannot be verified') from error


def verify_mpn_workbook(run, path):
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if read_mpn_proof(wb) != _proof(run):
            raise ValueError('MPN workbook differs from the selected capture')
        return set(wb.sheetnames)
    finally:
        wb.close()


def resume_mpn_declaration(run):
    if 'mpn' in run:
        return {'filename': mpn_filename(run), 'status': run['mpn']['status'],
                'scope': [run['npwp'], run['year'], run['year_basis']]}
    value = run.get('retained_mpn')
    if value is not None:
        _validate_resume(value, run)
        return deepcopy(value)
    return None


def _validate_resume(value, run):
    if (not isinstance(value, dict) or set(value) != {'filename', 'status', 'scope'}
            or value['filename'] != f"MPN_{run['year']}.xlsx"
            or value['status'] not in ('PENDING', 'PARTIAL', 'COMPLETE')
            or value['scope'] != [run['npwp'], run['year'], run['year_basis']]):
        raise ValueError('MPN resume declaration is invalid')


def apply_resume_mpn(workbook, run, row):
    from coretax_exporter import put
    value = resume_mpn_declaration(run)
    if value is None:
        return
    sheet = workbook['Ringkasan']
    for col, content in enumerate(('MPN', value['status'], value['filename']), 1):
        cell = put(sheet, row, col, content)
        # Reuse safe adjacent reference presentation, including the status slot.
        from copy import copy
        for key in ('font', 'fill', 'border', 'alignment', 'protection'):
            setattr(cell, key, copy(getattr(sheet.cell(9, min(col, 2)), key)))
    cell = sheet.cell(row, 3)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=value['filename'], location="'MPN'!A1")
    cell.font = Font(name=cell.font.name, size=cell.font.sz, color='2463A0', underline='single')
    audit = workbook['Audit']
    ar = audit.max_row + 2
    from coretax_resume_layout import _style
    for col, content in enumerate(('Status / Objek', 'Induk dibaca', 'Detail dibaca',
                                    'Pembayaran dipilih', 'Hasil', 'Berkas / Keterangan'), 1):
        cell = put(audit, ar, col, content)
        _style(cell, audit.cell(3, col))
    captured = run.get('mpn')
    counts = (captured['headers'], captured['details'], len(captured['rows'])) if captured else (None,) * 3
    for col, content in enumerate(('MPN', *counts, value['status'], value['filename']), 1):
        cell = put(audit, ar + 1, col, content)
        _style(cell, audit.cell(4, col))
    cell = audit.cell(ar + 1, 6)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=value['filename'], location="'MPN'!A1")
    cell.font = Font(name=cell.font.name, size=cell.font.sz, color='2463A0', underline='single')
    for offset, item in enumerate(captured['audit'] if captured else (), ar + 2):
        for col, content in ((1, 'MPN'), (5, 'PARTIAL'), (6, item['message'])):
            cell = put(audit, offset, col, content)
            _style(cell, audit.cell(4, col))


def read_resume_mpn(workbook):
    scope = summary_scope(workbook)
    sheet = workbook['Ringkasan']
    rows = [row[0].row for row in sheet.iter_rows(min_row=14, min_col=1, max_col=1) if row[0].value == 'MPN']
    if not rows:
        return None
    if len(rows) != 1:
        raise ValueError('MPN resume declaration is duplicated')
    row = rows[0]
    cell = sheet.cell(row, 3)
    value = {'filename': cell.value, 'status': sheet.cell(row, 2).value,
             'scope': [scope['npwp'], scope['year'], scope['year_basis']]}
    _validate_resume(value, scope)
    link = cell.hyperlink
    if (not link or link.target != value['filename'] or link.location != "'MPN'!A1"
            or cell.data_type != 's'):
        raise ValueError('MPN resume link is missing or changed')
    return value
