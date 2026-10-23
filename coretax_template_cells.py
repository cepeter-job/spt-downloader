"""Shared style-preserving primitives for the three approved Masa templates."""
from copy import copy
import re

from openpyxl import load_workbook


def load_template_workbook(resource, expected_names, reserved_titles):
    workbook = load_workbook(resource)
    if workbook.sheetnames != expected_names:
        workbook.close()
        raise ValueError('Resource format SPT Masa tidak sesuai.')
    templates = {}
    reserved = {title.casefold() for title in reserved_titles}
    counter = 0
    for sheet in workbook.worksheets:
        templates[sheet.title.casefold()] = sheet
        while f'_reference_{counter}'.casefold() in reserved:
            counter += 1
        sheet.title = f'_reference_{counter}'
        reserved.add(sheet.title.casefold())
        counter += 1
    return workbook, templates


def _key(value):
    return re.sub(r'[\W_]+', '', str(value or '').casefold())


def _write(sheet, row, column, value):
    cell = sheet.cell(row, column)
    cell.value = value if value is not None else ''
    if isinstance(cell.value, str):
        # Excel-injection protection must not replace the user's number format.
        cell.data_type = 's'


def _header_key(value):
    # DOM text can include a date-format hint or a tooltip, neither is a column.
    # Raw source headers are retained; only the matching key is transformed.
    text = re.split(r'\bTooltip\s*:', str(value or ''), maxsplit=1, flags=re.I)[0]
    text = re.sub(r'\s*\(dd-mm-yyyy\)\s*$', '', text, flags=re.I)
    return _key(text)


def _mapping(table, labels, aliases=None):
    headers = table.get('headers', [])
    aliases = {_header_key(k): _header_key(v) for k, v in (aliases or {}).items()}
    target = {_header_key(value): index + 1 for index, value in enumerate(labels) if value is not None}
    keys = [_header_key(value) for value in headers]
    # A canonical reference header wins over an alias used by another form.
    keys = [key if key in target else aliases.get(key, key) for key in keys]
    if (not keys or any(not value for value in keys) or len(set(keys)) != len(keys)
            or any(value not in target for value in keys)
            or any(len(row) > len(headers) for row in table.get('rows', []))):
        return None
    return [target[value] for value in keys]


def _body_rows(table):
    """Do not turn footer totals or collapsed physical cells into detail rows."""
    rows = table.get('rows', [])
    count = table.get('body_row_count', len(rows))
    if type(count) is not int or not 0 <= count <= len(rows):
        return [], rows
    width = len(table.get('headers', []))
    body, pending = [], []
    for index, values in enumerate(rows):
        if index < count and len(values) == width and any(v is not None and str(v).strip() for v in values):
            body.append(values)
        else:
            pending.append(values)
    return body, pending


def _table_projection(table, labels, aliases=None, extra_headers=()):
    """Project explicitly supported DOM-only columns, retaining the original raw table."""
    headers = table.get('headers', [])
    target = {_header_key(v) for v in labels if v is not None}
    extras = {_header_key(v) for v in extra_headers} - target
    indices = [i for i, value in enumerate(headers) if _header_key(value) not in extras]
    projected = {**table, 'headers': [headers[i] for i in indices],
                 'rows': [[values[i] for i in indices if i < len(values)] for values in table.get('rows', [])]}
    if any(len(values) != len(headers) for values in table.get('rows', [])):
        return table, None, False
    changed = projected['headers'] != headers
    return projected, _mapping(projected, labels, aliases), changed


def _copy_row_style(sheet, source_row, target_row, width):
    for column in range(1, width + 1):
        sheet.cell(target_row, column)._style = copy(sheet.cell(source_row, column)._style)
    if source_row in sheet.row_dimensions:
        dimension = copy(sheet.row_dimensions[source_row])
        dimension.index = target_row
        sheet.row_dimensions[target_row] = dimension


def _grow(sheet, boundary, count, style_row):
    if not count:
        return
    # openpyxl shifts cells, not row dimensions or merged ranges.
    merged = [copy(region) for region in sheet.merged_cells.ranges]
    if any(region.min_row < boundary <= region.max_row for region in merged):
        raise ValueError('Baris data SPT Masa memotong merge template.')
    moving_merges = [region for region in merged if region.min_row >= boundary]
    for region in moving_merges:
        sheet.unmerge_cells(str(region))
    sheet.insert_rows(boundary, count)
    for row in sorted(list(sheet.row_dimensions), reverse=True):
        if row >= boundary:
            dimension = sheet.row_dimensions.pop(row)
            dimension.index = row + count
            sheet.row_dimensions[row + count] = dimension
    for region in moving_merges:
        region.shift(row_shift=count)
        sheet.merge_cells(str(region))
    for row in range(boundary, boundary + count):
        _copy_row_style(sheet, style_row, row, sheet.max_column)


def _compact_rows(sheet, deleted):
    """Remove only declared empty rows, preserving merges and row dimensions."""
    ordered = sorted(set(deleted))
    for row in ordered:
        if row < 1 or any(c.value is not None and str(c.value).strip() for c in sheet[row]):
            raise ValueError('Pemadatan template menargetkan baris berisi data.')
    runs = []
    for row in ordered:
        if runs and runs[-1][0] + runs[-1][1] == row:
            runs[-1][1] += 1
        else:
            runs.append([row, 1])
    for first, count in reversed(runs):
        last = first + count - 1
        moving = [copy(m) for m in sheet.merged_cells.ranges if m.max_row >= first]
        if any(m.min_row < first <= m.max_row or m.min_row <= last < m.max_row for m in moving):
            raise ValueError('Pemadatan body memotong merge template.')
        for region in moving:
            sheet.unmerge_cells(str(region))
        dimensions = [(r, sheet.row_dimensions.pop(r)) for r in list(sheet.row_dimensions) if r >= first]
        sheet.delete_rows(first, count)
        for row, dimension in dimensions:
            if row > last:
                dimension.index = row - count
                sheet.row_dimensions[row - count] = dimension
        for region in moving:
            if region.min_row > last:
                region.shift(row_shift=-count)
                sheet.merge_cells(str(region))


def _whole_accounting(sheet):
    # Change display precision only, never stored amounts or percentage formats.
    for row in sheet:
        for cell in row:
            if '"Rp"' in cell.number_format and '* ' in cell.number_format:
                cell.number_format = cell.number_format.replace('#,##0.00', '#,##0')


def _raw_evidence(sheet, sections, fields, tables, style_row):
    unresolved = [t for t in tables if t.get('rows') or not t.get('complete', True)]
    report = {'complete': not fields and not unresolved, 'unmapped_fields': len(fields),
              'unmapped_tables': len(unresolved),
              'unmapped_rows': sum(len(t.get('rows', [])) for t in unresolved)}
    texts = [section.get('text', '') for section in sections if section.get('text')]
    if not fields and not tables and not texts:
        return report
    row = sheet.max_row + 2
    original_width = sheet.max_column
    def append(values):
        nonlocal row
        width = max(original_width, len(values))
        for column in range(1, width + 1):
            sheet.cell(row, column)._style = copy(sheet.cell(style_row, min(column, original_width))._style)
        if style_row in sheet.row_dimensions:
            dimension = copy(sheet.row_dimensions[style_row])
            dimension.index = row
            sheet.row_dimensions[row] = dimension
        for column, value in enumerate(values, 1):
            _write(sheet, row, column, value)
        row += 1
    # Unsupported or ambiguous content is still visible source evidence. It must
    # not be guessed into the wrong BPA/BP group, lost, or duplicated as amounts.
    append(['Data sumber tambahan / belum dipetakan'])
    for field in fields:
        append([field.get('label') or field.get('field', ''), field.get('value', '')])
    for table in tables:
        append([table.get('name', 'Tabel')])
        if table.get('headers'):
            append(table['headers'])
        for values in table.get('rows', []):
            append(values)
        if not table.get('complete', True):
            append(['PARTIAL: paginasi tidak selesai'])
    for text in texts:
        append(['Teks halaman / konteks sumber'])
        for start in range(0, len(text), 30000):
            append([text[start:start + 30000]])
    return report
